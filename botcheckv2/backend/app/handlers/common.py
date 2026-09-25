"""Tách từ app/bot.py (refactor 2026-09-25) — giữ nguyên 100% logic, chỉ chia module theo tính năng."""

import asyncio
import html
import logging
import os
import random
import time
from typing import Optional

import httpx
import re
from aiogram import Bot, Dispatcher, Router, F
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.aiohttp import AiohttpSession
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandStart, CommandObject, StateFilter
from aiogram.types import (
    BotCommand, Message, URLInputFile, FSInputFile, BufferedInputFile,
    ReplyKeyboardMarkup, KeyboardButton,
    InlineKeyboardMarkup, InlineKeyboardButton, CallbackQuery,
)
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup



from .. import db
from .. import perms as _perms
from .. import notify_bot as _notify_bot
from .. import config as _config
from ..persist import SQLiteStorage, check_cache_get, check_cache_set, init_cache_db
from ..util import now, parse_check_args, vnd, vn_time_str
from ..tiktok import parse_username, fetch_tiktok_info, fmt_num, build_info_caption
from ..ig import (
    parse_ig_username, parse_ig_post_id,
    fetch_ig_info, fetch_ig_post_info,
    build_ig_info_caption, build_ig_video_caption
)
from ..fb import check_uid, build_fb_caption
from ..poller import poller

from .core import _SHOP_WALLET_HINT, _sub_active, log, manager

class TrackMenuState(StatesGroup):
    waiting_fb_uid = State()
    waiting_tiktok_username = State()
    waiting_tiktok_video = State()
    waiting_ig_username = State()
    waiting_ig_post = State()
    waiting_zalo_phone = State()
    waiting_list_name = State()
    waiting_list_add = State()
    waiting_alert_add = State()

class AccShopState(StatesGroup):
    waiting_for_stock_file = State()
    waiting_for_update_file = State()
    waiting_for_cover = State()
    waiting_for_review_comment = State()
    waiting_for_custom_qty = State()
    waiting_for_pick_uid = State()
    waiting_for_edit_uid = State()
    waiting_for_edit_value = State()

_STOCK_DIE_STATUSES = {"dead", "disabled", "checkpoint", "checkpoint_282", "checkpoint_956"}

DEFAULT_MAIL_APP_LINK = "https://www.swisstransfer.com/d/3ec3521b-a0b7-4dc9-bd78-53298af278ac"

async def _run_sheet_import(msg, cat_id: int, ncc_id: int, cost: int):
    """Chạy nhập kho từ Google Sheet (dùng chung cho lệnh /nhapkhosheet và nút bấm)."""
    c = db.acc_category_get(cat_id)
    if not c:
        await msg.answer("❌ Không có loại acc này. Xem: /kho")
        return
    if ncc_id and not db.supplier_get(ncc_id):
        await msg.answer("❌ Không có NCC này. Xem: /ncc")
        return
    sid = db.get_setting("sheet_import_id") or ""
    default_tab = db.get_setting("sheet_import_tab") or "NhapKho"
    if not sid:
        await msg.answer("📊 Chưa cài đặt sheet nhập kho. Dùng: <code>/setsheet &lt;link&gt;</code>",
                         parse_mode="HTML")
        return
    from .. import sheet_import as _si
    try:
        stall = (c["stall"] if "stall" in c.keys() else "") or "Acc Facebook"
    except Exception:
        stall = "Acc Facebook"
    tab = _si.tab_for_stall(stall, default_tab)
    wait = await msg.answer("⏳ Đang đọc Google Sheet...")
    try:
        if not await _si.ensure_tab(sid, tab):
            await wait.edit_text("❌ Không tạo/kiểm tra được tab Sheet. Thử lại sau.")
            return
        sheet_rows = await _si.read_unmarked(sid, tab)
    except Exception as e:
        await wait.edit_text(f"❌ Không đọc được sheet: {html.escape(str(e)[:200])}")
        return
    if not sheet_rows:
        await wait.edit_text("📭 Sheet không có dòng mới nào.\n"
                             "<i>Các dòng đã có đánh dấu ở cột Trạng thái sẽ bị bỏ qua.</i>",
                             parse_mode="HTML")
        return
    rows = []
    for rnum, cells in sheet_rows:
        rows.append({
            "uid": cells[0], "password": cells[1], "created_date": cells[2],
            "backup_mail": cells[3], "note": cells[4], "totp": cells[5],
            "cookie": cells[6], "token": cells[7], "_sheet_row": rnum,
        })
    extra = ""
    if ncc_id:
        extra += f" NCC #{ncc_id}"
    if cost:
        extra += f" vốn {vnd(cost)}/acc"
    try:
        await wait.edit_text(f"⏳ Đọc được <b>{len(rows)}</b> dòng mới từ sheet.{extra}\nĐang nhập kho...",
                             parse_mode="HTML")
    except Exception:
        pass
    await _import_stock_rows(rows, cat_id, ncc_id, cost, c, msg, wait,
                            sheet_ctx={"sheet_id": sid, "tab": tab})

async def process_tiktok_check(msg: Message, username: str):
    wait = await msg.answer(f"⏳ Đang kiểm tra <b>@{username}</b>...")
    try:
        info    = await fetch_tiktok_info(username)
        caption = build_info_caption(info)
        if info["avatar"]:
            try:
                await msg.answer_photo(
                    photo=URLInputFile(info["avatar"], filename="avatar.jpg"),
                    caption=caption,
                )
                await wait.delete()
                return
            except Exception:
                pass
        await wait.edit_text(caption, disable_web_page_preview=True)
    except ValueError as e:
        await wait.edit_text(f"❌ {e}")
    except httpx.TimeoutException:
        await wait.edit_text("⏰ Timeout! Thử lại sau.")
    except Exception as e:
        log.exception("Lỗi check @%s", username)
        await wait.edit_text(f"❌ Lỗi: {e}")

async def process_ig_check(msg: Message, username: str):
    wait = await msg.answer(f"⏳ Đang kiểm tra IG <b>@{username}</b>...")
    try:
        info    = await fetch_ig_info(username)
        caption = build_ig_info_caption(info)
        if info.get("avatar"):
            try:
                await msg.answer_photo(
                    photo=URLInputFile(info["avatar"], filename="avatar.jpg"),
                    caption=caption,
                )
                await wait.delete()
                return
            except Exception:
                pass
        await wait.edit_text(caption, disable_web_page_preview=True)
    except ValueError as e:
        await wait.edit_text(f"❌ {e}")
    except httpx.TimeoutException:
        await wait.edit_text("⏰ Timeout! Thử lại sau.")
    except Exception as e:
        log.exception("Lỗi check IG @%s", username)
        await wait.edit_text(f"❌ Lỗi: {e}")

async def process_fb_check(msg: Message, uid: str):
    wait = await msg.answer(f"⏳ Đang kiểm tra FB UID <b>{uid}</b>...")
    try:
        res = await check_uid(uid)
        try:
            db.log_check_stat(msg.from_user.id, "fb", res.get("uid") or uid,
                              res.get("status", ""), res.get("via", ""))
            db.add_check_history(msg.from_user.id, "fb", res.get("uid") or uid,
                                 res.get("status", ""))
        except Exception:
            pass
        caption = build_fb_caption(res)

        from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="📝 Ghi chú", callback_data=f"fb_note_{res['uid']}"),
                InlineKeyboardButton(text="👁 Theo dõi", callback_data=f"fb_track_{res['uid']}")
            ]
        ])

        avatar = res.get("avatar_url")
        # Chỉ gửi ảnh khi là avatar THẬT (acc dùng ảnh mặc định thì không gửi ảnh trắng)
        if res.get("has_real_avatar") and avatar:
            # Acc LIVE có avatar thật => gửi kèm ảnh avatar của acc
            try:
                await wait.delete()
            except Exception:
                pass
            try:
                await msg.answer_photo(photo=avatar, caption=caption, reply_markup=kb)
            except Exception:
                # Telegram không tải được ảnh => fallback về tin nhắn chữ
                await msg.answer(caption, disable_web_page_preview=True, reply_markup=kb)
        else:
            await wait.edit_text(caption, disable_web_page_preview=True, reply_markup=kb)
    except Exception as e:
        log.exception("Lỗi check FB %s", uid)
        await wait.edit_text(f"❌ Lỗi: {e}")

async def _notify_super_subadmin_action(actor_id: int, actor_name: str, text: str):
    """Báo chủ shop khi admin phụ thực hiện thao tác nhạy cảm (duyệt/từ chối rút)."""
    try:
        if _perms.is_super(actor_id):
            return
        sid = _perms.super_id()
        if sid and manager.bot:
            await manager.bot.send_message(
                sid,
                f"👤 <b>Admin phụ {html.escape(actor_name or '')}</b> "
                f"({actor_id}) vừa {text}",
                parse_mode="HTML")
    except Exception:
        pass

def status_caption(status: str, note: str, price, header: str = "") -> str:
    icon = "🟢" if status == "live" else "🔴"
    word = "LIVE" if status == "live" else "DIE"
    lines = []
    if header:
        lines.append(header)
    lines.append(f"{icon} Tài khoản đang <b>{word}</b>")
    if note:
        lines.append(f"Ghi chú: {note}")
    if price:
        lines.append(f"Giá: {vnd(price)}")
    return "\n".join(lines)

async def _send_card(bot: Bot, chat_id: int, uid: str, status: str, note, price,
                     avatar: str, header: str = ""):
    caption = status_caption(status, note, price, header)
    if avatar:
        try:
            await bot.send_photo(chat_id, photo=URLInputFile(avatar), caption=caption)
            return
        except Exception:
            pass
    await bot.send_message(chat_id, caption)

def _esc_list_name(s: str) -> str:
    return html.escape(s or "", quote=False)

async def _run_scanlist(chat_id: int, user_id: int, name: str, answer, bot):
    """Logic quét danh sách dùng chung cho /scanlist và nút trackmenu:lists_scan."""
    items = db.get_list_items(chat_id, name)
    if not items:
        await answer(f"❌ Danh sách <b>{_esc_list_name(name)}</b> trống hoặc không tồn tại.", parse_mode="HTML")
        return
    uids = [str(dict(i).get("value", "")) for i in items if dict(i).get("value")]
    if not uids:
        await answer(f"❌ Danh sách <b>{_esc_list_name(name)}</b> trống.", parse_mode="HTML")
        return
    wait = await answer(f"⏳ Đang check {len(uids)} UID trong <b>{_esc_list_name(name)}</b>...", parse_mode="HTML")
    if not await _ensure_bulk_credits_simple(chat_id, user_id, wait, len(uids), bot):
        return
    live, die, err = await check_uids_batch(uids, user_id=user_id)
    lines = [f"⚡ <b>KẾT QUẢ SCAN — {len(uids)} UID:</b>\n",
             f"🟢 Live: <b>{len(live)}</b>  |  🔴 Die: <b>{len(die)}</b>  |  ❓ Lỗi: <b>{len(err)}</b>\n"]
    if live:
        lines.append("🟢 <b>Live:</b>")
        lines += [f"• <code>{_esc_list_name(u)}</code>" for u in live[:30]]
        if len(live) > 30:
            lines.append(f"  …và {len(live) - 30} UID live nữa")
    if die:
        lines.append("\n🔴 <b>Die:</b>")
        lines += [f"• <code>{_esc_list_name(u)}</code>" for u in die[:30]]
        if len(die) > 30:
            lines.append(f"  …và {len(die) - 30} UID die nữa")
    if err:
        lines.append(f"\n❓ <b>Lỗi kiểm tra:</b> {len(err)} UID (thử lại sau)")
    await wait.edit_text("\n".join(lines), parse_mode="HTML")

async def check_uids_batch(uids: list, concurrency: int = 15, user_id: int = None):
    """Quét hàng loạt UID Facebook bằng Async Batching."""
    sem = asyncio.Semaphore(concurrency)
    live_list = []
    die_list = []
    error_list = []

    async def _check_one(uid):
        async with sem:
            try:
                res = await check_uid(uid)
                status = res.get("status", "error")
                if user_id:
                    try:
                        db.log_check_stat(user_id, "fb", uid, status, res.get("via", ""))
                    except Exception:
                        pass
                if status == "live":
                    live_list.append(uid)
                elif status == "die":
                    die_list.append(uid)
                else:
                    error_list.append(uid)
            except Exception:
                error_list.append(uid)

    tasks = [_check_one(u) for u in uids]
    await asyncio.gather(*tasks, return_exceptions=True)

    if user_id and (live_list or die_list):
        try:
            summary_msg = f"Check file: {len(live_list)} Live, {len(die_list)} Die"
            db.add_check_history(user_id, "fb", summary_msg, "live" if live_list else "die")
        except Exception:
            pass

    return live_list, die_list, error_list

async def _ensure_bulk_credits_simple(chat_id: int, user_id: int, wait: Message, n_uids: int, bot) -> bool:
    """Bản gọn của _ensure_bulk_credits dùng cho nút bấm (không cần Message đầy đủ)."""
    bulk_cost = int(db.get_setting("bulk_credit_cost", "0") or 0)
    raw_u = db.get_user(user_id)
    if bulk_cost > 0 and not _sub_active(dict(raw_u) if raw_u else None):
        need = n_uids * bulk_cost
        have = db.get_credits(user_id)
        if have < need:
            await wait.edit_text(
                f"❌ <b>Không đủ credits!</b>\n\n"
                f"Cần: <b>{need}</b> credits ({n_uids} UID × {bulk_cost})\n"
                f"Bạn có: <b>{have}</b> credits\n\n"
                f"Mua thêm bằng /muacredit",
                parse_mode="HTML",
            )
            return False
        db.consume_credits(user_id, need)
        await _maybe_low_credit_warn(bot, user_id, chat_id)
    return True

def _build_excel_bytes(sheets: dict) -> bytes:
    """Tạo file .xlsx từ dict {tên_sheet: (headers, rows)}. Trả về bytes."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    first = True
    for sheet_name, (headers, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet(title=sheet_name[:31])
        if first:
            ws.title = sheet_name[:31]
            first = False
        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill("solid", fgColor="4472C4")
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(horizontal="center")
        for r, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 28
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

async def _do_doitien(msg, tg_id: int):
    """Đổi hoa hồng sang số dư — tái sử dụng logic on_doitien."""
    c = db.get_conn()
    user = c.execute("SELECT ref_earnings, ref_withdrawn FROM tg_users WHERE tg_id=?", (tg_id,)).fetchone()
    if not user:
        await msg.answer("❌ Bạn chưa có tài khoản.")
        return
    d = dict(user)
    available = (d.get("ref_earnings") or 0) - (d.get("ref_withdrawn") or 0)
    if available <= 0:
        await msg.answer("❌ Bạn chưa có hoa hồng khả dụng để đổi.")
        return
    bonus = int(available * 0.10)
    total = available + bonus
    with db._lock:
        c.execute("UPDATE tg_users SET ref_withdrawn = ref_withdrawn + ?, balance = balance + ? WHERE tg_id=?",
                  (available, total, tg_id))
        c.commit()
    db.add_log("doitien", f"Đổi {vnd(available)} hoa hồng + bonus {vnd(bonus)}", tg_id)
    await msg.answer(
        f"✅ <b>ĐỔI HOA HỒNG THÀNH CÔNG!</b>\n\n"
        f"💰 Hoa hồng đổi: <b>{vnd(available)}</b>\n"
        f"🎁 Bonus +10%: <b>{vnd(bonus)}</b>\n"
        f"💵 <b>Đã cộng {vnd(total)} vào số dư.</b>",
        parse_mode="HTML")

async def _show_history(msg, tg_id: int, kind):
    rows = db.get_user_logs(tg_id, kind=kind, limit=15)
    if not rows:
        await msg.answer("📭 Bạn chưa có lịch sử check nào.")
        return
    label = {"fb": "FB", "tiktok": "TikTok", "ig": "IG", "zalo": "Zalo"}.get(kind or "", "tất cả")
    lines = [f"📜 <b>LỊCH SỬ CHECK — {label.upper()}</b>", "━━━━━━━━━━━━━━━"]
    for r in rows:
        ts = vn_time_str("%d/%m %H:%M", r["ts"])
        uid = html.escape(str(r["uid"] or ""))
        info = html.escape(str(r["message"] or ""))[:50]
        lines.append(f"• <code>{uid}</code> [{r['kind']}] {info} <i>({ts})</i>")
    await msg.answer("\n".join(lines), parse_mode="HTML")

async def _maybe_low_credit_warn(bot, tg_id: int, chat_id: int):
    """Sau khi trừ credits bulk: nếu còn dưới ngưỡng thì nhắc nạp + nút mua nhanh."""
    try:
        threshold = int(db.get_setting("low_credit_warn", "50") or 50)
    except Exception:
        threshold = 50
    if threshold <= 0:
        return
    try:
        left = db.get_credits(tg_id)
    except Exception:
        return
    if left >= threshold:
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⚡ Mua thêm credits", callback_data="buymore")]
    ])
    try:
        await bot.send_message(
            chat_id,
            f"⚠️ <b>Credits sắp hết!</b>\n\n"
            f"💠 Còn lại: <b>{left}</b> credits\n"
            f"<i>Nạp thêm để không bị gián đoạn nhé 👇</i>",
            parse_mode="HTML", reply_markup=kb,
        )
    except Exception:
        pass

def _trackmenu_main_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📘 Facebook UID", callback_data="trackmenu:fb"),
         InlineKeyboardButton(text="🎵 TikTok", callback_data="trackmenu:tiktok")],
        [InlineKeyboardButton(text="📸 Instagram", callback_data="trackmenu:ig"),
         InlineKeyboardButton(text="💬 Zalo", callback_data="trackmenu:zalo")],
        [InlineKeyboardButton(text="📁 Danh sách UID", callback_data="trackmenu:lists"),
         InlineKeyboardButton(text="🔔 Cảnh báo", callback_data="trackmenu:alerts")],
    ])

def _trackmenu_sub_kb(prefix: str):
    """Submenu chung: Thêm mới + Danh sách + Quay lại."""
    rows = []
    if prefix == "fb":
        rows = [
            [InlineKeyboardButton(text="➕ Thêm UID FB", callback_data="trackmenu:fb_add")],
            [InlineKeyboardButton(text="📋 Đang theo dõi", callback_data="trackmenu:fb_list")],
        ]
    elif prefix == "tiktok":
        rows = [
            [InlineKeyboardButton(text="➕ Theo dõi tài khoản", callback_data="trackmenu:tiktok_add"),
             InlineKeyboardButton(text="🎬 Theo dõi video", callback_data="trackmenu:tiktok_video_add")],
            [InlineKeyboardButton(text="📋 DS tài khoản", callback_data="trackmenu:tiktok_list"),
             InlineKeyboardButton(text="📋 DS video", callback_data="trackmenu:tiktok_vlist")],
        ]
    elif prefix == "ig":
        rows = [
            [InlineKeyboardButton(text="➕ Theo dõi tài khoản", callback_data="trackmenu:ig_add"),
             InlineKeyboardButton(text="🎞️ Theo dõi bài viết", callback_data="trackmenu:ig_post_add")],
            [InlineKeyboardButton(text="📋 DS tài khoản", callback_data="trackmenu:ig_list"),
             InlineKeyboardButton(text="📋 DS bài viết", callback_data="trackmenu:ig_plist")],
        ]
    elif prefix == "zalo":
        rows = [
            [InlineKeyboardButton(text="➕ Theo dõi SĐT", callback_data="trackmenu:zalo_add")],
            [InlineKeyboardButton(text="📋 DS Zalo", callback_data="trackmenu:zalo_list")],
        ]
    elif prefix == "lists":
        rows = [
            [InlineKeyboardButton(text="➕ Tạo danh sách", callback_data="trackmenu:lists_new")],
            [InlineKeyboardButton(text="📋 Xem danh sách", callback_data="trackmenu:lists_view")],
            [InlineKeyboardButton(text="➕ Thêm UID vào DS", callback_data="trackmenu:lists_add")],
        ]
    elif prefix == "alerts":
        rows = [
            [InlineKeyboardButton(text="➕ Thêm cảnh báo", callback_data="trackmenu:alert_add")],
            [InlineKeyboardButton(text="📋 DS cảnh báo", callback_data="trackmenu:alert_list")],
        ]
    rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:main")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

def _vip_limit_for(user) -> int:
    vip_level = 0
    try:
        vip_level = dict(user).get("vip_level", 0) if user else 0
    except Exception:
        pass
    try:
        return int(db.get_setting(f"vip{vip_level}_limit", [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]))
    except Exception:
        return [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]

def _is_cancel(text: str) -> bool:
    t = (text or "").strip().lower()
    return t in ("/huy", "/cancel", "hủy", "huỷ")

def _purchase_alert_text(buyer, orders: list) -> str:
    """Tin báo admin: thông tin khách + chi tiết acc đã mua."""
    e = html.escape
    uname = (getattr(buyer, "username", "") or "").strip()
    fname = (getattr(buyer, "full_name", "") or "").strip()
    lines = ["🛒 <b>ĐƠN MUA ACC MỚI</b>", "━━━━━━━━━━━━",
             f"👤 Khách: <b>{e(fname)}</b>"
             + (f" @{e(uname)}" if uname else "")
             + f" — <code>{buyer.id}</code>"]
    for o in orders:
        lines += ["",
                  f"🆔 Mã đơn: <code>{o['id']}</code>",
                  f"📦 {e(o.get('cat_name') or '')} | 💰 {vnd(o.get('price') or 0)}",
                  f"👤 UID: <code>{e(o.get('uid') or '')}</code>",
                  f"🔑 MK: <code>{e(o.get('password') or '')}</code>"]
        if o.get("created_date"):
            lines.append(f"📅 Ngày tạo: <code>{e(o['created_date'])}</code>")
        if o.get("backup_mail"):
            lines.append(f"📧 Mail thay: <code>{e(o['backup_mail'])}</code>")
        if o.get("totp"):
            lines.append(f"🔐 2FA: <code>{e(o['totp'])}</code>")
        if o.get("note"):
            lines.append(f"📝 Ghi chú: {e(o['note'])}")
        if o.get("cookie"):
            lines.append(f"🍪 Cookie: <code>{e(o['cookie'])}</code>")
        if o.get("token"):
            lines.append(f"🎫 Token: <code>{e(o['token'])}</code>")
    txt = "\n".join(lines)
    return txt[:3900]

def _buyer_shim(tg_id: int):
    """Object buyer toi thieu (id/username/full_name) khi khong co message User."""
    username, full_name = "", ""
    try:
        u = db.get_user(tg_id)
        if u:
            username = u["username"] or ""
            try:
                full_name = u["full_name"] or ""
            except Exception:
                full_name = ""
    except Exception:
        pass
    return type("BuyerShim", (), {"id": tg_id, "username": username,
                                  "full_name": full_name})()

async def _notify_purchase_admin(bot, buyer, orders: list):
    """Báo admin khi khách mua acc: qua bot thông báo riêng nếu có,
    fallback về bot chính. Admin phụ có quyền Đơn hàng cũng nhận tin."""
    if not orders:
        return
    try:
        text = _purchase_alert_text(buyer, orders)
        sent = await _notify_bot.manager.send_to_privileged(text)
        if not sent:
            await _notify_shop_admin(bot, text)
        try:
            priv = set(_notify_bot.privileged_ids())
        except Exception:
            priv = set()
        for tid in _perms.notify_extra_ids("orders"):
            if tid in priv:
                continue
            try:
                await bot.send_message(tid, text, parse_mode="HTML")
            except Exception:
                pass
    except Exception as e:
        log.warning("Báo admin đơn mua acc lỗi: %s", e)

async def _check_stock_live(stock_rows):
    """Check song song 1 loạt acc trong kho.
    Trả (live_ids, die_ids, unknown_ids) theo stock id.
    - live: status == "live"
    - die: dead/disabled/checkpoint... → cách ly
    - unknown: cookie_invalid/error/... → lỗi hạ tầng check, không kết luận được"""
    sem = asyncio.Semaphore(10)
    live_ids, die_ids, unknown_ids = [], [], []

    async def _one(r):
        async with sem:
            try:
                res = await check_uid(str(r["uid"]))
                st = str(res.get("status") or "").lower()
            except Exception:
                st = "error"
            return r["id"], st

    for sid, st in await asyncio.gather(*[_one(r) for r in stock_rows]):
        if st == "live":
            live_ids.append(sid)
        elif st in _STOCK_DIE_STATUSES:
            die_ids.append(sid)
        else:
            unknown_ids.append(sid)
    return live_ids, die_ids, unknown_ids

async def _notify_die_quarantine(bot, die_rows):
    """Báo admin các acc vừa bị cách ly vì check DIE trước khi giao."""
    if not die_rows:
        return
    try:
        by_cat = {}
        for r in die_rows:
            by_cat.setdefault(r.get("cat_id"), []).append(str(r.get("uid") or ""))
        lines = ["🗑 <b>CÁCH LY ACC DIE</b>",
                 "Shop vừa check trước khi giao — đã loại khỏi kho bán:"]
        for cid, uids in by_cat.items():
            try:
                c = db.acc_category_get(cid)
                name = c["name"] if c else f"#{cid}"
            except Exception:
                name = f"#{cid}"
            show = ", ".join(f"<code>{html.escape(u)}</code>" for u in uids[:10])
            more = f" <i>(+{len(uids) - 10})</i>" if len(uids) > 10 else ""
            lines.append(f"📦 <b>{html.escape(name)}</b>: {show}{more}")
        lines.append("\n<i>Dọn hẳn: /xoadie [id_loại]</i>")
        text = "\n".join(lines)
        sent = await _notify_bot.manager.send_to_privileged(text)
        if not sent:
            await _notify_shop_admin(bot, text)
    except Exception as e:
        log.warning("Báo admin cách ly acc die lỗi: %s", e)

async def _sell_live_stock(bot, tg_id: int, price_total: int, qty: int, pick_fn,
                          check_live: bool = True):
    """Bán qty acc nhưng CHỈ giao acc đã check LIVE (nếu check_live=True).
    pick_fn(exclude_ids) -> list[{"id","uid","cat_id"}] (ứng viên AVAILABLE).
    - check_live=True: Acc check DIE → cách ly khỏi kho + báo admin.
    - check_live=False (gian hàng không phải FB): giao thẳng, không check.
    - Check lỗi hạ tầng (không kết luận được) → không giao, không cách ly.
    Trả (orders, fail): orders = list[dict order] hoặc None;
    fail ∈ {"short" (thiếu hàng live), "infra" (lỗi check), "race" (bị mua mất)}."""
    need = max(1, int(qty))
    seen, live_rows = set(), []
    infra = False
    for _ in range(4):
        cands = [r for r in pick_fn(seen) if r["id"] not in seen]
        if not cands:
            break
        seen.update(r["id"] for r in cands)
        if check_live:
            l_ids, d_ids, u_ids = await _check_stock_live(cands)
            if d_ids:
                db.acc_stock_quarantine(d_ids)
                dset = set(d_ids)
                await _notify_die_quarantine(bot, [r for r in cands if r["id"] in dset])
            if u_ids:
                infra = True
            idmap = {r["id"]: r for r in cands}
            for sid in l_ids:
                if len(live_rows) < need:
                    live_rows.append(idmap[sid])
        else:
            for r in cands:
                if len(live_rows) < need:
                    live_rows.append(r)
        if len(live_rows) >= need:
            break
    if len(live_rows) < need:
        return None, ("infra" if infra else "short")
    # Bán đúng các acc đã check live (nguyên tử theo từng loại)
    orders = []
    groups = {}
    for r in live_rows[:need]:
        groups.setdefault(r["cat_id"], []).append(r["id"])
    per = price_total // need if need else price_total
    left = price_total
    items = list(groups.items())
    for gi, (cid, sids) in enumerate(items):
        share = left if gi == len(items) - 1 else per * len(sids)
        left -= share
        sold = db.acc_sell_stock_ids(cid, tg_id, share, sids)
        if not sold:
            return None, "race"
        orders += [dict(db.acc_get_order(oid)) for oid, _ in sold]
    return orders, None

async def _sell_uid_items(bot, tg_id, priced):
    """Bán các UID cụ thể khách đã chọn trong giỏ (check LIVE từng acc nếu sạp FB).
    priced: list[(cat_dict, stock_dict, price)].
    Trả (delivered_orders, failed[(cat, stock, price)])."""
    delivered, failed = [], []
    groups = {}
    for cat, stock, price in priced:
        groups.setdefault(cat["id"], (cat, []))[1].append((stock, price))
    for cat_id, (cat, lst) in groups.items():
        rows = [s for s, _ in lst]
        price_of = {s["id"]: p for s, p in lst}
        bad = set()
        if _cat_live_check(cat_id):
            l_ids, d_ids, u_ids = await _check_stock_live(rows)
            if d_ids:
                db.acc_stock_quarantine(d_ids)
                try:
                    await _notify_die_quarantine(
                        bot, [s for s in rows if s["id"] in set(d_ids)])
                except Exception:
                    pass
            bad = set(d_ids) | set(u_ids)
        for s, p in lst:
            if s["id"] in bad:
                failed.append((cat, s, p))
        ok_rows = [s for s in rows if s["id"] not in bad]
        if not ok_rows:
            continue
        total = sum(price_of[s["id"]] for s in ok_rows)
        sold = db.acc_sell_stock_ids(cat_id, tg_id, total,
                                     [s["id"] for s in ok_rows])
        if not sold:
            for s, p in lst:
                if s["id"] not in bad:
                    failed.append((cat, s, p))
            continue
        delivered += [dict(db.acc_get_order(oid)) for oid, _ in sold]
    return delivered, failed

async def _notify_shop_admin(bot, text: str):
    """Báo cho admin qua bot chính."""
    try:
        aid = db.get_setting("admin_tg_id", "")
        if aid:
            await bot.send_message(int(aid), text, parse_mode="HTML")
    except Exception:
        pass

async def _notify_admin_photo(bot, file_id: str, caption: str):
    """Gui anh bang chung cho admin qua bot chinh."""
    targets = []
    try:
        aid = db.get_setting("admin_tg_id", "")
        if aid:
            targets.append(int(aid))
    except Exception:
        pass
    try:
        gid = db.get_setting("admin_tg_group_id", "")
        if gid:
            targets.append(int(gid))
    except Exception:
        pass
    for t in targets:
        try:
            await bot.send_photo(t, file_id, caption=caption, parse_mode="HTML")
        except Exception:
            pass

async def _notify_admin_smart(bot, text: str, perm: str = None):
    """Báo admin: ưu tiên bot thông báo riêng, fallback về bot chính.
    perm: nếu có, gửi thêm cho các admin phụ đang giữ quyền đó
    (tránh trùng người đã nhận qua kênh privileged)."""
    try:
        sent = await _notify_bot.manager.send_to_privileged(text)
        if sent:
            pass
        else:
            await _notify_shop_admin(bot, text)
    except Exception:
        try:
            await _notify_shop_admin(bot, text)
        except Exception:
            pass
    if perm:
        try:
            priv = set(_notify_bot.privileged_ids())
        except Exception:
            priv = set()
        for tid in _perms.notify_extra_ids(perm):
            if tid in priv:
                continue
            try:
                await bot.send_message(tid, text, parse_mode="HTML")
            except Exception:
                pass

def _totp_now(secret: str, digits: int = 6, period: int = 30) -> str:
    """Sinh mã TOTP 6 số hiện tại từ 2FA secret (RFC 6238, chu kỳ 30s)."""
    import base64
    import hashlib
    import hmac
    import struct
    import time
    s = re.sub(r"[^A-Za-z2-7]", "", (secret or "").upper())
    if not s:
        return ""
    try:
        key = base64.b32decode(s + "=" * (-len(s) % 8))
    except Exception:
        return ""
    counter = int(time.time() // period)
    mac = hmac.new(key, struct.pack(">Q", counter), hashlib.sha1).digest()
    off = mac[-1] & 0x0F
    code = struct.unpack(">I", mac[off:off + 4])[0] & 0x7FFFFFFF
    return str(code % (10 ** digits)).zfill(digits)

def _mail_app_link() -> str:
    """Link tải app mail ảo (mail của acc là mail ảo). Admin đổi bằng /setmailapp."""
    return (db.get_setting("mail_app_link", DEFAULT_MAIL_APP_LINK) or "").strip()

def _mail_app_line() -> str:
    link = _mail_app_link()
    if not link:
        return ""
    return (
        "📧 <b>Mail của acc là mail ảo</b> — tải app xem mail tại đây: "
        f"<a href=\"{html.escape(link)}\">Tải app mail</a>"
    )

def _pickup_suffix() -> str:
    """Dòng chọn cách nhận acc + link app mail (nếu có)."""
    line = _mail_app_line()
    return (line + "\n" if line else "") + "👇 <b>Chọn cách nhận acc:</b>"

def _acc_delivery_caption(o) -> str:
    e = html.escape
    parts = [
        f"🎉 <b>MUA THÀNH CÔNG!</b>",
        f"{_cat_icon(o['cat_name'])} <b>{e(o['cat_name'])}</b>",
        f"🧾 Đơn hàng: <b>#{o['id']}</b>  •  💰 <b>{vnd(o['price'])}</b>",
        "━━━━━━━━━━━━━━",
        f"👤 UID: <code>{e(o['uid'] or '')}</code>",
        f"🔑 Mật khẩu: <code>{e(o['password'] or '')}</code>",
    ]
    if o["created_date"]:
        parts.append(f"📅 Ngày tạo: <code>{e(o['created_date'])}</code>")
    if o["backup_mail"]:
        parts.append(f"📧 Mail thay: <code>{e(o['backup_mail'])}</code>")
    if o["totp"]:
        _code = _totp_now(o["totp"])
        _code_txt = f" — mã hiện tại: <code>{_code}</code>" if _code else ""
        parts.append(f"🔐 2FA: <code>{e(o['totp'])}</code>{_code_txt}")
        parts.append("   <i>Mã đổi 30s/lần — lấy mã mới trong /damua.</i>")
    if o["note"]:
        parts.append(f"📝 Ghi chú: {e(o['note'])}")
    if o["cookie"]:
        parts.append(f"🍪 Cookie: <code>{e(o['cookie'])}</code>")
    if o["token"]:
        parts.append(f"🎫 Token: <code>{e(o['token'])}</code>")
    parts += [
        "",
        "⚠️ <b>Đổi mật khẩu ngay</b> sau khi đăng nhập.",
        f"🛡 Bảo hành <b>{_fmt_warranty(o['warranty_hours'])}</b> — acc die thì vào /damua bấm \"Bảo hành\".",
        "📌 <b>Chỉ bảo hành trường hợp đăng nhập báo sai mật khẩu.</b>",
    ]
    mail_line = _mail_app_line()
    if mail_line:
        parts += ["", mail_line]
    return "\n".join(parts)

def _cat_icon(name: str) -> str:
    """Emoji theo tên loại acc cho giao diện shop đẹp hơn."""
    n = (name or "").lower()
    if "bm" in n: return "🏢"
    if "page" in n: return "📄"
    if "clone" in n: return "👥"
    if "via" in n: return "🌐"
    if "xmdt" in n or "xác minh" in n: return "✅"
    if "mail" in n: return "📧"
    return "📦"

async def _shop_stall_picker(msg: Message):
    """Màn hình chọn gian hàng (chỉ hiện khi có >1 gian hàng)."""
    stalls = [s for s in db.acc_stall_list() if int(s["n"] or 0) > 0]
    lines = ["🏪 <b>CHỌN GIAN HÀNG</b>", "━━━━━━━━━━━━━━", ""]
    kb_rows = []
    for s in stalls:
        st = s["stall"]
        total = db.acc_stock_count_stall(st)
        lines.append(f"🏪 <b>{html.escape(st)}</b> — 📦 {total} còn bán • {s['n']} loại hàng")
        kb_rows.append([InlineKeyboardButton(
            text=f"🏪 {st} ({total})", callback_data=f"shopstall:{st}")])
    kb_rows.append([InlineKeyboardButton(text="🛒 Xem giỏ hàng", callback_data="cartview")])
    await msg.answer("\n".join(lines), parse_mode="HTML",
                     reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))

async def _shop_list(msg: Message, stall=None, edit=False, user_id=None):
    cats = db.acc_category_list()
    if stall:
        cats = [c for c in cats
                if ((c["stall"] if "stall" in c.keys() else "") or "Acc Facebook") == stall]
    if not cats:
        txt = "🛒 Gian hàng này hiện chưa có mặt hàng nào. Quay lại sau nhé!"
        if edit:
            await msg.edit_text(txt)
        else:
            await msg.answer(txt)
        return
    _is_fb = (stall or "Acc Facebook") == "Acc Facebook"
    kb_rows = []
    now = int(time.time())
    _uid = user_id or (msg.from_user.id if getattr(msg, "from_user", None) else 0)
    _u = db.get_user(_uid) if _uid else None
    _shop_bal = int(_u["shop_balance"] or 0) if _u and "shop_balance" in _u.keys() else 0
    lines = [
        f"🏪 <b>{html.escape(stall)}</b>" if stall else "🛒 <b>SHOP TÀI KHOẢN</b>",
        "━━━━━━━━━━━━━━",
        f"👛 Ví shop: <b>{vnd(_shop_bal)}</b>",
        "",
    ]
    # 4.10 Giờ vàng giảm giá
    hh_on, hh_pct, hh_cat = db.happy_hour_active()
    if hh_on:
        scope = "toàn shop" if hh_cat == 0 else f"loại #{hh_cat}"
        lines += [f"⚡ <b>GIỜ VÀNG</b> — giảm <b>{hh_pct}%</b> {scope}!",
              ""]
    # gom số liệu từng loại
    infos = []
    for c in cats:
        c = dict(c)
        up = db.shop_unit_price(c)
        n = up["stock"]
        sold = db.acc_sold_count(c["id"])
        avg, rcnt = db.acc_review_avg(c["id"])
        sample = db.get_conn().execute(
            "SELECT uid FROM acc_stock WHERE cat_id=? AND status='AVAILABLE' ORDER BY id LIMIT 1",
            (c["id"],)).fetchone()
        sample_uid = (sample["uid"] or "") if sample else ""
        infos.append((c, up, n, sold, avg, rcnt, sample_uid))
    # 💡 GỢI Ý CHO BẠN: bán chạy / rẻ nhất / đánh giá cao
    in_stock = [i for i in infos if i[2] > 0]
    sugg = []
    if in_stock:
        best = max(in_stock, key=lambda i: i[3])
        cheap = min(in_stock, key=lambda i: i[1]["price"])
        rated = [i for i in in_stock if i[5] >= 2]
        top = max(rated, key=lambda i: i[4]) if rated else None
        seen = set()
        for label, item in (("🔥 Bán chạy", best), ("💰 Rẻ nhất", cheap), ("⭐ Đánh giá cao", top)):
            if item and item[0]["id"] not in seen:
                seen.add(item[0]["id"])
                sugg.append(InlineKeyboardButton(
                    text=f"{label}: {item[0]['name']}",
                    callback_data=f"accbuy:{item[0]['id']}"))
    if sugg and any(i[3] > 0 for i in in_stock):
        lines.append("💡 <b>GỢI Ý CHO BẠN</b>")
        lines.append("")
        kb_rows.append(sugg[:2])
        if len(sugg) > 2:
            kb_rows.append(sugg[2:])
    for c, up, n, sold, avg, rcnt, sample_uid in infos:
        icon = _cat_icon(c["name"])
        price_txt = vnd(up["price"])
        badges = []
        if n == 0:
            badges.append("⛔ HẾT HÀNG")
        else:
            if sold >= 10:
                badges.append("🔥 BÁN CHẠY")
            try:
                is_new = now - int(c.get("created_at") or 0) < 7 * 86400
            except Exception:
                is_new = False
            if is_new:
                badges.append("✨ MỚI")
            if n <= 5:
                badges.append("⚠️ SẮP HẾT")
        if up["happy"]:
            badges.append(f"⚡ giờ vàng −{up['happy_pct']}%")
        badge_txt = (" " + " ".join(badges)) if badges else ""
        stock_txt = f"📦 Còn <b>{n}</b>" if n else "📦 <b>Hết hàng</b>"
        sold_txt = f"  •  🔁 Đã bán <b>{sold}</b>" if sold else ""
        lines.append(f"{icon} <b>{html.escape(c['name'])}</b>{badge_txt}")
        lines.append(f"┌ 💰 <b>{price_txt}</b>/acc")
        lines.append(f"└ {stock_txt}{sold_txt}")
        if rcnt:
            stars = "⭐" * int(round(avg)) if avg else "⭐"
            lines.append(f"    {stars} <b>{avg}/5</b> <i>({rcnt} đánh giá)</i>")
        if c["description"]:
            lines.append(f"    <i>{html.escape(c['description'])}</i>")
        if sample_uid and _is_fb:
            lines.append(f'    🔗 Check thử: <a href="https://facebook.com/{html.escape(sample_uid)}">facebook.com/{html.escape(sample_uid)}</a>')
        lines.append("")
        kb_rows.append([InlineKeyboardButton(
            text=f"{icon} {c['name']} • {price_txt}",
            callback_data=f"accbuy:{c['id']}")])
    # 4.8 Hộp mù acc
    try:
        m_price = int(db.get_setting("mystery_price", "0") or 0)
    except Exception:
        m_price = 0
    if m_price > 0 and _is_fb:
        m_cats = [dict(x) for x in db.acc_category_list()
                  if int(x["mystery_eligible"] or 0) and db.acc_stock_count(x["id"]) > 0]
        if m_cats:
            lines += ["", f"🎁 <b>HỘP MÙ ACC — {vnd(m_price)}</b>: mở ngẫu nhiên 1 acc "
                         f"từ {len(m_cats)} loại, rẻ hơn giá lẻ!"]
            kb_rows.append([InlineKeyboardButton(
                text=f"🎁 Mua hộp mù — {vnd(m_price)}",
                callback_data="accmystery")])
    lines += ["👉 <i>Chạm vào từng loại để xem chi tiết & mua.</i>",
              "",
              "🎡 <i>/quay</i> vòng quay may mắn  •  🏅 <i>/hang</i> hạng thành viên  •  🎁 <i>/doiqua</i> đổi điểm",
              "🛒 <i>/giohang</i> giỏ hàng của bạn"]
    kb_rows.append([InlineKeyboardButton(text="🛒 Xem giỏ hàng", callback_data="cartview"),
                    InlineKeyboardButton(text="💰 Nạp ví shop", callback_data="wal_nap:shop")])
    if stall:
        kb_rows.append([InlineKeyboardButton(text="◀️ Gian hàng khác",
                                             callback_data="shopstall:__all")])
    _kb = InlineKeyboardMarkup(inline_keyboard=kb_rows)
    _text = "\n".join(lines)
    if edit:
        await msg.edit_text(_text, parse_mode="HTML", reply_markup=_kb,
                            disable_web_page_preview=True)
    else:
        await msg.answer(_text, parse_mode="HTML", reply_markup=_kb,
                         disable_web_page_preview=True)

def _shop_bulk_pcts() -> tuple[int, int, int]:
    """% giảm giá khi mua 5 / 10 / 20 acc (giá sỉ). Đổi qua settings shop_bulk_5_pct, shop_bulk_10_pct, shop_bulk_20_pct."""
    def _g(key, default):
        try:
            return max(0, int(db.get_setting(key, str(default)) or default))
        except Exception:
            return default
    return _g("shop_bulk_5_pct", 5), _g("shop_bulk_10_pct", 10), _g("shop_bulk_20_pct", 15)

async def _acc_after_purchase(bot, msg, from_user, c: dict, cat_id: int, tg_id: int,
                              final: int, qty: int, delivered: list,
                              want_upsell: bool = False, upsell_pct_cfg: int = 0,
                              wmin: int = 10):
    """Giao acc + qua tang + bao admin sau khi ban thanh cong (dung chung cho
    mua thuong va mua theo UID cu the)."""
    # Giao từng acc — khách chọn cách nhận: hiện thông tin hoặc tải file
    for order in delivered:
        order_id = order["id"]
        await msg.answer(
            f"🎉 <b>MUA THÀNH CÔNG!</b>\n"
            f"━━━━━━━━━━━━━━\n"
            f"{_cat_icon(order['cat_name'])} <b>{html.escape(order['cat_name'])}</b>\n"
            f"🧾 Đơn hàng: <b>#{order_id}</b>\n"
            f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
            f"💰 Đã thanh toán: <b>{vnd(order['price'])}</b>\n"
            f"{_live_line(cat_id)}"
            f"━━━━━━━━━━━━━━\n\n"
            f"{_pickup_suffix()}",
            parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
    # Báo admin: thông tin khách + acc đã mua
    await _notify_purchase_admin(bot, from_user, delivered)
    # Quà tặng kèm: credits + vé quay
    extras = []
    bonus_per = int(c.get("credit_bonus") or 0)
    if bonus_per > 0:
        new_credits = db.add_credits(tg_id, bonus_per * qty, f"combo_mua_acc:{cat_id}")
        extras.append(f"🎁 Tặng <b>{bonus_per * qty} credits</b> (số dư credits: {new_credits})")
    tickets = db.spin_add_tickets(tg_id, qty)
    extras.append(f"🎡 Nhận <b>{qty} vé quay</b> may mắn — gõ /quay để thử vận may (đang có {tickets} vé)")
    # Tích điểm loyalty: 1 điểm / loyalty_per_vnd
    try:
        per = int(db.get_setting("loyalty_per_vnd", "100000") or 100000)
    except Exception:
        per = 100000
    pts = final // per if per > 0 else 0
    if pts > 0:
        new_pts = db.loyalty_add(tg_id, pts, f"mua_acc:{cat_id}")
        extras.append(f"⭐ Tích <b>{pts} điểm</b> loyalty (tổng: <b>{new_pts}</b> điểm) — "
                      f"đủ điểm đổi acc miễn phí bằng /doiqua")
    if extras:
        rkey = f"o{min(o['id'] for o in delivered)}" if delivered else None
        await msg.answer("\n".join(extras), parse_mode="HTML",
                                reply_markup=_spin_kb(rkey, tg_id, final))
    # 4.6 Upsell: gợi ý mua thêm trong 10 phút được giảm thêm
    try:
        if upsell_pct_cfg > 0 and not want_upsell:
            kb = InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(
                    text=f"⚡ Mua thêm 1 acc giảm {upsell_pct_cfg}% (trong {wmin} phút)",
                    callback_data=f"accconfirm:{cat_id}:1:upsell")],
            ])
            await msg.answer(
                f"⚡ <b>ƯU ĐÃI NÓNG:</b> bạn vừa mua <b>{html.escape(c['name'])}</b> — "
                f"mua thêm acc cùng loại trong <b>{wmin} phút</b> được "
                f"<b>giảm thêm {upsell_pct_cfg}%</b>!",
                parse_mode="HTML", reply_markup=kb)
    except Exception:
        pass
    # 5.12 Hướng dẫn sau mua tự động
    try:
        guide = db.get_setting("postbuy_guide", "") or (
            "📋 <b>HƯỚNG DẪN SAU KHI MUA ACC</b>\n"
            "1️⃣ Đổi mật khẩu ngay sau khi đăng nhập.\n"
            "2️⃣ Bật xác thực 2 lớp (2FA) cho acc.\n"
            "3️⃣ 3 ngày đầu: đừng đổi tên/avatar vội, lướt newsfeed nhẹ nhàng.\n"
            "4️⃣ Không đăng nhập nhiều acc cùng 1 IP/proxy lạ.\n"
            "💡 Làm đúng các bước trên giúp acc sống lâu, ít bị checkpoint!")
        await msg.answer(guide, parse_mode="HTML")
    except Exception:
        pass
    # Hoa hồng cho người giới thiệu (F1)
    try:
        f1_id, comm = db.ref_shop_commission(tg_id, final)
        if f1_id and comm:
            try:
                await bot.send_message(
                    f1_id,
                    f"🎁 <b>Hoa hồng shop acc!</b>\n"
                    f"Người bạn giới thiệu vừa mua {qty} acc "
                    f"<b>{html.escape(c['name'])}</b> — bạn nhận <b>{vnd(comm)}</b>.",
                    parse_mode="HTML")
            except Exception:
                pass
    except Exception:
        pass
    # Cảnh báo hết hàng cho admin
    try:
        warn_at = int(db.get_setting("acc_low_stock_warn", "20") or 20)
    except Exception:
        warn_at = 20
    left = db.acc_stock_count(cat_id)
    if left <= warn_at:
        await _notify_admin_smart(
            bot,
            f"⚠️ <b>Sắp hết hàng:</b> {html.escape(c['name'])} chỉ còn <b>{left}</b> acc. "
            f"Nhập thêm bằng /themacc {cat_id}",
            perm="kho",
        )
    # Mua nhiều acc 1 lúc: tự động gửi 1 file gộp
    await _send_merged_acc_file(msg, delivered)

async def _send_merged_acc_file(msg, delivered):
    """Khách mua nhiều acc 1 lúc: tự động gửi 1 file txt + 1 file xlsx
    gộp tất cả acc vừa giao (mỗi dòng 1 acc, dòng cuối chú thích cột)."""
    if len(delivered) < 2:
        return
    import datetime as _dt
    F = ["uid", "password", "created_date", "backup_mail", "note",
         "totp", "cookie", "token"]
    rows = [[(o.get(k) or "") for k in F] for o in delivered]
    header = ["UID", "Mật khẩu", "Ngày tạo", "Mail thay", "Ghi chú",
              "2FA", "Cookie", "Token"]
    ts = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    n = len(delivered)
    try:
        txt_data = ("\n".join("|".join(r) for r in rows)
                    + "\n" + "|".join(header)).encode("utf-8")
        await msg.answer_document(
            BufferedInputFile(txt_data,
                              filename=f"acc_gop_{n}acc_{ts}.txt"),
            caption=f"📦 <b>File gộp {n} acc</b> vừa mua — mỗi dòng 1 acc, "
                    f"dòng cuối là chú thích cột.",
            parse_mode="HTML")
    except Exception:
        pass
    try:
        import openpyxl
        import io as _io
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "acc"
        ws.append(header)
        for r in rows:
            ws.append(r)
        ws.append([])
        ws.append(header)
        buf = _io.BytesIO()
        wb.save(buf)
        await msg.answer_document(
            BufferedInputFile(buf.getvalue(),
                              filename=f"acc_gop_{n}acc_{ts}.xlsx"),
            caption=f"📊 <b>File Excel gộp {n} acc</b> vừa mua.",
            parse_mode="HTML")
    except Exception:
        pass

def _uid_final_price(tg_id: int, c: dict):
    """Giá mua 1 acc theo UID: giữ nguyên công thức mua lẻ qty=1
    (giá đơn vị + giảm hạng TV + giảm mua thêm, chưa áp promo)."""
    up = db.shop_unit_price(c)
    price = up["price"]
    tier = db.member_tier_info(tg_id)
    tier_pct = int(tier["pct"])
    try:
        wmin = int(db.get_setting("upsell_window_min", "10") or 10)
        upsell_pct_cfg = int(db.get_setting("upsell_pct", "5") or 5)
    except Exception:
        wmin, upsell_pct_cfg = 10, 5
    upsell_pct = 0
    try:
        last_at = db.acc_last_order_at(tg_id)
        if upsell_pct_cfg > 0 and last_at and now() - last_at <= wmin * 60:
            upsell_pct = upsell_pct_cfg
    except Exception:
        pass
    final = price * (100 - tier_pct) // 100
    final = final * (100 - upsell_pct) // 100
    return final, price, tier, tier_pct, upsell_pct, wmin, upsell_pct_cfg

def _line_price(cat: dict, qty: int, tg_id: int) -> dict:
    """Giá 1 dòng giỏ/mua: giữ nguyên công thức mua lẻ
    (giảm bulk theo số lượng của dòng + giảm hạng TV)."""
    up = db.shop_unit_price(cat)
    price = up["price"]
    p5, p10, p20 = _shop_bulk_pcts()
    bulk_pct = p20 if qty >= 20 else (p10 if qty >= 10 else (p5 if qty >= 5 else 0))
    tier = db.member_tier_info(tg_id)
    tier_pct = int(tier["pct"])
    total = price * qty
    final = total * (100 - bulk_pct) // 100
    final = final * (100 - tier_pct) // 100
    return {"unit": price, "bulk_pct": bulk_pct, "tier_pct": tier_pct,
            "tier_name": tier.get("tier", ""), "total": total, "final": final}

def _cart_render(tg_id: int):
    """Trả (text, keyboard) cho giỏ hàng, hoặc None nếu giỏ trống."""
    items = [it for it in db.cart_list(tg_id) if it["active"]]
    uid_items = [it for it in db.cart_uid_list(tg_id) if it["active"]]
    if not items and not uid_items:
        return None
    lines = ["🛒 <b>GIỎ HÀNG CỦA BẠN</b>", "━━━━━━━━━━━━━━"]
    kb_rows = []
    grand = 0
    saved = 0
    for i, it in enumerate(items, 1):
        cat = db.acc_category_get(it["cat_id"])
        if not cat:
            continue
        lp = _line_price(dict(cat), it["qty"], tg_id)
        grand += lp["final"]
        disc = []
        if lp["bulk_pct"]:
            disc.append(f"−{lp['bulk_pct']}% mua nhiều")
        if lp["tier_pct"]:
            disc.append(f"−{lp['tier_pct']}% {lp['tier_name']}")
        # tiền tiết kiệm so với giá gốc
        try:
            base_total = int(lp["unit"]) * int(it["qty"])
            saved += max(0, base_total - int(lp["final"]))
        except Exception:
            pass
        dtxt = f"\n   <i>🎉 {', '.join(disc)}</i>" if disc else ""
        lines.append(
            f"<b>{i}. {_cat_icon(it['name'])} {html.escape(it['name'])}</b>\n"
            f"   {it['qty']} × {vnd(lp['unit'])} = <b>{vnd(lp['final'])}</b>{dtxt}")
        kb_rows.append([
            InlineKeyboardButton(text="➖", callback_data=f"cartdec:{it['cat_id']}"),
            InlineKeyboardButton(text=f"×{it['qty']}", callback_data="cartnoop"),
            InlineKeyboardButton(text="➕", callback_data=f"cartinc:{it['cat_id']}"),
            InlineKeyboardButton(text="❌", callback_data=f"cartdel:{it['cat_id']}"),
        ])
    if uid_items:
        lines.append(f"<b>🎯 UID đã chọn ({len(uid_items)}):</b>")
        for it in uid_items:
            cat = db.acc_category_get(it["cat_id"])
            if not cat:
                continue
            lp = _line_price(dict(cat), 1, tg_id)
            grand += lp["final"]
            lines.append(
                f"👤 <code>{html.escape(it['uid'] or '')}</code> — "
                f"<i>{html.escape(it['name'])}</i> = <b>{vnd(lp['final'])}</b>")
            kb_rows.append([
                InlineKeyboardButton(
                    text=f"❌ Xóa UID {it['uid']}",
                    callback_data=f"cartuiddel:{it['stock_id']}"),
            ])
    u = db.get_user(tg_id)
    balance = int(u["shop_balance"] or 0) if u else 0
    grand_promo, promo_code = db.preview_user_promo(tg_id, grand, wallet="shop")
    lines += ["━━━━━━━━━━━━━━",
              f"🧾 <b>Tổng cộng: {vnd(grand)}</b>"]
    if saved > 0:
        lines.append(f"🎉 <i>Bạn tiết kiệm được {vnd(saved)}</i>")
    if promo_code:
        lines.append(f"🎟️ Mã <b>{promo_code}</b> sẽ được áp khi thanh toán → còn <b>{vnd(grand_promo)}</b>")
    lines.append(f"👛 Ví shop: <b>{vnd(balance)}</b>"
                 + ("" if balance >= grand_promo else f"  <i>(thiếu {vnd(grand_promo - balance)})</i>"))
    kb_rows.append([InlineKeyboardButton(
        text=f"💳 Thanh toán — {vnd(grand_promo)}", callback_data="cartcheckout")])
    kb_rows.append([
        InlineKeyboardButton(text="🗑 Xóa giỏ", callback_data="cartclear"),
        InlineKeyboardButton(text="🛍 Tiếp tục mua", callback_data="accshop_back"),
    ])
    return "\n".join(lines), InlineKeyboardMarkup(inline_keyboard=kb_rows)

def _loyalty_random_cfg():
    """(min, max, cap_user, min_order, daily_max) cho điểm ngẫu nhiên sau mua.
    None = đang tắt. min_order: đơn phải >= X tiền mới được random (0 = không
    giới hạn). daily_max: mỗi user tối đa N lượt random/ngày (0 = không giới hạn)."""
    try:
        lo = int(db.get_setting("loyalty_random_min", "1") or 1)
        hi = int(db.get_setting("loyalty_random_max", "5") or 5)
        cap = int(db.get_setting("loyalty_random_cap", "0") or 0)
        min_order = int(db.get_setting("loyalty_random_min_order", "0") or 0)
        daily_max = int(db.get_setting("loyalty_random_daily_max", "0") or 0)
    except Exception:
        return None
    if cap <= 0 or lo <= 0 or hi < lo:
        return None
    return (lo, hi, cap, max(0, min_order), max(0, daily_max))

def _loyalty_random_received(tg_id: int) -> int:
    """Tổng điểm ngẫu nhiên user đã nhận (để so với trần cap_user)."""
    r = db.get_conn().execute(
        "SELECT COALESCE(SUM(delta),0) AS s FROM loyalty_history "
        "WHERE tg_id=? AND reason LIKE 'randpts:%'", (tg_id,)).fetchone()
    try:
        return int(r["s"] or 0)
    except Exception:
        return 0

def _loyalty_random_today(tg_id: int) -> int:
    """Số lượt random user đã bấm trong hôm nay (chống cày điểm)."""
    try:
        import datetime as _dt
        start = int(_dt.datetime.now().replace(
            hour=0, minute=0, second=0, microsecond=0).timestamp())
        r = db.get_conn().execute(
            "SELECT COUNT(*) AS c FROM loyalty_history "
            "WHERE tg_id=? AND reason LIKE 'randpts:%' AND created_at>=?",
            (tg_id, start)).fetchone()
        return int(r["c"] or 0)
    except Exception:
        return 0

def _order_total_for_randkey(key: str) -> int:
    """Tổng tiền của đơn từ rand key (o<order_id> / c<order_id>). 0 nếu không rõ."""
    try:
        oid = int((key or "")[1:])
        r = db.get_conn().execute(
            "SELECT price FROM acc_orders WHERE id=?", (oid,)).fetchone()
        return int(r["price"] or 0) if r else 0
    except Exception:
        return 0

def _spin_kb(order_key=None, tg_id: int = 0, order_total: int = 0):
    """Bàn phím tin quà tặng sau mua: nút quay + nút điểm ngẫu nhiên (nếu bật
    và đơn/user đủ điều kiện chống cày điểm). Nhãn nút hiện tiến độ (x/y)."""
    rows = [[InlineKeyboardButton(text="🎡 Quay ngay", callback_data="spin_now")]]
    if order_key:
        cfg = _loyalty_random_cfg()
        if cfg:
            lo, hi, cap, min_order, daily_max = cfg
            ok = True
            if min_order > 0 and order_total < min_order:
                ok = False
            if ok and daily_max > 0 and tg_id:
                ok = _loyalty_random_today(tg_id) < daily_max
            if ok:
                prog = ""
                if tg_id and cap > 0:
                    prog = f" ({_loyalty_random_received(tg_id)}/{cap})"
                rows.append([InlineKeyboardButton(
                    text=f"🎲 Nhận điểm ngẫu nhiên{prog}",
                    callback_data=f"randpts:{order_key}")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

async def _do_spin(msg, tg_id: int, tickets: int):
    if not db.spin_consume_ticket(tg_id):
        await msg.answer("😅 Vé vừa được dùng hết. Mua acc ở /shop để nhận thêm vé nhé!")
        return
    prize = db.spin_roll()
    db.spin_award(tg_id, prize)
    left = db.spin_get_tickets(tg_id)
    label = prize.get("label", "")
    kind = prize.get("kind", "none")
    value = int(prize.get("value", 0) or 0)
    if kind == "none" or value <= 0:
        txt = f"🎡 <b>VÒNG QUAY MAY MẮN</b>\n\n{label}\n\nCòn <b>{left}</b> vé — chúc bạn may mắn lần sau!"
    else:
        txt = (f"🎡 <b>VÒNG QUAY MAY MẮN</b>\n\n"
               f"🎉 <b>TRÚNG: {label}!</b>\n"
               f"Quà đã cộng vào tài khoản của bạn.\n\n"
               f"Còn <b>{left}</b> vé.")
    kb = None
    if left > 0:
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🎡 Quay tiếp", callback_data="spin_now")],
        ])
    await msg.answer(txt, parse_mode="HTML", reply_markup=kb)

def _acc_delivery_kb(order_id: int) -> InlineKeyboardMarkup:
    """Bàn phím chọn cách nhận acc: hiện thông tin để sao chép / tải file."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📋 Hiện thông tin (sao chép)",
                              callback_data=f"accshow:{order_id}")],
        [InlineKeyboardButton(text="📄 Tải file .txt",
                              callback_data=f"accfile:{order_id}:txt"),
         InlineKeyboardButton(text="📊 Tải file .xlsx",
                              callback_data=f"accfile:{order_id}:xlsx")],
    ])

def _strip_acc(s: str) -> str:
    """Bỏ dấu tiếng Việt để so khớp đơn vị."""
    import unicodedata as _ud
    t = _ud.normalize("NFD", s or "")
    t = "".join(c for c in t if _ud.category(c) != "Mn")
    return t.replace("đ", "d").replace("Đ", "D")

def _parse_warranty(text: str):
    """Parse th\u1eddi gian b\u1ea3o h\u00e0nh: 30p | 24h | 2 ng\u00e0y | 1 tu\u1ea7n | 24 (m\u1eb7c \u0111\u1ecbnh = gi\u1edd). Tr\u1ea3 v\u1ec1 s\u1ed1 PH\u00daT, None n\u1ebfu sai."""
    t = _strip_acc((text or "").strip().lower()).replace(",", ".")
    m = re.match(r"^(\d+(?:\.\d+)?)\s*([a-z ]*)$", t)
    if not m:
        return None
    num = float(m.group(1))
    unit = m.group(2).strip()
    if unit in ("", "h", "g", "gio", "tieng"):
        mins = num * 60
    elif unit in ("p", "m", "ph", "phut", "min", "mins"):
        mins = num
    elif unit in ("d", "ngay", "day", "days"):
        mins = num * 1440
    elif unit in ("w", "tuan", "week", "weeks"):
        mins = num * 10080
    else:
        return None
    mins = int(round(mins))
    return mins if mins > 0 else None

def _fmt_warranty(mins) -> str:
    """Hi\u1ec3n th\u1ecb th\u1eddi gian b\u1ea3o h\u00e0nh (ph\u00fat) d\u1ea1ng \u0111\u1eb9p: 30 ph\u00fat | 24h | 1h30."""
    try:
        m = int(mins or 0)
    except Exception:
        m = 0
    if m <= 0:
        return "0"
    if m < 60:
        return f"{m} ph\u00fat"
    if m % 60 == 0:
        return f"{m // 60}h"
    return f"{m // 60}h{m % 60}"

def _cat_live_check(cat_id: int) -> bool:
    """Loại acc này có cần check LIVE trước khi giao không?
    Chỉ sạp 'Acc Facebook' (live_check=1) mới check."""
    try:
        c = db.acc_category_get(cat_id)
        if c and "live_check" in c.keys():
            return int(c["live_check"] or 0) == 1
    except Exception:
        pass
    return True

def _live_line(cat_id: int) -> str:
    """Dòng 'Đã kiểm tra LIVE' trong tin giao hàng — chỉ hiện khi có check thật."""
    return "🟢 <i>Đã kiểm tra LIVE trước khi giao</i>\n" if _cat_live_check(cat_id) else ""

def _smart_stock_fields(fields: list) -> dict:
    """Nhận diện trường acc từ 1 dòng .txt theo MẪU nội dung (không phụ thuộc thứ tự cột):
    link FB | email | mã 2FA (base32) | ngày tạo | cookie | token EAAG | UID số.
    Trường còn lại: trường đầu -> mật khẩu, các trường sau -> ghi chú.
    Riêng bộ 4 mail Outlook (mail|mật khẩu|refresh_token M.C...|client_id UUID)
    được gộp lại thành 1 chuỗi "liền mạch" vào backup_mail."""
    import re as _re
    # Pre-pass: gộp bộ 4 Outlook về 1 trường trước khi nhận diện.
    _fields = [(f or "").strip() for f in fields]
    _UUID = _re.compile(
        r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-"
        r"[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")
    merged: list = []
    _i, _n = 0, len(_fields)
    while _i < _n:
        _f = _fields[_i]
        if (_f and _re.search(r"\S+@\S+\.\S+", _f)
                and _i + 3 < _n
                and _fields[_i + 1]
                and _fields[_i + 2].startswith("M.C")
                and len(_fields[_i + 2]) > 50
                and _UUID.match(_fields[_i + 3])):
            _j = _i + 4
            # Bundle mail co the kem them email phu lien sau (vd mail fviainboxes)
            while _j < _n and _re.search(r"\S+@\S+\.\S+", _fields[_j]):
                _j += 1
            merged.append("|".join(_fields[_i:_j]))
            _i = _j
        else:
            merged.append(_f)
            _i += 1
    fields = merged
    r = {"uid": "", "password": "", "created_date": "", "backup_mail": "",
         "note": "", "totp": "", "cookie": "", "token": ""}
    others = []
    for _f in fields:
        f = (_f or "").strip()
        if not f:
            continue
        fl = f.lower()
        if "facebook.com" in fl or "fb.com" in fl or "fb.watch" in fl:
            if not r["uid"]:
                r["uid"] = f
            else:
                others.append(f)
        elif _re.search(r"\S+@\S+\.\S+", f):
            if not r["backup_mail"]:
                r["backup_mail"] = f
            else:
                others.append(f)
        elif (f.startswith("EAAG") or f.startswith("EAAB")) and not r["token"]:
            r["token"] = f
        elif ("c_user" in fl or ("=" in f and len(f) > 80)) and not r["cookie"]:
            r["cookie"] = f
        elif (_re.fullmatch(r"[A-Z2-7][A-Z2-7 ]+", f.upper())
              and not r["totp"]
              and ((" " in f and len(_re.sub(r"\s+", "", f)) >= 12)
                   or len(f) >= 16)):
            r["totp"] = _re.sub(r"\s+", "", f).upper()
        elif (_re.search(r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2}", f)
              and not r["created_date"]):
            r["created_date"] = f
        elif not r["uid"]:
            # UID số thuần (độ dài >= 6). Chấp nhận cả dạng float do Excel
            # ("61593959792972.0") -> cắt đuôi .0, tránh check die oan.
            _m = _re.fullmatch(r"(\d+)(?:\.0+)?", f)
            if _m and len(_m.group(1)) >= 6:
                r["uid"] = _m.group(1)
            else:
                others.append(f)
        else:
            others.append(f)
    if others:
        r["password"] = others[0]
        if len(others) > 1:
            r["note"] = " | ".join(others[1:])
    return r

async def _resolve_stock_links(rows, wait, action_text="nhập kho"):
    """Giải link FB ở cột đầu -> UID số (dùng chung cho nhập kho và cập nhật).
    Trả (resolved, failed, failed_lines, uid_to_link, link_idx)."""
    link_idx = [i for i, r in enumerate(rows)
                if re.search(r'facebook\.com|fb\.com|fb\.watch|^https?://', (r.get("uid") or ""), re.I)]
    resolved = failed = 0
    failed_lines = []
    if link_idx:
        await wait.edit_text(f"⏳ Đang {action_text}... 🔗 Giải {len(link_idx)} link FB → UID...")
        from ..fb import resolve_fb_uid
        _cache = {}
        _sem = asyncio.Semaphore(4)
        async def _res(i):
            nonlocal resolved, failed
            link = (rows[i].get("uid") or "").strip()
            rows[i]["_orig_link"] = link  # giữ link gốc để vòng thử lại dùng
            if link in _cache:
                uid = _cache[link]
            else:
                async with _sem:
                    try:
                        uid, _nm, _md = await resolve_fb_uid(link)
                    except Exception:
                        uid = ""
                _cache[link] = uid
            if uid and str(uid).isdigit():
                rows[i]["uid"] = str(uid)
                resolved += 1
            else:
                rows[i]["uid"] = ""
                failed += 1
                failed_lines.append(f"Dòng {i + 1}: {link[:60]}")
        await asyncio.gather(*[_res(i) for i in link_idx])
        if failed:
            # Vòng 2: thử lại TUẦN TỰ các link lỗi, cách nhau vài giây.
            # (Vòng 1 quét song song dễ bị API giới hạn tốc độ -> báo lỗi oan.)
            try:
                await wait.edit_text(
                    f"🔁 Thử lại {failed} link chưa giải được (chậm để tránh bị giới hạn)...")
            except Exception:
                pass
            still_failed = []
            for i in link_idx:
                if rows[i].get("uid"):
                    continue
                link = (rows[i].get("_orig_link") or "").strip()
                if not link:
                    continue
                try:
                    uid, _nm, _md = await resolve_fb_uid(link)
                except Exception:
                    uid = ""
                if uid and str(uid).isdigit():
                    rows[i]["uid"] = str(uid)
                    resolved += 1
                    failed -= 1
                else:
                    still_failed.append(f"Dòng {i + 1}: {link[:60]}")
                await asyncio.sleep(2)
            failed_lines = still_failed
    # Map uid -> link gốc (hiển thị danh sách acc die sau khi quét)
    uid_to_link = {}
    for i in link_idx:
        if rows[i]["uid"] and str(rows[i]["uid"]).isdigit():
            uid_to_link[rows[i]["uid"]] = (rows[i].get("_orig_link") or "")
    return resolved, failed, failed_lines, uid_to_link, link_idx

async def _import_stock_rows(rows, cat_id, ncc_id, cost, c, msg, wait, sheet_ctx=None):
    """Pipeline nhập kho dùng chung cho /themacc (file) và /nhapkhosheet (Google Sheet).
    rows: list dict {uid,password,created_date,backup_mail,note,totp,cookie,token}.
    sheet_ctx: None hoặc {"sheet_id","tab"} — ghi trạng thái ngược vào sheet."""
    # Cột mk trống -> dùng mặc định "khai2006" (các cột khác giữ đúng vị trí)
    for r in rows:
        if not (r.get("password") or "").strip():
            r["password"] = "khai2006"
    resolved, failed, failed_lines, uid_to_link, link_idx = await _resolve_stock_links(rows, wait)
    # Chống trùng: bỏ dòng trùng trong file và UID đang còn trong kho
    # (chỉ tính acc AVAILABLE/DIE; acc đã bán SOLD được nhập lại bình thường)
    _seen = set()
    try:
        _existing = {r[0] for r in db.get_conn().execute(
            "SELECT uid FROM acc_stock WHERE cat_id=? AND status IN ('AVAILABLE','DIE')",
            (cat_id,)).fetchall()}
    except Exception:
        _existing = set()
    _uniq, empty, dup_in_file, dup_in_stock = [], 0, 0, 0
    all_rows = rows  # giữ để ghi trạng thái sheet cho cả dòng bị loại
    for r in rows:
        uid = (r.get("uid") or "").strip()
        if not uid:
            empty += 1
            if sheet_ctx is not None:
                r["_sheet_status"] = "🔗 LỖI LINK" if r.get("_orig_link") else "⏭ BỎ QUA"
            continue
        if uid in _seen:
            dup_in_file += 1
            if sheet_ctx is not None:
                r["_sheet_status"] = "🗑 TRÙNG"
            continue
        if uid in _existing:
            dup_in_stock += 1
            if sheet_ctx is not None:
                r["_sheet_status"] = "🗑 TRÙNG"
            continue
        _seen.add(uid)
        _uniq.append(r)
    rows = _uniq
    if sheet_ctx is not None:
        for r in _uniq:
            r["_sheet_status"] = "✅ OK"
            # Giữ lại vị trí dòng trên Sheet để sau này đẩy dấu "đã bán" ngược
            if r.get("_sheet_row"):
                r["_sheet_ref"] = f"{sheet_ctx['tab']}:{r['_sheet_row']}"
        sheet_ctx["uid_to_row"] = {
            str(r["uid"]): r.get("_sheet_row")
            for r in _uniq if str(r.get("uid") or "").isdigit() and r.get("_sheet_row")
        }
    skipped = empty + dup_in_file + dup_in_stock
    stock_since = int(time.time())
    added, _ = db.acc_stock_add_batch(
        cat_id, rows,
        batch="Lô " + time.strftime("%d/%m %H:%M", time.localtime()),
        supplier_id=ncc_id, cost_per_acc=cost)
    total = db.acc_stock_count(cat_id)
    extra_info = ""
    if ncc_id:
        extra_info += f"\n🏭 NCC: <b>#{ncc_id}</b>"
    if cost:
        extra_info += f" — 💰 Vốn: <b>{vnd(cost)}</b>/acc"
    dup_txt = ""
    if dup_in_file or dup_in_stock:
        dup_txt = f"🗑 Trùng loại: {dup_in_file} trong file, {dup_in_stock} đã có trong kho\n"
    link_txt = ""
    if link_idx:
        link_txt = f"\n🔗 Giải link FB: <b>{resolved}</b> OK" + (f", <b>{failed}</b> lỗi" if failed else "")
    await wait.edit_text(
        f"✅ <b>Nhập kho xong:</b> {html.escape(c['name'])}\n"
        f"➕ Thêm: <b>{added}</b> acc | ⏭ Bỏ qua: <b>{skipped}</b> dòng{extra_info}\n"
        f"{dup_txt}"
        f"📊 Tồn kho hiện tại: <b>{total}</b> acc"
        f"{link_txt}"
        "\n\n🔍 Đang tự quét kiểm tra chất lượng...",
        parse_mode="HTML",
    )
    if sheet_ctx is not None:
        # Ghi đánh dấu ngược vào sheet: cột I trạng thái + ghi đè link bằng UID (cách 1)
        # + cột L loại/gian hàng
        try:
            from .. import sheet_import as _si
            try:
                _wb_stall = (c["stall"] if "stall" in c.keys() else "") or "Acc Facebook"
            except Exception:
                _wb_stall = "Acc Facebook"
            _wb_label = _si.cat_label(cat_id, c["name"], _wb_stall)
            _updates = []
            for r in all_rows:
                sr = r.get("_sheet_row")
                if not sr:
                    continue
                _updates.append((sr, 9, r.get("_sheet_status") or "⏭ BỎ QUA"))
                _updates.append((sr, _si.CAT_COL, _wb_label))
                if r.get("_orig_link") and str(r.get("uid") or "").isdigit():
                    _updates.append((sr, 1, str(r["uid"])))
            await _si.write_updates(sheet_ctx["sheet_id"], sheet_ctx["tab"], _updates)
        except Exception as e:
            log.warning("sheet write-back: %s", e)
    if failed_lines:
        try:
            txt = ("⚠️ <b>Các link không giải được UID</b> (đã thử lại nhiều lần, đã bỏ qua):\n"
                   "<i>Thường do link đã chết hoặc bài viết không để công khai — kiểm tra lại các link này.</i>\n"
                   ) + "\n".join(
                html.escape(x) for x in failed_lines[:15])
            if len(failed_lines) > 15:
                txt += f"\n...và {len(failed_lines) - 15} dòng nữa"
            await msg.answer(txt, parse_mode="HTML")
        except Exception:
            pass
    # Quét chất lượng nền: acc die -> loại khỏi kho bán
    async def _scan(stock_since: int):
        from .. import fb as fb_mod
        import asyncio
        ids = db.get_conn().execute(
            "SELECT id, uid FROM acc_stock WHERE cat_id=? AND status='AVAILABLE' "
            "AND added_at >= ? ORDER BY id", (cat_id, stock_since)).fetchall()
        sem = asyncio.Semaphore(10)
        dead = 0
        dead_list = []
        dead_uids = []
        live_uids = []
        async def _one(sid, uid):
            nonlocal dead
            async with sem:
                try:
                    r = await fb_mod.check_uid(str(uid))
                    st = str(r.get("status", "")).lower()
                    if st in ("die", "dead"):
                        db.acc_mark_status(sid, "DEAD")
                        dead += 1
                        dead_uids.append(str(uid))
                        lk = (uid_to_link.get(str(uid)) or "").strip()
                        dead_list.append(f"• {html.escape(str(uid))}" + (f" — {html.escape(lk)}" if lk else ""))
                    elif st == "live":
                        live_uids.append(str(uid))
                except Exception:
                    pass
        await asyncio.gather(*[_one(r["id"], r["uid"]) for r in ids])
        left = db.acc_stock_count(cat_id)
        dead_txt = ""
        if dead_list:
            dead_txt = "\n⚠️ <b>Danh sách acc die:</b>\n" + "\n".join(dead_list[:20])
            if len(dead_list) > 20:
                dead_txt += f"\n...và {len(dead_list) - 20} acc nữa"
        if sheet_ctx is not None and (dead_uids or live_uids):
            # Đánh dấu tình trạng acc ngược vào cột K (cột I giữ nguyên trạng thái nhập kho)
            try:
                from .. import sheet_import as _si
                _u2r = sheet_ctx.get("uid_to_row", {}) or {}
                _u = [(1, _si.HEALTH_COL, "Tình trạng")]
                _u += [(r, _si.HEALTH_COL, "☠️ DIE")
                       for u, r in _u2r.items() if u in dead_uids]
                _u += [(r, _si.HEALTH_COL, "🟢 LIVE")
                       for u, r in _u2r.items() if u in live_uids]
                await _si.write_updates(sheet_ctx["sheet_id"], sheet_ctx["tab"], _u)
            except Exception as e:
                log.warning("sheet health mark: %s", e)
        await msg.answer(
            f"🔍 <b>Quét xong {len(ids)} acc mới nhập:</b> "
            f"<b>{dead}</b> acc die đã loại khỏi kho.\n"
            f"📊 Tồn kho bán được: <b>{left}</b> acc.{dead_txt}",
            parse_mode="HTML",
        )
        # Báo / tự giao cho những ai đã đăng ký "có hàng nhắn tôi"
        if left > 0:
            try:
                await _fulfill_restock_subs(cat_id, msg.bot, dict(c))
            except Exception as e:
                log.warning("fulfill restock subs: %s", e)
            # 4.9 Giao hàng cho người đã đặt cọc (FIFO)
            try:
                await _fulfill_deposits(cat_id, msg.bot, dict(c))
            except Exception as e:
                log.warning("fulfill deposits: %s", e)
    asyncio.create_task(_scan(int(now()) - 600))

def _parse_stock_file(raw: bytes, file_name: str):
    """Đọc file .txt/.xlsx chứa acc → (rows, err).
    rows: list dict {uid,password,created_date,backup_mail,note,totp,cookie,token}.
    err: chuỗi lỗi hoặc None."""
    fn = (file_name or "").lower()
    if not (fn.endswith(".txt") or fn.endswith(".xlsx")):
        return None, "❌ Chỉ nhận file .txt hoặc .xlsx."
    rows = []
    if fn.endswith(".xlsx"):
        # .xlsx: 8 cột đầu = uid/link|mk|ngày tạo|mail thay|ghi chú|2fa|cookie|token
        from openpyxl import load_workbook
        import io as _io
        try:
            wb = load_workbook(_io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
        except Exception as e:
            return None, f"❌ Không đọc được file Excel: {e}"

        def _cell_str(c):
            # Excel hay convert UID dài thành số float (61593959792972.0):
            # ép về số nguyên để không lưu UID lỗi vào kho (gây check die oan).
            if c is None:
                return ""
            if isinstance(c, float) and c.is_integer():
                return str(int(c))
            return str(c).strip()

        first_row = True
        for vals in ws.iter_rows(values_only=True):
            cells = [_cell_str(c) for c in vals[:8]]
            while len(cells) < 8:
                cells.append("")
            if not any(cells):
                continue
            c0 = cells[0].lower()
            if first_row:
                first_row = False
                # Bỏ dòng tiêu đề (UID, Link, Url, STT...)
                if not (c0.isdigit() or "http" in c0 or "facebook.com" in c0 or "fb.com" in c0):
                    continue
            rows.append({
                "uid": cells[0], "password": cells[1], "created_date": cells[2],
                "backup_mail": cells[3], "note": cells[4], "totp": cells[5],
                "cookie": cells[6], "token": cells[7],
            })
    else:
        text = raw.decode("utf-8", errors="ignore")
        for line in text.splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            p = [x.strip() for x in line.split("|")]
            rows.append(_smart_stock_fields(p))
    if not rows:
        return None, "❌ File không có dòng acc nào hợp lệ."
    return rows, None

def _price_warn_pct() -> int:
    try:
        return max(1, int(db.get_setting("price_warn_pct", "50") or 50))
    except Exception:
        return 50

def _price_change_ok(old: int, new: int) -> bool:
    """True nếu mức đổi giá bình thường (không cần xác nhận)."""
    if old <= 0 or new < 0:
        return True
    pct = _price_warn_pct()
    lo = old * (100 - pct) / 100
    hi = old * (100 + pct) / 100
    return lo <= new <= hi

def _loyalty_gift_desc() -> str:
    """Mô tả quà đổi điểm hiện tại (acc hoặc tiền về ví)."""
    mode = db.get_setting("loyalty_redeem_mode", "acc") or "acc"
    pts = db.get_setting("loyalty_redeem_points", "10") or "10"
    if mode == "money":
        amt = db.get_setting("loyalty_redeem_amount", "0") or "0"
        wallet = db.get_setting("loyalty_redeem_wallet", "main") or "main"
        w = db.wallet_label(wallet)
        return f"<b>{db.wallet_amount_text(wallet, int(amt))}</b> vào {w} — <b>{html.escape(str(pts))}</b> điểm"
    scope = db.get_setting("loyalty_redeem_scope", "cat") or "cat"
    if scope == "stall":
        return f"1 acc <b>bất kỳ trong gian hàng Acc Facebook</b> — <b>{html.escape(str(pts))}</b> điểm"
    cid = db.loyalty_redeem_cat()
    c = db.acc_category_get(cid) if cid else None
    return f"1 acc <b>{html.escape(c['name']) if c else 'chưa cài'}</b> — <b>{html.escape(str(pts))}</b> điểm"

def _fb_stall_gift_cats() -> list:
    """Các loại acc còn hàng trong gian hàng Acc Facebook (để đổi quà)."""
    out = []
    for c in db.acc_category_list():
        c = dict(c)
        if (c.get("stall") or "Acc Facebook") != "Acc Facebook":
            continue
        n = db.acc_stock_count(c["id"])
        if n > 0:
            out.append((c, n))
    return out

async def _doiqua_redeem_acc(bot, tg_id: int, cat_id: int, need: int):
    """Trừ điểm + giao 1 acc live. Trả về (ok, order|None, reason)."""
    c = db.acc_category_get(cat_id)
    if not c or not c["active"] or db.acc_stock_count(cat_id) <= 0:
        return False, None, "out_of_stock"
    if not db.loyalty_consume(tg_id, need):
        return False, None, "points_changed"
    sold_orders, _sell_fail = await _sell_live_stock(
        bot, tg_id, 0, 1,
        lambda seen: db.acc_stock_pick_candidates(cat_id, 5, seen),
        check_live=_cat_live_check(cat_id))
    if not sold_orders:
        db.loyalty_add(tg_id, need, "hoan_diem_khong_du_hang_live")
        return False, None, "no_live"
    return True, sold_orders[0], ""

async def _fulfill_restock_subs(cat_id: int, bot, cat: dict):
    """Hàng về -> xử lý đăng ký báo hàng theo FIFO:
    - auto_buy=1: tự trừ ví shop, check LIVE rồi giao acc ngay.
    - auto_buy=0: chỉ nhắn báo hàng về.
    Sub nào ví không đủ tiền / chưa giao được thì GIỮ LẠI để thử đợt sau."""
    subs = db.acc_restock_subscribers(cat_id)
    if not subs:
        return
    up = db.shop_unit_price(cat)
    unit = up["price"]
    p5, p10, p20 = _shop_bulk_pcts()
    done_auto = 0
    for s in subs:
        tg_id = int(s["tg_id"])
        qty = max(1, int(s.get("qty") or 1))
        auto = int(s.get("auto_buy") or 0)
        if not auto:
            try:
                await bot.send_message(
                    tg_id,
                    f"🔔 <b>CÓ HÀNG RỒI!</b>\n\n"
                    f"📦 <b>{html.escape(cat['name'])}</b> vừa về "
                    f"<b>{db.acc_stock_count(cat_id)}</b> acc — "
                    f"giá {vnd(unit)}/acc.\n"
                    f"👉 Vào /shop mua ngay kẻo hết!",
                    parse_mode="HTML")
            except Exception:
                pass
            db.acc_unsub_restock(tg_id, cat_id)
            continue
        # Tự mua: giá như shop (giảm mua nhiều theo số lượng thực mua)
        stock = db.acc_stock_count(cat_id)
        buy_qty = min(qty, stock)
        if buy_qty <= 0:
            try:
                await bot.send_message(
                    tg_id,
                    f"😔 <b>Hàng bạn đặt trước:</b> <b>{html.escape(cat['name'])}</b> "
                    f"vừa về nhưng đã hết trước khi bot kịp giao. "
                    f"Bot giữ đơn đặt trước cho đợt hàng sau nhé!",
                    parse_mode="HTML")
            except Exception:
                pass
            continue  # giữ sub cho đợt sau
        bulk_pct = p20 if buy_qty >= 20 else (p10 if buy_qty >= 10 else (p5 if buy_qty >= 5 else 0))
        final = unit * buy_qty * (100 - bulk_pct) // 100
        u = db.get_user(tg_id)
        bal = int(u["shop_balance"] or 0) if u else 0
        if bal < final:
            try:
                await bot.send_message(
                    tg_id,
                    f"🔔 <b>HÀNG BẠN ĐẶT TRƯỚC ĐÃ VỀ!</b>\n\n"
                    f"📦 <b>{html.escape(cat['name'])}</b> — {buy_qty} acc = "
                    f"<b>{vnd(final)}</b> (ví shop bạn: {vnd(bal)}).\n"
                    f"👉 Nạp thêm bằng /napshop để bot tự giao acc ngay nhé! "
                    f"Bot giữ đơn đặt trước cho bạn.",
                    parse_mode="HTML")
            except Exception:
                pass
            continue  # giữ sub cho đợt sau
        if not db.adjust_shop_balance(tg_id, -final, f"dat_truoc:{cat_id}x{buy_qty}"):
            continue
        sold_orders, _sell_fail = await _sell_live_stock(
            bot, tg_id, final, buy_qty,
            lambda seen: db.acc_stock_pick_candidates(cat_id, buy_qty + 4, seen),
            check_live=_cat_live_check(cat_id))
        if not sold_orders:
            db.add_shop_balance_only(tg_id, final, "hoan_tien_dat_truoc")
            try:
                await bot.send_message(
                    tg_id,
                    "😔 <b>Hàng bạn đặt trước:</b> shop vừa kiểm tra lại, hiện chưa có "
                    "acc <b>LIVE</b> để giao. Tiền đã được hoàn vào ví shop — "
                    "có hàng live bot sẽ giao ngay cho bạn!",
                    parse_mode="HTML")
            except Exception:
                pass
            continue  # giữ sub cho đợt sau
        db.acc_unsub_restock(tg_id, cat_id)
        done_auto += 1
        for order in sold_orders:
            order_id = order["id"]
            try:
                await bot.send_message(
                    tg_id,
                    f"⚡ <b>ĐẶT TRƯỚC THÀNH CÔNG — ĐÃ GIAO ACC</b>\n\n"
                    f"📦 Loại: <b>{html.escape(cat['name'])}</b>\n"
                    f"🧾 Đơn hàng: <b>#{order_id}</b> | 💰 {vnd(order['price'])}\n"
                    f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
                    f"{_live_line(cat_id)}\n"
                    f"{_pickup_suffix()}",
                    parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
            except Exception:
                pass
        try:
            tickets = db.spin_add_tickets(tg_id, buy_qty)
            per = int(db.get_setting("loyalty_per_vnd", "100000") or 100000)
            pts = final // per if per > 0 else 0
            if pts > 0:
                db.loyalty_add(tg_id, pts, f"dat_truoc:{cat_id}")
            await bot.send_message(
                tg_id,
                f"🎁 Đặt trước {buy_qty} acc: nhận <b>{buy_qty} vé quay</b> /quay"
                + (f" + <b>{pts} điểm</b> loyalty" if pts > 0 else "")
                + f" (đang có {tickets} vé).",
                parse_mode="HTML")
        except Exception:
            pass
        try:
            await _notify_purchase_admin(bot, _buyer_shim(tg_id), sold_orders)
        except Exception:
            pass
    if done_auto:
        try:
            await _notify_admin_smart(
                bot,
                f"⚡ <b>Tự giao hàng đặt trước:</b> {html.escape(cat['name'])} — "
                f"đã giao <b>{done_auto}</b> đơn đặt trước.")
        except Exception:
            pass

async def _fulfill_deposits(cat_id: int, bot, cat: dict):
    """4.9 Hàng về -> giao acc cho người đặt cọc theo FIFO, trừ nốt tiền."""
    deps = db.acc_deposit_waiting(cat_id)
    if not deps:
        return
    up = db.shop_unit_price(cat)
    unit = up["price"]
    done = 0
    for d in deps:
        d = dict(d)
        tg_id = int(d["tg_id"])
        if db.acc_stock_count(cat_id) <= 0:
            break
        rest = unit - int(d["amount"])
        u = db.get_user(tg_id)
        bal = int(u["shop_balance"] or 0) if u else 0
        if rest > 0 and bal < rest:
            try:
                await bot.send_message(
                    tg_id,
                    f"🔔 <b>HÀNG BẠN ĐẶT CỌC ĐÃ VỀ!</b>\n\n"
                    f"📦 <b>{html.escape(cat['name'])}</b> — còn thiếu <b>{vnd(rest)}</b> "
                    f"(ví shop bạn: {vnd(bal)}).\n"
                    f"👉 Nạp thêm bằng /napshop để bot tự giao acc ngay nhé!",
                    parse_mode="HTML")
            except Exception:
                pass
            continue
        if rest > 0:
            if not db.adjust_shop_balance(tg_id, -rest, f"dat_coc_giao:{cat_id}"):
                continue
        sold_orders, _sell_fail = await _sell_live_stock(
            bot, tg_id, unit, 1,
            lambda seen: db.acc_stock_pick_candidates(cat_id, 5, seen),
            check_live=_cat_live_check(cat_id))
        if not sold_orders:
            if rest > 0:
                db.add_shop_balance_only(tg_id, rest, "hoan_tien_dat_coc")
            try:
                await bot.send_message(
                    tg_id,
                    "😔 <b>Hàng bạn đặt cọc:</b> shop vừa kiểm tra lại, hiện chưa có "
                    "acc <b>LIVE</b> để giao. Tiền cọc vẫn được giữ — có hàng live "
                    "bot sẽ giao ngay cho bạn!",
                    parse_mode="HTML")
            except Exception:
                pass
            break
        order = sold_orders[0]
        order_id = order["id"]
        db.acc_deposit_set_status(d["id"], "DONE", order_id)
        done += 1
        try:
            await bot.send_message(
                tg_id,
                f"✅ <b>ĐẶT CỌC THÀNH CÔNG — ĐÃ GIAO ACC</b>\n\n"
                f"📦 Loại: <b>{html.escape(cat['name'])}</b>\n"
                f"💰 Đã cọc: {vnd(d['amount'])} + trừ thêm {vnd(rest)}\n"
                f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
                f"{_live_line(cat_id)}\n"
                f"{_pickup_suffix()}",
                parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
        except Exception:
            pass
        await _notify_purchase_admin(bot, _buyer_shim(tg_id), [order])
    if done:
        await _notify_admin_smart(
            bot,
            f"💰 <b>Tự giao hàng đặt cọc:</b> {html.escape(cat['name'])} — "
            f"đã giao <b>{done}</b> acc cho khách đặt cọc.")

async def _do_deposit(tg_id: int, cat_id: int, msg, bot):
    c = db.acc_category_get(cat_id)
    if not c:
        await msg.answer("❌ Loại acc không tồn tại.")
        return
    c = dict(c)
    if db.acc_stock_count(cat_id) > 0:
        await msg.answer("✅ Loại này đang còn hàng — bạn mua luôn ở /shop nhé!")
        return
    try:
        dep_pct = int(db.get_setting("deposit_pct", "30") or 30)
    except Exception:
        dep_pct = 30
    amt = int(c["price"]) * dep_pct // 100
    amt, promo_code = db.preview_user_promo(tg_id, amt, wallet="shop")
    u = db.get_user(tg_id)
    bal = int(u["shop_balance"] or 0) if u else 0
    if bal < amt:
        await msg.answer(
            f"❌ Ví shop không đủ đặt cọc.\nCọc {dep_pct}% = <b>{vnd(amt)}</b>, "
            f"ví shop bạn: <b>{vnd(bal)}</b>.\nNạp thêm bằng /napshop nhé."
            + _SHOP_WALLET_HINT,
            parse_mode="HTML")
        return
    if not db.adjust_shop_balance(tg_id, -amt, f"dat_coc:{cat_id}" + (f" (promo {promo_code})" if promo_code else "")):
        await msg.answer("❌ Ví shop không đủ đặt cọc." + _SHOP_WALLET_HINT, parse_mode="HTML")
        return
    dep_id = db.acc_deposit_create(tg_id, cat_id, amt)
    db.finalize_user_promo(tg_id, promo_code)  # trừ lượt promo SAU khi đặt cọc thành công
    await msg.answer(
        f"✅ <b>ĐẶT CỌC THÀNH CÔNG #{dep_id}</b>\n\n"
        f"📦 Loại: <b>{html.escape(c['name'])}</b>\n"
        f"💰 Đã cọc: <b>{vnd(amt)}</b> ({dep_pct}% giá {vnd(c['price'])})\n\n"
        f"👉 Khi có hàng về, bot <b>tự giao acc ngay</b> và trừ nốt "
        f"<b>{vnd(int(c['price']) - amt)}</b> từ ví shop của bạn.\n"
        f"Xem/hủy cọc: /coclist",
        parse_mode="HTML")

__all__ = [
    "_STOCK_DIE_STATUSES",
    "DEFAULT_MAIL_APP_LINK",
    "TrackMenuState",
    "AccShopState",
    "_run_sheet_import",
    "process_tiktok_check",
    "process_ig_check",
    "process_fb_check",
    "_notify_super_subadmin_action",
    "status_caption",
    "_send_card",
    "_esc_list_name",
    "_run_scanlist",
    "check_uids_batch",
    "_ensure_bulk_credits_simple",
    "_build_excel_bytes",
    "_do_doitien",
    "_show_history",
    "_maybe_low_credit_warn",
    "_trackmenu_main_kb",
    "_trackmenu_sub_kb",
    "_vip_limit_for",
    "_is_cancel",
    "_purchase_alert_text",
    "_buyer_shim",
    "_notify_purchase_admin",
    "_check_stock_live",
    "_notify_die_quarantine",
    "_sell_live_stock",
    "_sell_uid_items",
    "_notify_shop_admin",
    "_notify_admin_photo",
    "_notify_admin_smart",
    "_totp_now",
    "_mail_app_link",
    "_mail_app_line",
    "_pickup_suffix",
    "_acc_delivery_caption",
    "_cat_icon",
    "_shop_stall_picker",
    "_shop_list",
    "_shop_bulk_pcts",
    "_acc_after_purchase",
    "_send_merged_acc_file",
    "_uid_final_price",
    "_line_price",
    "_cart_render",
    "_loyalty_random_cfg",
    "_loyalty_random_received",
    "_loyalty_random_today",
    "_order_total_for_randkey",
    "_spin_kb",
    "_do_spin",
    "_acc_delivery_kb",
    "_strip_acc",
    "_parse_warranty",
    "_fmt_warranty",
    "_cat_live_check",
    "_live_line",
    "_smart_stock_fields",
    "_resolve_stock_links",
    "_import_stock_rows",
    "_parse_stock_file",
    "_price_warn_pct",
    "_price_change_ok",
    "_loyalty_gift_desc",
    "_fb_stall_gift_cats",
    "_doiqua_redeem_acc",
    "_fulfill_restock_subs",
    "_fulfill_deposits",
    "_do_deposit",
]