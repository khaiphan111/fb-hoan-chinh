import asyncio
import html
import logging
import os
import time
from typing import Optional

import httpx
import re
from aiogram import Bot, Dispatcher, Router, F
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.aiohttp import AiohttpSession
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandStart, CommandObject, StateFilter
from aiogram.types import (
    BotCommand, Message, URLInputFile, FSInputFile, BufferedInputFile,
    ReplyKeyboardMarkup, KeyboardButton,
    InlineKeyboardMarkup, InlineKeyboardButton, CallbackQuery,
)
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup

class WarrantyClaimState(StatesGroup):
    waiting_for_evidence = State()


class BankState(StatesGroup):
    waiting_for_amount = State()

class PayOSState(StatesGroup):
    waiting_for_amount = State()
    waiting_for_shop_amount = State()

class FBNoteState(StatesGroup):
    waiting_for_note = State()
    uid = None


class TienIchState(StatesGroup):
    """States cho menu tiện ích /tienich — nhập liệu từng bước thay vì gõ lệnh tay."""
    waiting_for_giftcode = State()
    waiting_for_promo = State()
    waiting_for_birthday = State()
    waiting_for_refcode = State()
    waiting_for_tiktok = State()
    waiting_for_ig = State()
    waiting_for_fb = State()
    waiting_for_getuid = State()
    # /ruttien 3 bước
    waiting_withdraw_amount = State()
    waiting_withdraw_bank = State()
    waiting_withdraw_stk = State()
    # /chuyentien 2 bước
    waiting_transfer_uid = State()
    waiting_transfer_amount = State()

from . import db
from . import perms as _perms
from . import notify_bot as _notify_bot
from . import config as _config
from .persist import SQLiteStorage, check_cache_get, check_cache_set, init_cache_db
from .util import now, parse_check_args, vnd, vn_time_str
DAY = 86400
from .tiktok import parse_username, fetch_tiktok_info, fmt_num, build_info_caption
from .ig import (
    parse_ig_username, parse_ig_post_id,
    fetch_ig_info, fetch_ig_post_info,
    build_ig_info_caption, build_ig_video_caption
)
from .fb import check_uid, build_fb_caption
from .poller import poller

log = logging.getLogger(__name__)
router = Router()


@router.message.middleware()
async def _perm_middleware(handler, event, data):
    """Chặn admin phụ gõ lệnh tay thuộc nhóm quyền chưa được cấp.

    Chỉ áp dụng cho admin phụ (không phải super admin). Super admin và
    khách thường đi qua bình thường. Menu nút (/adm, /shopadm) tự lọc riêng.
    """
    try:
        from aiogram.types import Message as _Msg
        if isinstance(event, _Msg) and event.from_user:
            uid = event.from_user.id
            if _perms.is_admin(uid) and not _perms.is_super(uid):
                need = _perms.cmd_perm_for_text(event.text or "")
                if need and not _perms.has_perm(uid, need):
                    await event.answer(
                        f"🚫 Bạn không có quyền <b>{_perms.perm_label(need)}</b>.\n"
                        f"Liên hệ chủ shop để được cấp thêm quyền.",
                        parse_mode="HTML")
                    return
            if _perms.is_admin(uid):
                # Nhật ký: admin (kể cả chủ shop) gõ lệnh tay thuộc diện quản trị
                _need2 = _perms.cmd_perm_for_text(event.text or "")
                _head = ((event.text or "").strip().split() or [""])[0].lower()
                if _need2 or _head in ("/adm", "/shopadm"):
                    try:
                        db.admin_audit_add(
                            uid, event.from_user.full_name, "lenh_tay",
                            (event.text or "")[:200])
                    except Exception:
                        pass
    except Exception:
        pass
    return await handler(event, data)

from .admin_bot import register_adm_menu as _register_adm_menu
_register_adm_menu(router)  # menu nút /adm — đăng ký sớm, trước on_other
from .shop_menu import register_shop_menu as _register_shop_menu
_register_shop_menu(router)  # menu nút /shopadm — đăng ký sớm, trước on_other
from aiogram import BaseMiddleware
from aiogram.types import Message

_user_last_cmd = {}
_user_spam_count = {}
_user_muted_until = {}

_spam_cleanup_ts = 0

class AntiSpamMiddleware(BaseMiddleware):
    async def __call__(self, handler, event, data):
        global _spam_cleanup_ts
        if isinstance(event, Message) and event.text:
            # Key theo (chat, user) để 1 người spam không mute cả group
            uid = event.from_user.id if event.from_user else 0
            key = (event.chat.id, uid)
            now_t = time.time()

            # Dọn dict chống rò rỉ bộ nhớ mỗi 10 phút
            if now_t - _spam_cleanup_ts > 600:
                _spam_cleanup_ts = now_t
                for k in [k for k, v in _user_muted_until.items() if v < now_t]:
                    _user_muted_until.pop(k, None)
                for k in [k for k, v in _user_last_cmd.items() if v < now_t - 3600]:
                    _user_last_cmd.pop(k, None)
                if len(_user_spam_count) > 5000:
                    _user_spam_count.clear()

            # Exempt admin completely from anti-spam / muting
            admin_id = db.get_setting("admin_tg_id")
            if admin_id and str(uid) == str(admin_id):
                return await handler(event, data)

            # Check if muted
            if key in _user_muted_until:
                if now_t < _user_muted_until[key]:
                    return  # Ignore silently
                else:
                    del _user_muted_until[key]
                    if key in _user_spam_count:
                        del _user_spam_count[key]

            # Anti flood (1 cmd per 3 seconds)
            last_cmd_time = _user_last_cmd.get(key, 0)
            if now_t - last_cmd_time < 3:
                _user_spam_count[key] = _user_spam_count.get(key, 0) + 1
                if _user_spam_count[key] >= 5:
                    _user_muted_until[key] = now_t + 900  # Mute 15 mins
                    await event.answer("🚫 Bạn đã gửi lệnh quá nhanh liên tục. Hệ thống tạm khóa bạn trong 15 phút để chống spam!")
                    return
                elif _user_spam_count[key] == 3:
                    await event.answer("⚠️ Cảnh báo: Vui lòng gửi lệnh chậm lại (mỗi 3 giây 1 lệnh). Nếu tiếp tục bạn sẽ bị khóa 15 phút!")
                    return
                return # Ignore too fast cmds without warning if count < 3
            else:
                _user_spam_count[key] = 0

            _user_last_cmd[key] = now_t
            
            # Sub check (original logic)
            cmd = event.text.split()[0].lower()
            if "@" in cmd:
                cmd = cmd.split("@")[0]
                
            main_group_id = db.get_setting("main_tg_group_id")
            is_main_group = main_group_id and str(event.chat.id) == str(main_group_id)
            
            if is_main_group:
                user = db.get_user(event.chat.id)
                if not user:
                    db.upsert_user(event.chat.id, "admin_group", "Admin Group")
                    db.set_sub_until(event.chat.id, now() + 3650*24*3600)
                return await handler(event, data)
                
            # Whitelist public / free commands
            free_cmds = (
                "/start", "/help", "/balance", "/sub", "/bank", "/nap", "/ref", "/web",
                "/refcode", "/ruttien", "/doitien", "/vip", "/trial", "/mycodes",
                "/hdcookie", "/ping2", "/alert", "/alertlist", "/alertoff",
                "/muagoi", "/checkfile", "/checkcookie", "/dailyreport", "/code",
                "/stats", "/history", "/top", "/adm", "/daily", "/chuyentien",
                "/shop", "/damua", "/quay", "/hang", "/doiqua", "/giasi",
                "/coc", "/huycoc", "/coclist", "/napshop", "/giohang"
            )
            if cmd not in free_cmds:
                user = db.get_user(event.chat.id)
                if not user:
                    await event.answer("Bạn chưa /start. Gõ /start trước nhé.")
                    return
                    
                # Check daily limit for all tracking and checking cmds
                if cmd in ("/check", "/tiktok", "/ig", "/track", "/trackv", "/trackig", "/trackvig", "/trackfb"):
                    can_check, err_msg = db.check_daily_limit(event.chat.id)
                    if not can_check:
                        await event.answer(f"❌ {err_msg}")
                        return
                        
                has_sub = user["sub_until"] and user["sub_until"] > now()
                has_balance = user["balance"] and user["balance"] > 0
                if not has_sub and not has_balance:
                    await event.answer("⚠️ Lỗi: Bạn cần được Admin cấp tiền hoặc cấp gói ngày sử dụng để dùng các chức năng này.\n👉 Gõ /balance để kiểm tra số dư, gõ /sub để mua gói.")
                    return
        return await handler(event, data)

router.message.middleware(AntiSpamMiddleware())


MENU = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="/muagoi"), KeyboardButton(text="/checkfile"), KeyboardButton(text="/checkcookie")],
        [KeyboardButton(text="/theodoi"), KeyboardButton(text="/tiktok"), KeyboardButton(text="/ig")],
        [KeyboardButton(text="/check"), KeyboardButton(text="/list"), KeyboardButton(text="/balance"), KeyboardButton(text="/sub")],
        [KeyboardButton(text="/vip"), KeyboardButton(text="/ref"), KeyboardButton(text="/bank")],
        [KeyboardButton(text="/tienich"), KeyboardButton(text="/shop"), KeyboardButton(text="/help")],
        [KeyboardButton(text="/web"), KeyboardButton(text="/huongdan")],
    ],
    resize_keyboard=True,
)

# Bàn phím riêng cho admin: thêm nút nhập kho từ Google Sheet luôn hiển thị
ADMIN_MENU = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="/muagoi"), KeyboardButton(text="/checkfile"), KeyboardButton(text="/checkcookie")],
        [KeyboardButton(text="/theodoi"), KeyboardButton(text="/tiktok"), KeyboardButton(text="/ig")],
        [KeyboardButton(text="/check"), KeyboardButton(text="/list"), KeyboardButton(text="/balance"), KeyboardButton(text="/sub")],
        [KeyboardButton(text="/vip"), KeyboardButton(text="/ref"), KeyboardButton(text="/bank")],
        [KeyboardButton(text="/tienich"), KeyboardButton(text="/shop"), KeyboardButton(text="/help")],
        [KeyboardButton(text="/web"), KeyboardButton(text="/adm"), KeyboardButton(text="/huongdan")],
        [KeyboardButton(text="/shopadm")],
        [KeyboardButton(text="\U0001f4ca Nhập kho Sheet")],
    ],
    resize_keyboard=True,
)

ADMIN_COMMANDS_EXTRA = [
    BotCommand(command="setsheet", description="Cài đặt Google Sheet nhập kho"),
    BotCommand(command="nhapkhosheet", description="Nhập kho từ Google Sheet: /nhapkhosheet <id_loại>"),
    BotCommand(command="adm", description="Bảng lệnh admin"),
    BotCommand(command="shopadm", description="🛒 Menu shop acc & NCC (nút bấm)"),
    BotCommand(command="kho", description="Xem tồn kho"),
    BotCommand(command="themacc", description="Nhập kho từ file: /themacc <id_loại>"),
]


class SheetImportState(StatesGroup):
    waiting_ncc_cost = State()


async def _run_sheet_import(msg, cat_id: int, ncc_id: int, cost: int):
    """Chạy nhập kho từ Google Sheet (dùng chung cho lệnh /nhapkhosheet và nút bấm)."""
    c = db.acc_category_get(cat_id)
    if not c:
        await msg.answer("❌ Không có loại acc này. Xem: /kho")
        return
    if ncc_id and not db.supplier_get(ncc_id):
        await msg.answer("❌ Không có NCC này. Xem: /ncc")
        return
    sid = db.get_setting("sheet_import_id") or ""
    default_tab = db.get_setting("sheet_import_tab") or "NhapKho"
    if not sid:
        await msg.answer("📊 Chưa cài đặt sheet nhập kho. Dùng: <code>/setsheet &lt;link&gt;</code>",
                         parse_mode="HTML")
        return
    from . import sheet_import as _si
    try:
        stall = (c["stall"] if "stall" in c.keys() else "") or "Acc Facebook"
    except Exception:
        stall = "Acc Facebook"
    tab = _si.tab_for_stall(stall, default_tab)
    wait = await msg.answer("⏳ Đang đọc Google Sheet...")
    try:
        if not await _si.ensure_tab(sid, tab):
            await wait.edit_text("❌ Không tạo/kiểm tra được tab Sheet. Thử lại sau.")
            return
        sheet_rows = await _si.read_unmarked(sid, tab)
    except Exception as e:
        await wait.edit_text(f"❌ Không đọc được sheet: {html.escape(str(e)[:200])}")
        return
    if not sheet_rows:
        await wait.edit_text("📭 Sheet không có dòng mới nào.\n"
                             "<i>Các dòng đã có đánh dấu ở cột Trạng thái sẽ bị bỏ qua.</i>",
                             parse_mode="HTML")
        return
    rows = []
    for rnum, cells in sheet_rows:
        rows.append({
            "uid": cells[0], "password": cells[1], "created_date": cells[2],
            "backup_mail": cells[3], "note": cells[4], "totp": cells[5],
            "cookie": cells[6], "token": cells[7], "_sheet_row": rnum,
        })
    extra = ""
    if ncc_id:
        extra += f" NCC #{ncc_id}"
    if cost:
        extra += f" vốn {vnd(cost)}/acc"
    try:
        await wait.edit_text(f"⏳ Đọc được <b>{len(rows)}</b> dòng mới từ sheet.{extra}\nĐang nhập kho...",
                             parse_mode="HTML")
    except Exception:
        pass
    await _import_stock_rows(rows, cat_id, ncc_id, cost, c, msg, wait,
                            sheet_ctx={"sheet_id": sid, "tab": tab})


@router.message(F.text == "\U0001f4ca Nhập kho Sheet")
async def on_sheet_button(msg: Message, state: FSMContext):
    """Nút bấm cố định trên bàn phím admin: mở flow nhập kho từ Sheet."""
    if not _is_admin(msg.from_user.id):
        return
    if not _perms.has_perm(msg.from_user.id, "kho"):
        await msg.answer("🚫 Bạn không có quyền 📦 Kho & loại acc.")
        return
    await state.clear()
    sid = db.get_setting("sheet_import_id") or ""
    tab = db.get_setting("sheet_import_tab") or "NhapKho"
    cats = db.acc_category_list(active_only=False, include_hidden=True)
    if not cats:
        await msg.answer("❌ Chưa có loại acc nào. Tạo trước bằng /themloai")
        return
    kb_rows, row = [], []
    for cc in cats:
        dd = dict(cc)
        name = (dd.get("name") or f"Loại {dd['id']}")[:18]
        row.append(InlineKeyboardButton(text=f"{dd['id']}. {name}",
                                       callback_data=f"sheetpick_{dd['id']}"))
        if len(row) == 2:
            kb_rows.append(row)
            row = []
    if row:
        kb_rows.append(row)
    txt = "📊 <b>NHẬP KHO TỪ GOOGLE SHEET</b>\n"
    if sid:
        txt += (f'🔗 <a href="https://docs.google.com/spreadsheets/d/{sid}/edit">Mở sheet</a> '
                f"(tab {html.escape(tab)})\n")
    txt += ("\n<i>Chỉ quét dòng chưa có đánh dấu ở cột Trạng thái.</i>\n"
            "\n👇 <b>Chọn loại acc để nhập:</b>")
    await msg.answer(txt, reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows),
                     parse_mode="HTML", disable_web_page_preview=True)


@router.callback_query(F.data.startswith("sheetpick_"))
async def on_sheet_pick(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("🚫")
        return
    if not _perms.has_perm(cb.from_user.id, "kho"):
        await cb.answer("🚫 Không có quyền.", show_alert=True)
        return
    try:
        cat_id = int(cb.data.split("_", 1)[1])
    except Exception:
        await cb.answer("❌")
        return
    c = db.acc_category_get(cat_id)
    if not c:
        await cb.answer("❌ Loại không tồn tại", show_alert=True)
        return
    await state.update_data(sheet_cat_id=cat_id)
    await state.set_state(SheetImportState.waiting_ncc_cost)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Nhập luôn (không NCC/vốn)",
                              callback_data=f"sheetgo_{cat_id}_0_0")],
    ])
    await cb.message.answer(
        f"📦 Loại: <b>{html.escape(c['name'])}</b> (#{cat_id})\n\n"
        "Nhập <code>ncc_id giá_vốn</code> (vd: <code>3 15000</code>)\n"
        "hoặc bấm nút để nhập luôn:",
        reply_markup=kb, parse_mode="HTML")
    await cb.answer()


@router.message(SheetImportState.waiting_ncc_cost)
async def on_sheet_ncc_cost(msg: Message, state: FSMContext):
    if not _is_admin(msg.from_user.id):
        await state.clear()
        return
    if not _perms.has_perm(msg.from_user.id, "kho"):
        await state.clear()
        await msg.answer("🚫 Bạn không có quyền 📦 Kho & loại acc.")
        return
    data = await state.get_data()
    cat_id = data.get("sheet_cat_id")
    await state.clear()
    if not cat_id:
        return
    parts = (msg.text or "").split()
    try:
        ncc_id = int(parts[0]) if len(parts) > 0 else 0
        cost = int(parts[1].replace(".", "").replace(",", "")) if len(parts) > 1 else 0
    except Exception:
        await msg.answer("❌ Nhập số thôi, vd: <code>3 15000</code> — hoặc bấm lại nút 📊 để làm lại.",
                         parse_mode="HTML")
        return
    await _run_sheet_import(msg, cat_id, ncc_id, cost)


@router.callback_query(F.data.startswith("sheetgo_"))
async def on_sheet_go(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("🚫")
        return
    if not _perms.has_perm(cb.from_user.id, "kho"):
        await cb.answer("🚫 Không có quyền.", show_alert=True)
        return
    await state.clear()
    try:
        _, cat_id, ncc_id, cost = cb.data.split("_")
    except Exception:
        await cb.answer("❌")
        return
    await cb.answer("⏳ Đang nhập kho...")
    await _run_sheet_import(cb.message, int(cat_id), int(ncc_id), int(cost))


COMMANDS = [
    BotCommand(command="start",       description="Bắt đầu sử dụng bot"),
    BotCommand(command="muagoi",      description="Mua gói theo ngày (1, 3, 5, 7 ngày)"),
    BotCommand(command="sub",         description="Xem gói tháng & mua gói"),
    BotCommand(command="bank",        description="Nạp tiền / Lấy thông tin chuyển khoản"),
    BotCommand(command="nap",         description="Nạp tiền tự động qua PayOS (chọn ví)"),
    BotCommand(command="napshop",     description="Nạp tiền vào ví shop mua acc"),
    BotCommand(command="balance",     description="Xem số dư hiện tại"),
    BotCommand(command="sodu",        description="Xem tất cả số dư: ví, credits, điểm"),
    BotCommand(command="huongdan",    description="Hướng dẫn nhanh theo từng mục"),
    BotCommand(command="vip",         description="Xem cấp độ VIP và đặc quyền"),
    BotCommand(command="checkfile",   description="Check UID Facebook qua file .txt/.xlsx"),
    BotCommand(command="checkcookie", description="Check dàn Cookie format UID|PASS|COOKIE|2FA"),
    BotCommand(command="muacredit",   description="Mua gói Credits (lượt check)"),
    BotCommand(command="lichsu",      description="Lịch sử check 7 ngày"),
    BotCommand(command="app",         description="Mở Mini App check UID"),
    BotCommand(command="accuracy",    description="Thống kê độ chính xác check"),
    BotCommand(command="theodoi",     description="👁️ Trung tâm theo dõi (menu nút gọn)"),
    BotCommand(command="tienich",      description="🧰 Tiện ích: giftcode, mã giảm giá, ngày sinh, rút/chuyển tiền, lịch sử, đặt cọc (menu nút)"),
    BotCommand(command="dailyreport", description="Cài đặt báo cáo tự động hằng ngày"),
    BotCommand(command="stats",       description="Thống kê cá nhân của bạn"),
    BotCommand(command="daily",       description="Điểm danh nhận credits mỗi ngày"),
    BotCommand(command="tiktok",      description="Check info TikTok: /tiktok <username>"),
    BotCommand(command="ig",          description="Check info Instagram: /ig <username>"),
    BotCommand(command="fb",          description="Check Facebook Live/Die: /fb <uid>"),
    BotCommand(command="shop",        description="Shop tài khoản Facebook"),
    BotCommand(command="damua",       description="Acc FB đã mua + yêu cầu bảo hành"),
    BotCommand(command="quay",        description="Vòng quay may mắn (vé từ mua acc)"),
    BotCommand(command="hang",        description="Hạng thành viên shop acc"),
    BotCommand(command="giasi",       description="Bảng giá sỉ mua nhiều"),
    BotCommand(command="doiqua",      description="Đổi điểm loyalty lấy acc miễn phí"),
    BotCommand(command="coclist",     description="Xem các khoản đặt cọc của bạn"),
    BotCommand(command="ref",         description="Lấy link giới thiệu kiếm tiền"),
    BotCommand(command="help",        description="Hướng dẫn sử dụng"),
    BotCommand(command="web",         description="Đăng nhập Bảng điều khiển Web"),
]


@router.message(Command("web"))
async def cmd_web(msg: Message):
    tg_id = msg.chat.id
    web_domain = (db.get_setting("web_domain", "") or "").strip().rstrip("/")
    if not web_domain:
        await msg.answer(
            "🌐 <b>WEB DASHBOARD</b>\n\n"
            "Web chưa được cấu hình public URL.\n"
            "<i>Admin: set setting <code>web_domain</code> thành địa chỉ public của backend "
            "(VD: https://bot.ban.com) rồi thử lại.</i>",
            parse_mode="HTML",
        )
        return
    token = db.create_magic_link(tg_id)
    url = f"{web_domain}/auth?token={token}"
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🌐 Đăng nhập Web", url=url)
    ]])
    await msg.answer("🔗 Bấm vào nút bên dưới để tự động đăng nhập vào Web:", reply_markup=kb)

# ─── PROCESS TIKTOK CHECK ─────────────────────────────────────
async def process_tiktok_check(msg: Message, username: str):
    wait = await msg.answer(f"⏳ Đang kiểm tra <b>@{username}</b>...")
    try:
        info    = await fetch_tiktok_info(username)
        caption = build_info_caption(info)
        if info["avatar"]:
            try:
                await msg.answer_photo(
                    photo=URLInputFile(info["avatar"], filename="avatar.jpg"),
                    caption=caption,
                )
                await wait.delete()
                return
            except Exception:
                pass
        await wait.edit_text(caption, disable_web_page_preview=True)
    except ValueError as e:
        await wait.edit_text(f"❌ {e}")
    except httpx.TimeoutException:
        await wait.edit_text("⏰ Timeout! Thử lại sau.")
    except Exception as e:
        log.exception("Lỗi check @%s", username)
        await wait.edit_text(f"❌ Lỗi: {e}")

async def process_ig_check(msg: Message, username: str):
    wait = await msg.answer(f"⏳ Đang kiểm tra IG <b>@{username}</b>...")
    try:
        info    = await fetch_ig_info(username)
        caption = build_ig_info_caption(info)
        if info.get("avatar"):
            try:
                await msg.answer_photo(
                    photo=URLInputFile(info["avatar"], filename="avatar.jpg"),
                    caption=caption,
                )
                await wait.delete()
                return
            except Exception:
                pass
        await wait.edit_text(caption, disable_web_page_preview=True)
    except ValueError as e:
        await wait.edit_text(f"❌ {e}")
    except httpx.TimeoutException:
        await wait.edit_text("⏰ Timeout! Thử lại sau.")
    except Exception as e:
        log.exception("Lỗi check IG @%s", username)
        await wait.edit_text(f"❌ Lỗi: {e}")

async def process_fb_check(msg: Message, uid: str):
    wait = await msg.answer(f"⏳ Đang kiểm tra FB UID <b>{uid}</b>...")
    try:
        res = await check_uid(uid)
        try:
            db.log_check_stat(msg.from_user.id, "fb", res.get("uid") or uid,
                              res.get("status", ""), res.get("via", ""))
            db.add_check_history(msg.from_user.id, "fb", res.get("uid") or uid,
                                 res.get("status", ""))
        except Exception:
            pass
        caption = build_fb_caption(res)

        from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="📝 Ghi chú", callback_data=f"fb_note_{res['uid']}"),
                InlineKeyboardButton(text="👁 Theo dõi", callback_data=f"fb_track_{res['uid']}")
            ]
        ])

        avatar = res.get("avatar_url")
        # Chỉ gửi ảnh khi là avatar THẬT (acc dùng ảnh mặc định thì không gửi ảnh trắng)
        if res.get("has_real_avatar") and avatar:
            # Acc LIVE có avatar thật => gửi kèm ảnh avatar của acc
            try:
                await wait.delete()
            except Exception:
                pass
            try:
                await msg.answer_photo(photo=avatar, caption=caption, reply_markup=kb)
            except Exception:
                # Telegram không tải được ảnh => fallback về tin nhắn chữ
                await msg.answer(caption, disable_web_page_preview=True, reply_markup=kb)
        else:
            await wait.edit_text(caption, disable_web_page_preview=True, reply_markup=kb)
    except Exception as e:
        log.exception("Lỗi check FB %s", uid)
        await wait.edit_text(f"❌ Lỗi: {e}")


# ─── HANDLERS ────────────────────────────────────────────────
# Tin chào mừng khách mới (chỉ hiện lần đầu /start)
_WELCOME_NEW = (
    "🎉 <b>Chào mừng bạn đến với shop!</b>\n"
    "━━━━━━━━━━━━━━━\n"
    "🛒 <b>Mua acc FB:</b> gõ /shop — chọn loại acc, thanh toán là nhận acc ngay\n"
    "💳 <b>Nạp tiền:</b> gõ /nap &lt;số tiền&gt; — quét mã QR là tiền vào ví\n"
    "🎁 <b>Tân thủ</b> được tặng ngày dùng thử miễn phí (nếu đang bật)\n"
    "❓ Cần hỗ trợ? Gõ /huongdan để xem hướng dẫn theo từng mục\n"
    "━━━━━━━━━━━━━━━\n\n"
)

# Gợi ý ví đúng lúc: chỉ hiện khi thiếu tiền ví shop (đỡ rối cho khách mới)
_SHOP_WALLET_HINT = (
    "\n<i>💡 Ví shop là ví riêng để mua acc (khác ví chính dùng để check UID/mua gói)."
    " Nạp bằng /napshop nhé.</i>"
)


@router.message(CommandStart())
async def on_start(msg: Message):
    u = msg.from_user
    parts = (msg.text or "").split(maxsplit=1)
    ref_id = 0
    if len(parts) > 1:
        code_or_id = parts[1].strip()
        try: 
            ref_id = int(code_or_id)
        except: 
            c = db.get_conn()
            row = c.execute("SELECT tg_id FROM tg_users WHERE ref_code=?", (code_or_id,)).fetchone()
            if row:
                ref_id = row["tg_id"]
        
    if ref_id and int(ref_id) == int(u.id):
        ref_id = 0  # chống tự giới thiệu chính mình để ăn hoa hồng
    user = db.get_user(u.id)
    is_new = False
    if not user:
        is_new = True
        user = db.upsert_user(u.id, u.username or "", u.full_name or "", ref_id)
        
        # New User Notification for Admin
        admin_msg = (
            "🎉 <b>CÓ NGƯỜI DÙNG MỚI!</b>\n"
            "━━━━━━━━━━━━━━━━━━━━\n"
            f"👤 <b>Tên:</b> {html.escape(u.full_name or '')}\n"
            f"🔗 <b>Username:</b> @{html.escape(u.username) if u.username else 'Không có'}\n"
            f"🆔 <b>ID:</b> <code>{u.id}</code>\n"
            f"🕒 <b>Thời gian:</b> {vn_time_str('%H:%M:%S %d/%m/%Y')}\n"
        )
        if ref_id > 0 and ref_id != u.id:
            admin_msg += f"🤝 <b>Mời bởi:</b> <code>{ref_id}</code>\n"
            try:
                await msg.bot.send_message(ref_id, f"🎉 <b>Tin vui!</b>\nNgười dùng <b>{html.escape(u.full_name or "")}</b> vừa tham gia Bot qua link giới thiệu của bạn!\nKhi họ nạp tiền bạn sẽ nhận được 10% hoa hồng.", parse_mode="HTML")
            except: pass
            
        # Send to admin tg
        admins = []
        try:
            if db.get_setting("admin_tg_id"): admins.append(int(db.get_setting("admin_tg_id")))
        except: pass
        try:
            if db.get_setting("admin_tg_group_id"): admins.append(int(db.get_setting("admin_tg_group_id")))
        except: pass
        
        admin_tg_token = db.get_setting("admin_bot_token", "")
        if admin_tg_token:
            from .admin_bot import manager as admin_manager
            admin_sender_bot = admin_manager.bot
        else:
            admin_sender_bot = msg.bot
            
        if admin_sender_bot:
            for admin_id in admins:
                try:
                    await admin_sender_bot.send_message(admin_id, admin_msg, parse_mode="HTML")
                except Exception as e:
                    log.error("Failed to notify TG admin %s of new user: %s", admin_id, e)
                
        # Send to admin zalo
        admin_zalo = db.get_setting("admin_zalo_id", "")
        if admin_zalo:
            try:
                if zalo_manager.running:
                    asyncio.create_task(zalo_manager.send_message(admin_zalo, admin_msg))
            except: pass
    else:
        user = db.upsert_user(u.id, u.username or "", u.full_name or "", 0)
    db.add_log("system", f"/start {u.id} @{u.username}", u.id)
    
    # Auto Trial Logic
    trial_msg = ""
    if db.get_setting("enable_free_trial", "1") == "1":
        try:
            days = int(db.get_setting("free_trial_days", "3"))
        except ValueError:
            days = 3
        if db.activate_trial(u.id, days):
            db.add_log("trial", f"Auto trial {days} ngày", u.id)
            user = db.get_user(u.id) # refresh user data
            trial_msg = f"🎁 <b>Quà tặng tân thủ:</b> Bạn đã được hệ thống tự động tặng <b>{days} ngày</b> dùng thử miễn phí!\n\n"
    
    await msg.answer(
        f"👋 Xin chào <b>{msg.from_user.full_name}</b>!\n\n"
        "📱 Bot <b>TikTok/IG/FB Checker V2</b> sẵn sàng!\n\n"
        f"💰 Số dư: <b>{vnd(user['balance'])}</b>\n"
        f"Gói FB: <b>{_sub_text(user)}</b>\n\n"
        f"{trial_msg}"
        + (_WELCOME_NEW if is_new else "") +
        "Gõ /huongdan để xem hướng dẫn theo từng mục.\n"
        "Gõ /ref để lấy link giới thiệu nhận 10% hoa hồng.",
        reply_markup=(ADMIN_MENU if _is_admin(u.id) else MENU),
    )
    # Khách mới: gửi thêm tin hướng dẫn 3 bước mua acc đầu tiên (nút bấm)
    if is_new:
        try:
            kb = InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="🛒 Mua acc đầu tiên (3 bước)",
                                      callback_data="guide_firstbuy:1")],
                [InlineKeyboardButton(text="📖 Hướng dẫn nhanh",
                                      callback_data="guide_hd")],
            ])
            await msg.answer(
                "🆕 <b>Bạn là khách mới?</b>\n"
                "Bấm nút bên dưới để được dẫn từng bước mua acc đầu tiên nhé 👇",
                parse_mode="HTML", reply_markup=kb,
            )
        except Exception:
            pass

# ─── Hướng dẫn 3 bước mua acc đầu tiên (cho khách mới) ───
_GUIDE_STEPS = {
    "1": (
        "🛒 <b>MUA ACC ĐẦU TIÊN — Bước 1/3: Nạp tiền</b>\n"
        "━━━━━━━━━━━━━━━\n"
        "Acc mua bằng <b>ví shop</b> (riêng với ví chính).\n\n"
        "👉 Gõ <b>/napshop &lt;số tiền&gt;</b> (VD: /napshop 50000)\n"
        "→ quét mã QR → tiền vào ví shop ngay.\n\n"
        "Nạp xong bấm Tiếp nhé 👇"
    ),
    "2": (
        "🛒 <b>MUA ACC ĐẦU TIÊN — Bước 2/3: Chọn acc</b>\n"
        "━━━━━━━━━━━━━━━\n"
        "👉 Gõ <b>/shop</b> → chọn loại acc → nhập số lượng\n"
        "→ bấm xác nhận.\n\n"
        "Bot tự <b>kiểm tra LIVE</b> từng acc trước khi giao —\n"
        "acc chết được đổi/cách ly, bạn chỉ nhận acc sống.\n\n"
        "Xong bước này bấm Tiếp nhé 👇"
    ),
    "3": (
        "🛒 <b>MUA ACC ĐẦU TIÊN — Bước 3/3: Nhận acc</b>\n"
        "━━━━━━━━━━━━━━━\n"
        "Sau khi mua, bạn nhận acc theo 1 trong 3 cách:\n"
        "• 📋 <b>Hiện thông tin</b> — xem ngay trong chat\n"
        "• 📄 <b>Tải file .txt</b> — lưu về máy\n"
        "• 📊 <b>Tải file .xlsx</b> — mở bằng Excel\n\n"
        "Acc lỗi trong thời gian bảo hành? Gõ <b>/damua</b>\n"
        "→ chọn đơn → yêu cầu bảo hành.\n\n"
        "🎉 Xong! Chúc bạn mua sắm vui vẻ."
    ),
}


@router.callback_query(F.data.startswith("guide_firstbuy:"))
async def on_guide_firstbuy(cb: CallbackQuery):
    await cb.answer()
    step = cb.data.split(":", 1)[1]
    text = _GUIDE_STEPS.get(step, _GUIDE_STEPS["1"])
    buttons = []
    if step == "1":
        buttons.append([InlineKeyboardButton(text="Tiếp: chọn acc →",
                                             callback_data="guide_firstbuy:2")])
    elif step == "2":
        buttons.append([InlineKeyboardButton(text="Tiếp: nhận acc →",
                                             callback_data="guide_firstbuy:3")])
    else:
        buttons.append([InlineKeyboardButton(text="← Xem lại từ đầu",
                                             callback_data="guide_firstbuy:1")])
    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    try:
        await cb.message.edit_text(text, parse_mode="HTML", reply_markup=kb)
    except Exception:
        pass


@router.callback_query(F.data == "guide_hd")
async def on_guide_hd(cb: CallbackQuery):
    await cb.answer()
    await _send_huongdan(cb.message, edit=True)


@router.message(Command("ref"))
async def on_ref(msg: Message):
    try:
        bot_info = await msg.bot.get_me()
        user = db.get_user(msg.chat.id)
        
        # Ensure we don't hit KeyError for ref_earnings or ref_withdrawn by converting to dict
        user_dict = dict(user) if user else {}
        
        earnings = user_dict.get("ref_earnings") or 0
        withdrawn = user_dict.get("ref_withdrawn") or 0
        available = earnings - withdrawn
        
        ref_code = user_dict.get("ref_code") or f"REF{msg.chat.id}"
        if user and not user_dict.get("ref_code"):
            with db._lock:
                try:
                    db.get_conn().execute("UPDATE tg_users SET ref_code=? WHERE tg_id=?", (ref_code, msg.chat.id))
                    db.get_conn().commit()
                except Exception as db_e:
                    log.error("Could not update ref_code: %s", db_e)
                
        ref_link = f"https://t.me/{bot_info.username}?start={msg.chat.id}"
        ref_link_code = f"https://t.me/{bot_info.username}?start={ref_code}"

        rates = db.get_ref_rates()
        _pct = lambda x: ("%g" % x) + "%"
        
        f1_count = 0
        f2_count = 0
        if user:
            try:
                c = db.get_conn()
                f1_res = c.execute("SELECT COUNT(*) FROM tg_users WHERE referrer_id=?", (msg.chat.id,)).fetchone()
                f1_count = f1_res[0] if f1_res else 0
                f2_res = c.execute("SELECT COUNT(*) FROM tg_users WHERE referrer_id IN (SELECT tg_id FROM tg_users WHERE referrer_id=?)", (msg.chat.id,)).fetchone()
                f2_count = f2_res[0] if f2_res else 0
            except Exception as cnt_e:
                log.error("Could not count refs: %s", cnt_e)

        text = (
            "🎁 <b>HỆ THỐNG GIỚI THIỆU - KIẾM TIỀN</b>\n"
            "━━━━━━━━━━━━━━━━━━━━\n\n"
            "🔗 <b>Link giới thiệu của bạn:</b>\n"
            f"👉 <code>{ref_link}</code>\n"
            f"Hoặc mã: <code>{ref_link_code}</code>\n\n"
            "💰 <b>Hoa hồng nhận được (tùy cấp độ):</b>\n"
            f"- F1 Hạng Đồng: <b>{_pct(rates['f1_pct'])}</b>\n"
            f"- F1 Hạng Bạc (F1 nạp từ {vnd(int(rates['f1_silver_min']))}): <b>{_pct(rates['f1_silver_pct'])}</b>\n"
            f"- F1 Hạng Vàng (F1 nạp từ {vnd(int(rates['f1_gold_min']))}): <b>{_pct(rates['f1_gold_pct'])}</b>\n"
            f"- F2 (gián tiếp): <b>{_pct(rates['f2_pct'])}</b>\n\n"
            "📊 <b>Thống kê của bạn:</b>\n"
            f"• Đã mời F1: <b>{f1_count} người</b>\n"
            f"• Đã mời F2: <b>{f2_count} người</b>\n"
            f"• Hoa hồng tổng: <b>{vnd(earnings)}</b>\n"
            f"• Đã rút: <b>{vnd(withdrawn)}</b>\n"
            f"• Khả dụng: <b>{vnd(available)}</b>\n\n"
            "💡 Lệnh hỗ trợ:\n"
            "<code>/refcode &lt;mã&gt;</code> - Đổi mã giới thiệu tùy chỉnh\n"
            "<code>/doitien &lt;số_tiền&gt;</code> - Đổi hoa hồng thành số dư (+10% Bonus)\n"
            "<code>/ruttien &lt;số_tiền&gt; &lt;Tên_NH&gt; &lt;STK&gt;</code> - Rút hoa hồng (Min 50k)"
        )
        await msg.answer(text, parse_mode="HTML")
    except Exception as e:
        import traceback
        err_msg = traceback.format_exc()
        await msg.answer(f"Lỗi ref: {e}\n\n{err_msg[:3000]}")

@router.message(Command("daily"))
async def on_daily(msg: Message):
    success, streak, reward = db.checkin_daily(msg.from_user.id)
    if success:
        await msg.answer(f"🎉 <b>Điểm danh thành công!</b>\n\nBạn nhận được <b>{reward} credits</b> check UID.\n🔥 Chuỗi điểm danh: <b>{streak} ngày</b>.", parse_mode="HTML")
    else:
        await msg.answer(f"⚠️ Bạn đã điểm danh hôm nay rồi!\n🔥 Chuỗi hiện tại: <b>{streak} ngày</b>.\nQuay lại vào ngày mai nhé.", parse_mode="HTML")

@router.message(Command("scan_all"))
async def on_scan_all(msg: Message):
    wait = await msg.answer("⏳ Đang tổng hợp dữ liệu toàn bộ dàn tài khoản của bạn...")
    try:
        tg_id = msg.from_user.id
        fb_tracks = db.user_fb_tracks(tg_id)
        tk_tracks = db.user_tracks(tg_id)
        ig_tracks = db.user_ig_tracks(tg_id)
        
        import csv
        import io
        
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Platform", "Target", "Status/Followers"])
        
        fb_die = 0
        for t in fb_tracks:
            writer.writerow(["FB", t["fb_uid"], t["last_status"]])
            if t.get("last_status") == "die":
                fb_die += 1
                
        tk_total = 0
        for t in tk_tracks:
            writer.writerow(["TikTok", t["tiktok_username"], t.get("last_followers", 0)])
            tk_total += t.get("last_followers", 0)
            
        for t in ig_tracks:
            writer.writerow(["IG", t["ig_username"], t.get("last_followers", 0)])
            
        output.seek(0)
        from aiogram.types import BufferedInputFile
        file = BufferedInputFile(output.getvalue().encode('utf-8'), filename="scan_all_report.csv")
        
        summary = (
            f"📊 <b>BÁO CÁO TỔNG QUAN</b>\n\n"
            f"🔹 <b>Facebook:</b> {len(fb_tracks)} nick ({fb_die} chết)\n"
            f"🔹 <b>TikTok:</b> {len(tk_tracks)} nick (Tổng {fmt_num(tk_total)} follow)\n"
            f"🔹 <b>Instagram:</b> {len(ig_tracks)} nick\n\n"
            "Tải file CSV đính kèm để xem chi tiết."
        )
        
        await msg.answer_document(document=file, caption=summary, parse_mode="HTML")
        await wait.delete()
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi khi quét: {e}")

@router.message(Command("refcode"))
async def on_refcode(msg: Message):
    parts = msg.text.split(maxsplit=1)
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: /refcode &lt;mã_mới&gt;\nLưu ý: Mã chỉ gồm chữ và số, không dấu cách.")
        return
        
    code = parts[1].strip()
    if not code.isalnum():
        await msg.answer("❌ Mã giới thiệu chỉ được chứa chữ cái và số!")
        return
        
    c = db.get_conn()
    exists = c.execute("SELECT tg_id FROM tg_users WHERE ref_code=?", (code,)).fetchone()
    if exists and exists["tg_id"] != msg.chat.id:
        await msg.answer("❌ Mã này đã có người sử dụng. Vui lòng chọn mã khác.")
        return
        
    with db._lock:
        c.execute("UPDATE tg_users SET ref_code=? WHERE tg_id=?", (code, msg.chat.id))
        c.commit()
        
    await msg.answer(f"✅ Đã đổi mã giới thiệu thành công: <code>{code}</code>")

@router.message(Command("ruttien"))
async def on_ruttien(msg: Message):
    parts = msg.text.split(maxsplit=3)
    if len(parts) < 4:
        await msg.answer("⚠️ Cú pháp: /ruttien &lt;số tiền&gt; &lt;Ngân hàng&gt; &lt;STK&gt;\nVí dụ: /ruttien 50000 MBBank 123456789")
        return
        
    try:
        amount = int(parts[1].replace(",", "").replace(".", "").replace("k", "000").replace("K", "000").strip())
    except:
        await msg.answer("❌ Số tiền không hợp lệ!")
        return
        
    bank_name = parts[2].strip()
    bank_account = parts[3].strip()
    bank_info = f"{bank_name} - {bank_account}"

    if amount < 50000:
        await msg.answer("❌ Số tiền rút tối thiểu là 50,000 VNĐ.")
        return
        
    c = db.get_conn()
    user = c.execute("SELECT ref_earnings, ref_withdrawn FROM tg_users WHERE tg_id=?", (msg.chat.id,)).fetchone()
    user_dict = dict(user) if user else {}
    
    earnings = user_dict.get("ref_earnings") or 0
    withdrawn = user_dict.get("ref_withdrawn") or 0
    available = earnings - withdrawn
    
    import datetime
    current_month_start = int(datetime.datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0).timestamp())
    count_withdrawals = c.execute("SELECT COUNT(*) as c FROM withdrawal_requests WHERE tg_id=? AND created_at >= ?", (msg.chat.id, current_month_start)).fetchone()["c"]

    fee = 0
    if count_withdrawals >= 2:
        fee = 10000

    if available < amount + fee:
        await msg.answer(f"❌ Số dư khả dụng không đủ! (Khả dụng: {vnd(available)}, Cần: {vnd(amount + fee)} bao gồm phí {vnd(fee)} nếu có)")
        return
        
    req_id = db.create_withdrawal_request(msg.chat.id, amount, bank_info, fee)
    await notify_admin_withdrawal_request(req_id, msg.chat.id, amount, bank_info, fee)
    await msg.answer(f"✅ Đã gửi yêu cầu rút <b>{vnd(amount)}</b>.\nVui lòng chờ Admin kiểm tra và duyệt chuyển khoản!", parse_mode="HTML")

async def notify_admin_withdrawal_request(req_id: int, tg_id: int, amount: int, bank_info: str, fee: int):
    admin_tg_token = db.get_setting("admin_bot_token", "")
    admin_zalo = db.get_setting("admin_zalo_id", "")
    
    c = db.get_conn()
    user = c.execute("SELECT username, name FROM tg_users WHERE tg_id=?", (tg_id,)).fetchone()
    username_str = f"@{user['username']}" if user and user['username'] else (user['name'] if user and user['name'] else str(tg_id))
    
    actual_amount = amount - fee
    admin_msg = (
        "🔔 <b>CÓ YÊU CẦU RÚT TIỀN HOA HỒNG!</b>\n\n"
        f"🆔 Đơn rút: <b>#{req_id}</b>\n"
        f"👤 Khách: {html.escape(username_str)} (ID: <code>{tg_id}</code>)\n"
        f"💰 Số tiền yêu cầu: <b>{vnd(amount)}</b>\n"
        f"💸 Phí rút: <b>{vnd(fee)}</b>\n"
        f"💵 <b>Thực nhận chuyển khoản: {vnd(actual_amount)}</b>\n"
        f"🏦 Ngân hàng / STK: <b>{html.escape(bank_info or '')}</b>\n\n"
        "👉 Bấm nút bên dưới để Duyệt hoặc Từ chối:"
    )
    
    if admin_tg_token and admin_tg_token.strip() != (db.get_setting("bot_token") or "").strip():
        from .admin_bot import manager as admin_manager
        admin_sender_bot = admin_manager.bot or manager.bot
    else:
        admin_sender_bot = manager.bot

    if admin_sender_bot:
        admins = _withdraw_admin_ids()
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="✅ Duyệt & Đã Chuyển", callback_data=f"tg_admin_withdraw_approve_{req_id}_{tg_id}_{amount}"),
                InlineKeyboardButton(text="❌ Từ Chối", callback_data=f"tg_admin_withdraw_reject_{req_id}_{tg_id}_{amount}")
            ]
        ])
        for admin_id in admins:
            try:
                await admin_sender_bot.send_message(admin_id, admin_msg, parse_mode="HTML", reply_markup=kb)
            except Exception as e:
                log.error("Failed to notify TG admin %s for withdrawal: %s", admin_id, e)
                
    if admin_zalo and zalo_manager.running:
        zalo_kb = {
            "inline_keyboard": [
                [
                    {"text": "✅ Đã Chuyển Tiền", "callback_data": f"zalo_withdraw_approve_{req_id}"},
                    {"text": "❌ Từ Chối", "callback_data": f"zalo_withdraw_reject_{req_id}"}
                ]
            ]
        }
        asyncio.create_task(zalo_manager.send_message(admin_zalo, admin_msg, reply_markup=zalo_kb))

def _withdraw_admin_ids() -> list:
    """ID được phép duyệt/từ chối rút tiền: chủ shop + group +
    admin phụ đang có quyền Tiền tệ."""
    ids = []
    try:
        if db.get_setting("admin_tg_id"): ids.append(int(db.get_setting("admin_tg_id")))
    except: pass
    try:
        if db.get_setting("admin_tg_group_id"): ids.append(int(db.get_setting("admin_tg_group_id")))
    except: pass
    try:
        for tid in _perms.notify_extra_ids("tien"):
            if tid not in ids:
                ids.append(tid)
    except: pass
    return ids


async def _notify_super_subadmin_action(actor_id: int, actor_name: str, text: str):
    """Báo chủ shop khi admin phụ thực hiện thao tác nhạy cảm (duyệt/từ chối rút)."""
    try:
        if _perms.is_super(actor_id):
            return
        sid = _perms.super_id()
        if sid and manager.bot:
            await manager.bot.send_message(
                sid,
                f"👤 <b>Admin phụ {html.escape(actor_name or '')}</b> "
                f"({actor_id}) vừa {text}",
                parse_mode="HTML")
    except Exception:
        pass


@router.callback_query(F.data.startswith("tg_admin_withdraw_approve_"))
async def on_admin_withdraw_approve(cb: CallbackQuery):
    admins = _withdraw_admin_ids()
    if cb.message.chat.id not in admins and cb.from_user.id not in admins:
        await cb.answer("❌ Bạn không có quyền duyệt!", show_alert=True)
        return
        
    parts = cb.data.split("_")
    # tg_admin_withdraw_approve_{req_id}_{tg_id}_{amount}
    req_id = int(parts[4])
    tg_id = int(parts[5])
    amount = int(parts[6])
    
    c = db.get_conn()
    if req_id == 0:
        req = c.execute("SELECT * FROM withdrawal_requests WHERE tg_id=? AND status='pending' ORDER BY id DESC LIMIT 1", (tg_id,)).fetchone()
    else:
        req = c.execute("SELECT * FROM withdrawal_requests WHERE id=?", (req_id,)).fetchone()
        
    if not req or req["status"] != "pending":
        await cb.answer("⚠️ Đơn này đã được xử lý từ trước!", show_alert=True)
        return
        
    actual_req_id = req["id"]
    with db._lock:
        c.execute("UPDATE withdrawal_requests SET status='approved', updated_at=? WHERE id=?", (int(time.time()), actual_req_id))
        c.execute("UPDATE tg_users SET ref_withdrawn = ref_withdrawn + ? WHERE tg_id=?", (amount, tg_id))
        c.commit()
        
    await cb.answer("✅ Đã duyệt đơn rút tiền thành công!", show_alert=True)
    try:
        await cb.message.edit_text(f"{cb.message.text}\n\n✅ <b>ĐÃ DUYỆT BỞI {cb.from_user.full_name} ({vnd(amount)})</b>", parse_mode="HTML")
    except: pass
    
    # Notify customer
    try:
        cust_msg = (
            "🎉 <b>RÚT TIỀN HOA HỒNG THÀNH CÔNG!</b>\n\n"
            f"Yêu cầu rút tiền <b>#{actual_req_id}</b> của bạn đã được Admin duyệt và chuyển tiền.\n"
            f"💰 Số tiền: <b>{vnd(amount)}</b>\n"
            f"🏦 Ngân hàng / STK: <b>{html.escape(req.get('bank_info', '') or '')}</b>\n\n"
            "Cảm ơn bạn đã đồng hành và phát triển cùng hệ thống! ❤️"
        )
        if manager.bot:
            await manager.bot.send_message(tg_id, cust_msg, parse_mode="HTML")
    except Exception as notify_err:
        log.error("Could not notify user of approved withdrawal: %s", notify_err)
    db.admin_audit_add(cb.from_user.id, cb.from_user.full_name, "duyet_rut_tien",
                       f"#{actual_req_id} {vnd(amount)} cho user {tg_id}")
    await _notify_super_subadmin_action(
        cb.from_user.id, cb.from_user.full_name,
        f"duyệt đơn rút <b>#{actual_req_id}</b> ({vnd(amount)})")

@router.callback_query(F.data.startswith("tg_admin_withdraw_reject_"))
async def on_admin_withdraw_reject(cb: CallbackQuery):
    admins = _withdraw_admin_ids()
    if cb.message.chat.id not in admins and cb.from_user.id not in admins:
        await cb.answer("❌ Bạn không có quyền từ chối!", show_alert=True)
        return
        
    parts = cb.data.split("_")
    req_id = int(parts[4])
    tg_id = int(parts[5])
    amount = int(parts[6])
    
    c = db.get_conn()
    if req_id == 0:
        req = c.execute("SELECT * FROM withdrawal_requests WHERE tg_id=? AND status='pending' ORDER BY id DESC LIMIT 1", (tg_id,)).fetchone()
    else:
        req = c.execute("SELECT * FROM withdrawal_requests WHERE id=?", (req_id,)).fetchone()
        
    if not req or req["status"] != "pending":
        await cb.answer("⚠️ Đơn này đã được xử lý từ trước!", show_alert=True)
        return
        
    actual_req_id = req["id"]
    with db._lock:
        c.execute("UPDATE withdrawal_requests SET status='rejected', updated_at=? WHERE id=?", (int(time.time()), actual_req_id))
        c.commit()
        
    await cb.answer("❌ Đã từ chối đơn rút tiền.", show_alert=True)
    try:
        await cb.message.edit_text(f"{cb.message.text}\n\n❌ <b>ĐÃ TỪ CHỐI BỞI {cb.from_user.full_name}</b>", parse_mode="HTML")
    except: pass
    
    # Notify customer
    try:
        cust_msg = (
            "❌ <b>YÊU CẦU RÚT TIỀN BỊ TỪ CHỐI</b>\n\n"
            f"Yêu cầu rút tiền hoa hồng <b>#{actual_req_id}</b> ({vnd(amount)}) của bạn đã bị Admin từ chối.\n"
            "Số dư hoa hồng của bạn vẫn được giữ nguyên.\n"
            "Vui lòng kiểm tra lại thông tin Ngân hàng / STK hoặc liên hệ Admin để được hỗ trợ."
        )
        if manager.bot:
            await manager.bot.send_message(tg_id, cust_msg, parse_mode="HTML")
    except Exception as notify_err:
        log.error("Could not notify user of rejected withdrawal: %s", notify_err)
    db.admin_audit_add(cb.from_user.id, cb.from_user.full_name, "tu_choi_rut_tien",
                       f"#{actual_req_id} {vnd(amount)} của user {tg_id}")
    await _notify_super_subadmin_action(
        cb.from_user.id, cb.from_user.full_name,
        f"từ chối đơn rút <b>#{actual_req_id}</b> ({vnd(amount)})")

@router.message(Command("doitien"))
async def on_doitien(msg: Message):
    parts = msg.text.split(maxsplit=1)
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: /doitien &lt;số_tiền&gt;\nVí dụ: /doitien 50000")
        return
        
    try:
        amount = int(parts[1].replace(",", "").replace(".", "").replace("k", "000").replace("K", "000").strip())
    except:
        await msg.answer("❌ Số tiền không hợp lệ!")
        return
        
    if amount <= 0:
        await msg.answer("❌ Số tiền không hợp lệ!")
        return

    c = db.get_conn()
    user = c.execute("SELECT ref_earnings, ref_withdrawn FROM tg_users WHERE tg_id=?", (msg.chat.id,)).fetchone()
    user_dict = dict(user) if user else {}
    
    earnings = user_dict.get("ref_earnings") or 0
    withdrawn = user_dict.get("ref_withdrawn") or 0
    available = earnings - withdrawn
    
    if amount > available:
        await msg.answer(f"❌ Số dư khả dụng không đủ! (Khả dụng: {vnd(available)})")
        return
        
    bonus_amount = int(amount * 1.1)
    
    with db._lock:
        c.execute("UPDATE tg_users SET ref_withdrawn = ref_withdrawn + ?, balance = balance + ? WHERE tg_id=?", (amount, bonus_amount, msg.chat.id))
        c.execute("INSERT INTO txns(ts, tg_id, amount, reason) VALUES(?,?,?,?)", (int(time.time()), msg.chat.id, bonus_amount, 'Đổi hoa hồng sang số dư (+10% Bonus)'))
        c.commit()
        
    await msg.answer(f"✅ Đã đổi <b>{vnd(amount)}</b> hoa hồng sang <b>{vnd(bonus_amount)}</b> số dư thành công!", parse_mode="HTML")

@router.message(Command("trial"))
async def on_trial(msg: Message):
    user = db.get_user(msg.chat.id)
    if not user:
        await msg.answer("Bạn chưa /start. Gõ /start trước nhé.")
        return
    
    if db.get_setting("enable_free_trial", "1") != "1":
        await msg.answer("❌ Rất tiếc, chương trình dùng thử hiện đang đóng.")
        return
        
    try:
        days = int(db.get_setting("free_trial_days", "3"))
    except ValueError:
        days = 3
        
    if db.activate_trial(msg.chat.id, days):
        db.add_log("trial", f"User nhận trial {days} ngày", msg.chat.id)
        u2 = db.get_user(msg.chat.id)
        await msg.answer(
            f"🎉 <b>Chúc mừng!</b>\n\nBạn đã nhận được <b>{days} ngày</b> dùng thử miễn phí full tính năng!\n"
            f"Hạn sử dụng mới: <b>{_sub_text(u2)}</b>\n\n"
            "Hãy trải nghiệm các lệnh theo dõi nhé!"
        )
    else:
        await msg.answer("⚠️ Bạn đã nhận gói dùng thử rồi hoặc gói VIP của bạn đã từng được kích hoạt!")

@router.message(Command("bank"))
async def on_bank(msg: Message):
    import json
    banks_list_str = db.get_setting("banks_list", "")
    banks = []
    if banks_list_str:
        try: banks = json.loads(banks_list_str)
        except: pass
    if not banks:
        bank_name = db.get_setting("bank_name", "")
        if bank_name:
            banks = [{"name": bank_name, "account": db.get_setting("bank_account", ""), "owner": db.get_setting("bank_owner", "")}]
            
    if not banks:
        await msg.answer("⚠️ Admin chưa thiết lập thông tin ngân hàng.")
        return
        
    transfer_content = msg.from_user.username if msg.from_user.username else msg.chat.id
    
    lines = ["🏦 <b>THÔNG TIN CHUYỂN KHOẢN</b>\n"]
    for i, b in enumerate(banks, 1):
        lines.append(f"<b>{i}. {b.get('name', '')}</b>")
        lines.append(f"• Số tài khoản: <code>{b.get('account', '')}</code>")
        lines.append(f"• Chủ tài khoản: <b>{b.get('owner', '')}</b>\n")
        
    lines.append(f"📝 <b>Nội dung CK bắt buộc:</b> <code>{transfer_content}</code>\n")
    lines.append("<i>Sau khi chuyển khoản thành công, hãy bấm nút bên dưới để xác nhận!</i>")
    text = "\n".join(lines)
    
    parts = msg.text.split(maxsplit=1)
    amount = 0
    if len(parts) > 1:
        try:
            amount = int(parts[1].replace(",", "").replace(".", "").replace("k", "000").replace("K", "000").strip())
        except: pass

    if amount > 0:
        cb_data = f"bank_confirm_{amount}"
        btn_text = f"✅ Tôi đã chuyển {vnd(amount)}"
    else:
        cb_data = "bank_confirm"
        btn_text = "✅ Tôi đã chuyển tiền"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=btn_text, callback_data=cb_data)],
        [InlineKeyboardButton(text="⚡ Nạp tự động (PayOS)", callback_data="payos_auto")],
    ])
    
    import os
    img_dir = os.path.join(os.path.dirname(__file__), "..", "data", "images")
    if os.path.exists(img_dir):
        for f in os.listdir(img_dir):
            if f.startswith("qr_"):
                try:
                    await msg.answer_photo(photo=FSInputFile(os.path.join(img_dir, f)))
                except: pass
                
    await msg.answer(text, reply_markup=kb)

@router.callback_query(F.data.startswith("use_code_"))
async def on_use_code(cb: CallbackQuery):
    code = cb.data.replace("use_code_", "")
    success, amount, msg_text = db.use_code(code, cb.fromuser.id) if hasattr(cb, 'fromuser') else db.use_code(code, cb.from_user.id)
    if success:
        db.adjust_balance(cb.from_user.id, amount, f"Sử dụng Giftcode: {code}")
        try:
            msg_text_resp = (
                f"✅ <b>NẠP TIỀN THÀNH CÔNG!</b>\n\n"
                f"Bạn đã sử dụng mã <code>{code}</code> và được cộng <b>{vnd(amount)}</b> vào tài khoản.\n"
                f"Cảm ơn bạn đã tin tưởng dịch vụ!"
            )
            upgraded, new_vip, is_lifetime = db.check_vip_upgrade(cb.from_user.id)
            if upgraded and new_vip > 0:
                limit = db.get_setting(f"vip{new_vip}_limit", "10")
                msg_text_resp += (
                    f"\n\n🎉 <b>CHÚC MỪNG BẠN ĐÃ LÊN VIP {new_vip}!</b> 🎉\n\n"
                    f"🎁 <b>Đặc quyền mới:</b>\n"
                    f"- Mức độ theo dõi tối đa: <b>{limit} mục</b>/nền tảng\n"
                )
                if is_lifetime:
                    msg_text_resp += "- Hạn sử dụng: <b>VĨNH VIỄN</b>\n\n"
                else:
                    msg_text_resp += "\n"
            await cb.message.edit_text(msg_text_resp, parse_mode="HTML")
        except: pass
        
        admin_zalo = db.get_setting("admin_zalo_id", "")
        if admin_zalo and zalo_manager.running:
            username_str = f"@{cb.from_user.username}" if cb.from_user.username else cb.from_user.full_name
            asyncio.create_task(zalo_manager.send_message(admin_zalo, f"💵 Khách {username_str} ({cb.from_user.id}) đã sử dụng thành công mã {code} ({vnd(amount)})"))
            
        await cb.answer("Nạp tiền thành công!")
    else:
        await cb.answer(f"❌ {msg_text}", show_alert=True)

@router.callback_query(F.data.startswith("save_code_"))
async def on_save_code(cb: CallbackQuery):
    code = cb.data.replace("save_code_", "")
    success, msg_text = db.save_code_for_user(cb.from_user.id, code)
    if success:
        await cb.answer("✅ Đã lưu mã vào ví của bạn! Dùng lệnh /mycodes để xem lại nhé.", show_alert=True)
    else:
        await cb.answer(f"❌ {msg_text}", show_alert=True)

@router.message(Command("mycodes"))
async def on_mycodes(msg: Message):
    codes = db.get_user_saved_codes(msg.chat.id)
    if not codes:
        await msg.answer("📭 Bạn không có mã lưu trữ nào chưa sử dụng.")
        return
        
    text = "📥 <b>KHO MÃ LƯU TRỮ CỦA BẠN</b>\n\n"
    keyboard = []
    
    import datetime
    for c in codes:
        code_str = c["code"]
        amount = c["amount"]
        expire_at = c["expire_at"]
        
        expire_text = "Vĩnh viễn"
        if expire_at > 0:
            expire_text = vn_time_str('%H:%M %d/%m', expire_at)
            
        text += f"• <code>{code_str}</code>: <b>{vnd(amount)}</b> (Hạn: {expire_text})\n"
        keyboard.append([InlineKeyboardButton(text=f"🎁 Dùng mã {vnd(amount)}", callback_data=f"use_code_{code_str}")])
        
    text += "\n<i>Bấm nút bên dưới để sử dụng:</i>"
    await msg.answer(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=keyboard), parse_mode="HTML")

async def _ask_bank_wallet(msg, amount: int):
    """Hỏi user cộng tiền bank vào ví nào (sau khi admin duyệt)."""
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="💰 Ví chính", callback_data=f"bank_wallet:main:{amount}")],
        [InlineKeyboardButton(text="🛒 Ví shop", callback_data=f"bank_wallet:shop:{amount}")],
    ])
    await msg.answer(f"💳 Cộng <b>{vnd(amount)}</b> vào ví nào (sau khi admin duyệt)?",
                     parse_mode="HTML", reply_markup=kb)


@router.callback_query(F.data.startswith("bank_confirm"))
async def on_bank_confirm(cb: CallbackQuery, state: FSMContext):
    await cb.message.edit_reply_markup(reply_markup=None)
    
    if cb.data == "bank_confirm":
        await cb.message.answer("✍️ Vui lòng nhập <b>số tiền</b> bạn đã chuyển khoản (ví dụ: 50000):")
        await state.set_state(BankState.waiting_for_amount)
        await cb.answer()
    else:
        try:
            amount = int(cb.data.split("_")[2])
        except:
            amount = 0
            
        if amount > 0:
            await _ask_bank_wallet(cb.message, amount)
        await cb.answer()


@router.callback_query(F.data.startswith("bank_wallet:"))
async def on_bank_wallet(cb: CallbackQuery):
    await cb.answer()
    try:
        _, target, amount_s = cb.data.split(":")
        amount = int(amount_s)
        assert target in ("main", "shop") and amount > 0
    except Exception:
        await cb.message.answer("❌ Yêu cầu không hợp lệ.")
        return
    await process_bank_amount(cb.message, cb.from_user, amount, target)

async def process_bank_amount(msg: Message, user, amount: int, target: str = "main"):
    admin_tg_token = db.get_setting("admin_bot_token", "")
    admin_zalo = db.get_setting("admin_zalo_id", "")
    
    username_str = f"@{user.username}" if user.username else user.full_name
    admin_msg = (
        "🔔 <b>CÓ KHÁCH BÁO CHUYỂN KHOẢN!</b>\n\n"
        f"👤 Khách: {html.escape(username_str)}\n"
        f"🆔 ID Telegram: {user.id}\n"
        f"💰 Số tiền: {vnd(amount)}\n"
        f"👛 Ví: {'🛒 Ví shop' if target == 'shop' else '💰 Ví chính'}\n\n"
        f"👉 ĐỂ TẠO & PHÁT CODE, gửi lệnh:\n/phatcode {user.id} {amount}\n\n"
        f"👉 Cú pháp cộng thẳng: /topup {user.id} {amount}\n"
        "👉 Hoặc cộng thủ công trên trang Quản lý."
    )
    
    notified = False
    if admin_tg_token:
        from .admin_bot import manager as admin_manager
        admin_sender_bot = admin_manager.bot
    else:
        from .bot import manager as main_manager
        admin_sender_bot = main_manager.bot

    if admin_sender_bot:
        admins = []
        try:
            if db.get_setting("admin_tg_id"): admins.append(int(db.get_setting("admin_tg_id")))
        except: pass
        try:
            if db.get_setting("admin_tg_group_id"): admins.append(int(db.get_setting("admin_tg_group_id")))
        except: pass
        
        kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="✅ Xác nhận + Cộng tiền", callback_data=f"tg_admin_confirm_{user.id}_{amount}_{target}")]])
        for admin_id in admins:
            try:
                await admin_sender_bot.send_message(admin_id, admin_msg, parse_mode="HTML", reply_markup=kb)
                notified = True
            except Exception as e:
                log.error("Failed to notify TG admin %s: %s", admin_id, e)
                    
    if admin_zalo and zalo_manager.running:
        zalo_kb = {
            "inline_keyboard": [
                [{"text": "✅ Đã nhận tiền (Phát Code)", "callback_data": f"zalo_confirm_{user.id}_{amount}"}]
            ]
        }
        asyncio.create_task(zalo_manager.send_message(admin_zalo, admin_msg, reply_markup=zalo_kb))
        notified = True
        
    if notified:
        _w = "🛒 <b>Ví shop</b>" if target == "shop" else "💰 <b>Ví chính</b>"
        await msg.answer(f"✅ Đã gửi thông báo cho Admin xác nhận khoản nạp <b>{vnd(amount)}</b> vào {_w}.\nTiền sẽ được cộng vào tài khoản của bạn sau khi Admin kiểm tra xong (thường trong vòng 1-5 phút)!")
    else:
        await msg.answer(f"✅ Đã ghi nhận báo cáo <b>{vnd(amount)}</b>.\nTiền sẽ được cộng vào tài khoản của bạn sau khi Admin kiểm tra xong!")

@router.callback_query(F.data.startswith("tg_admin_confirm_"))
async def on_main_admin_confirm(cb: CallbackQuery):
    admins = []
    try:
        if db.get_setting("admin_tg_id"): admins.append(int(db.get_setting("admin_tg_id")))
    except: pass
    try:
        if db.get_setting("admin_tg_group_id"): admins.append(int(db.get_setting("admin_tg_group_id")))
    except: pass
    
    if cb.message.chat.id not in admins and cb.from_user.id not in admins:
        await cb.answer("❌ Bạn không có quyền duyệt!", show_alert=True)
        return
        
    parts = cb.data.split("_")
    user_id = int(parts[3])
    amount = int(parts[4])
    target = parts[5] if len(parts) > 5 and parts[5] in ("main", "shop") else "main"

    try:
        if target == "shop":
            db.credit_topup(user_id, amount, reason="bank_transfer", wallet="shop")
        else:
            db.adjust_balance(user_id, amount, reason="bank_transfer")
        success = True
    except Exception as e:
        log.error(f"Lỗi cộng tiền: {e}")
        success = False
    if success:
        await cb.answer("✅ Đã cộng tiền thành công!", show_alert=True)
        try:
            await cb.message.edit_text(f"{cb.message.text}\n\n✅ <b>Đã duyệt {amount:,.0f} VNĐ bởi {cb.from_user.full_name}</b>", parse_mode="HTML")
        except: pass
        
        # Notify user via main bot
        try:
            _w = "🛒 <b>Ví shop</b>" if target == "shop" else "💰 <b>Ví chính</b>"
            msg_text = (
                f"✅ <b>NẠP TIỀN THÀNH CÔNG</b>\n\n"
                f"Bạn vừa được cộng <b>{amount:,.0f} VNĐ</b> vào {_w}.\n"
                f"Cảm ơn bạn đã sử dụng dịch vụ!"
            )
            try:
                _u = db.get_user(user_id)
                _bal = (_u["shop_balance"] if target == "shop" else _u["balance"]) if _u is not None else 0
                msg_text += f"\n\n{_w} số dư hiện tại: <b>{_bal:,.0f} VNĐ</b>"
            except Exception:
                pass
            
            # Kiem tra VIP upgrade
            upgraded, new_vip, is_lifetime = db.check_vip_upgrade(user_id)
            if upgraded and new_vip > 0:
                limit = db.get_setting(f"vip{new_vip}_limit", "10")
                msg_text += (
                    f"\n\n🎉 <b>CHÚC MỪNG BẠN ĐÃ LÊN VIP {new_vip}!</b> 🎉\n\n"
                    f"🎁 <b>Đặc quyền mới:</b>\n"
                    f"- Mức độ theo dõi tối đa: <b>{limit} mục</b>/nền tảng\n"
                )
                if is_lifetime:
                    msg_text += "- Hạn sử dụng: <b>VĨNH VIỄN</b>\n\n"
                else:
                    msg_text += "\n"
                    
            await cb.bot.send_message(
                user_id,
                msg_text,
                parse_mode="HTML"
            )
        except: pass
    else:
        await cb.answer("❌ Lỗi khi cộng tiền!", show_alert=True)

@router.message(BankState.waiting_for_amount)
async def on_bank_amount(msg: Message, state: FSMContext):
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("\U0001f6ab \u0110\u00e3 h\u1ee7y thao t\u00e1c \u0111ang nh\u1eadp. B\u1ea1n g\u00f5 l\u1ea1i l\u1ec7nh v\u1eeba r\u1ed3i nh\u00e9.")
        return
    amount_str = msg.text.strip()
    try:
        amount = int(amount_str.replace(",", "").replace(".", ""))
        if amount <= 0: raise ValueError()
    except:
        await msg.answer("❌ Số tiền không hợp lệ. Vui lòng nhập lại số tiền (ví dụ: 50000):")
        return
        
    await state.clear()
    await _ask_bank_wallet(msg, amount)

# ============================ PayOS — nạp tiền tự động ============================
def _parse_payos_amount(raw: str) -> int:
    s = (raw or "").strip().replace(",", "").replace(".", "").replace(" ", "")
    s = s.replace("k", "000").replace("K", "000")
    amount = int(s)
    if amount <= 0:
        raise ValueError()
    return amount


async def _make_payos_order(tg_id: int, amount: int, target: str = "main"):
    """Tạo (hoặc dùng lại) đơn PayOS. Trả (order_code, checkout_url, qr_code, reused)."""
    if target not in ("main", "shop"):
        target = "main"
    from . import payos as payos_mod
    if not payos_mod.is_configured():
        raise payos_mod.PayOSError("Admin chưa cấu hình PayOS, vui lòng nạp bằng /bank.")
    if amount < payos_mod.PAYOS_MIN_AMOUNT:
        raise payos_mod.PayOSError(
            f"Số tiền tối thiểu là {vnd(payos_mod.PAYOS_MIN_AMOUNT)}.")
    existing = db.get_user_pending_payos_order(tg_id, payos_mod.ORDER_TTL, target)
    if existing and existing["checkout_url"] and int(existing["amount"]) == int(amount):
        # Đơn dùng lại: lấy QR đã lưu lúc tạo đơn (API đọc trạng thái không trả qrCode)
        qr_code = ""
        try:
            qr_code = existing["qr_code"] or ""
        except Exception:
            pass
        return int(existing["order_code"]), existing["checkout_url"], qr_code, True
    if existing and int(existing["amount"]) != int(amount):
        # Số tiền khác -> hủy đơn cũ, tạo đơn mới (tránh link thanh toán sai tiền)
        try:
            await payos_mod.cancel_payment_link(int(existing["order_code"]))
        except Exception:
            pass
        db.mark_payos_status(int(existing["order_code"]), "CANCELLED")
    import sqlite3
    description = f"NAP{tg_id}"[-25:]
    return_url, cancel_url = payos_mod.get_return_urls()
    data = None
    for _ in range(5):
        order_code = payos_mod.new_order_code(tg_id)
        data = await payos_mod.create_payment_link(
            order_code, amount, description, return_url, cancel_url)
        try:
            db.create_payos_order(order_code, tg_id, amount,
                                  data.get("paymentLinkId", ""), data.get("checkoutUrl", ""),
                                  data.get("qrCode", "") or "", target)
            break
        except sqlite3.IntegrityError:
            continue  # mã đơn trùng (cực hiếm) -> sinh mã mới, tạo link mới
    return order_code, data.get("checkoutUrl", ""), data.get("qrCode", "") or "", False


def _qr_code_png(qr_text: str):
    """Dựng ảnh PNG mã QR VietQR từ chuỗi PayOS trả về."""
    import io
    import qrcode
    qr = qrcode.QRCode(border=2)
    qr.add_data(qr_text)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white").convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


async def _send_payos_invoice(msg: Message, tg_id: int, amount: int,
                              order_code: int, checkout_url: str,
                              qr_code: str, reused: bool, target: str = "main"):
    kb = InlineKeyboardMarkup(inline_keyboard=[
        *([[InlineKeyboardButton(text="🌐 Mở trang thanh toán", url=checkout_url)]]
          if checkout_url else []),
        [InlineKeyboardButton(text="❌ Hủy đơn", callback_data=f"payos_cancel:{order_code}")],
    ])
    wallet_txt = "🛒 <b>Ví shop</b> <i>(mua tài khoản)</i>" if target == "shop" else "💰 <b>Ví chính</b> <i>(check UID • mua gói)</i>"
    caption = (
        "💳 <b>NẠP TIỀN TỰ ĐỘNG</b>" + (" <i>(dùng lại đơn đang chờ)</i>" if reused else "") + "\n"
        "━━━━━━━━━━━━━━\n"
        f"💵 Số tiền: <b>{vnd(amount)}</b>\n"
        f"👛 Nạp vào: {wallet_txt}\n"
        f"🧾 Mã đơn: <code>{order_code}</code>\n"
        "━━━━━━━━━━━━━━\n"
        "① Mở app ngân hàng, quét <b>mã QR</b> bên dưới\n"
        "② Thanh toán đúng số tiền\n"
        "③ Tiền <b>tự động cộng</b> vào ví (thường dưới 1 phút) ⚡\n\n"
        "<i>⏰ Đơn tự hủy sau 45 phút nếu chưa thanh toán.</i>"
    )
    if qr_code:
        try:
            photo = BufferedInputFile(_qr_code_png(qr_code).read(),
                                      filename=f"payos_{order_code}.png")
            await msg.answer_photo(photo, caption=caption,
                                   parse_mode="HTML", reply_markup=kb)
            return
        except Exception:
            pass
    await msg.answer(caption, parse_mode="HTML", reply_markup=kb)


async def _ask_payos_wallet(msg, amount: int):
    """Hỏi user nạp vào ví nào trước khi tạo đơn PayOS."""
    u = db.get_user(msg.chat.id if hasattr(msg, "chat") else msg.from_user.id)
    main_bal = int(u["balance"] or 0) if u else 0
    shop_bal = int(u["shop_balance"] or 0) if u and "shop_balance" in u.keys() else 0
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(
            text=f"💰 Ví chính — đang có {vnd(main_bal)}",
            callback_data=f"payos_wallet:main:{amount}")],
        [InlineKeyboardButton(
            text=f"🛒 Ví shop — đang có {vnd(shop_bal)}",
            callback_data=f"payos_wallet:shop:{amount}")],
    ])
    await msg.answer(
        f"💳 <b>NẠP {vnd(amount)}</b>\n"
        f"━━━━━━━━━━━━━━\n"
        f"Chọn ví muốn nạp vào:",
        parse_mode="HTML", reply_markup=kb)


_NAP_QUICK_AMOUNTS = [50000, 100000, 200000, 500000, 1000000]


async def _nap_amount_picker(msg: Message, target: str = "main"):
    """Bàn phím chọn nhanh số tiền nạp."""
    rows = []
    row = []
    for a in _NAP_QUICK_AMOUNTS:
        row.append(InlineKeyboardButton(
            text=f"{a // 1000}K", callback_data=f"nap_amt:{target}:{a}"))
        if len(row) == 3:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    rows.append([InlineKeyboardButton(text="✏️ Nhập số khác",
                                      callback_data=f"nap_custom:{target}")])
    wallet_txt = "🛒 <b>Ví shop</b>" if target == "shop" else "💰 <b>Ví chính</b>"
    await msg.answer(
        f"💳 <b>NẠP TIỀN</b> → {wallet_txt}\n"
        f"━━━━━━━━━━━━━━\n"
        f"Chọn nhanh số tiền bên dưới, hoặc nhập số khác:",
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))


@router.callback_query(F.data.startswith("nap_amt:"))
async def on_nap_amt(cb: CallbackQuery):
    await cb.answer()
    try:
        _, target, amount_s = cb.data.split(":")
        amount = int(amount_s)
        assert target in ("main", "shop") and amount > 0
    except Exception:
        await cb.message.answer("❌ Yêu cầu không hợp lệ, gõ /nap lại nhé.")
        return
    if target == "shop":
        try:
            order_code, checkout_url, qr_code, reused = await _make_payos_order(
                cb.from_user.id, amount, "shop")
        except Exception as e:
            await cb.message.answer(f"❌ {e}")
            return
        await _send_payos_invoice(cb.message, cb.from_user.id, amount,
                                  order_code, checkout_url, qr_code, reused, "shop")
    else:
        await _ask_payos_wallet(cb.message, amount)


@router.callback_query(F.data.startswith("nap_custom:"))
async def on_nap_custom(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    target = cb.data.split(":", 1)[1]
    await cb.message.answer("✍️ Nhập <b>số tiền</b> muốn nạp (ví dụ: 50000):",
                            parse_mode="HTML")
    if target == "shop":
        await state.set_state(PayOSState.waiting_for_shop_amount)
    else:
        await state.set_state(PayOSState.waiting_for_amount)


@router.message(Command("nap"))
async def on_nap(msg: Message, state: FSMContext):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) > 1:
        try:
            amount = _parse_payos_amount(parts[1])
        except Exception:
            await msg.answer("❌ Số tiền không hợp lệ. Ví dụ: <code>/nap 50000</code>")
            return
        await _ask_payos_wallet(msg, amount)
    else:
        await _nap_amount_picker(msg, "main")


@router.callback_query(F.data.startswith("payos_wallet:"))
async def on_payos_wallet(cb: CallbackQuery):
    await cb.answer()
    try:
        _, target, amount_s = cb.data.split(":")
        amount = int(amount_s)
        assert target in ("main", "shop") and amount > 0
    except Exception:
        await cb.message.answer("❌ Yêu cầu không hợp lệ, gõ /nap lại nhé.")
        return
    try:
        order_code, checkout_url, qr_code, reused = await _make_payos_order(
            cb.from_user.id, amount, target)
    except Exception as e:
        await cb.message.answer(f"❌ {e}")
        return
    await _send_payos_invoice(cb.message, cb.from_user.id, amount,
                              order_code, checkout_url, qr_code, reused, target)


@router.message(Command("napshop"))
async def on_napshop(msg: Message):
    """Nạp thẳng vào ví shop (mua acc)."""
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) > 1:
        try:
            amount = _parse_payos_amount(parts[1])
        except Exception:
            await msg.answer("❌ Số tiền không hợp lệ. Ví dụ: <code>/napshop 50000</code>")
            return
        try:
            order_code, checkout_url, qr_code, reused = await _make_payos_order(
                msg.from_user.id, amount, "shop")
        except Exception as e:
            await msg.answer(f"❌ {e}")
            return
        await _send_payos_invoice(msg, msg.from_user.id, amount,
                                  order_code, checkout_url, qr_code, reused, "shop")
    else:
        await _nap_amount_picker(msg, "shop")


@router.message(PayOSState.waiting_for_amount)
async def on_nap_amount(msg: Message, state: FSMContext):
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("\U0001f6ab \u0110\u00e3 h\u1ee7y thao t\u00e1c \u0111ang nh\u1eadp. B\u1ea1n g\u00f5 l\u1ea1i l\u1ec7nh v\u1eeba r\u1ed3i nh\u00e9.")
        return
    try:
        amount = _parse_payos_amount(msg.text)
    except Exception:
        await msg.answer("❌ Số tiền không hợp lệ. Vui lòng nhập lại (ví dụ: 50000):")
        return
    await state.clear()
    await _ask_payos_wallet(msg, amount)


@router.message(PayOSState.waiting_for_shop_amount)
async def on_napshop_amount(msg: Message, state: FSMContext):
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("\U0001f6ab \u0110\u00e3 h\u1ee7y thao t\u00e1c \u0111ang nh\u1eadp. B\u1ea1n g\u00f5 l\u1ea1i l\u1ec7nh v\u1eeba r\u1ed3i nh\u00e9.")
        return
    try:
        amount = _parse_payos_amount(msg.text)
    except Exception:
        await msg.answer("❌ Số tiền không hợp lệ. Vui lòng nhập lại (ví dụ: 50000):")
        return
    await state.clear()
    try:
        order_code, checkout_url, qr_code, reused = await _make_payos_order(
            msg.from_user.id, amount, "shop")
    except Exception as e:
        await msg.answer(f"❌ {e}")
        return
    await _send_payos_invoice(msg, msg.from_user.id, amount,
                              order_code, checkout_url, qr_code, reused, "shop")


@router.callback_query(F.data == "payos_auto")
async def on_payos_auto(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    await cb.message.answer("✍️ Nhập <b>số tiền</b> muốn nạp tự động (ví dụ: 50000):")
    await state.set_state(PayOSState.waiting_for_amount)


@router.callback_query(F.data.startswith("payos_cancel:"))
async def on_payos_cancel(cb: CallbackQuery):
    from . import payos as payos_mod
    try:
        order_code = int(cb.data.split(":")[1])
    except Exception:
        await cb.answer("❌ Đơn không hợp lệ.", show_alert=True)
        return
    order = db.get_payos_order(order_code)
    if not order or int(order["tg_id"]) != cb.from_user.id:
        await cb.answer("❌ Không tìm thấy đơn.", show_alert=True)
        return
    if str(order["status"]).upper() != "PENDING":
        await cb.answer("Đơn này đã được xử lý rồi.", show_alert=True)
        return
    try:
        await payos_mod.cancel_payment_link(order_code, "User hủy trên bot")
    except Exception:
        pass
    db.mark_payos_status(order_code, "CANCELLED")
    await cb.answer("✅ Đã hủy đơn.", show_alert=True)
    try:
        await cb.message.edit_text(
            f"❌ <b>ĐÃ HỦY ĐƠN</b>\n\n🧾 Mã đơn: <code>{order_code}</code>",
            parse_mode="HTML", reply_markup=None)
    except Exception:
        pass


@router.message(Command("payos"))
async def on_payos_status(msg: Message):
    """Admin: xem trạng thái cấu hình + thống kê PayOS hôm nay."""
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    from . import payos as payos_mod
    cid, _, _ = payos_mod.get_creds()
    cfg = "✅ Đã cấu hình" if payos_mod.is_configured() else "❌ Chưa cấu hình"
    masked = (cid[:4] + "…" + cid[-4:]) if cid and len(cid) > 8 else ("(trống)" if not cid else "…")
    st = db.payos_stats_today()
    ret_url, cancel_url = payos_mod.get_return_urls()
    await msg.answer(
        "💳 <b>PAYOS — TRẠNG THÁI</b>\n\n"
        f"• Cấu hình: <b>{cfg}</b>\n"
        f"• Client ID: <code>{masked}</code>\n"
        f"• Return URL: <code>{ret_url}</code>\n\n"
        f"📊 <b>Hôm nay:</b> {st['paid']} đơn thành công "
        f"({vnd(st['paid_amount'])}), {st['pending']} đơn đang chờ\n\n"
        "<i>Cấu hình key: </i><code>/payosset &lt;client_id&gt; &lt;api_key&gt; &lt;checksum_key&gt;</code>",
        parse_mode="HTML",
    )


@router.message(Command("payosset"))
async def on_payosset(msg: Message):
    """Chủ shop: lưu key PayOS (tin nhắn chứa key sẽ bị xóa ngay)."""
    if not _perms.is_super(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    try:
        await msg.delete()
    except Exception:
        pass
    if len(parts) != 4:
        await msg.answer(
            "⚠️ Cú pháp: <code>/payosset &lt;client_id&gt; &lt;api_key&gt; &lt;checksum_key&gt;</code>\n"
            "<i>Tin nhắn chứa key đã được xóa. Lấy 3 key này trong trang payos (Kênh thanh toán → Thông tin tích hợp).</i>",
            parse_mode="HTML")
        return
    _, cid, akey, ckey = parts
    db.set_setting("payos_client_id", cid.strip())
    db.set_setting("payos_api_key", akey.strip())
    db.set_setting("payos_checksum_key", ckey.strip())
    await msg.answer("✅ Đã lưu cấu hình PayOS (tin nhắn chứa key đã bị xóa).\n"
                     "Dùng /payos để kiểm tra trạng thái.")

@router.message(Command("hdcookie"))
async def on_hdcookie(msg: Message):
    text = (
        "🍪 <b>HƯỚNG DẪN LẤY COOKIE CÁC NỀN TẢNG</b> 🍪\n\n"
        "<b>1. ZALO (Lấy Cookie & IMEI)</b>\n"
        "• Truy cập: <code>chat.zalo.me</code> trên máy tính (F12 hoặc Chuột phải -> Kiểm tra)\n"
        "• Chọn tab <b>Network</b> (Mạng), F5 tải lại trang. Bấm vào một yêu cầu (request) bất kỳ, kéo xuống phần <b>Request Headers</b>, copy toàn bộ đoạn <code>Cookie: ...</code>\n"
        "• Hoặc chọn tab <b>Application</b> (Ứng dụng) -> Cookies. Tìm khóa <code>zpw_sek</code> và copy giá trị.\n"
        "• <b>Lấy IMEI:</b> Ở tab <b>Application</b> -> Local Storage -> tìm khóa <code>z_uuid</code>, đó chính là IMEI.\n\n"
        "<b>2. FACEBOOK (Cookie)</b>\n"
        "• Đăng nhập tài khoản Clone FB trên trình duyệt.\n"
        "• F12 -> tab <b>Network</b> -> F5. Bấm vào request đầu tiên, kéo xuống <b>Request Headers</b> -> copy toàn bộ dòng <code>Cookie: c_user=...</code>\n"
        "• Hoặc dùng tiện ích mở rộng (Extension) như <b>J2TEAM Security</b> hoặc <b>Get Token by Ninja</b> để copy nhanh.\n\n"
        "<b>3. INSTAGRAM</b>\n"
        "• <b>Cách 1 (Dễ nhất - Khuyên dùng):</b> Bạn chỉ cần vào trang <b>Cài đặt</b> trên Web Dashboard, chọn mục Phương thức check Instagram là <code>Instaloader</code>, sau đó điền trực tiếp <b>Tên đăng nhập (Username)</b> và <b>Mật khẩu</b> của acc Clone IG vào. Tool sẽ tự động đăng nhập ngầm và lấy Cookie cho bạn!\n"
        "• <b>Cách 2 (Thủ công):</b> Đăng nhập IG trên Web -> F12 -> <b>Application</b> -> Cookies -> tìm khóa <code>sessionid</code> và copy dán vào mục IG Session Cookie.\n\n"
        "⚠️ <i>Lưu ý: Tuyệt đối KHÔNG sử dụng tài khoản CHÍNH cho việc này để tránh rủi ro bị khóa. Chỉ nên dùng acc Clone/Phụ để nạp vào Tool!</i>"
    )
    await msg.answer(text, parse_mode="HTML")


@router.message(Command("ping2"))
async def on_ping2(msg: Message):
    await msg.answer("pong version 2 - Đã cập nhật code mới thành công!")

@router.message(Command("tiktok"))
async def on_tiktok(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /tiktok &lt;username&gt;\nVí dụ: /tiktok cristiano")
        return
    username = parse_username(parts[1].strip())
    if not username:
        await msg.answer("❌ Không nhận diện được username.")
        return
    await process_tiktok_check(msg, username)
@router.message(Command("fb"))
async def on_fb(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /fb &lt;uid&gt;\nVí dụ: /fb 100089260699193")
        return
    from .fb import extract_uid
    uid = extract_uid(parts[1].strip())
    if not uid:
        await msg.answer("❌ Không nhận diện được UID.")
        return
    await process_fb_check(msg, uid)


@router.message(Command("getuid"))
async def on_getuid(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer(
            "⚠️ Cú pháp: /getuid &lt;link facebook&gt;\n"
            "Ví dụ: /getuid https://www.facebook.com/zuck"
        )
        return
    link = parts[1].strip()
    wait = await msg.answer("⏳ Đang lấy UID từ link...")
    try:
        from .fb import resolve_fb_uid
        from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton
        import html as html_lib
        uid, fb_name, method = await resolve_fb_uid(link)
        if uid:
            kb = InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="🔍 Check Live/Die ngay", callback_data=f"fb_quickcheck_{uid}")
            ]])
            name_line = f"👤 Tên: <b>{html_lib.escape(fb_name)}</b>\n" if fb_name else ""
            await wait.edit_text(
                f"✅ <b>Lấy UID thành công!</b>\n\n"
                f"🆔 UID: <code>{uid}</code>\n"
                f"{name_line}"
                f"🔎 Lấy bằng: {method}\n\n"
                f"👆 Bấm vào UID để copy, hoặc bấm nút bên dưới để check luôn.",
                reply_markup=kb,
            )
        else:
            await wait.edit_text(
                "❌ Không lấy được UID từ link này.\n"
                "Hãy kiểm tra lại link (cần là link trang cá nhân hoặc page công khai)."
            )
    except Exception as e:
        log.exception("Lỗi getuid %s", link)
        await wait.edit_text(f"❌ Lỗi: {e}")


@router.callback_query(F.data.startswith("fb_quickcheck_"))
async def on_fb_quickcheck(cb: CallbackQuery):
    uid = cb.data.replace("fb_quickcheck_", "")
    await cb.answer()
    try:
        await cb.message.delete()
    except Exception:
        pass
    await process_fb_check(cb.message, uid)


# ─── PROCESS FB POST CHECK ────────────────────────────────────
async def process_fb_post_check(msg: Message, url: str):
    wait = await msg.answer("⏳ Đang lấy thông tin bài viết Facebook...")
    from .fb import fetch_fb_post_info, build_fb_post_caption
    info = await fetch_fb_post_info(url)
    if not info or not info.get("post_id"):
        await wait.edit_text("❌ Không lấy được thông tin bài viết FB. Vui lòng kiểm tra lại link.")
        return
    caption = build_fb_post_caption(info)
    await wait.edit_text(caption, disable_web_page_preview=True)


@router.message(Command("track"))
async def on_track(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /track &lt;username&gt;\nVí dụ: /track cristiano")
        return
    username = parse_username(parts[1].strip())
    if not username:
        await msg.answer("❌ Không nhận diện được username.")
        return

    wait = await msg.answer(f"⏳ Đang thêm theo dõi <b>@{username}</b>...")
    
    # VIP Limit check
    user = db.get_user(msg.chat.id)
    vip_level = dict(user).get("vip_level", 0) if user else 0
    try:
        max_limit = int(db.get_setting(f"vip{vip_level}_limit", [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]))
    except:
        max_limit = [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]
    with db._lock:
        count = db.get_conn().execute("SELECT COUNT(*) FROM tracks WHERE tg_user_id=?", (msg.chat.id,)).fetchone()[0]
    if count >= max_limit:
        await wait.edit_text(f"❌ <b>Giới hạn hạng VIP!</b>\nHạng của bạn chỉ cho phép theo dõi tối đa <b>{max_limit}</b> mục.\nVui lòng /untrack các mục cũ hoặc nâng cấp VIP.")
        return
    try:
        info = await fetch_tiktok_info(username)
        u = msg.from_user
        result = db.add_track(
            u.id, u.username or u.full_name,
            info["username"],
            info["followers"], info["following"], info["videos"]
        )
        if result == -1:
            await wait.edit_text(f"⚠️ Bạn đã theo dõi <b>@{info['username']}</b> rồi!")
            return
        db.add_log("track_add", f"Thêm theo dõi @{info['username']}", u.id, info["username"])
        
        caption = (
            f"✅ <b>Đã thêm theo dõi tài khoản!</b>\n\n"
            f"📱 Kênh: <b><a href='https://www.tiktok.com/@{info['username']}'>@{info['username']}</a></b>\n"
            f"👥 Followers hiện tại: <b>{fmt_num(info['followers'])}</b>\n"
            f"➡️ Đang follow: <b>{fmt_num(info['following'])}</b>\n"
            f"🎬 Tổng videos: <b>{fmt_num(info['videos'])}</b>\n\n"
            f"📩 <i>Bot sẽ thông báo khi có thay đổi follower.</i>"
        )
        if info.get("avatar"):
            try:
                await msg.answer_photo(photo=URLInputFile(info["avatar"], filename="avatar.jpg"), caption=caption)
                await wait.delete()
            except Exception:
                await wait.edit_text(caption, disable_web_page_preview=False)
        else:
            await wait.edit_text(caption, disable_web_page_preview=False)
    except ValueError as e:
        await wait.edit_text(f"❌ {e}")
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi: {e}")


@router.message(Command("untrack"))
async def on_untrack(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /untrack &lt;username&gt;\nVí dụ: /untrack cristiano")
        return
    username = parse_username(parts[1].strip())
    if not username:
        await msg.answer("❌ Không nhận diện được username.")
        return
    ok = db.remove_track(msg.chat.id, username)
    if ok:
        db.add_log("track_remove", f"Huỷ theo dõi @{username}", msg.chat.id, username)
        await msg.answer(f"✅ Đã huỷ theo dõi <b>@{username}</b>.")
    else:
        await msg.answer(f"❌ Không tìm thấy <b>@{username}</b> trong danh sách của bạn.")


@router.message(Command("tracklist"))
async def on_tracklist(msg: Message):
    tracks = db.user_tracks(msg.chat.id)
    if not tracks:
        await msg.answer(
            "📋 Bạn chưa theo dõi tài khoản nào.\n\n"
            "Dùng /track &lt;username&gt; để thêm."
        )
        return
    lines = ["📋 <b>Danh sách đang theo dõi:</b>\n"]
    for i, t in enumerate(tracks, 1):
        lines.append(
            f"{i}. <b>@{t['tiktok_username']}</b>\n"
            f"   👥 {fmt_num(t['last_followers'])} followers"
        )
    await msg.answer("\n".join(lines))


@router.message(Command("trackvlist"))
async def on_trackvlist(msg: Message):
    vtracks = db.user_video_tracks(msg.chat.id)
    if not vtracks:
        await msg.answer("📋 Bạn chưa theo dõi video nào.\n\nDùng /trackv &lt;link_video&gt; để thêm.")
        return
    lines = ["🎬 <b>Video đang theo dõi:</b>\n"]
    for i, v in enumerate(vtracks, 1):
        interval_min = v["check_interval"] // 60
        desc = (v["video_desc"][:50] + "...") if len(v.get("video_desc","")) > 50 else v.get("video_desc","")
        lines.append(
            f"{i}. <a href=\"{v['video_url']}\">@{v['tiktok_username']}</a>\n"
            f"   📝 {desc or 'Khong co mo ta'}\n"
            f"   ▶️ {v['last_plays']:,}  ❤️ {v['last_likes']:,}  💬 {v['last_comments']:,}  🔁 {v['last_shares']:,}  ⭐ {v.get('last_favorites', 0):,}\n"
            f"   ⏱ Check mỗi {interval_min} phút"
        )
    await msg.answer("\n\n".join(lines))


@router.message(Command("trackv"))
async def on_trackv(msg: Message):
    parts = (msg.text or "").split(maxsplit=2)
    if len(parts) < 2:
        await msg.answer(
            "⚠️ Cú pháp: /trackv &lt;link_video&gt; [phút]\n\n"
            "Ví dụ:\n"
            "  /trackv https://tiktok.com/@user/video/123\n"
            "  /trackv https://tiktok.com/@user/video/123 30   (check mỗi 30 phút)"
        )
        return

    from .tiktok import fetch_video_info, parse_video_id
    video_url = parts[1].strip()
    interval_min = 60
    if len(parts) >= 3:
        try:
            interval_min = max(1, int(parts[2]))
        except ValueError:
            pass

    if not parse_video_id(video_url):
        await msg.answer("❌ Link video không hợp lệ. Cần dạng: tiktok.com/@user/video/ID")
        return

    wait = await msg.answer(f"⏳ Đang lấy thông tin video...")
    try:
        info = await fetch_video_info(video_url)
        u = msg.from_user
        r = db.add_video_track(
            u.id, u.username or u.full_name,
            video_url, info["id"], info.get("username",""),
            info.get("desc",""), info.get("cover",""),
            interval_min * 60,
            info["plays"], info["likes"], info["comments"], info["shares"], info.get("favorites", 0)
        )
        if r == -1:
            await wait.edit_text("⚠️ Bạn đã theo dõi video này rồi!")
            return
        db.add_log("video_track_add", f"Them video @{info.get('username','')}", u.id, info.get("username",""))
        desc = (info.get("desc", "")[:80]+"...") if len(info.get("desc","")) > 80 else info.get("desc","")
        caption = (
            f"✅ <b>Đã thêm theo dõi video!</b>\n\n"
            f"📱 Kênh: <b><a href='https://www.tiktok.com/@{info.get('username','')}'>@{info.get('username','')}</a></b>\n"
            f"📝 Mô tả: <i>{desc or 'Không có mô tả'}</i>\n\n"
            f"📊 <b>Thống kê hiện tại:</b>\n"
            f" ┣ ▶️ Lượt xem: <b>{info['plays']:,}</b>\n"
            f" ┣ ❤️ Lượt thích: <b>{info['likes']:,}</b>\n"
            f" ┣ 💬 Bình luận: <b>{info['comments']:,}</b>\n"
            f" ┣ 🔁 Chia sẻ: <b>{info['shares']:,}</b>\n"
            f" ┗ ⭐ Yêu thích: <b>{info.get('favorites', 0):,}</b>\n\n"
            f"⏱ <i>Tự động check mỗi <b>{interval_min} phút</b></i>\n"
            f"📩 <i>Bot sẽ thông báo khi có tương tác mới!</i>"
        )
        if info.get("cover"):
            try:
                await msg.answer_photo(photo=URLInputFile(info["cover"], filename="cover.jpg"), caption=caption)
                await wait.delete()
            except Exception:
                await wait.edit_text(caption, disable_web_page_preview=False)
        else:
            await wait.edit_text(caption, disable_web_page_preview=False)
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi: {e}")


@router.message(Command("untrackv"))
async def on_untrackv(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: /untrackv &lt;link_video&gt;")
        return
    from .tiktok import parse_video_id
    vid_id = parse_video_id(parts[1].strip())
    if not vid_id:
        await msg.answer("❌ Không nhận diện được Video ID.")
        return
    ok = db.remove_video_track(msg.chat.id, vid_id)
    if ok:
        db.add_log("video_track_remove", f"Huy video ID {vid_id}", msg.chat.id)
        await msg.answer("✅ Đã huỷ theo dõi video.")
    else:
        await msg.answer("❌ Không tìm thấy video này trong danh sách của bạn.")

# ─── INSTAGRAM COMMANDS ──────────────────────────────────────
@router.message(Command("ig"))
async def on_ig(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /ig &lt;username&gt;\nVí dụ: /ig cristiano")
        return
    username = parse_ig_username(parts[1].strip())
    if not username:
        await msg.answer("❌ Không nhận diện được username IG.")
        return
    await process_ig_check(msg, username)

@router.message(Command("trackig"))
async def on_trackig(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /trackig &lt;username&gt;\nVí dụ: /trackig cristiano")
        return
    username = parse_ig_username(parts[1].strip())
    if not username:
        await msg.answer("❌ Không nhận diện được username IG.")
        return

    wait = await msg.answer(f"⏳ Đang thêm theo dõi IG <b>@{username}</b>...")
    try:
        info = await fetch_ig_info(username)
        u = msg.from_user
        result = db.add_ig_track(
            u.id, u.username or u.full_name,
            info["username"],
            info["followers"], info["following"], info["posts"]
        )
        if result == -1:
            await wait.edit_text(f"⚠️ Bạn đã theo dõi IG <b>@{info['username']}</b> rồi!")
            return
        db.add_log("track_add", f"Thêm theo dõi IG @{info['username']}", u.id, info["username"])
        caption = (
            f"✅ <b>Đã thêm theo dõi tài khoản IG!</b>\n\n"
            f"📸 Kênh: <b><a href='https://www.instagram.com/{info['username']}'>@{info['username']}</a></b>\n"
            f"👥 Followers hiện tại: <b>{fmt_num(info['followers'])}</b>\n"
            f"➡️ Đang follow: <b>{fmt_num(info['following'])}</b>\n"
            f"🖼 Bài viết: <b>{fmt_num(info['posts'])}</b>\n\n"
            f"📩 <i>Bot sẽ thông báo khi có thay đổi follower.</i>"
        )
        if info.get("avatar"):
            try:
                await msg.answer_photo(photo=URLInputFile(info["avatar"], filename="avatar.jpg"), caption=caption)
                await wait.delete()
            except Exception:
                await wait.edit_text(caption, disable_web_page_preview=False)
        else:
            await wait.edit_text(caption, disable_web_page_preview=False)
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi: {e}")

@router.message(Command("untrackig"))
async def on_untrackig(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /untrackig &lt;username&gt;")
        return
    username = parse_ig_username(parts[1].strip())
    if not username: return
    ok = db.remove_ig_track(msg.chat.id, username)
    if ok:
        db.add_log("track_remove", f"Huỷ theo dõi IG @{username}", msg.chat.id, username)
        await msg.answer(f"✅ Đã huỷ theo dõi IG <b>@{username}</b>.")
    else:
        await msg.answer(f"❌ Không tìm thấy IG <b>@{username}</b> trong danh sách của bạn.")

@router.message(Command("trackiglist"))
async def on_trackiglist(msg: Message):
    tracks = db.user_ig_tracks(msg.chat.id)
    if not tracks:
        await msg.answer("📋 Bạn chưa theo dõi tài khoản IG nào.")
        return
    lines = ["📸 <b>Danh sách IG đang theo dõi:</b>\n"]
    for i, t in enumerate(tracks, 1):
        lines.append(
            f"{i}. <b>@{t['ig_username']}</b>\n"
            f"   👥 {fmt_num(t['last_followers'])} followers"
        )
    await msg.answer("\n".join(lines))

@router.message(Command("trackvig"))
async def on_trackvig(msg: Message):
    parts = (msg.text or "").split(maxsplit=2)
    if len(parts) < 2:
        await msg.answer(
            "⚠️ Cú pháp: /trackvig &lt;link_bài_viết_ig&gt; [phút]\n\n"
            "Ví dụ:\n"
            "  /trackvig https://www.instagram.com/p/C123456/\n"
            "  /trackvig https://www.instagram.com/p/C123456/ 30"
        )
        return

    post_url = parts[1].strip()
    interval_min = 60
    if len(parts) >= 3:
        try: interval_min = max(1, int(parts[2]))
        except ValueError: pass

    post_id = parse_ig_post_id(post_url)
    if not post_id:
        await msg.answer("❌ Link bài viết IG không hợp lệ.")
        return

    wait = await msg.answer(f"⏳ Đang lấy thông tin bài viết IG...")
    try:
        info = await fetch_ig_post_info(post_url)
        u = msg.from_user
        r = db.add_ig_video_track(
            u.id, u.username or u.full_name,
            post_url, info["id"], info.get("username",""),
            info.get("desc",""), info.get("cover",""),
            interval_min * 60,
            info["likes"], info["comments"], info.get("views", 0)
        )
        if r == -1:
            await wait.edit_text("⚠️ Bạn đã theo dõi bài viết IG này rồi!")
            return
        db.add_log("video_track_add", f"Them IG post {info['id']}", u.id, info.get("username",""))
        desc = (info.get("desc","")[:80]+"...") if len(info.get("desc","")) > 80 else info.get("desc","")
        caption = (
            f"✅ <b>Đã thêm theo dõi bài viết IG!</b>\n\n"
            f"📸 Kênh: <b><a href='https://www.instagram.com/{info.get('username','')}'>@{info.get('username','')}</a></b>\n"
            f"📝 Mô tả: <i>{desc or 'Không có mô tả'}</i>\n\n"
            f"📊 <b>Thống kê hiện tại:</b>\n"
            f" ┣ ❤️ Lượt thích: <b>{info['likes']:,}</b>\n"
            f" ┣ 💬 Bình luận: <b>{info['comments']:,}</b>\n"
            f" ┗ 👁️ Lượt xem: <b>{info.get('views', 0):,}</b>\n\n"
            f"⏱ <i>Tự động check mỗi <b>{interval_min} phút</b></i>\n"
            f"📩 <i>Bot sẽ thông báo khi có tương tác mới!</i>"
        )
        if info.get("cover"):
            try:
                await msg.answer_photo(photo=URLInputFile(info["cover"], filename="cover.jpg"), caption=caption)
                await wait.delete()
            except Exception:
                await wait.edit_text(caption, disable_web_page_preview=False)
        else:
            await wait.edit_text(caption, disable_web_page_preview=False)
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi: {e}")

@router.message(Command("untrackvig"))
async def on_untrackvig(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2: return
    post_id = parse_ig_post_id(parts[1].strip())
    if not post_id: return
    ok = db.remove_ig_video_track(msg.chat.id, post_id)
    if ok:
        await msg.answer("✅ Đã huỷ theo dõi bài viết IG.")
    else:
        await msg.answer("❌ Không tìm thấy bài viết IG này.")

@router.message(Command("trackviglist"))
async def on_trackviglist(msg: Message):
    vtracks = db.user_ig_video_tracks(msg.chat.id)
    if not vtracks:
        await msg.answer("📋 Bạn chưa theo dõi bài viết IG nào.")
        return
    lines = ["🎬 <b>Bài viết IG đang theo dõi:</b>\n"]
    for i, v in enumerate(vtracks, 1):
        interval_min = v["check_interval"] // 60
        desc = (v["post_desc"][:50] + "...") if len(v.get("post_desc","")) > 50 else v.get("post_desc","")
        lines.append(
            f"{i}. <a href=\"{v['post_url']}\">@{v['ig_username']}</a>\n"
            f"   📝 {desc or 'Khong co mo ta'}\n"
            f"   ❤️ {v['last_likes']:,}  💬 {v['last_comments']:,}\n"
            f"   ⏱ Check mỗi {interval_min} phút"
        )
    await msg.answer("\n\n".join(lines))


def _sub_active(user) -> bool:
    return user and user["sub_until"] and user["sub_until"] > now()

def _sub_text(user) -> str:
    if _sub_active(user):
        if user["sub_until"] >= 9999999999:
            return "VĨNH VIỄN"
        days_left = (user["sub_until"] - now()) // DAY
        return f"Còn hạn ({days_left} ngày)"
    return "Chưa có / Đã hết hạn"

def status_caption(status: str, note: str, price, header: str = "") -> str:
    icon = "🟢" if status == "live" else "🔴"
    word = "LIVE" if status == "live" else "DIE"
    lines = []
    if header:
        lines.append(header)
    lines.append(f"{icon} Tài khoản đang <b>{word}</b>")
    if note:
        lines.append(f"Ghi chú: {note}")
    if price:
        lines.append(f"Giá: {vnd(price)}")
    return "\n".join(lines)

async def _send_card(bot: Bot, chat_id: int, uid: str, status: str, note, price,
                     avatar: str, header: str = ""):
    caption = status_caption(status, note, price, header)
    if avatar:
        try:
            await bot.send_photo(chat_id, photo=URLInputFile(avatar), caption=caption)
            return
        except Exception:
            pass
    await bot.send_message(chat_id, caption)

@router.message(Command("balance"))
async def on_balance(msg: Message):
    user = db.get_user(msg.chat.id)
    if not user:
        await msg.answer("Bạn chưa /start. Gõ /start trước nhé.")
        return
    credits = db.get_credits(msg.from_user.id)
    shop_bal = int(user["shop_balance"] or 0) if "shop_balance" in user.keys() else 0
    main_bal = int(user["balance"] or 0)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="💰 Nạp ví chính", callback_data="wal_nap:main"),
         InlineKeyboardButton(text="🛒 Nạp ví shop", callback_data="wal_nap:shop")],
        [InlineKeyboardButton(text="📜 Lịch sử giao dịch", callback_data="wal_hist")],
    ])
    await msg.answer(
        f"👛 <b>VÍ CỦA BẠN</b>\n"
        f"━━━━━━━━━━━━━━\n"
        f"💰 <b>Ví chính</b>\n"
        f"      <b>{vnd(main_bal)}</b>\n"
        f"      <i>Check UID • mua gói • credits</i>\n\n"
        f"🛒 <b>Ví shop</b>\n"
        f"      <b>{vnd(shop_bal)}</b>\n"
        f"      <i>Mua tài khoản tự động</i>\n\n"
        f"⚡ <b>Credits:</b> {credits} lượt\n"
        f"👑 <b>Gói:</b> {_sub_text(user)}\n"
        f"━━━━━━━━━━━━━━",
        parse_mode="HTML", reply_markup=kb,
    )


@router.callback_query(F.data.startswith("wal_nap:"))
async def on_wal_nap(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    target = cb.data.split(":", 1)[1]
    await _nap_amount_picker(cb.message, target if target == "shop" else "main")


@router.callback_query(F.data == "wal_hist")
async def on_wal_hist(cb: CallbackQuery):
    await cb.answer()
    await cb.message.answer("📜 Xem lịch sử giao dịch bằng lệnh /lichsu nhé.")

# ─── /sodu: tất cả số dư + giải thích từng loại (cho khách mới đỡ rối) ───
@router.message(Command("sodu"))
async def on_sodu(msg: Message):
    user = db.get_user(msg.chat.id)
    if not user:
        await msg.answer("Bạn chưa /start. Gõ /start trước nhé.")
        return
    credits = db.get_credits(msg.from_user.id)
    shop_bal = int(user["shop_balance"] or 0) if "shop_balance" in user.keys() else 0
    main_bal = int(user["balance"] or 0)
    points = db.loyalty_get(msg.from_user.id)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="💰 Nạp ví chính", callback_data="wal_nap:main"),
         InlineKeyboardButton(text="🛒 Nạp ví shop", callback_data="wal_nap:shop")],
    ])
    await msg.answer(
        "👛 <b>SỐ DƯ CỦA BẠN</b>\n"
        "━━━━━━━━━━━━━━\n"
        f"💰 <b>Ví chính:</b> {vnd(main_bal)}\n"
        "      <i>→ Check UID, mua gói/VIP, mua credits</i>\n\n"
        f"🛒 <b>Ví shop:</b> {vnd(shop_bal)}\n"
        "      <i>→ Mua acc, đặt cọc, hộp mù (nạp bằng /napshop)</i>\n\n"
        f"⚡ <b>Credits:</b> {credits} lượt\n"
        "      <i>→ Check UID hàng loạt (/checkfile, /muacredit để mua thêm)</i>\n\n"
        f"🎁 <b>Điểm:</b> {points} điểm\n"
        "      <i>→ Nhận sau mỗi đơn mua acc, đổi quà bằng /doiqua</i>\n"
        "━━━━━━━━━━━━━━",
        parse_mode="HTML", reply_markup=kb,
    )

# ─── /huongdan: menu nút gọn theo từng mục ───
_HUONGDAN_TEXTS = {
    "shop": (
        "🛒 <b>MUA ACC FACEBOOK</b>\n"
        "━━━━━━━━━━━━━━━\n"
        "1️⃣ Gõ <b>/shop</b> → chọn loại acc\n"
        "2️⃣ Nhập số lượng (hoặc bấm nút nhập số tùy ý)\n"
        "3️⃣ Xác nhận → bot <b>tự check LIVE</b> trước khi giao\n"
        "4️⃣ Nhận acc: 📋 hiện thông tin / 📄 file .txt / 📊 file .xlsx\n\n"
        "💡 Mua bằng <b>ví shop</b> — nạp bằng /napshop.\n"
        "🛒 Mua nhiều loại 1 lần: dùng /giohang."
    ),
    "nap": (
        "💳 <b>NẠP TIỀN</b>\n"
        "━━━━━━━━━━━━━━━\n"
        "• <b>/nap &lt;số tiền&gt;</b> — nạp tự động (quét QR), chọn ví chính/shop\n"
        "• <b>/napshop &lt;số tiền&gt;</b> — nạp thẳng vào ví shop để mua acc\n"
        "• <b>/bank &lt;số tiền&gt;</b> — lấy thông tin chuyển khoản tay\n"
        "• <b>/sodu</b> — xem tất cả số dư: ví chính, ví shop, credits, điểm\n\n"
        "💡 <b>Ví chính</b>: check UID, mua gói/VIP.\n"
        "💡 <b>Ví shop</b>: mua acc, đặt cọc, hộp mù."
    ),
    "bh": (
        "🛡️ <b>BẢO HÀNH ACC</b>\n"
        "━━━━━━━━━━━━━━━\n"
        "1️⃣ Gõ <b>/damua</b> → chọn đơn hàng\n"
        "2️⃣ Bấm yêu cầu bảo hành, ghi rõ lỗi\n"
        "3️⃣ Admin duyệt → được đổi acc hoặc hoàn tiền\n\n"
        "⏱ Mỗi loại acc có thời gian bảo hành riêng\n"
        "(xem trong chi tiết loại acc ở /shop)."
    ),
    "check": (
        "🔍 <b>KIỂM TRA ACC (LIVE/DIE)</b>\n"
        "━━━━━━━━━━━━━━━\n"
        "• <b>/fb &lt;uid&gt;</b> — check 1 acc Facebook\n"
        "• <b>/checkfile</b> — check hàng loạt từ file .txt/.xlsx\n"
        "• <b>/tiktok &lt;username&gt;</b> / <b>/ig &lt;username&gt;</b> — check TikTok/IG\n"
        "• <b>/theodoi</b> — theo dõi UID tự động, die báo ngay\n\n"
        "💡 Check lẻ trừ ví chính, check hàng loạt dùng credits."
    ),
}

_HUONGDAN_MENU_KB = InlineKeyboardMarkup(inline_keyboard=[
    [InlineKeyboardButton(text="🛒 Mua acc", callback_data="hd:shop"),
     InlineKeyboardButton(text="💳 Nạp tiền", callback_data="hd:nap")],
    [InlineKeyboardButton(text="🛡️ Bảo hành", callback_data="hd:bh"),
     InlineKeyboardButton(text="🔍 Kiểm tra acc", callback_data="hd:check")],
])


async def _send_huongdan(target, edit: bool = False):
    text = "📖 <b>HƯỚNG DẪN NHANH</b>\nChọn mục bạn cần 👇"
    if edit:
        try:
            await target.edit_text(text, parse_mode="HTML",
                                   reply_markup=_HUONGDAN_MENU_KB)
        except Exception:
            pass
    else:
        await target.answer(text, parse_mode="HTML",
                            reply_markup=_HUONGDAN_MENU_KB)


@router.message(Command("huongdan"))
async def on_huongdan(msg: Message):
    await _send_huongdan(msg)


@router.callback_query(F.data.startswith("hd:"))
async def on_hd(cb: CallbackQuery):
    await cb.answer()
    topic = cb.data.split(":", 1)[1]
    text = _HUONGDAN_TEXTS.get(topic)
    if not text:
        await _send_huongdan(cb.message, edit=True)
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="← Quay lại", callback_data="hd:back")],
    ])
    try:
        await cb.message.edit_text(text, parse_mode="HTML", reply_markup=kb)
    except Exception:
        pass

from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, CallbackQuery

@router.message(Command("sub"))
async def on_sub(msg: Message):
    p1 = int(db.get_setting("price_1m", "0") or 0)
    bonus = _sub_credit_bonus(p1)
    kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text=f"1 tháng - {vnd(p1)}", callback_data="sub:1")],
            [InlineKeyboardButton(text=f"2 tháng - {vnd(p1 * 2)}", callback_data="sub:2")],
            [InlineKeyboardButton(text=f"3 tháng - {vnd(p1 * 3)}", callback_data="sub:3")],
        ]
    )
    extra = f"\n🎁 <i>Mỗi tháng được tặng kèm <b>{bonus}</b> credits (90% so với mua gói credit thẳng).</i>" if bonus > 0 else ""
    await msg.answer("Chọn gói muốn mua / gia hạn:" + extra, parse_mode="HTML", reply_markup=kb)

@router.callback_query(F.data.startswith("sub:"))
async def on_sub_pick(cb: CallbackQuery):
    months = int(cb.data.split(":")[1])
    p1 = int(db.get_setting("price_1m", "0") or 0)
    cost = p1 * months
    user = db.get_user(cb.from_user.id)
    if not user:
        await cb.answer("Bạn chưa /start.", show_alert=True)
        return
    if user["balance"] < cost:
        await cb.answer(f"Số dư không đủ. Cần {vnd(cost)}, bạn có {vnd(user['balance'])}.", show_alert=True)
        return
    if not db.adjust_balance(cb.from_user.id, -cost, f"Mua gói {months} tháng"):
        await cb.answer("Số dư không đủ.", show_alert=True)
        return
    base = max(now(), user["sub_until"] or 0)
    db.set_sub_until(cb.from_user.id, base + months * 30 * DAY)
    db.add_log("sub", f"Mua {months} tháng (-{cost})", cb.from_user.id)
    # Tặng credits kèm gói tháng: 90% so với mua gói credit thẳng
    bonus = _sub_credit_bonus(p1) * months
    bonus_txt = ""
    if bonus > 0:
        total_cr = db.add_credits(cb.from_user.id, bonus, f"Tặng kèm gói {months} tháng")
        bonus_txt = f"\n🎁 Tặng kèm: <b>{bonus}</b> credits (tổng: {total_cr})"
    u2 = db.get_user(cb.from_user.id)
    await cb.message.answer(
        f"Đã kích hoạt gói <b>{months} tháng</b>.\n"
        f"Số dư còn: <b>{vnd(u2['balance'])}</b>\nGói: <b>{_sub_text(u2)}</b>{bonus_txt}",
        parse_mode="HTML",
    )
    await cb.answer("Thành công")

@router.message(Command("list", "trackfblist"))
async def on_list(msg: Message):
    rows = db.user_watches(msg.chat.id)
    if not rows:
        await msg.answer("Bạn chưa theo dõi UID nào. Dùng /check để thêm.")
        return
    lines = ["<b>Danh sách đang theo dõi</b>"]
    for w in rows:
        st = w["last_status"] or "?"
        icon = "🟢" if st == "live" else ("🔴" if st == "die" else "⚪")
        extra = f" — {w['note']}" if w["note"] else ""
        lines.append(f"{icon} {w['uid']}{extra}")
    await msg.answer("\n".join(lines))

@router.message(Command("remove", "untrackfb"))
async def on_remove(msg: Message):
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("Cú pháp: /remove {uid}")
        return
    n = db.remove_watch(msg.chat.id, parts[1].strip())
    await msg.answer("Đã bỏ theo dõi." if n else "Không tìm thấy UID này.")

@router.message(Command("check", "trackfb"))
async def on_check(msg: Message):
    user = db.get_user(msg.chat.id)
    if not user:
        await msg.answer("Bạn chưa /start. Gõ /start trước nhé.")
        return
    if not _sub_active(user):
        await msg.answer("Bạn cần có gói còn hạn để dùng /check. Gõ /sub để mua gói.")
        return

    uid, note, price, days = parse_check_args(msg.text or "")
    if not uid:
        await msg.answer("Cú pháp: /check {uid} [ghi chú] [giá] [số ngày]")
        return

    from .fb import check_uid, avatar_url
    res = await check_uid(uid)
    status = "live" if res["alive"] else "die"
    avatar = res["avatar_url"] or avatar_url(uid)
    expire_at = now() + days * DAY if days else 0
    wid, is_new = db.add_watch(msg.chat.id, res["uid"], note or "", price or 0, expire_at)
    db.update_watch_status(wid, status, avatar)
    if is_new:
        db.add_log("add", f"Thêm UID {res['uid']} ({status})", msg.chat.id, res["uid"])

    header = "Đã thêm theo dõi:" if is_new else "UID này đã trong danh sách theo dõi:"
    if days:
        header += f" trong {days} ngày"
    await _send_card(msg.bot, msg.chat.id, res["uid"], status, note, price, avatar, header)


@router.callback_query(F.data.startswith("fb_track_"))
async def on_fb_track_btn(cb: CallbackQuery):
    uid = cb.data.replace("fb_track_", "")
    user = db.get_user(cb.from_user.id)
    if not user or not _sub_active(user):
        await cb.answer("❌ Bạn cần có gói còn hạn để dùng tính năng theo dõi.", show_alert=True)
        return
        
    from .fb import check_uid, avatar_url
    res = await check_uid(uid)
    status = "live" if res["alive"] else "die"
    avatar = res.get("avatar_url") or avatar_url(uid)
    
    wid, is_new = db.add_watch(cb.from_user.id, res["uid"], "", 0, 0)
    db.update_watch_status(wid, status, avatar)
    if is_new:
        db.add_log("add", f"Thêm UID {res['uid']} ({status})", cb.from_user.id, res["uid"])

    await cb.answer("✅ Đã thêm vào danh sách theo dõi!" if is_new else "ℹ️ UID này đã được theo dõi rồi.", show_alert=True)
    await _send_card(cb.bot, cb.message.chat.id, res["uid"], status, "", 0, avatar, "Đã thêm theo dõi:")

@router.callback_query(F.data.startswith("fb_note_"))
async def on_fb_note_btn(cb: CallbackQuery, state: FSMContext):
    uid = cb.data.replace("fb_note_", "")
    user = db.get_user(cb.from_user.id)
    if not user or not _sub_active(user):
        await cb.answer("❌ Bạn cần có gói còn hạn để dùng tính năng này.", show_alert=True)
        return
        
    await state.set_state(FBNoteState.waiting_for_note)
    await state.update_data(uid=uid)
    await cb.message.answer(f"✍️ Vui lòng nhập nội dung ghi chú cho UID <b>{uid}</b>:")
    await cb.answer()

@router.message(FBNoteState.waiting_for_note)
async def on_fb_note_input(msg: Message, state: FSMContext):
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("\U0001f6ab \u0110\u00e3 h\u1ee7y thao t\u00e1c \u0111ang nh\u1eadp. B\u1ea1n g\u00f5 l\u1ea1i l\u1ec7nh v\u1eeba r\u1ed3i nh\u00e9.")
        return
    data = await state.get_data()
    uid = data.get("uid")
    note = msg.text.strip()
    await state.clear()
    
    from .fb import check_uid, avatar_url
    res = await check_uid(uid)
    status = "live" if res["alive"] else "die"
    avatar = res.get("avatar_url") or avatar_url(uid)
    
    wid, is_new = db.add_watch(msg.chat.id, res["uid"], note, 0, 0)
    db.update_watch_status(wid, status, avatar)
    if is_new:
        db.add_log("add", f"Thêm UID {res['uid']} kèm ghi chú", msg.chat.id, res["uid"])

    await msg.answer("✅ Đã lưu ghi chú và thêm vào danh sách theo dõi!" if is_new
                     else "✅ Đã lưu ghi chú (UID này đã được theo dõi từ trước).")
    await _send_card(msg.bot, msg.chat.id, res["uid"], status, note, 0, avatar, "Đã thêm theo dõi (Có ghi chú):")



# ---- Nhập liệu cho TRUNG TÂM THEO DÕI /theodoi (đặt trước on_other để không bị nuốt) ----

class TrackMenuState(StatesGroup):
    waiting_fb_uid = State()
    waiting_tiktok_username = State()
    waiting_tiktok_video = State()
    waiting_ig_username = State()
    waiting_ig_post = State()
    waiting_zalo_phone = State()
    waiting_list_name = State()
    waiting_list_add = State()
    waiting_alert_add = State()



@router.message(TrackMenuState.waiting_fb_uid)
async def on_trackmenu_fb_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    uid = (msg.text or "").strip().split()[0]
    if not uid:
        await msg.answer("❌ UID không hợp lệ. Gửi lại hoặc /huy để hủy.")
        return
    user = db.get_user(msg.chat.id)
    if not user or not _sub_active(user):
        await state.clear()
        await msg.answer("❌ Bạn cần có gói còn hạn để dùng tính năng theo dõi.")
        return
    wait = await msg.answer("⏳ Đang kiểm tra UID...")
    try:
        from .fb import check_uid, avatar_url
        res = await check_uid(uid)
        status = "live" if res.get("alive") else "die"
        avatar = res.get("avatar_url") or avatar_url(uid)
        wid, is_new = db.add_watch(msg.chat.id, res.get("uid") or uid, "", 0, 0)
        db.update_watch_status(wid, status, avatar)
        await state.clear()
        await wait.edit_text(
            f"{'✅ <b>Đã thêm theo dõi!</b>' if is_new else 'ℹ️ <b>UID này đã được theo dõi rồi.</b>'}\n"
            f"🆔 UID: <code>{html.escape(res.get('uid') or uid)}</code> — <b>{status.upper()}</b>\n\n"
            f"Mở /theodoi để quản lý.",
            parse_mode="HTML", reply_markup=_trackmenu_main_kb())
    except Exception as e:
        await state.clear()
        await wait.edit_text(f"❌ Lỗi: {html.escape(str(e))}")


@router.message(TrackMenuState.waiting_tiktok_username)
async def on_trackmenu_tiktok_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    username = parse_username((msg.text or "").strip())
    if not username:
        await msg.answer("❌ Không nhận diện được username. Gửi lại hoặc /huy để hủy.")
        return
    user = db.get_user(msg.chat.id)
    max_limit = _vip_limit_for(user)
    with db._lock:
        count = db.get_conn().execute("SELECT COUNT(*) FROM tracks WHERE tg_user_id=?", (msg.chat.id,)).fetchone()[0]
    if count >= max_limit:
        await state.clear()
        await msg.answer(f"❌ Giới hạn VIP: tối đa <b>{max_limit}</b> mục.", parse_mode="HTML")
        return
    wait = await msg.answer(f"⏳ Đang thêm theo dõi <b>@{html.escape(username)}</b>...", parse_mode="HTML")
    try:
        info = await fetch_tiktok_info(username)
        u = msg.from_user
        result = db.add_track(u.id, u.username or u.full_name, info["username"],
                              info["followers"], info["following"], info["videos"])
        await state.clear()
        if result == -1:
            await wait.edit_text(f"⚠️ Bạn đã theo dõi <b>@{html.escape(info['username'])}</b> rồi!", parse_mode="HTML")
        else:
            await wait.edit_text(f"✅ Đã thêm theo dõi <b>@{html.escape(info['username'])}</b>!\nMở /theodoi để quản lý.",
                                 parse_mode="HTML", reply_markup=_trackmenu_main_kb())
    except Exception as e:
        await state.clear()
        await wait.edit_text(f"❌ Lỗi: {html.escape(str(e))}")


@router.message(TrackMenuState.waiting_tiktok_video)
async def on_trackmenu_tiktokv_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    from .tiktok import fetch_video_info, parse_video_id
    video_url = (msg.text or "").strip().split()[0]
    if not parse_video_id(video_url):
        await msg.answer("❌ Link video không hợp lệ. Gửi lại hoặc /huy để hủy.")
        return
    wait = await msg.answer("⏳ Đang lấy thông tin video...")
    try:
        info = await fetch_video_info(video_url)
        u = msg.from_user
        r = db.add_video_track(u.id, u.username or u.full_name, video_url, info["id"],
                               info.get("username", ""), info.get("desc", ""), info.get("cover", ""),
                               3600, info["plays"], info["likes"], info["comments"], info["shares"],
                               info.get("favorites", 0))
        await state.clear()
        if r == -1:
            await wait.edit_text("⚠️ Bạn đã theo dõi video này rồi!")
        else:
            await wait.edit_text("✅ Đã thêm theo dõi video!\nMở /theodoi để quản lý.",
                                 reply_markup=_trackmenu_main_kb())
    except Exception as e:
        await state.clear()
        await wait.edit_text(f"❌ Lỗi: {html.escape(str(e))}")


@router.message(TrackMenuState.waiting_ig_username)
async def on_trackmenu_ig_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    username = parse_ig_username((msg.text or "").strip())
    if not username:
        await msg.answer("❌ Không nhận diện được username IG. Gửi lại hoặc /huy để hủy.")
        return
    wait = await msg.answer(f"⏳ Đang thêm theo dõi IG <b>@{html.escape(username)}</b>...", parse_mode="HTML")
    try:
        info = await fetch_ig_info(username)
        u = msg.from_user
        result = db.add_ig_track(u.id, u.username or u.full_name, info["username"],
                                 info["followers"], info["following"], info["posts"])
        await state.clear()
        if result == -1:
            await wait.edit_text(f"⚠️ Bạn đã theo dõi IG <b>@{html.escape(info['username'])}</b> rồi!", parse_mode="HTML")
        else:
            await wait.edit_text(f"✅ Đã thêm theo dõi IG <b>@{html.escape(info['username'])}</b>!\nMở /theodoi để quản lý.",
                                 parse_mode="HTML", reply_markup=_trackmenu_main_kb())
    except Exception as e:
        await state.clear()
        await wait.edit_text(f"❌ Lỗi: {html.escape(str(e))}")


@router.message(TrackMenuState.waiting_ig_post)
async def on_trackmenu_igp_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    post_url = (msg.text or "").strip().split()[0]
    post_id = parse_ig_post_id(post_url)
    if not post_id:
        await msg.answer("❌ Link bài viết IG không hợp lệ. Gửi lại hoặc /huy để hủy.")
        return
    wait = await msg.answer("⏳ Đang lấy thông tin bài viết IG...")
    try:
        info = await fetch_ig_post_info(post_url)
        u = msg.from_user
        r = db.add_ig_video_track(u.id, u.username or u.full_name, post_url, info["id"],
                                  info.get("username", ""), info.get("desc", ""), info.get("cover", ""),
                                  3600, info["likes"], info["comments"], info.get("views", 0))
        await state.clear()
        if r == -1:
            await wait.edit_text("⚠️ Bạn đã theo dõi bài viết này rồi!")
        else:
            await wait.edit_text("✅ Đã thêm theo dõi bài viết IG!\nMở /theodoi để quản lý.",
                                 reply_markup=_trackmenu_main_kb())
    except Exception as e:
        await state.clear()
        await wait.edit_text(f"❌ Lỗi: {html.escape(str(e))}")


@router.message(TrackMenuState.waiting_zalo_phone)
async def on_trackmenu_zalo_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    phone = "".join(ch for ch in (msg.text or "") if ch.isdigit() or ch == "+").strip()
    if len(phone) < 9:
        await msg.answer("❌ SĐT không hợp lệ. Gửi lại hoặc /huy để hủy.")
        return
    user = db.get_user(msg.chat.id)
    max_limit = _vip_limit_for(user)
    with db._lock:
        c1 = db.get_conn().execute("SELECT COUNT(*) FROM tracks WHERE tg_user_id=?", (msg.chat.id,)).fetchone()[0]
        c2 = db.get_conn().execute("SELECT COUNT(*) FROM zalo_tracks WHERE tg_user_id=?", (msg.chat.id,)).fetchone()[0]
    if c1 + c2 >= max_limit:
        await state.clear()
        await msg.answer(f"❌ Giới hạn VIP: tối đa <b>{max_limit}</b> mục.", parse_mode="HTML")
        return
    wait = await msg.answer("⏳ Đang xử lý theo dõi SĐT Zalo...")
    try:
        ok, err = db.check_daily_limit(msg.chat.id)
        if not ok:
            await state.clear()
            await wait.edit_text(f"❌ {html.escape(err)}")
            return
        from app.zalo_checker import check_zalo_phone
        cookie = db.get_setting("zalo_cookie", "")
        imei = db.get_setting("zalo_imei", "")
        res = await check_zalo_phone(phone, cookie, imei)
        status = "LIVE" if res.get("live") else "DIE"
        db.add_zalo_track(msg.chat.id, msg.from_user.username or msg.from_user.full_name,
                          phone, res.get("name", ""), res.get("avatar", ""), status)
        await state.clear()
        await wait.edit_text(f"✅ Đã thêm SĐT <b>{html.escape(phone)}</b> ({status})!\nMở /theodoi để quản lý.",
                             parse_mode="HTML", reply_markup=_trackmenu_main_kb())
    except Exception as e:
        await state.clear()
        await wait.edit_text(f"❌ Lỗi: {html.escape(str(e))}")


@router.message(TrackMenuState.waiting_list_name)
async def on_trackmenu_listname_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    name = (msg.text or "").strip()
    if not name:
        await msg.answer("❌ Tên không hợp lệ. Gửi lại hoặc /huy để hủy.")
        return
    ok, reason = db.create_user_list(msg.chat.id, name)
    await state.clear()
    if ok:
        await msg.answer(f"✅ Đã tạo danh sách <b>{html.escape(name)}</b>.", parse_mode="HTML",
                         reply_markup=_trackmenu_main_kb())
    else:
        await msg.answer(f"❌ {html.escape(reason)}")


@router.message(TrackMenuState.waiting_list_add)
async def on_trackmenu_listadd_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    raw = (msg.text or "").strip()
    if "|" not in raw:
        await msg.answer("❌ Sai dạng. Gửi: <code>tên danh sách | uid</code>\nVD: <code>khach-vip | 1000123456789</code>\n\n/huy để hủy.",
                         parse_mode="HTML")
        return
    name, uid_raw = [p.strip() for p in raw.split("|", 1)]
    try:
        from .fb import extract_uid
        uid = extract_uid(uid_raw) or uid_raw
    except Exception:
        uid = uid_raw
    ok, reason = db.add_to_user_list(msg.chat.id, name, uid)
    await state.clear()
    if ok:
        await msg.answer(f"✅ Đã thêm <code>{html.escape(uid)}</code> vào <b>{html.escape(name)}</b>.",
                         parse_mode="HTML", reply_markup=_trackmenu_main_kb())
    else:
        await msg.answer(f"❌ {html.escape(reason)}")


@router.message(TrackMenuState.waiting_alert_add)
async def on_trackmenu_alert_input(msg: Message, state: FSMContext):
    if _is_cancel(msg.text):
        await state.clear()
        await msg.answer("Đã hủy.", reply_markup=_trackmenu_main_kb())
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("❌ Sai dạng. Gửi: <code>platform target</code>\nVD: <code>fb_watch 1000123456789</code>\n\n/huy để hủy.",
                         parse_mode="HTML")
        return
    rule_id = db.create_alert_rule(str(msg.chat.id), parts[0], parts[1])
    await state.clear()
    await msg.answer(f"✅ Đã thêm cảnh báo [{html.escape(parts[0])}] {html.escape(parts[1])} (ID: {rule_id}).",
                     parse_mode="HTML", reply_markup=_trackmenu_main_kb())
# ─── FSM handlers: nhập liệu từng bước cho /tienich ──────────────────────────

@router.message(TienIchState.waiting_for_giftcode)
async def on_tienich_giftcode(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy nhập giftcode.")
        return
    code = (msg.text or "").strip()
    await state.clear()
    if not code:
        await msg.answer("❌ Mã trống, thử lại nhé.")
        return
    success, amount, msg_text = db.use_code(code, msg.from_user.id)
    if success:
        db.adjust_balance(msg.from_user.id, amount, f"Sử dụng Giftcode: {code}")
        db.check_vip_upgrade(msg.from_user.id)
        await msg.answer(
            f"✅ <b>NẠP TIỀN THÀNH CÔNG!</b>\n\n"
            f"Bạn đã dùng mã <code>{html.escape(code)}</code> và được cộng <b>{vnd(amount)}</b> vào tài khoản.",
            parse_mode="HTML")
    else:
        await msg.answer(f"❌ {html.escape(msg_text)}", parse_mode="HTML")


@router.message(TienIchState.waiting_for_promo)
async def on_tienich_promo(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy áp mã giảm giá.")
        return
    code = (msg.text or "").strip().upper()
    await state.clear()
    if not code:
        await msg.answer("❌ Mã trống, thử lại nhé.")
        return
    ok, why, row = db.promo_valid(code)
    if not ok:
        await msg.answer(f"❌ {why}")
        return
    db.set_user_promo(msg.from_user.id, code)
    exp_txt = f"\n⏳ Hết hạn: {time.strftime('%d/%m/%Y %H:%M', time.localtime(row['expires_at']))}" if row["expires_at"] else ""
    left_txt = f"\n🎫 Còn lại: <b>{int(row['max_uses']) - int(row['used_count'])}</b> lượt" if row["max_uses"] else ""
    await msg.answer(
        f"✅ <b>Áp mã thành công!</b>\n\n"
        f"🎟️ Mã: <b>{code}</b>\n"
        f"💸 Giảm: <b>{int(row['pct'])}%</b> cho lần mua gói credit tiếp theo{exp_txt}{left_txt}\n\n"
        f"<i>Mở /muacredit để mua ngay.</i>",
        parse_mode="HTML")


@router.message(TienIchState.waiting_for_birthday)
async def on_tienich_birthday(msg: Message, state: FSMContext):
    import datetime as _dt
    import re as _re
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy nhập ngày sinh.")
        return
    arg = (msg.text or "").strip()
    await state.clear()
    m = _re.match(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\s*$", arg)
    if not m:
        await msg.answer("❌ Sai định dạng. Ví dụ: <code>25/12/2000</code>", parse_mode="HTML")
        return
    try:
        d = _dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        await msg.answer("❌ Ngày không hợp lệ. Ví dụ: <code>25/12/2000</code>", parse_mode="HTML")
        return
    today = _dt.date.today()
    if d > today or d.year < 1920:
        await msg.answer("❌ Ngày sinh không hợp lệ.")
        return
    db.set_dob(msg.from_user.id, d.isoformat())
    await msg.answer(f"✅ Đã lưu ngày sinh <b>{d.strftime('%d/%m/%Y')}</b>!\n🎁 Bot sẽ tự tặng quà vào đúng ngày sinh nhật hằng năm.", parse_mode="HTML")


@router.message(TienIchState.waiting_for_refcode)
async def on_tienich_refcode(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy đổi mã giới thiệu.")
        return
    code = (msg.text or "").strip()
    await state.clear()
    if not code.isalnum():
        await msg.answer("❌ Mã giới thiệu chỉ được chứa chữ cái và số!")
        return
    c = db.get_conn()
    exists = c.execute("SELECT tg_id FROM tg_users WHERE ref_code=?", (code,)).fetchone()
    if exists and exists["tg_id"] != msg.chat.id:
        await msg.answer("❌ Mã này đã có người sử dụng. Vui lòng chọn mã khác.")
        return
    with db._lock:
        c.execute("UPDATE tg_users SET ref_code=? WHERE tg_id=?", (code, msg.chat.id))
        c.commit()
    await msg.answer(f"✅ Đã đổi mã giới thiệu thành công: <code>{html.escape(code)}</code>", parse_mode="HTML")


@router.message(TienIchState.waiting_for_getuid)
async def on_tienich_getuid(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy lấy UID.")
        return
    link = (msg.text or "").strip()
    await state.clear()
    if not link:
        await msg.answer("❌ Link trống, thử lại nhé.")
        return
    wait = await msg.answer("⏳ Đang lấy UID từ link...")
    try:
        from .fb import resolve_fb_uid
        uid, fb_name, method = await resolve_fb_uid(link)
        if uid:
            kb = InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="🔍 Check Live/Die ngay", callback_data=f"fb_quickcheck_{uid}")
            ]])
            name_line = f"👤 Tên: <b>{html.escape(fb_name)}</b>\n" if fb_name else ""
            await wait.edit_text(
                f"✅ <b>Lấy UID thành công!</b>\n\n"
                f"🆔 UID: <code>{uid}</code>\n"
                f"{name_line}"
                f"🔎 Lấy bằng: {method}\n\n"
                f"👆 Bấm vào UID để copy, hoặc bấm nút bên dưới để check luôn.",
                parse_mode="HTML", reply_markup=kb)
        else:
            await wait.edit_text("❌ Không lấy được UID từ link này.\nHãy kiểm tra lại link (cần là link trang cá nhân hoặc page công khai).")
    except Exception as e:
        log.exception("tienich getuid %s", link)
        await wait.edit_text(f"❌ Lỗi: {e}")


# ── /ruttien 3 bước ──
@router.message(TienIchState.waiting_withdraw_amount)
async def on_tienich_wd_amount(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy rút tiền.")
        return
    try:
        amount = int((msg.text or "").replace(",", "").replace(".", "").replace("k", "000").replace("K", "000").strip())
    except Exception:
        await msg.answer("❌ Số tiền không hợp lệ! Gửi lại số tiền (VD: 50000).")
        return
    if amount < 50000:
        await msg.answer("❌ Số tiền rút tối thiểu là 50,000 VNĐ.")
        return
    await state.update_data(wd_amount=amount)
    await state.set_state(TienIchState.waiting_withdraw_bank)
    await msg.answer("💸 <b>RÚT HOA HỒNG</b> (bước 2/3)\n\nGửi <b>tên ngân hàng</b> (VD: MBBank, Vietcombank).\nGõ /huy để hủy.", parse_mode="HTML")


@router.message(TienIchState.waiting_withdraw_bank)
async def on_tienich_wd_bank(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy rút tiền.")
        return
    bank = (msg.text or "").strip()
    if not bank:
        await msg.answer("❌ Tên ngân hàng trống, gửi lại nhé.")
        return
    await state.update_data(wd_bank=bank)
    await state.set_state(TienIchState.waiting_withdraw_stk)
    await msg.answer("💸 <b>RÚT HOA HỒNG</b> (bước 3/3)\n\nGửi <b>số tài khoản</b> ngân hàng.\nGõ /huy để hủy.", parse_mode="HTML")


@router.message(TienIchState.waiting_withdraw_stk)
async def on_tienich_wd_stk(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy rút tiền.")
        return
    stk = (msg.text or "").strip()
    if not stk:
        await msg.answer("❌ STK trống, gửi lại nhé.")
        return
    data = await state.get_data()
    await state.clear()
    amount = data.get("wd_amount")
    bank_info = f"{data.get('wd_bank')} - {stk}"
    # tái sử dụng logic kiểm tra của on_ruttien
    c = db.get_conn()
    user = c.execute("SELECT ref_earnings, ref_withdrawn FROM tg_users WHERE tg_id=?", (msg.chat.id,)).fetchone()
    d = dict(user) if user else {}
    available = (d.get("ref_earnings") or 0) - (d.get("ref_withdrawn") or 0)
    import datetime as _dt
    month_start = int(_dt.datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0).timestamp())
    cnt = c.execute("SELECT COUNT(*) as c FROM withdrawal_requests WHERE tg_id=? AND created_at >= ?", (msg.chat.id, month_start)).fetchone()["c"]
    fee = 10000 if cnt >= 2 else 0
    if available < amount + fee:
        await msg.answer(f"❌ Số dư khả dụng không đủ! (Khả dụng: {vnd(available)}, Cần: {vnd(amount + fee)} bao gồm phí {vnd(fee)} nếu có)")
        return
    req_id = db.create_withdrawal_request(msg.chat.id, amount, bank_info, fee)
    await notify_admin_withdrawal_request(req_id, msg.chat.id, amount, bank_info, fee)
    await msg.answer(f"✅ Đã gửi yêu cầu rút <b>{vnd(amount)}</b>.\n🏦 {html.escape(bank_info)}\nVui lòng chờ Admin kiểm tra và duyệt chuyển khoản!", parse_mode="HTML")


# ── /chuyentien 2 bước ──
@router.message(TienIchState.waiting_transfer_uid)
async def on_tienich_tf_uid(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy chuyển tiền.")
        return
    try:
        to_id = int((msg.text or "").strip())
    except ValueError:
        await msg.answer("❌ User ID phải là số. Gửi lại nhé.")
        return
    if to_id == msg.from_user.id:
        await msg.answer("❌ Không thể chuyển tiền cho chính mình.")
        return
    recv = db.get_user(to_id)
    if not recv:
        await msg.answer("❌ Không tìm thấy người nhận (user chưa từng dùng bot).")
        return
    await state.update_data(tf_to_id=to_id)
    await state.set_state(TienIchState.waiting_transfer_amount)
    await msg.answer(
        f"↔️ <b>CHUYỂN TIỀN</b> (bước 2/2)\n\nNgười nhận: <b>{html.escape(recv['name'] or '')}</b> (<code>{to_id}</code>)\n"
        f"Gửi <b>số tiền</b> muốn chuyển.\nGõ /huy để hủy.",
        parse_mode="HTML")


@router.message(TienIchState.waiting_transfer_amount)
async def on_tienich_tf_amount(msg: Message, state: FSMContext):
    if (msg.text or "").strip().lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("❌ Đã hủy chuyển tiền.")
        return
    try:
        amount = int((msg.text or "").replace(".", "").replace(",", "").replace("đ", "").strip())
    except ValueError:
        await msg.answer("❌ Số tiền phải là số. Gửi lại nhé.")
        return
    if amount <= 0:
        await msg.answer("❌ Số tiền phải lớn hơn 0.")
        return
    data = await state.get_data()
    await state.clear()
    to_id = data.get("tf_to_id")
    recv = db.get_user(to_id)
    if not recv:
        await msg.answer("❌ Không tìm thấy người nhận.")
        return
    _pending_transfer[(msg.chat.id, msg.from_user.id)] = (to_id, amount)
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Xác nhận chuyển", callback_data="tf_yes"),
        InlineKeyboardButton(text="❌ Hủy", callback_data="tf_no"),
    ]])
    await msg.answer(
        f"💸 Xác nhận chuyển <b>{vnd(amount)}</b> cho "
        f"{html.escape(recv['name'] or '')} (<code>{to_id}</code>)?",
        reply_markup=kb, parse_mode="HTML")


@router.message(F.text & ~F.text.startswith("/"))
async def on_other(msg: Message):
    username = parse_username(msg.text or "")
    if username:
        await process_tiktok_check(msg, username)
    else:
        await msg.answer(
            "💡 Gõ /tiktok &lt;username&gt; để check TikTok.\n"
            "Hoặc /help để xem hướng dẫn.",
            reply_markup=MENU,
        )


@router.message(Command("vip"))
async def on_vip(msg: Message):
    user = db.get_user(msg.chat.id)
    if not user:
        await msg.answer("Bạn chưa /start. Gõ /start trước nhé.")
        return
    
    vip_level = dict(user).get("vip_level", 0)
    vip_names = {0: "Thường (Free)", 1: "VIP 1", 2: "VIP 2", 3: "VIP 3"}
    try:
        limit = int(db.get_setting(f"vip{vip_level}_limit", [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]))
    except:
        limit = [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]
    
    name = vip_names.get(vip_level, "Thường")
    
    auto_renew_status = "Đang Bật 🟢" if dict(user).get("auto_renew") == 1 else "Đang Tắt 🔴"
    
    text = (
        f"👑 <b>THÔNG TIN HẠNG VIP</b>\n"
        f"━━━━━━━━━━━━━━━━━━\n\n"
        f"👤 Tài khoản: <b>{dict(user).get('username') or dict(user).get('full_name', '')}</b>\n"
        f"💎 Hạng hiện tại: <b>{name}</b>\n"
        f"⏳ Hạn sử dụng: <b>{_sub_text(user)}</b>\n"
        f"📊 Giới hạn theo dõi: <b>{limit} mục</b>\n"
        f"🔄 Gia hạn tự động: <b>{auto_renew_status}</b>\n\n"
        f"💡 <i>Mẹo: Nhấn nút bên dưới để Bật/Tắt tính năng tự động gia hạn gói khi hết hạn (cần đủ số dư ví).</i>\n"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🔄 Bật/Tắt Auto-renew", callback_data="toggle_autorenew")]])
    await msg.answer(text, reply_markup=kb)

@router.callback_query(F.data == "toggle_autorenew")
async def on_toggle_autorenew(cb: CallbackQuery):
    user = db.get_user(cb.from_user.id)
    if not user: return
    new_status = 0 if dict(user).get("auto_renew") == 1 else 1
    with db._lock:
        db.get_conn().execute("UPDATE tg_users SET auto_renew=? WHERE tg_id=?", (new_status, cb.from_user.id))
        db.get_conn().commit()
    status_text = "Bật" if new_status == 1 else "Tắt"
    await cb.answer(f"Đã {status_text} tự động gia hạn!", show_alert=True)
    # Edit msg text
    new_kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🔄 Bật/Tắt Auto-renew", callback_data="toggle_autorenew")]])
    try: await cb.message.edit_reply_markup(reply_markup=new_kb)
    except: pass

@router.callback_query(F.data.startswith("chart_"))
async def on_chart(cb: CallbackQuery):
    import time, json, urllib.parse
    parts = cb.data.split("_")
    if len(parts) < 4: return
    platform = parts[1]
    track_type = parts[2]
    try: track_id = int(parts[3])
    except: return
    
    c = db.get_conn()
    rows = c.execute("SELECT stat_value, created_at FROM track_history WHERE track_id=? AND platform=? AND track_type=? ORDER BY created_at ASC LIMIT 50", (track_id, platform, track_type)).fetchall()
    
    if len(rows) < 2:
        await cb.answer("Chưa đủ dữ liệu để vẽ biểu đồ (Cần ít nhất 2 lần quét).", show_alert=True)
        return
        
    labels = [vn_time_str("%d/%m %H:%M", r["created_at"]) for r in rows]
    data = [r["stat_value"] for r in rows]
    
    chart_config = {
        "type": "line",
        "data": {
            "labels": labels,
            "datasets": [{
                "label": f"Tăng trưởng {track_type} ({platform})",
                "data": data,
                "fill": False,
                "borderColor": "blue",
                "backgroundColor": "rgba(0,0,255,0.1)",
                "borderWidth": 2
            }]
        },
        "options": {
            "title": {"display": True, "text": f"Biểu đồ {track_type}"}
        }
    }
    url = "https://quickchart.io/chart?c=" + urllib.parse.quote(json.dumps(chart_config))
    await cb.message.answer_photo(URLInputFile(url), caption=f"📊 Biểu đồ lịch sử {track_type}")
    await cb.answer()

# ─── BOT MANAGER ─────────────────────────────────────────────
class BotManager:
    def __init__(self):
        self.bot: Optional[Bot] = None
        self.dp:  Optional[Dispatcher] = None
        self._task: Optional[asyncio.Task] = None

    @property
    def running(self) -> bool:
        return self._task is not None and not self._task.done()

    def _make_session(self):
        proxy = (
            os.environ.get("HTTPS_PROXY")
            or os.environ.get("https_proxy")
            or os.environ.get("HTTP_PROXY")
            or os.environ.get("http_proxy")
        )
        return AiohttpSession(proxy=proxy) if proxy else None

    async def start(self, token: str) -> bool:
        await self.stop()
        if not token:
            return False
        self.bot = Bot(token=token, default=DefaultBotProperties(parse_mode=ParseMode.HTML), session=self._make_session())
        try:
            me = await self.bot.get_me()
        except Exception as e:
            log.error("Loi get_me: %s", e)
            self.bot = None
            return False
        try:
            await self.bot.delete_webhook(drop_pending_updates=False)  # GIỮ tin nhắn đang chờ: restart giữa chừng không được nuốt tin user
        except Exception:
            pass
        init_cache_db(_config.DB_PATH)
        self.dp = Dispatcher(storage=SQLiteStorage(_config.DB_PATH))
        self.dp.include_router(router)
        await self.bot.set_my_commands(COMMANDS)
        try:
            _admin_id = int(db.get_setting("admin_tg_id") or 0)
            if _admin_id:
                from aiogram.types import BotCommandScopeChat
                await self.bot.set_my_commands(
                    COMMANDS + ADMIN_COMMANDS_EXTRA,
                    scope=BotCommandScopeChat(chat_id=_admin_id))
        except Exception as e:
            log.warning("admin commands: %s", e)
        poller.set_bot(self.bot)
        self._task = asyncio.create_task(
            self.dp.start_polling(self.bot, handle_signals=False)
        )
        db.add_log("system", f"Bot khởi động: @{me.username}")
        log.info("Bot @%s đang chạy.", me.username)
        return True

    async def stop(self):
        if self.dp:
            try:
                await self.dp.stop_polling()
            except Exception:
                pass
        if self._task:
            self._task.cancel()
            try:
                await self._task
            except (asyncio.CancelledError, Exception):
                pass
            self._task = None
        if self.bot:
            try:
                await self.bot.session.close()
            except Exception:
                pass
            self.bot = None
        self.dp = None

    async def verify_token(self, token: str) -> Optional[str]:
        b = Bot(token=token, session=self._make_session())
        try:
            me = await b.get_me()
            return me.username
        except Exception:
            return None
        finally:
            await b.session.close()


manager = BotManager()


# ─── ZALO BOT MANAGER ────────────────────────────────────────
class ZaloBotManager:
    def __init__(self):
        self.token: str = ""
        self.base_url: str = ""
        self._task: Optional[asyncio.Task] = None
        self._client: Optional[httpx.AsyncClient] = None
        self.offset = 0

    @property
    def running(self) -> bool:
        return self._task is not None and not self._task.done()

    async def start(self, token: str) -> bool:
        await self.stop()
        if not token:
            return False
        self.token = token
        self.base_url = f"https://bot-api.zaloplatforms.com/bot{token}"
        self._client = httpx.AsyncClient(timeout=35.0)

        # Test token validity via getMe (if exists, or just start polling)
        log.info("Zalo Bot starting polling...")
        poller.set_zalo_bot(self)
        self._task = asyncio.create_task(self.polling_loop())
        return True

    async def stop(self):
        if self._task:
            self._task.cancel()
            try:
                await self._task
            except asyncio.CancelledError:
                pass
            self._task = None
        if self._client:
            await self._client.aclose()
            self._client = None

    def _strip_html(self, html_text: str) -> str:
        """Convert HTML to plain text for Zalo (does not support HTML)."""
        # Replace links: <a href="URL">TEXT</a> -> TEXT (URL)
        def link_replacer(match):
            url = match.group(1)
            text = match.group(2)
            if url in text:
                return text
            return f"{text} ({url})"
        text = re.sub(r'<a[^>]*href=["\']([^"\']+)["\'][^>]*>(.*?)</a>', link_replacer, html_text)
        # Remove remaining tags: <b>, <i>, <code>, etc.
        text = re.sub(r'<[^>]+>', '', text)
        # Decode HTML entities
        text = text.replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&').replace('&quot;', '"')
        return text

    async def send_message(self, chat_id: str, text: str, reply_markup: dict = None):
        if not self._client: return
        text = self._strip_html(text)
        url = f"{self.base_url}/sendMessage"
        payload = {
            "chat_id": chat_id,
            "text": text
        }
        if reply_markup:
            payload["reply_markup"] = reply_markup
        try:
            resp = await self._client.post(url, json=payload)
            data = resp.json()
            if not data.get("ok"):
                log.warning("Zalo sendMessage failed: %s", data)
        except Exception as e:
            log.warning("Zalo sendMessage error: %s", e)

    async def send_photo(self, chat_id: str, photo_url: str, caption: str = ""):
        if not self._client: return
        caption = self._strip_html(caption)
        text = f"{caption}\n\n📷 Ảnh: {photo_url}" if photo_url else caption
        await self.send_message(chat_id, text)

    async def polling_loop(self):
        while True:
            try:
                url = f"{self.base_url}/getUpdates"
                payload = {"offset": self.offset, "timeout": 30}
                resp = await self._client.post(url, json=payload)
                if resp.status_code == 200:
                    data = resp.json()
                    if data.get("ok"):
                        result = data.get("result", [])
                        # Zalo co the tra ve dict (1 message) hoac list (nhieu message)
                        if isinstance(result, dict):
                            updates = [result]
                        elif isinstance(result, list):
                            updates = result
                        else:
                            updates = []

                        for u in updates:
                            if not isinstance(u, dict):
                                continue
                            # Cap nhat offset
                            upd_id = u.get("update_id")
                            if upd_id is not None:
                                self.offset = int(upd_id) + 1
                            # LOG CAU TRUC THUC TE DE DEBUG
                            log.warning("ZALO UPDATE STRUCT: keys=%s | data=%s", list(u.keys()), str(u)[:500])
                            # Xu ly tin nhan - thu nhieu key khac nhau
                            msg = u.get("message") or u.get("edited_message") or u.get("channel_post")
                            # Neu chinh u la tin nhan (Zalo co the tra truc tiep object)
                            if msg is None and u.get("text") is not None:
                                msg = u
                                
                            if "callback_query" in u:
                                await self.handle_callback(u["callback_query"])
                            elif isinstance(msg, dict):
                                await self.handle_message(msg)
                            elif isinstance(msg, str):
                                await self.handle_message({"text": msg, "chat": {}, "from": {}})

                    else:
                        err_code = data.get("error_code")
                        if err_code not in (408, 504):
                            log.warning("Zalo getUpdates error: %s", data)
                            await asyncio.sleep(2)
                else:
                    log.warning("Zalo HTTP %s: %s", resp.status_code, resp.text[:200])
                    await asyncio.sleep(2)
            except asyncio.CancelledError:
                break
            except Exception as e:
                log.error("Zalo polling error: %s", e, exc_info=True)
                await asyncio.sleep(5)



    async def handle_callback(self, cb: dict):
        data = cb.get("data", "")
        if not data: return
        
        chat_id = cb.get("message", {}).get("chat", {}).get("id")
        
        if data.startswith("zalo_confirm_"):
            parts = data.split("_")
            if len(parts) >= 4:
                tg_id = int(parts[2])
                amount = int(parts[3])
                
                code = db.get_unused_code(amount)
                
                try:
                    kb = InlineKeyboardMarkup(inline_keyboard=[[
                        InlineKeyboardButton(text="🎁 Sử dụng luôn", callback_data=f"use_code_{code}")
                    ]])
                    await manager.bot.send_message(
                        tg_id, 
                        f"🎉 <b>Thanh toán thành công!</b>\n\nĐây là mã code nạp tiền trị giá <b>{vnd(amount)}</b> của bạn:\n"
                        f"👉 <code>{code}</code>\n\n"
                        "<i>Nhấn nút bên dưới để sử dụng mã ngay lập tức!</i>",
                        reply_markup=kb
                    )
                    
                    if chat_id:
                        await self.send_message(str(chat_id), f"✅ Đã xác nhận và phát mã {code} ({vnd(amount)}) cho ID {tg_id}")
                except Exception as e:
                    log.error(f"Lỗi gửi code cho khách: {e}")
                    if chat_id:
                        await self.send_message(str(chat_id), f"❌ Lỗi gửi code cho khách: {e}")

    async def handle_message(self, msg: dict):
        chat_id = str(msg.get("chat", {}).get("id", ""))
        text = msg.get("text", "").strip()
        if not chat_id or not text: return

        username = msg.get("from", {}).get("display_name") or msg.get("from", {}).get("username") or msg.get("from", {}).get("first_name", chat_id)
        
        txt_lower = text.lower()
        if txt_lower.startswith("/start") or txt_lower.startswith("/id"):
            await self.cmd_help(chat_id, username)
        elif txt_lower.startswith("/topup"):
            await self.cmd_topup(chat_id, text)
        elif txt_lower.startswith("/phatcode"):
            await self.cmd_phatcode(chat_id, text)
        elif txt_lower.startswith("/web"):
            token = db.create_magic_link(int(chat_id))
            web_domain = db.get_setting("web_domain", "http://127.0.0.1:8000")
            url = f"{web_domain.rstrip('/')}/auth?token={token}"
            kb = InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="🌐 Đăng nhập Web", url=url)
            ]])
            await self.send_message(chat_id, "🔗 Bấm vào nút bên dưới để tự động đăng nhập vào Web:", reply_markup=kb)
        else:
            await self.send_message(chat_id, f"💡 Zalo Chat ID của bạn: {chat_id}\n\nLệnh có sẵn:\n- Copy Chat ID dán vào web để nhận thông báo\n- Phát Code: /phatcode <ID> <SỐ TIỀN>\n- Cộng thẳng: /topup <ID> <SỐ TIỀN>\n- Bảng điều khiển Web: /web")

    async def cmd_phatcode(self, chat_id, text):
        admin_id = db.get_setting("admin_zalo_id", "")
        if not admin_id or chat_id != admin_id:
            await self.send_message(chat_id, "⛔ Chỉ Admin mới được dùng lệnh này!")
            return
            
        parts = text.split()
        if len(parts) < 3:
            await self.send_message(chat_id, "⚠️ Cú pháp sai! Vui lòng gửi lệnh có sẵn trong thông báo.\nVD: /phatcode 12345 50000")
            return
            
        try:
            tg_id = int(parts[1])
            amount = int(parts[2])
            
            code = db.get_unused_code(amount)
            kb = InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="🎁 Sử dụng luôn", callback_data=f"use_code_{code}")
            ]])
            from .bot import manager
            await manager.bot.send_message(
                tg_id, 
                f"🎉 <b>Thanh toán thành công!</b>\n\nĐây là mã code nạp tiền trị giá <b>{vnd(amount)}</b> của bạn:\n"
                f"👉 <code>{code}</code>\n\n"
                "<i>Nhấn nút bên dưới để sử dụng mã ngay lập tức!</i>",
                reply_markup=kb
            )
            await self.send_message(chat_id, f"✅ Đã phát mã {code} ({vnd(amount)}) cho ID {tg_id}")
        except Exception as e:
            await self.send_message(chat_id, f"❌ Lỗi: {e}")

    async def cmd_help(self, chat_id, username):
        help_text = (
            f"👋 Xin chào Admin <b>{username}</b>!\n\n"
            f"🆔 Zalo Chat ID của bạn là: {chat_id}\n\n"
            "1. Copy Chat ID này và dán vào ô 'Admin Zalo Chat ID' trên Web Dashboard để nhận thông báo khách nạp tiền.\n"
            "2. Để nạp tiền cho khách trực tiếp từ đây, dùng lệnh:\n"
            "👉 /topup <ID_TELE> <SỐ TIỀN>\n"
            "(Ví dụ: /topup 123456789 50000)"
        )
        await self.send_message(chat_id, help_text)

    async def cmd_topup(self, chat_id, text):
        admin_id = db.get_setting("admin_zalo_id", "")
        if not admin_id or chat_id != admin_id:
            await self.send_message(chat_id, "⛔ Chỉ Admin (Zalo ID đã cài đặt) mới được dùng lệnh này!")
            return
            
        parts = text.split()
        if len(parts) < 3:
            await self.send_message(chat_id, "⚠️ Cú pháp sai!\nVí dụ: /topup 123456789 50000")
            return
            
        try:
            tg_id = int(parts[1])
            amount = int(parts[2].replace(",", "").replace(".", "").replace("k", "000").replace("K", "000"))
        except:
            await self.send_message(chat_id, "❌ ID Telegram hoặc Số tiền không hợp lệ!")
            return
            
        db.adjust_balance(tg_id, amount, "Nạp tiền qua Zalo Bot")
        await self.send_message(chat_id, f"✅ Đã cộng thành công {vnd(amount)} cho ID Telegram: {tg_id}")
        
        try:
            from .bot import manager
            msg_text_resp = f"🎉 <b>NẠP TIỀN THÀNH CÔNG!</b>\n\nAdmin vừa cộng cho bạn: <b>{vnd(amount)}</b>\n👉 Gõ /balance để kiểm tra số dư nhé."
            upgraded, new_vip, is_lifetime = db.check_vip_upgrade(tg_id)
            if upgraded and new_vip > 0:
                limit = db.get_setting(f"vip{new_vip}_limit", "10")
                msg_text_resp += (
                    f"\n\n🎉 <b>CHÚC MỪNG BẠN ĐÃ LÊN VIP {new_vip}!</b> 🎉\n\n"
                    f"🎁 <b>Đặc quyền mới:</b>\n"
                    f"- Mức độ theo dõi tối đa: <b>{limit} mục</b>/nền tảng\n"
                )
                if is_lifetime:
                    msg_text_resp += "- Hạn sử dụng: <b>VĨNH VIỄN</b>\n\n"
                else:
                    msg_text_resp += "\n"
            await manager.bot.send_message(tg_id, msg_text_resp, parse_mode="HTML")
        except:
            pass


zalo_manager = ZaloBotManager()


# --- ZALO TRACKING COMMANDS ---
from app.zalo_checker import check_zalo_phone

@router.message(Command("zalo"))
async def on_zalo(msg: Message, command: CommandObject):
    phone = command.args
    if not phone:
        await msg.answer("💡 Gõ /zalo &lt;sđt&gt; để kiểm tra nhanh SĐT Zalo.")
        return
        
    wait = await msg.answer("⏳ Đang kiểm tra Zalo...")
    try:
        cookie = db.get_setting("zalo_cookie", "")
        imei = db.get_setting("zalo_imei", "")
        res = await check_zalo_phone(phone, cookie, imei)
        if res.get("live"):
            await wait.edit_text(f"✅ <b>LIVE</b>\\nSĐT: {phone}\\nTên Zalo: <b>{res['name']}</b>", parse_mode="HTML")
        else:
            await wait.edit_text(f"❌ <b>DIE / KHÔNG TÌM THẤY</b>\\nSĐT: {phone}\\nLỗi: {res.get('error', '')}", parse_mode="HTML")
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi: {str(e)}")

@router.message(Command("trackzalo"))
async def on_trackzalo(msg: Message, command: CommandObject):
    phone = command.args
    if not phone:
        await msg.answer("💡 Gõ /trackzalo &lt;sđt&gt; để theo dõi biến động SĐT Zalo.")
        return
        
    user = db.get_user(msg.chat.id)
    vip_level = dict(user).get("vip_level", 0) if user else 0
    try: max_limit = int(db.get_setting(f"vip{vip_level}_limit", [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]))
    except: max_limit = [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]
    
    with db._lock: count = db.get_conn().execute("SELECT COUNT(*) FROM tracks WHERE tg_user_id=?", (msg.chat.id,)).fetchone()[0]
    with db._lock: z_count = db.get_conn().execute("SELECT COUNT(*) FROM zalo_tracks WHERE tg_user_id=?", (msg.chat.id,)).fetchone()[0]
    if count + z_count >= max_limit:
        await msg.answer(f"❌ <b>Giới hạn hạng VIP!</b>\\nHạng của bạn chỉ cho phép tối đa <b>{max_limit}</b> mục.", parse_mode="HTML")
        return
        
    wait = await msg.answer("⏳ Đang xử lý theo dõi SĐT Zalo...")
    try:
        ok, err = db.check_daily_limit(msg.chat.id)
        if not ok:
            await wait.edit_text(f"❌ {err}")
            return
            
        cookie = db.get_setting("zalo_cookie", "")
        imei = db.get_setting("zalo_imei", "")
        res = await check_zalo_phone(phone, cookie, imei)
        
        if res.get("live"):
            status = "LIVE"
            name = res.get("name", "")
            avatar = res.get("avatar", "")
        else:
            status = "DIE"
            name = ""
            avatar = ""
            
        db.add_zalo_track(msg.chat.id, msg.from_user.username or msg.from_user.full_name, phone, name, avatar, status)
        db.add_log("track_add", f"Thêm Zalo {phone}", msg.chat.id, phone)
        await wait.edit_text(f"✅ Đã thêm SĐT Zalo <b>{phone}</b> vào danh sách theo dõi!\\nTrạng thái hiện tại: {status}", parse_mode="HTML")
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi: {str(e)}")

# --- ALERTS ---
@router.message(Command("alert"))
async def on_alert_cmd(msg: Message):
    parts = msg.text.split()
    if len(parts) < 3:
        await msg.answer("⚠️ Cú pháp: /alert <platform> <target>\nVD: /alert fb_watch 123456789")
        return
    platform = parts[1]
    target = parts[2]
    rule_id = db.create_alert_rule(str(msg.chat.id), platform, target)
    await msg.answer(f"✅ Đã thêm cảnh báo cho {platform} mục {target} (ID: {rule_id})")


# ─── 📁 DANH SÁCH NGƯỜI DÙNG (/newlist, /lists, /addtolist, /scanlist, /deletelist) ───

def _esc_list_name(s: str) -> str:
    return html.escape(s or "", quote=False)


@router.message(Command("newlist"))
async def on_newlist_cmd(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /newlist &lt;tên danh sách&gt;")
        return
    name = parts[1].strip()
    ok, reason = db.create_user_list(msg.chat.id, name)
    if ok:
        await msg.answer(f"✅ Đã tạo danh sách <b>{_esc_list_name(name)}</b>.", parse_mode="HTML")
    else:
        await msg.answer(f"❌ {html.escape(reason, quote=False)}")


@router.message(Command("lists"))
async def on_lists_cmd(msg: Message):
    lists = db.get_user_lists(msg.chat.id)
    if not lists:
        await msg.answer("📁 Bạn chưa có danh sách nào.\nGõ <code>/newlist tên</code> để tạo.", parse_mode="HTML")
        return
    lines = ["📁 <b>Danh sách của bạn:</b>\n"]
    for l in lists:
        l = dict(l)
        d = time.strftime("%d/%m/%Y", time.localtime(l.get("created_at") or 0))
        lines.append(f"• <b>{_esc_list_name(l.get('name'))}</b> — {l.get('item_count', 0)} UID ({d})")
    lines.append("\nGõ <code>/addtolist tên uid</code> để thêm UID.")
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("addtolist"))
async def on_addtolist_cmd(msg: Message):
    parts = (msg.text or "").rsplit(maxsplit=2)
    if len(parts) < 3 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /addtolist &lt;tên danh sách&gt; &lt;uid&gt;")
        return
    name = parts[1].strip()
    raw_val = parts[2].strip()
    try:
        from .fb import extract_uid
        uid = extract_uid(raw_val) or raw_val
    except Exception:
        uid = raw_val
    ok, reason = db.add_to_user_list(msg.chat.id, name, uid)
    if ok:
        await msg.answer(f"✅ Đã thêm <code>{_esc_list_name(uid)}</code> vào <b>{_esc_list_name(name)}</b>.", parse_mode="HTML")
    else:
        await msg.answer(f"❌ {html.escape(reason, quote=False)}")


@router.message(Command("scanlist"))
async def on_scanlist_cmd(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /scanlist &lt;tên danh sách&gt;")
        return
    await _run_scanlist(msg.chat.id, msg.from_user.id, parts[1].strip(),
                        answer=msg.answer, bot=msg.bot)


async def _run_scanlist(chat_id: int, user_id: int, name: str, answer, bot):
    """Logic quét danh sách dùng chung cho /scanlist và nút trackmenu:lists_scan."""
    items = db.get_list_items(chat_id, name)
    if not items:
        await answer(f"❌ Danh sách <b>{_esc_list_name(name)}</b> trống hoặc không tồn tại.", parse_mode="HTML")
        return
    uids = [str(dict(i).get("value", "")) for i in items if dict(i).get("value")]
    if not uids:
        await answer(f"❌ Danh sách <b>{_esc_list_name(name)}</b> trống.", parse_mode="HTML")
        return
    wait = await answer(f"⏳ Đang check {len(uids)} UID trong <b>{_esc_list_name(name)}</b>...", parse_mode="HTML")
    if not await _ensure_bulk_credits_simple(chat_id, user_id, wait, len(uids), bot):
        return
    live, die, err = await check_uids_batch(uids, user_id=user_id)
    lines = [f"⚡ <b>KẾT QUẢ SCAN — {len(uids)} UID:</b>\n",
             f"🟢 Live: <b>{len(live)}</b>  |  🔴 Die: <b>{len(die)}</b>  |  ❓ Lỗi: <b>{len(err)}</b>\n"]
    if live:
        lines.append("🟢 <b>Live:</b>")
        lines += [f"• <code>{_esc_list_name(u)}</code>" for u in live[:30]]
        if len(live) > 30:
            lines.append(f"  …và {len(live) - 30} UID live nữa")
    if die:
        lines.append("\n🔴 <b>Die:</b>")
        lines += [f"• <code>{_esc_list_name(u)}</code>" for u in die[:30]]
        if len(die) > 30:
            lines.append(f"  …và {len(die) - 30} UID die nữa")
    if err:
        lines.append(f"\n❓ <b>Lỗi kiểm tra:</b> {len(err)} UID (thử lại sau)")
    await wait.edit_text("\n".join(lines), parse_mode="HTML")


@router.message(Command("deletelist"))
async def on_deletelist_cmd(msg: Message):
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("⚠️ Cú pháp: /deletelist &lt;tên danh sách&gt;")
        return
    name = parts[1].strip()
    if db.delete_user_list(msg.chat.id, name):
        await msg.answer(f"🗑 Đã xóa danh sách <b>{_esc_list_name(name)}</b>.", parse_mode="HTML")
    else:
        await msg.answer(f"❌ Không tìm thấy danh sách <b>{_esc_list_name(name)}</b>.", parse_mode="HTML")


@router.message(Command("alertlist"))
async def on_alertlist_cmd(msg: Message):
    rules = db.get_alert_rules(tg_id=str(msg.chat.id))
    if not rules:
        await msg.answer("Bạn không có cảnh báo nào.")
        return
    lines = ["🚨 <b>Danh sách Cảnh báo</b>\n"]
    for r in rules:
        lines.append(f"• ID {r['id']}: [{r['platform']}] {r['target']} ({r['condition']})")
    await msg.answer("\n".join(lines), parse_mode="HTML")

@router.message(Command("alertoff"))
async def on_alertoff_cmd(msg: Message):
    parts = msg.text.split(maxsplit=1)
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: /alertoff &lt;target_hoặc_id&gt;")
        return
    target = parts[1]
    rules = db.get_alert_rules(tg_id=str(msg.chat.id))
    deleted = False
    for r in rules:
        if str(r["id"]) == target or r["target"] == target:
            db.delete_alert_rule(r["id"])
            deleted = True
    if deleted:
        await msg.answer(f"✅ Đã xoá cảnh báo cho {target}.")
    else:
        await msg.answer("❌ Không tìm thấy cảnh báo phù hợp.")

@router.callback_query(lambda c: c.data and c.data.startswith("camp_giveaway_"))
async def on_camp_giveaway(cb: CallbackQuery):
    camp_id = int(cb.data.split("_")[-1])
    tg_id = cb.from_user.id
    
    if db.check_campaign_participation(camp_id, tg_id):
        await cb.answer("Bạn đã nhận lì xì từ chiến dịch này rồi!", show_alert=True)
        return
        
    camp = db.get_campaign(camp_id)
    if not camp or camp["status"] != "finished":
        await cb.answer("Chiến dịch này không khả dụng!", show_alert=True)
        return
        
    import json
    import random
    try: config = json.loads(camp.get("config", "{}"))
    except: config = {}
    
    min_reward = int(config.get("min_reward") or 1000)
    max_reward = int(config.get("max_reward") or 5000)
    max_winners = int(config.get("max_winners") or 0)
    
    if min_reward > max_reward:
        min_reward, max_reward = max_reward, min_reward
    
    # check max winners
    try: stats = json.loads(camp.get("stats", "{}"))
    except: stats = {}
    claims = stats.get("claims", 0)
    
    if max_winners > 0 and claims >= max_winners:
        await cb.answer("Rất tiếc, số lượng lì xì đã hết!", show_alert=True)
        return
        
    reward = random.randint(min_reward, max_reward)
    
    db.adjust_balance(tg_id, reward, f"Lì xì chiến dịch #{camp_id}")
    db.add_campaign_participant(camp_id, tg_id, "claimed", f"{reward}")
    
    stats["claims"] = claims + 1
    db.update_campaign_stats(camp_id, json.dumps(stats))
    
    await cb.answer(f"🎉 Chúc mừng bạn nhận được {fmt_num(reward)}đ lì xì!", show_alert=True)

@router.callback_query(lambda c: c.data and c.data.startswith("camp_bounty_"))
async def on_camp_bounty(cb: CallbackQuery):
    camp_id = int(cb.data.split("_")[-1])
    tg_id = cb.from_user.id
    
    if db.check_campaign_participation(camp_id, tg_id):
        await cb.answer("Bạn đã tham gia nhiệm vụ này rồi! (Đang chờ duyệt)", show_alert=True)
        return
        
    db.add_campaign_participant(camp_id, tg_id, "pending_review", "")
    await cb.answer("Hệ thống đã ghi nhận yêu cầu tham gia của bạn! Admin sẽ kiểm tra và duyệt sau.", show_alert=True)

# ─── BẢNG HẰNG SỐ & STATE NÂNG CAO ───────────────────────────────────────────

# Cache ket qua check file/cookie: xem app/persist.py (luu RAM + DB, song sot qua restart)

class CookieCheckState(StatesGroup):
    waiting_for_file = State()

class FileCheckState(StatesGroup):
    waiting_for_file = State()


# ─── ⚡ QUÉT SIÊU TỐC ASYNC BATCHING HELPER ────────────────────────────────────

async def check_uids_batch(uids: list, concurrency: int = 15, user_id: int = None):
    """Quét hàng loạt UID Facebook bằng Async Batching."""
    sem = asyncio.Semaphore(concurrency)
    live_list = []
    die_list = []
    error_list = []

    async def _check_one(uid):
        async with sem:
            try:
                res = await check_uid(uid)
                status = res.get("status", "error")
                if user_id:
                    try:
                        db.log_check_stat(user_id, "fb", uid, status, res.get("via", ""))
                    except Exception:
                        pass
                if status == "live":
                    live_list.append(uid)
                elif status == "die":
                    die_list.append(uid)
                else:
                    error_list.append(uid)
            except Exception:
                error_list.append(uid)

    tasks = [_check_one(u) for u in uids]
    await asyncio.gather(*tasks, return_exceptions=True)

    if user_id and (live_list or die_list):
        try:
            summary_msg = f"Check file: {len(live_list)} Live, {len(die_list)} Die"
            db.add_check_history(user_id, "fb", summary_msg, "live" if live_list else "die")
        except Exception:
            pass

    return live_list, die_list, error_list


# ─── 1. CHECK UID FACEBOOK BẰNG FILE TEXT (.txt) ──────────────────────────────

@router.message(Command("checkfile"))
@router.message(Command("scanfile"))
@router.message(Command("fbfile"))
async def on_checkfile_cmd(msg: Message, state: FSMContext):
    """Lệnh check UID FB bằng file .txt."""
    if msg.document:
        await _process_file_check(msg, state)
        return

    await state.set_state(FileCheckState.waiting_for_file)
    await msg.answer(
        "📁 <b>CHECK UID FACEBOOK BẰNG FILE (.txt / .xlsx)</b>\n"
        "━━━━━━━━━━━━━━━━━━━━\n\n"
        "Vui lòng đính kèm và <b>gửi file</b> chứa danh sách UID Facebook:\n"
        "• File <b>.txt</b>: mỗi UID hoặc đường link 1 dòng.\n"
        "• File <b>.xlsx</b>: mọi cấu trúc cột — ô nào chứa link Facebook "
        "(kể cả link chưa có ID) sẽ <b>tự động giải thành UID</b> rồi mới check. "
        "Xong sẽ trả file mới giữ nguyên cột, ô link được thay bằng UID.\n\n"
        "💡 <i>Gửi file đính kèm ngay tại đây. Nếu muốn huỷ, hãy gõ /cancel</i>",
        parse_mode="HTML"
    )


@router.message(FileCheckState.waiting_for_file, F.document)
async def on_file_received(msg: Message, state: FSMContext):
    await state.clear()
    await _process_file_check(msg, state)


async def _ensure_bulk_credits(msg: Message, wait: Message, n_uids: int) -> bool:
    """Kiểm tra + trừ credits cho check bulk (VIP còn hạn được miễn).
    Trả về True nếu được phép tiếp tục."""
    bulk_cost = int(db.get_setting("bulk_credit_cost", "0") or 0)
    raw_u = db.get_user(msg.from_user.id)
    if bulk_cost > 0 and not _sub_active(dict(raw_u) if raw_u else None):
        need = n_uids * bulk_cost
        have = db.get_credits(msg.from_user.id)
        if have < need:
            await wait.edit_text(
                f"❌ <b>Không đủ credits!</b>\n\n"
                f"Cần: <b>{need}</b> credits ({n_uids} UID × {bulk_cost})\n"
                f"Bạn có: <b>{have}</b> credits\n\n"
                f"Mua thêm bằng /muacredit",
                parse_mode="HTML",
            )
            return False
        db.consume_credits(msg.from_user.id, need)
        await _maybe_low_credit_warn(msg.bot, msg.from_user.id, msg.chat.id)
    return True


async def _ensure_bulk_credits_simple(chat_id: int, user_id: int, wait: Message, n_uids: int, bot) -> bool:
    """Bản gọn của _ensure_bulk_credits dùng cho nút bấm (không cần Message đầy đủ)."""
    bulk_cost = int(db.get_setting("bulk_credit_cost", "0") or 0)
    raw_u = db.get_user(user_id)
    if bulk_cost > 0 and not _sub_active(dict(raw_u) if raw_u else None):
        need = n_uids * bulk_cost
        have = db.get_credits(user_id)
        if have < need:
            await wait.edit_text(
                f"❌ <b>Không đủ credits!</b>\n\n"
                f"Cần: <b>{need}</b> credits ({n_uids} UID × {bulk_cost})\n"
                f"Bạn có: <b>{have}</b> credits\n\n"
                f"Mua thêm bằng /muacredit",
                parse_mode="HTML",
            )
            return False
        db.consume_credits(user_id, need)
        await _maybe_low_credit_warn(bot, user_id, chat_id)
    return True


async def _process_file_check(msg: Message, state: FSMContext):
    doc = msg.document
    if not doc:
        await msg.answer("❌ Không tìm thấy file đính kèm!")
        return

    file_name = doc.file_name or "list.txt"
    is_xlsx = file_name.lower().endswith(".xlsx")
    if not (file_name.endswith(".txt") or file_name.endswith(".csv") or file_name.endswith(".log") or is_xlsx):
        await msg.answer("❌ Hệ thống chỉ hỗ trợ file .txt, .csv, .log, .xlsx. Vui lòng gửi lại!")
        return

    wait = await msg.answer(f"⏳ Đang đọc file <b>{file_name}</b>...", parse_mode="HTML")

    try:
        file_info = await msg.bot.get_file(doc.file_id)
        file_bytes = await msg.bot.download_file(file_info.file_path)

        if is_xlsx:
            await _process_xlsx_check(msg, wait, file_name, file_bytes.getvalue())
            return

        raw_content = file_bytes.getvalue().decode("utf-8", errors="ignore")

        from .fb import extract_uid
        lines = [l.strip() for l in raw_content.splitlines() if l.strip()]
        uids = []
        for line in lines:
            uid = extract_uid(line)
            if not uid or not uid.isalnum():
                m = re.search(r'\d{6,}', line)
                if m: uid = m.group(0)
            if uid and uid not in uids:
                uids.append(uid)

        if not uids:
            await wait.edit_text("❌ Không tìm thấy UID hợp lệ nào trong file! Vui lòng kiểm tra lại nội dung file.")
            return

        # Trừ credits cho check bulk (user VIP còn hạn được miễn; mặc định cost=0 = miễn phí)
        if not await _ensure_bulk_credits(msg, wait, len(uids)):
            return

        if len(uids) > 1000:
            await wait.edit_text("⚠️ Số lượng UID vượt quá giới hạn (Tối đa 1,000 UID/lần). Hệ thống sẽ lấy 1,000 UID đầu tiên.")
            uids = uids[:1000]

        await wait.edit_text(f"⚡ Đang quét siêu tốc (Async Batching) <b>{len(uids)}</b> UID từ file <b>{file_name}</b>...", parse_mode="HTML")

        live_list, die_list, error_list = await check_uids_batch(uids, concurrency=15, user_id=msg.from_user.id)

        check_cache_set("file", msg.chat.id, msg.from_user.id, {
            "file_name": file_name,
            "uids": uids,
            "live": live_list,
            "die": die_list,
            "error": error_list,
            "time": time.time()
        })

        text = (
            f"📊 <b>KẾT QUẢ CHECK FILE UID FACEBOOK</b>\n"
            f"━━━━━━━━━━━━━━━━━━━━\n"
            f"📂 File: <b>{file_name}</b>\n"
            f"🔢 Tổng UID: <b>{len(uids)}</b>\n"
            f"🟢 LIVE: <b>{len(live_list)}</b> tài khoản\n"
            f"🔴 DIE: <b>{len(die_list)}</b> tài khoản\n"
        )
        if error_list:
            text += f"⚠️ Lỗi: <b>{len(error_list)}</b> tài khoản\n"
        text += "\n"

        if live_list:
            text += "🟢 <b>LIVE (Xem trước tối đa 15 UID):</b>\n" + "\n".join(f"• <code>{u}</code>" for u in live_list[:15])
            if len(live_list) > 15:
                text += f"\n<i>...và {len(live_list)-15} UID Live khác</i>"
            text += "\n\n"

        if die_list:
            text += "🔴 <b>DIE (Xem trước tối đa 10 UID):</b>\n" + "\n".join(f"• <code>{u}</code>" for u in die_list[:10])
            if len(die_list) > 10:
                text += f"\n<i>...và {len(die_list)-10} UID Die khác</i>"

        buttons = [
            [InlineKeyboardButton(text=f"📁 Thêm tất cả vào Danh Sách ({len(uids)})", callback_data="fc_addlist")],
            [InlineKeyboardButton(text=f"🔔 Theo Dõi Tự Động UID Live ({len(live_list)})", callback_data="fc_tracklive")],
            [InlineKeyboardButton(text="📥 Tải File CSV Kết Quả", callback_data="fc_exportcsv")],
            [InlineKeyboardButton(text="📊 Xuất File Excel (.xlsx)", callback_data="fc_exportxlsx")],
        ]
        # Upsell: die nhiều (>=5 và >=20%) -> gợi ý mua acc thay thế 1 chạm
        if len(die_list) >= 5 and len(die_list) * 5 >= len(uids):
            text += (f"\n💡 <b>Cần acc thay thế?</b> Bạn vừa check {len(uids)} acc, "
                     f"die tới <b>{len(die_list)}</b> — bấm 1 chạm để vào shop mua ngay!\n")
            buttons.insert(0, [InlineKeyboardButton(text="🛒 Mua acc thay thế ngay", callback_data="accshop_back")])
        kb = InlineKeyboardMarkup(inline_keyboard=buttons)

        await wait.delete()
        await msg.answer(text, parse_mode="HTML", reply_markup=kb, disable_web_page_preview=True)

    except Exception as e:
        log.error(f"Error processing file check: {e}")
        await wait.edit_text(f"❌ Lỗi xử lý file: {e}")


async def _process_xlsx_check(msg: Message, wait: Message, file_name: str, content: bytes):
    """Check UID từ file .xlsx với mọi cấu trúc cột: link FB tự resolve thành UID,
    check từng acc, rồi trả file mới giữ nguyên cấu trúc cột (ô link -> UID)."""
    from .fb import extract_uids_from_xlsx, build_xlsx_result, check_uid

    await wait.edit_text(
        f"🔗 Đang đọc file <b>{file_name}</b>, tự giải link Facebook thành UID...",
        parse_mode="HTML")

    try:
        parsed = await extract_uids_from_xlsx(content)
    except Exception as e:
        log.error(f"extract_uids_from_xlsx failed: {e}")
        await wait.edit_text(f"❌ Không đọc được file Excel: {e}")
        return

    uids = parsed["uids"]
    if not uids:
        extra = ""
        if parsed["unresolved"]:
            extra = f"\n⚠️ Có {len(parsed['unresolved'])} link không giải được thành UID."
        await wait.edit_text(
            "❌ Không tìm thấy UID hoặc link Facebook hợp lệ nào trong file!" + extra)
        return

    # Trừ credits cho check bulk (VIP còn hạn được miễn)
    if not await _ensure_bulk_credits(msg, wait, len(uids)):
        return

    if len(uids) > 1000:
        await wait.edit_text("⚠️ Số lượng UID vượt quá giới hạn (tối đa 1.000 UID/lần). Hệ thống sẽ lấy 1.000 UID đầu tiên.")
        uids = uids[:1000]

    await wait.edit_text(
        f"⚡ Đang quét siêu tốc <b>{len(uids)}</b> UID từ file <b>{file_name}</b>...",
        parse_mode="HTML")

    # Check từng UID, giữ lại tên để ghi vào file kết quả
    sem = asyncio.Semaphore(15)
    results: dict = {}

    async def _one(uid: str):
        async with sem:
            try:
                res = await check_uid(uid)
                st = res.get("status", "error")
                try:
                    db.log_check_stat(msg.from_user.id, "fb", uid, st, res.get("via", ""))
                except Exception:
                    pass
                results[uid] = {"status": st, "name": res.get("name") or ""}
            except Exception:
                results[uid] = {"status": "error", "name": ""}

    await asyncio.gather(*[_one(u) for u in uids])

    live_list = [u for u in uids if results[u]["status"] == "live"]
    die_list = [u for u in uids if results[u]["status"] == "die"]
    error_list = [u for u in uids if u not in live_list and u not in die_list]
    if live_list or die_list:
        try:
            db.add_check_history(
                msg.from_user.id, "fb",
                f"Check file xlsx: {len(live_list)} Live, {len(die_list)} Die",
                "live" if live_list else "die")
        except Exception:
            pass

    try:
        out_bytes = build_xlsx_result(parsed, results)
    except Exception as e:
        log.error(f"build_xlsx_result failed: {e}")
        out_bytes = None

    check_cache_set("file", msg.chat.id, msg.from_user.id, {
        "file_name": file_name,
        "uids": uids,
        "live": live_list,
        "die": die_list,
        "error": error_list,
        "xlsx_out": out_bytes,
        "time": time.time()
    })

    text = (
        f"📊 <b>KẾT QUẢ CHECK FILE EXCEL</b>\n"
        f"━━━━━━━━━━━━━━━━━━━━\n"
        f"📂 File: <b>{file_name}</b>\n"
        f"🔗 Link đã giải thành UID: <b>{len(parsed['link_cells'])}</b>\n"
        f"🔢 Tổng UID: <b>{len(uids)}</b>\n"
        f"🟢 LIVE: <b>{len(live_list)}</b> tài khoản\n"
        f"🔴 DIE: <b>{len(die_list)}</b> tài khoản\n"
    )
    if error_list:
        text += f"⚠️ Lỗi: <b>{len(error_list)}</b> tài khoản\n"
    if parsed["unresolved"]:
        text += f"❓ Link không giải được: <b>{len(parsed['unresolved'])}</b>\n"
        show_bad = parsed["unresolved"][:10]
        for lk in show_bad:
            short = lk if len(lk) <= 60 else lk[:57] + "..."
            text += f"• <code>{short}</code>\n"
        if len(parsed["unresolved"]) > 10:
            text += f"<i>...và {len(parsed['unresolved']) - 10} link khác (xem cột Trạng thái trong file Excel)</i>\n"
    text += "\n"

    if live_list:
        text += "🟢 <b>LIVE (xem trước tối đa 15):</b>\n" + "\n".join(f"• <code>{u}</code>" for u in live_list[:15])
        if len(live_list) > 15:
            text += f"\n<i>...và {len(live_list)-15} UID Live khác</i>"
        text += "\n\n"

    if die_list:
        text += "🔴 <b>DIE (xem trước tối đa 10):</b>\n" + "\n".join(f"• <code>{u}</code>" for u in die_list[:10])
        if len(die_list) > 10:
            text += f"\n<i>...và {len(die_list)-10} UID Die khác</i>"

    buttons = [
        [InlineKeyboardButton(text=f"📁 Thêm tất cả vào Danh Sách ({len(uids)})", callback_data="fc_addlist")],
        [InlineKeyboardButton(text=f"🔔 Theo Dõi Tự Động UID Live ({len(live_list)})", callback_data="fc_tracklive")],
    ]
    if out_bytes:
        buttons.append([InlineKeyboardButton(text="📥 Tải file Excel", callback_data="fc_export_uidxlsx")])
    # Upsell: die nhiều (>=5 và >=20%) -> gợi ý mua acc thay thế 1 chạm
    if len(die_list) >= 5 and len(die_list) * 5 >= len(uids):
        text += (f"\n💡 <b>Cần acc thay thế?</b> Bạn vừa check {len(uids)} acc, "
                 f"die tới <b>{len(die_list)}</b> — bấm 1 chạm để vào shop mua ngay!\n")
        buttons.insert(0, [InlineKeyboardButton(text="🛒 Mua acc thay thế ngay", callback_data="accshop_back")])
    kb = InlineKeyboardMarkup(inline_keyboard=buttons)

    await wait.delete()
    await msg.answer(text, parse_mode="HTML", reply_markup=kb, disable_web_page_preview=True)


@router.callback_query(F.data == "fc_export_uidxlsx")
async def on_fc_export_uidxlsx(cb: CallbackQuery):
    """Gửi file Excel kết quả với cấu trúc cột cố định:
    uid | mk | tình trạng | gmail | mail thay | 2fa | Ghi chú | Cookie | Token."""
    cache = check_cache_get("file", cb.message.chat.id, cb.from_user.id)
    if not cache or not cache.get("xlsx_out"):
        await cb.answer("❌ Đã hết phiên lưu trữ, vui lòng gửi lại file!", show_alert=True)
        return
    raw_name = (cache.get("file_name") or "ketqua.xlsx").rsplit(".", 1)[0]
    file = BufferedInputFile(cache["xlsx_out"], filename=f"ketqua_{raw_name}.xlsx")
    await cb.message.answer_document(
        document=file,
        caption="📥 File Excel kết quả: uid | mk | tình trạng | gmail | mail thay | 2fa | Ghi chú | Cookie | Token.")
    await cb.answer()


@router.callback_query(F.data == "fc_addlist")
async def on_fc_addlist(cb: CallbackQuery):
    cache = check_cache_get("file", cb.message.chat.id, cb.from_user.id)
    if not cache:
        await cb.answer("❌ Đã hết phiên lưu trữ, vui lòng gửi lại file!", show_alert=True)
        return

    uids = cache["uids"]
    list_name = f"File_{time.strftime('%d%m_%H%M')}"
    db.create_user_list(cb.from_user.id, list_name)
    for uid in uids:
        db.add_to_user_list(cb.from_user.id, list_name, uid)

    await cb.answer("✅ Đã lưu danh sách!", show_alert=False)
    await cb.message.answer(
        f"✅ <b>ĐÃ THÊM VÀO DANH SÁCH!</b>\n\n"
        f"📁 Đã lưu <b>{len(uids)}</b> UID vào danh sách: <b>{list_name}</b>\n"
        f"👉 Gõ <code>/scanlist {list_name}</code> để kiểm tra lại bất cứ lúc nào.",
        parse_mode="HTML"
    )


@router.callback_query(F.data == "fc_tracklive")
async def on_fc_tracklive(cb: CallbackQuery):
    cache = check_cache_get("file", cb.message.chat.id, cb.from_user.id)
    if not cache:
        await cb.answer("❌ Đã hết phiên lưu trữ, vui lòng gửi lại file!", show_alert=True)
        return

    live_uids = cache["live"]
    if not live_uids:
        await cb.answer("❌ Không có UID Live nào để theo dõi!", show_alert=True)
        return

    added = 0
    for uid in live_uids:
        try:
            db.create_alert_rule(str(cb.from_user.id), "fb", uid, "status_change")
            added += 1
        except Exception:
            pass

    await cb.answer(f"🔔 Đã bật theo dõi cho {added} UID Live!", show_alert=True)
    await cb.message.answer(
        f"🔔 <b>ĐÃ BẬT THEO DÕI TỰ ĐỘNG!</b>\n\n"
        f"Đã thêm <b>{added}</b> UID Live vào danh sách Cảnh Báo Tự Động.\n"
        f"Bot sẽ thông báo ngay lập tức nếu có UID bị Die trong tương lai!",
        parse_mode="HTML"
    )


def _build_excel_bytes(sheets: dict) -> bytes:
    """Tạo file .xlsx từ dict {tên_sheet: (headers, rows)}. Trả về bytes."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    first = True
    for sheet_name, (headers, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet(title=sheet_name[:31])
        if first:
            ws.title = sheet_name[:31]
            first = False
        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill("solid", fgColor="4472C4")
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(horizontal="center")
        for r, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 28
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.callback_query(F.data == "fc_exportxlsx")
async def on_fc_exportxlsx(cb: CallbackQuery):
    cache = check_cache_get("file", cb.message.chat.id, cb.from_user.id)
    if not cache:
        await cb.answer("❌ Đã hết phiên lưu trữ!", show_alert=True)
        return
    try:
        data = _build_excel_bytes({
            "Ket qua": (
                ["UID", "Trạng thái"],
                [[u, "LIVE"] for u in cache["live"]]
                + [[u, "DIE"] for u in cache["die"]]
                + [[u, "ERROR"] for u in cache["error"]],
            )
        })
    except Exception as e:
        await cb.answer(f"❌ Lỗi tạo Excel: {e}", show_alert=True)
        return
    from aiogram.types import BufferedInputFile
    raw_name = cache['file_name'].rsplit('.', 1)[0]
    file = BufferedInputFile(data, filename=f"result_{raw_name}.xlsx")
    await cb.answer()
    await cb.message.answer_document(document=file, caption="📊 File Excel tổng hợp kết quả check UID.")


@router.callback_query(F.data == "fc_exportcsv")
async def on_fc_exportcsv(cb: CallbackQuery):
    cache = check_cache_get("file", cb.message.chat.id, cb.from_user.id)
    if not cache:
        await cb.answer("❌ Đã hết phiên lưu trữ!", show_alert=True)
        return

    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["UID", "Status"])
    for u in cache["live"]: writer.writerow([u, "LIVE"])
    for u in cache["die"]: writer.writerow([u, "DIE"])
    for u in cache["error"]: writer.writerow([u, "ERROR"])
    output.seek(0)

    from aiogram.types import BufferedInputFile
    raw_name = cache['file_name'].rsplit('.', 1)[0]
    out_name = f"result_{raw_name}.csv"
    file = BufferedInputFile(output.getvalue().encode('utf-8-sig'), filename=out_name)
    await cb.answer()
    await cb.message.answer_document(document=file, caption="📥 File CSV tổng hợp kết quả check UID.")


# ─── 2. CHECK FORMAT UID|PASS|COOKIE|2FA ─────────────────────────────────────

@router.message(Command("checkcookie"))
async def on_checkcookie_cmd(msg: Message, state: FSMContext):
    """Check dàn Cookie format UID|PASS|COOKIE|2FA."""
    parts = msg.text.split(maxsplit=1)
    if len(parts) > 1 and parts[1].strip():
        await _process_cookie_text(msg, parts[1].strip())
        return

    if msg.document:
        await _process_cookie_file(msg, state)
        return

    await state.set_state(CookieCheckState.waiting_for_file)
    await msg.answer(
        "🍪 <b>CHECK COOKIE FACEBOOK (UID|PASS|COOKIE|2FA)</b>\n"
        "━━━━━━━━━━━━━━━━━━━━\n\n"
        "Vui lòng gửi file <b>.txt</b> đính kèm chứa danh sách Cookie/Nick.\n"
        "Hệ thống hỗ trợ các định dạng:\n"
        "• <code>UID|PASS|COOKIE|2FA</code>\n"
        "• <code>COOKIE</code> (Nguyên chuỗi c_user=...)\n"
        "• <code>UID|COOKIE</code>\n\n"
        "💡 <i>Gửi file đính kèm tại đây. Hoặc gõ <code>/checkcookie &lt;đoạn_text&gt;</code></i>",
        parse_mode="HTML"
    )


@router.message(CookieCheckState.waiting_for_file, F.document)
async def on_cookie_file_received(msg: Message, state: FSMContext):
    await state.clear()
    await _process_cookie_file(msg, state)


async def _process_cookie_file(msg: Message, state: FSMContext):
    doc = msg.document
    if not doc:
        await msg.answer("❌ Vui lòng gửi file đính kèm!")
        return

    wait = await msg.answer(f"⏳ Đang tải và kiểm tra file cookie <b>{doc.file_name}</b>...", parse_mode="HTML")
    try:
        file_info = await msg.bot.get_file(doc.file_id)
        file_bytes = await msg.bot.download_file(file_info.file_path)
        raw_text = file_bytes.getvalue().decode("utf-8", errors="ignore")
        await _process_cookie_text(msg, raw_text, wait_msg=wait, file_name=doc.file_name)
    except Exception as e:
        await wait.edit_text(f"❌ Lỗi đọc file: {e}")


async def _process_cookie_text(msg: Message, raw_text: str, wait_msg=None, file_name: str = "cookie.txt"):
    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
    if not lines:
        if wait_msg: await wait_msg.edit_text("❌ Nội dung trống!")
        else: await msg.answer("❌ Nội dung trống!")
        return

    # Trừ credits cho check cookie bulk (user VIP còn hạn được miễn; mặc định cost=0 = miễn phí)
    bulk_cost = int(db.get_setting("bulk_credit_cost", "0") or 0)
    raw_u2 = db.get_user(msg.from_user.id)
    if bulk_cost > 0 and not _sub_active(dict(raw_u2) if raw_u2 else None):
        need = len(lines) * bulk_cost
        have = db.get_credits(msg.from_user.id)
        if have < need:
            txt = (
                f"❌ <b>Không đủ credits!</b>\n\n"
                f"Cần: <b>{need}</b> credits ({len(lines)} dòng × {bulk_cost})\n"
                f"Bạn có: <b>{have}</b> credits\n\n"
                f"Mua thêm bằng /muacredit"
            )
            if wait_msg:
                await wait_msg.edit_text(txt, parse_mode="HTML")
            else:
                await msg.answer(txt, parse_mode="HTML")
            return
        db.consume_credits(msg.from_user.id, need)
        await _maybe_low_credit_warn(msg.bot, msg.from_user.id, msg.chat.id)

    if not wait_msg:
        wait_msg = await msg.answer(f"⚡ Đang kiểm tra siêu tốc <b>{len(lines)}</b> cookie...", parse_mode="HTML")
    else:
        await wait_msg.edit_text(f"⚡ Đang kiểm tra siêu tốc <b>{len(lines)}</b> cookie...", parse_mode="HTML")

    from .fb import _check_with_cookie, extract_uid
    sem = asyncio.Semaphore(15)

    live_items = []
    cp282_items = []
    cp956_items = []
    checkpoint_items = []
    die_items = []

    async def _check_line(line):
        async with sem:
            cookie_str = ""
            uid = ""
            parts = line.split("|")
            for p in parts:
                p_clean = p.strip()
                if "c_user=" in p_clean or "xs=" in p_clean or "fr=" in p_clean:
                    cookie_str = p_clean
                elif not uid and p_clean.isdigit() and len(p_clean) >= 6:
                    uid = p_clean

            if not cookie_str:
                cookie_str = line

            if not uid:
                uid = extract_uid(line) or "unknown"

            res = await _check_with_cookie(uid, cookie_str)
            status = res.get("status", "dead")

            if status == "live":
                live_items.append(line)
            elif status == "checkpoint_282":
                cp282_items.append(line)
            elif status == "checkpoint_956":
                cp956_items.append(line)
            elif status == "checkpoint":
                checkpoint_items.append(line)
            else:
                die_items.append(line)

    tasks = [_check_line(line) for line in lines[:500]]
    await asyncio.gather(*tasks, return_exceptions=True)

    check_cache_set("cookie", msg.chat.id, msg.from_user.id, {
        "file_name": file_name,
        "live": live_items,
        "checkpoint_282": cp282_items,
        "checkpoint_956": cp956_items,
        "checkpoint": checkpoint_items,
        "die": die_items,
        "total": len(lines)
    })

    text = (
        f"🍪 <b>KẾT QUẢ CHECK COOKIE FACEBOOK</b>\n"
        f"━━━━━━━━━━━━━━━━━━━━\n"
        f"📂 File: <b>{file_name}</b>\n"
        f"🔢 Tổng số: <b>{len(lines)}</b> cookie\n"
        f"🟢 Cookie LIVE: <b>{len(live_items)}</b>\n"
        f"🟡 Checkpoint 282 (xác minh): <b>{len(cp282_items)}</b>\n"
        f"🟠 Checkpoint 956 (vi phạm): <b>{len(cp956_items)}</b>\n"
        f"⚪ Checkpoint khác: <b>{len(checkpoint_items)}</b>\n"
        f"🔴 DIE / Expired: <b>{len(die_items)}</b>\n\n"
        "<i>Bấm nút bên dưới để xuất file TXT riêng từng loại:</i>"
    )

    buttons = [
        [
            InlineKeyboardButton(text=f"🟢 Tải Nick LIVE ({len(live_items)})", callback_data="ck_export_live"),
            InlineKeyboardButton(text=f"🔴 Tải Nick DIE ({len(die_items)})", callback_data="fc_export_die_ck"),
        ],
        [
            InlineKeyboardButton(text=f"🟡 Tải CP 282 ({len(cp282_items)})", callback_data="ck_export_282"),
            InlineKeyboardButton(text=f"🟠 Tải CP 956 ({len(cp956_items)})", callback_data="ck_export_956"),
        ],
        [
            InlineKeyboardButton(text="📊 Xuất File Excel (.xlsx)", callback_data="ck_export_xlsx"),
        ],
    ]

    await wait_msg.delete()
    await msg.answer(text, parse_mode="HTML", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))


@router.callback_query(F.data == "ck_export_live")
async def on_ck_export_live(cb: CallbackQuery):
    cache = check_cache_get("cookie", cb.message.chat.id, cb.from_user.id)
    if not cache or not cache["live"]:
        await cb.answer("❌ Không có nick LIVE nào!", show_alert=True)
        return

    content = "\n".join(cache["live"])
    from aiogram.types import BufferedInputFile
    file = BufferedInputFile(content.encode("utf-8"), filename="live_cookies.txt")
    await cb.answer()
    await cb.message.answer_document(document=file, caption="🟢 File danh sách Cookie LIVE.")


@router.callback_query(F.data == "fc_export_die_ck")
async def on_fc_export_die_ck(cb: CallbackQuery):
    cache = check_cache_get("cookie", cb.message.chat.id, cb.from_user.id)
    if not cache or not cache["die"]:
        await cb.answer("❌ Không có nick DIE nào!", show_alert=True)
        return

    content = "\n".join(cache["die"])
    from aiogram.types import BufferedInputFile
    file = BufferedInputFile(content.encode("utf-8"), filename="die_cookies.txt")
    await cb.answer()
    await cb.message.answer_document(document=file, caption="🔴 File danh sách Cookie DIE / Expired.")


# ─── 3. BÁO CÁO TỰ ĐỘNG HẰNG NGÀY ──────────────────────────────────────────

@router.message(Command("sinhnhat"))
async def on_birthday_cmd(msg: Message):
    """Nhap ngay sinh de nhan qua sinh nhat tu dong moi nam."""
    import datetime as _dt
    import re as _re
    parts = (msg.text or "").split(None, 1)
    arg = parts[1].strip() if len(parts) > 1 else ""
    u = db.get_user(msg.from_user.id) or {}
    try:
        cur = dict(u).get("dob") or ""
    except Exception:
        cur = ""
    if not arg:
        if cur:
            d = _dt.date(*map(int, cur.split("-")))
            await msg.answer(f"🎂 Ngày sinh của bạn: <b>{d.strftime('%d/%m/%Y')}</b>\n"
                             "Bot sẽ tự tặng quà vào đúng ngày sinh nhật hằng năm.\n"
                             "Muốn đổi, gửi: <code>/sinhnhat ngày/tháng/năm</code>")
        else:
            await msg.answer("🎂 Nhập ngày sinh để nhận quà sinh nhật hằng năm:\n"
                             "<code>/sinhnhat 25/12/2000</code>")
        return
    m = _re.match(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\s*$", arg)
    if not m:
        await msg.answer("❌ Sai định dạng. Ví dụ: <code>/sinhnhat 25/12/2000</code>")
        return
    try:
        d = _dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        await msg.answer("❌ Ngày không hợp lệ. Ví dụ: <code>/sinhnhat 25/12/2000</code>")
        return
    today = _dt.date.today()
    if d > today or d.year < 1920:
        await msg.answer("❌ Ngày sinh không hợp lệ.")
        return
    db.set_dob(msg.from_user.id, d.isoformat())
    note = ""
    if (d.month, d.day) == (today.month, today.day):
        note = "\n<i>Lưu ý: cần nhập trước sinh nhật ít nhất 1 ngày mới được tặng quà năm nay.</i>"
    await msg.answer(f"✅ Đã lưu ngày sinh: <b>{d.strftime('%d/%m/%Y')}</b>\n"
                     f"🎁 Bot sẽ tự tặng giftcode vào sinh nhật của bạn mỗi năm.{note}")


@router.message(Command("dailyreport"))
@router.message(Command("report"))
async def on_daily_report_cmd(msg: Message):
    """Cấu hình giờ nhận báo cáo tự động định kỳ hằng ngày."""
    parts = msg.text.split(maxsplit=1)
    raw_user = db.get_user(msg.from_user.id)
    user = dict(raw_user) if raw_user else {}
    curr_hour = user.get("daily_report_hour", -1)

    if len(parts) < 2:
        status_str = "🔴 TẮT" if curr_hour < 0 else f"🟢 BẬT (Lúc {curr_hour:02d}:00 hàng ngày)"
        text = (
            "☀️ <b>CẤU HÌNH BÁO CÁO TỰ ĐỘNG HẰNG NGÀY</b>\n"
            "━━━━━━━━━━━━━━━━━━━━\n\n"
            f"Trạng thái hiện tại: <b>{status_str}</b>\n\n"
            "<b>Cú pháp cài đặt:</b>\n"
            "• <code>/dailyreport &lt;giờ&gt;</code> — Đặt giờ báo cáo (0 đến 23)\n"
            "• <code>/dailyreport off</code> — Tắt báo cáo tự động\n\n"
            "<i>Ví dụ: <code>/dailyreport 8</code> để nhận báo cáo vào 8h00 sáng mỗi ngày.</i>"
        )
        buttons = [
            [
                InlineKeyboardButton(text="☀️ 08:00 Sáng", callback_data="set_report_8"),
                InlineKeyboardButton(text="🌙 20:00 Tối", callback_data="set_report_20"),
            ],
            [
                InlineKeyboardButton(text="🔴 Tắt Báo Cáo", callback_data="set_report_off"),
            ]
        ]
        await msg.answer(text, parse_mode="HTML", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
        return

    arg = parts[1].strip().lower()
    if arg in ("off", "tat", "tắt", "stop"):
        db.set_daily_report_hour(msg.from_user.id, -1)
        await msg.answer("🔴 Đã <b>TẮT</b> báo cáo tự động định kỳ.", parse_mode="HTML")
        return

    try:
        hour = int(arg)
        if not (0 <= hour <= 23):
            raise ValueError()
    except ValueError:
        await msg.answer("❌ Giờ không hợp lệ! Vui lòng chọn số từ 0 đến 23 (Ví dụ: /dailyreport 8)")
        return

    db.set_daily_report_hour(msg.from_user.id, hour)
    await msg.answer(f"✅ Đã bật báo cáo tự động hằng ngày vào <b>{hour:02d}:00</b>!", parse_mode="HTML")


@router.callback_query(lambda c: c.data and c.data.startswith("set_report_"))
async def on_set_report_cb(cb: CallbackQuery):
    val = cb.data.split("_")[-1]
    if val == "off":
        db.set_daily_report_hour(cb.from_user.id, -1)
        await cb.answer("Đã tắt báo cáo tự động!", show_alert=True)
        await cb.message.edit_text("🔴 Đã <b>TẮT</b> báo cáo tự động định kỳ.", parse_mode="HTML")
    else:
        try:
            hour = int(val)
            db.set_daily_report_hour(cb.from_user.id, hour)
            await cb.answer(f"Đã bật báo cáo lúc {hour:02d}:00!", show_alert=True)
            await cb.message.edit_text(f"✅ Đã bật báo cáo tự động hằng ngày vào <b>{hour:02d}:00</b>!", parse_mode="HTML")
        except Exception:
            await cb.answer("Lỗi cài đặt!", show_alert=True)


# ─── 4. /muagoi — MUA GÓI NGÀY ───────────────────────────────────────────────

@router.message(Command("muagoi"))
async def on_muagoi(msg: Message):
    """Hiển thị bảng giá và mua gói theo ngày bằng số dư."""
    parts = msg.text.split(maxsplit=1)

    price_1d = int(db.get_setting("price_1d") or 0)
    price_3d = int(db.get_setting("price_3d") or 0)
    price_5d = int(db.get_setting("price_5d") or 0)
    price_7d = int(db.get_setting("price_7d") or 0)

    raw_user = db.get_user(msg.from_user.id)
    user = dict(raw_user) if raw_user else {}
    balance = user.get("balance") or 0
    pkgs = [(d, p) for d, p in [(1, price_1d), (3, price_3d), (5, price_5d), (7, price_7d)] if p > 0]

    if len(parts) < 2:
        # Hiện bảng giá kèm nút bấm
        if not pkgs:
            await msg.answer("❌ Hệ thống chưa cấu hình giá gói ngày. Vui lòng liên hệ Admin!")
            return
        sub_until = user.get("sub_until") or 0
        sub_text = "Không có" if not sub_until else (
            "Vĩnh viễn" if sub_until > 9000000000
            else time.strftime("%d/%m/%Y %H:%M", time.localtime(sub_until))
        )
        text = (
            "📦 <b>MUA GÓI SỬ DỤNG THEO NGÀY</b>\n"
            "━━━━━━━━━━━━━━━━━━━━\n\n"
            f"💳 Số dư hiện tại: <b>{vnd(balance)}</b>\n"
            f"⏳ Hạn hiện tại: <b>{sub_text}</b>\n\n"
            "<b>📋 Các gói:</b>\n"
        )
        buttons = []
        for days, price in pkgs:
            per_day = price // days
            ok = balance >= price
            status = "✅" if ok else "❌"
            text += f"{status} <b>{days} ngày</b> — {vnd(price)} (~{vnd(per_day)}/ngày)\n"
            buttons.append([InlineKeyboardButton(
                text=f"{status} {days} ngày — {vnd(price)}",
                callback_data=f"buy_pkg_{days}_{price}"
            )])
        text += "\n💡 <i>Gói cộng dồn nếu bạn đang còn hạn. Bấm nút để mua ngay:</i>"
        await msg.answer(text, parse_mode="HTML", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
        return

    # /muagoi <số_ngày>
    try:
        days = int(parts[1].strip())
        if days not in (1, 3, 5, 7):
            raise ValueError()
    except ValueError:
        await msg.answer("❌ Số ngày hợp lệ: 1, 3, 5, 7\nVD: /muagoi 7\nGõ /muagoi để xem bảng giá.")
        return

    price_map = {1: price_1d, 3: price_3d, 5: price_5d, 7: price_7d}
    price = price_map[days]
    if not price:
        await msg.answer(f"❌ Gói {days} ngày chưa được cấu hình. Liên hệ Admin!")
        return
    await _do_buy_package(msg, msg.from_user.id, days, price)


@router.callback_query(lambda c: c.data and c.data.startswith("buy_pkg_"))
async def on_buy_pkg_cb(cb: CallbackQuery):
    """Xử lý khi bấm nút mua gói inline."""
    try:
        parts = cb.data.split("_")
        days = int(parts[2])
        price = int(parts[3])
    except (ValueError, IndexError):
        await cb.answer("Lỗi dữ liệu!", show_alert=True)
        return
    raw_user = db.get_user(cb.from_user.id)
    user = dict(raw_user) if raw_user else {}
    if not user or (user.get("balance") or 0) < price:
        await cb.answer(f"❌ Số dư không đủ! Cần {vnd(price)}", show_alert=True)
        return
    await cb.answer()
    await _do_buy_package(cb.message, cb.from_user.id, days, price, edit_msg=True)


async def _do_buy_package(target_msg, tg_id: int, days: int, price: int, edit_msg: bool = False):
    """Logic mua gói chung — trừ số dư, gia hạn sub_until."""
    raw_user = db.get_user(tg_id)
    if not raw_user:
        await target_msg.answer("❌ Không tìm thấy tài khoản!")
        return
    user = dict(raw_user)
    balance = user.get("balance") or 0
    if balance < price:
        txt = (
            f"❌ <b>Số dư không đủ!</b>\n\n"
            f"💳 Số dư: <b>{vnd(balance)}</b>\n"
            f"💰 Cần: <b>{vnd(price)}</b>\n"
            f"🔺 Thiếu: <b>{vnd(price - balance)}</b>\n\n"
            "Nạp thêm bằng lệnh /bank"
        )
        await target_msg.answer(txt, parse_mode="HTML")
        return

    # Trừ tiền
    if not db.adjust_balance(tg_id, -price, f"Mua gói {days} ngày"):
        await target_msg.answer("❌ Số dư không đủ.", parse_mode="HTML")
        return

    # Tặng credits kèm gói (tỉ lệ theo số ngày, chuẩn 30 ngày = 1 tháng)
    monthly_ref = int(db.get_setting("price_1m", "0") or 0)
    if monthly_ref <= 0 and days > 0:
        monthly_ref = int(price * 30 / days)
    bonus = int(_sub_credit_bonus(monthly_ref) * days / 30) if monthly_ref > 0 else 0
    bonus_txt = ""
    if bonus > 0:
        total_cr = db.add_credits(tg_id, bonus, f"Tặng kèm gói {days} ngày")
        bonus_txt = f"\n🎁 Tặng kèm: <b>{bonus}</b> credits (tổng: {total_cr})"

    # Gia hạn sub_until — cộng dồn nếu còn hạn
    now_ts = int(time.time())
    current_sub = user.get("sub_until") or 0
    new_sub = max(now_ts, current_sub) + days * 86400

    # Đảm bảo ít nhất VIP 1 nếu user chưa có
    if (user.get("vip_level") or 0) < 1:
        db.admin_set_vip(tg_id, 1, 0)

    # Ghi sub_until
    with db._lock:
        c = db.get_conn()
        c.execute("UPDATE tg_users SET sub_until=? WHERE tg_id=?", (new_sub, tg_id))
        try:
            c.commit()
        except Exception:
            pass

    updated_user = db.get_user(tg_id)
    new_balance = (dict(updated_user).get("balance") or 0) if updated_user else 0
    expire_str = time.strftime("%d/%m/%Y %H:%M", time.localtime(new_sub))
    result = (
        f"✅ <b>MUA GÓI THÀNH CÔNG!</b>\n\n"
        f"📦 Gói: <b>{days} ngày</b>\n"
        f"💰 Đã trừ: <b>{vnd(price)}</b>\n"
        f"💳 Số dư còn: <b>{vnd(new_balance)}</b>\n"
        f"⏳ Hạn đến: <b>{expire_str}</b>{bonus_txt}\n\n"
        "🎉 Cảm ơn bạn đã sử dụng dịch vụ!"
    )
    if edit_msg:
        try:
            await target_msg.edit_text(result, parse_mode="HTML")
        except Exception:
            await target_msg.answer(result, parse_mode="HTML")
    else:
        await target_msg.answer(result, parse_mode="HTML")


# ─── 5. CẬP NHẬT /help TEXT ──────────────────────────────────────────────────

@router.message(Command("help"))
async def on_help(msg: Message):
    """Hướng dẫn đầy đủ Checker V2 — tất cả tính năng user."""
    help_text = (
        "📖 <b>HƯỚNG DẪN CHECKER V2 PRO</b>\n"
        "━━━━━━━━━━━━━━━━━━━━\n\n"

        "<b>💰 TÀI KHOẢN &amp; SỐ DƯ</b>\n"
        "• /balance — Xem số dư hiện tại\n"
        "• /bank — Xem thông tin nạp tiền &amp; QR\n"
        "• /bank &lt;số_tiền&gt; — Nạp nhanh (VD: /bank 50000)\n"
        "• /nap &lt;số_tiền&gt; — Nạp tự động qua PayOS, chọn ví chính/shop (VD: /nap 50000)\n"
        "• /napshop &lt;số_tiền&gt; — Nạp thẳng vào ví shop mua acc\n"
        "• /ref — Lấy link giới thiệu kiếm hoa hồng\n"
        "• /doitien &lt;số_tiền&gt; — Đổi hoa hồng → số dư (+10% Bonus)\n"
        "• /ruttien &lt;tiền&gt; &lt;ngân_hàng&gt; &lt;stk&gt; — Rút hoa hồng về bank\n"
        "• /chuyentien &lt;user_id&gt; &lt;số_tiền&gt; — Chuyển số dư cho người khác\n\n"

        "<b>🛒 SHOP ACC FACEBOOK</b>\n"
        "• /shop — Xem và mua tài khoản Facebook\n"
        "• /giohang — Giỏ hàng: gom nhiều loại acc, thanh toán 1 lần\n"
        "• /damua — Acc đã mua + yêu cầu bảo hành\n"
        "• /quay — Vòng quay may mắn (mua acc được vé)\n"
        "• /hang — Hạng thành viên & ưu đãi giảm giá\n"
        "• /doiqua — Đổi điểm loyalty lấy acc miễn phí\n"
        "• /giasi — Bảng giá sỉ khi mua nhiều\n"
        "• /coc &lt;id_loại&gt; — Đặt cọc giữ hàng hot khi hết hàng\n"
        "• /coclist — Xem các khoản đặt cọc của bạn\n"
        "• /huycoc &lt;id_cọc&gt; — Hủy đặt cọc (hoàn tiền)\n"
        "• <i>Mua acc/đặt cọc/hộp mù trừ <b>ví shop</b> riêng, không dùng ví chính</i>\n\n"

        "<b>🎫 GÓI &amp; VIP</b>\n"
        "• /muagoi — Bảng giá &amp; mua gói ngày (1, 3, 5, 7 ngày)\n"
        "• /muagoi &lt;số_ngày&gt; — Mua nhanh gói ngày (VD: /muagoi 7)\n"
        "• /sub — Xem gói tháng &amp; mua gói\n"
        "• /vip — Xem cấp độ VIP và đặc quyền\n"
        "• /mycodes — Xem kho mã quà tặng\n"
        "• /code &lt;mã&gt; — Nhập mã Giftcode nhận tiền\n"
        "• /muacredit — Mua gói Credits (lượt check hàng loạt)\n"
        "• /promo &lt;mã&gt; — Áp mã giảm giá Flash Sale\n"
        "• /lichsu — Xem lại lịch sử check 7 ngày\n"
        "• /app — Mở Mini App check UID\n\n"

        "<b>🎁 TÍNH NĂNG V2 PRO</b>\n"
        "• /daily — Điểm danh nhận thưởng mỗi ngày\n"
        "• /stats — Dashboard thống kê cá nhân\n"
        "• /history — Lịch sử check 7 ngày gần nhất\n"
        "• /history &lt;platform&gt; — Lọc lịch sử (fb/tiktok/ig/zalo)\n"
        "• /top — Bảng xếp hạng nạp tiền tháng\n"
        "• /top ref — Bảng xếp hạng giới thiệu\n"
        "• /dailyreport &lt;giờ|off&gt; — Hẹn giờ nhận báo cáo dàn nick hằng ngày\n"
        "• /scan_all — Báo cáo tổng hợp toàn bộ dàn tài khoản\n"
        "• /accuracy — Thống kê độ chính xác &amp; nhất quán khi check\n"
"• /sinhnhat &lt;ngày/tháng/năm&gt; — Nhập ngày sinh, nhận quà sinh nhật hằng năm\n\n"

        "<b>📁 CHECK FILE &amp; DANH SÁCH</b>\n"
        "• /checkfile — Check hàng loạt UID FB từ file .txt/.xlsx (link FB tự giải thành UID)\n"
        "• /checkcookie — Check dàn Cookie format UID|PASS|COOKIE|2FA từ file .txt\n"
        "• /newlist &lt;tên&gt; — Tạo danh sách tài khoản mới\n"
        "• /lists — Xem tất cả danh sách đã tạo\n"
        "• /addtolist &lt;tên&gt; &lt;uid&gt; — Thêm UID vào danh sách\n"
        "• /scanlist &lt;tên&gt; — Check siêu tốc Live/Die toàn bộ danh sách\n"
        "• /deletelist &lt;tên&gt; — Xóa danh sách\n\n"

        "<b>🔔 CẢNH BÁO TỰ ĐỘNG (ALERTS)</b>\n"
        "• /alert &lt;platform&gt; &lt;target&gt; — Bật cảnh báo tự động\n"
        "• /alertlist — Xem danh sách cảnh báo\n"
        "• /alertoff &lt;id&gt; — Tắt cảnh báo\n"
        "• /alertpause &lt;id&gt; — Tạm dừng (không xóa)\n"
        "• /alertsnooze &lt;id&gt; &lt;giờ&gt; — Tắt tạm N giờ rồi tự bật lại\n\n"

        "<b>📘 FACEBOOK</b>\n"
        "• /fb &lt;uid/link&gt; — Check Live/Die nhanh\\n"
        "• /getuid &lt;link&gt; — Lấy UID từ link FB\n"
        "• /trackfb &lt;uid&gt; — Theo dõi Live/Die\n"
        "• /untrackfb &lt;uid&gt; — Huỷ theo dõi\n"
        "• /trackfblist — Danh sách FB đang theo dõi\n"
        "• /mywatches — UID đang theo dõi + chế độ báo\n"
        "• /trackmode &lt;uid&gt; &lt;all|die&gt; — Chế độ báo (VD: /trackmode 1000123 die)\n\n"

        "<b>🎵 TIKTOK</b>\n"
        "• /tiktok &lt;user&gt; — Check nhanh\n"
        "• /track &lt;user&gt; — Theo dõi follower\n"
        "• /untrack &lt;user&gt; — Huỷ theo dõi\n"
        "• /tracklist — Danh sách TikTok đang theo dõi\n"
        "• /trackv &lt;link&gt; [phút] — Theo dõi video\n"
        "• /untrackv &lt;link&gt; — Huỷ video\n"
        "• /trackvlist — Danh sách video đang theo dõi\n\n"

        "<b>📷 INSTAGRAM</b>\n"
        "• /ig &lt;user&gt; — Check nhanh\n"
        "• /trackig &lt;user&gt; — Theo dõi follower\n"
        "• /untrackig &lt;user&gt; — Huỷ theo dõi\n"
        "• /trackiglist — Danh sách IG đang theo dõi\n"
        "• /trackvig &lt;link&gt; [phút] — Theo dõi bài viết\n"
        "• /untrackvig &lt;link&gt; — Huỷ bài viết\n"
        "• /trackviglist — Danh sách bài viết IG\n\n"

        "<b>💬 ZALO</b>\n"
        "• /zalo &lt;sđt&gt; — Check SĐT Zalo Live/Die\n"
        "• /trackzalo &lt;sđt&gt; — Theo dõi SĐT Zalo\n\n"

        "<b>💻 KHÁC</b>\n"
        "• /web — Đăng nhập Bảng điều khiển Web\n"
        "• /hdcookie — Hướng dẫn lấy Cookie các nền tảng\n"
        "• /trial — Nhận dùng thử\n"
        "• /refcode — Xem mã giới thiệu của bạn\n"
        "• /report — Báo cáo dàn nick hằng ngày\n\n"

        "💬 <b>Hỗ trợ:</b>\n"
        "• Telegram: @khaikhai998\n"
        "• Facebook: facebook.com/khaitradecoin"
    )
    await msg.answer(help_text, parse_mode="HTML")


# ─── Helpers cho /tienich ────────────────────────────────────────────────────

async def _do_doitien(msg, tg_id: int):
    """Đổi hoa hồng sang số dư — tái sử dụng logic on_doitien."""
    c = db.get_conn()
    user = c.execute("SELECT ref_earnings, ref_withdrawn FROM tg_users WHERE tg_id=?", (tg_id,)).fetchone()
    if not user:
        await msg.answer("❌ Bạn chưa có tài khoản.")
        return
    d = dict(user)
    available = (d.get("ref_earnings") or 0) - (d.get("ref_withdrawn") or 0)
    if available <= 0:
        await msg.answer("❌ Bạn chưa có hoa hồng khả dụng để đổi.")
        return
    bonus = int(available * 0.10)
    total = available + bonus
    with db._lock:
        c.execute("UPDATE tg_users SET ref_withdrawn = ref_withdrawn + ?, balance = balance + ? WHERE tg_id=?",
                  (available, total, tg_id))
        c.commit()
    db.add_log("doitien", f"Đổi {vnd(available)} hoa hồng + bonus {vnd(bonus)}", tg_id)
    await msg.answer(
        f"✅ <b>ĐỔI HOA HỒNG THÀNH CÔNG!</b>\n\n"
        f"💰 Hoa hồng đổi: <b>{vnd(available)}</b>\n"
        f"🎁 Bonus +10%: <b>{vnd(bonus)}</b>\n"
        f"💵 <b>Đã cộng {vnd(total)} vào số dư.</b>",
        parse_mode="HTML")


async def _show_history(msg, tg_id: int, kind):
    rows = db.get_user_logs(tg_id, kind=kind, limit=15)
    if not rows:
        await msg.answer("📭 Bạn chưa có lịch sử check nào.")
        return
    label = {"fb": "FB", "tiktok": "TikTok", "ig": "IG", "zalo": "Zalo"}.get(kind or "", "tất cả")
    lines = [f"📜 <b>LỊCH SỬ CHECK — {label.upper()}</b>", "━━━━━━━━━━━━━━━"]
    for r in rows:
        ts = vn_time_str("%d/%m %H:%M", r["ts"])
        uid = html.escape(str(r["uid"] or ""))
        info = html.escape(str(r["message"] or ""))[:50]
        lines.append(f"• <code>{uid}</code> [{r['kind']}] {info} <i>({ts})</i>")
    await msg.answer("\n".join(lines), parse_mode="HTML")


# ─── 6. /adm — LỆNH ADMIN TÍCH HỢP VÀO BOT CHÍNH ──────────────────────────────

@router.message(Command("adm"))
async def on_adm(msg: Message, state: FSMContext):
    """Lệnh admin tập trung — chỉ admin_tg_id được cấp phép mới dùng được."""
    from .admin_bot import _handle_adm_cmd, is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    await state.clear()
    await _handle_adm_cmd(msg, bot_instance=manager.bot)


# ─── 7. LỆNH MENU CÒN THIẾU: stats / history / top / chuyentien / code ────────

@router.message(Command("stats"))
async def on_stats(msg: Message):
    """Thống kê cá nhân của user."""
    u = db.get_user(msg.from_user.id)
    if not u:
        await msg.answer("❌ Bạn chưa có tài khoản. Gõ /start để bắt đầu.")
        return
    total_checks = db.get_conn().execute(
        "SELECT COUNT(*) AS c FROM logs WHERE tg_id=?", (msg.from_user.id,)
    ).fetchone()["c"]
    text = (
        "📊 <b>THỐNG KÊ CÁ NHÂN</b>\n"
        "━━━━━━━━━━━━━━━\n"
        f"👤 <b>Tên:</b> {html.escape(u['name'] or '')}\n"
        f"🆔 <b>ID:</b> <code>{u['tg_id']}</code>\n"
        f"💰 <b>Số dư:</b> {vnd(u['balance'] or 0)}\n"
        f"⭐ <b>VIP:</b> {u['vip_level'] or 0}\n"
        f"🔍 <b>Tổng lượt check:</b> {total_checks:,}\n"
        f"💵 <b>Tổng đã nạp:</b> {vnd(u['total_topup'] or 0)}\n"
        f"🤝 <b>Hoa hồng giới thiệu:</b> {vnd(u['ref_earnings'] or 0)}\n"
    )
    await msg.answer(text, parse_mode="HTML")


@router.message(Command("history"))
async def on_history(msg: Message, command: CommandObject):
    """Lịch sử check: /history [fb|tiktok|ig|zalo]"""
    arg = (command.args or "").strip().lower()
    valid = {"fb", "tiktok", "ig", "zalo"}
    if arg and arg not in valid:
        await msg.answer("❌ Nền tảng không hợp lệ. Dùng: <code>/history [fb|tiktok|ig|zalo]</code>", parse_mode="HTML")
        return
    rows = db.get_user_logs(msg.from_user.id, kind=(arg or None), limit=15)
    if not rows:
        await msg.answer("📭 Bạn chưa có lịch sử check nào.")
        return
    lines = ["📜 <b>LỊCH SỬ CHECK GẦN ĐÂY</b>", "━━━━━━━━━━━━━━━"]
    for r in rows:
        ts = vn_time_str("%d/%m %H:%M", r["ts"])
        uid = html.escape(str(r["uid"] or ""))
        info = html.escape(str(r["message"] or ""))[:50]
        lines.append(f"• <code>{uid}</code> [{r['kind']}] {info} <i>({ts})</i>")
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("top"))
async def on_top(msg: Message, command: CommandObject):
    """Bảng xếp hạng: /top (nạp tiền tháng) hoặc /top ref (giới thiệu)."""
    arg = (command.args or "").strip().lower()
    kind = "ref" if arg == "ref" else "topup"
    rows = db.get_leaderboard(kind=kind, limit=10)
    title = "🤝 <b>TOP GIỚI THIỆU</b>" if kind == "ref" else "🏆 <b>TOP NẠP TIỀN THÁNG</b>"
    if not rows:
        await msg.answer("📭 Chưa có dữ liệu xếp hạng.")
        return
    medals = ["🥇", "🥈", "🥉"]
    lines = [title, "━━━━━━━━━━━━━━━"]
    for i, r in enumerate(rows):
        medal = medals[i] if i < 3 else f"{i + 1}."
        name = html.escape(r["name"] or r["username"] or str(r["tg_id"]))
        val = r["ref_earnings"] if kind == "ref" else r["monthly_topup"]
        lines.append(f"{medal} {name} — <b>{vnd(val or 0)}</b>")
    lines.append("\n💡 Dùng <code>/top ref</code> để xem top giới thiệu.")
    await msg.answer("\n".join(lines), parse_mode="HTML")


_pending_transfer = {}

@router.message(Command("chuyentien"))
async def on_chuyentien(msg: Message, command: CommandObject):
    """/chuyentien <user_id> <số tiền> — có bước xác nhận chống nhập nhầm."""
    args = (command.args or "").split()
    if len(args) != 2:
        await msg.answer(
            "💸 <b>CHUYỂN TIỀN</b>\n"
            "Cú pháp: <code>/chuyentien &lt;user_id&gt; &lt;số_tiền&gt;</code>\n"
            "VD: <code>/chuyentien 123456789 50000</code>",
            parse_mode="HTML",
        )
        return
    try:
        to_id = int(args[0])
        amount = int(args[1].replace(".", "").replace(",", "").replace("đ", ""))
    except ValueError:
        await msg.answer("❌ User ID và số tiền phải là số.")
        return
    if amount <= 0:
        await msg.answer("❌ Số tiền phải lớn hơn 0.")
        return
    if to_id == msg.from_user.id:
        await msg.answer("❌ Không thể chuyển tiền cho chính mình.")
        return
    recv = db.get_user(to_id)
    if not recv:
        await msg.answer("❌ Không tìm thấy người nhận (user chưa từng dùng bot).")
        return
    _pending_transfer[(msg.chat.id, msg.from_user.id)] = (to_id, amount)
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Xác nhận chuyển", callback_data="tf_yes"),
        InlineKeyboardButton(text="❌ Hủy", callback_data="tf_no"),
    ]])
    await msg.answer(
        f"💸 Xác nhận chuyển <b>{vnd(amount)}</b> cho "
        f"{html.escape(recv['name'] or '')} (<code>{to_id}</code>)?",
        reply_markup=kb, parse_mode="HTML",
    )


@router.callback_query(F.data.in_({"tf_yes", "tf_no"}))
async def on_transfer_confirm(cb: CallbackQuery):
    key = (cb.message.chat.id, cb.from_user.id)
    pending = _pending_transfer.pop(key, None)
    if cb.data == "tf_no" or not pending:
        try:
            await cb.message.edit_text("❌ Đã hủy chuyển tiền.")
        except Exception:
            pass
        await cb.answer()
        return
    to_id, amount = pending
    ok, text = db.transfer_balance(cb.from_user.id, to_id, amount)
    try:
        await cb.message.edit_text(("✅ " if ok else "❌ ") + html.escape(text), parse_mode="HTML")
    except Exception:
        pass
    await cb.answer()


@router.message(Command("code"))
async def on_code(msg: Message, command: CommandObject):
    """/code <mã giftcode> — nhập mã nhận thưởng."""
    code = (command.args or "").strip()
    if not code:
        await msg.answer("🎁 Nhập mã giftcode: <code>/code &lt;mã&gt;</code>", parse_mode="HTML")
        return
    success, amount, msg_text = db.use_code(code, msg.from_user.id)
    if success:
        db.adjust_balance(msg.from_user.id, amount, f"Sử dụng Giftcode: {code}")
        db.check_vip_upgrade(msg.from_user.id)
        await msg.answer(
            f"✅ <b>NẠP TIỀN THÀNH CÔNG!</b>\n\n"
            f"Bạn đã dùng mã <code>{html.escape(code)}</code> và được cộng <b>{vnd(amount)}</b> vào tài khoản.",
            parse_mode="HTML",
        )
    else:
        await msg.answer(f"❌ {html.escape(msg_text)}", parse_mode="HTML")


@router.callback_query(F.data == "ck_export_282")
async def on_ck_export_282(cb: CallbackQuery):
    cache = check_cache_get("cookie", cb.message.chat.id, cb.from_user.id)
    if not cache or not cache.get("checkpoint_282"):
        await cb.answer("❌ Không có nick Checkpoint 282 nào!", show_alert=True)
        return
    content = "\n".join(cache["checkpoint_282"])
    from aiogram.types import BufferedInputFile
    file = BufferedInputFile(content.encode("utf-8"), filename="checkpoint_282.txt")
    await cb.answer()
    await cb.message.answer_document(document=file, caption="🟡 File danh sách Cookie Checkpoint 282 (xác minh danh tính).")


@router.callback_query(F.data == "ck_export_956")
async def on_ck_export_956(cb: CallbackQuery):
    cache = check_cache_get("cookie", cb.message.chat.id, cb.from_user.id)
    if not cache or not cache.get("checkpoint_956"):
        await cb.answer("❌ Không có nick Checkpoint 956 nào!", show_alert=True)
        return
    content = "\n".join(cache["checkpoint_956"])
    from aiogram.types import BufferedInputFile
    file = BufferedInputFile(content.encode("utf-8"), filename="checkpoint_956.txt")
    await cb.answer()
    await cb.message.answer_document(document=file, caption="🟠 File danh sách Cookie Checkpoint 956 (vi phạm/khóa).")


# ─── 8. CREDITS (lượt check) + RESELLER KEYS ─────────────────────────────────

def _credit_packs() -> list:
    """Đọc gói credits từ setting credit_packs (JSON). Mặc định 3 gói."""
    import json
    default = [
        {"credits": 500, "price": 50000, "label": "500 lượt"},
        {"credits": 900, "price": 200000, "label": "900 lượt"},
        {"credits": 2000, "price": 350000, "label": "2000 lượt"},
    ]
    try:
        raw = db.get_setting("credit_packs", "")
        if raw:
            packs = json.loads(raw)
            if isinstance(packs, list) and packs:
                return packs
    except Exception:
        pass
    return default


def _is_admin(tg_id: int) -> bool:
    """Kiểm tra quyền admin: super admin (setting admin_tg_id / admin_tg_group_id)
    hoặc admin phụ được cấp quyền. Phân quyền chi tiết xem app/perms.py."""
    try:
        if db.get_setting("admin_tg_id") and int(db.get_setting("admin_tg_id")) == int(tg_id):
            return True
    except Exception:
        pass
    try:
        if db.get_setting("admin_tg_group_id") and int(db.get_setting("admin_tg_group_id")) == int(tg_id):
            return True
    except Exception:
        pass
    try:
        return _perms.is_admin(tg_id)
    except Exception:
        return False


def _sub_credit_bonus(price_per_month: int) -> int:
    """Credits tặng kèm khi mua 1 tháng gói.

    = tỉ lệ (setting sub_credit_ratio, mặc định 0.9) × số lượt của gói credit
    có giá cao nhất mà vẫn <= giá 1 tháng. Nếu giá tháng thấp hơn mọi gói thì
    lấy gói rẻ nhất làm chuẩn.
    """
    try:
        ratio = float(db.get_setting("sub_credit_ratio", "0.9") or 0.9)
    except Exception:
        ratio = 0.9
    if price_per_month <= 0:
        return 0
    packs = [p for p in _credit_packs()
             if int(p.get("price", 0) or 0) > 0 and int(p.get("credits", 0) or 0) > 0]
    if not packs:
        return 0
    packs.sort(key=lambda p: int(p["price"]))
    chosen = packs[0]
    for p in packs:
        if int(p["price"]) <= price_per_month:
            chosen = p
    return int(int(chosen["credits"]) * ratio)


async def _send_credit_packs(target_msg, tg_id: int, edit: bool = False):
    """Gửi bảng gói credits (dùng chung cho /muacredit và nút mua nhanh)."""
    from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton
    packs = _credit_packs()
    credits = db.get_credits(tg_id)
    up = db.get_user_promo(tg_id)
    promo_line = ""
    if up:
        ok, _, row = db.promo_valid(up["code"])
        if ok:
            promo_line = f"\n🎟️ Mã <b>{up['code']}</b> đang áp dụng: <b>giảm {int(row['pct'])}%</b> cho lần mua tới!\n"
    lines = [
        "⚡ <b>MUA GÓI CREDITS (lượt check)</b>",
        "━━━━━━━━━━━━━━━━━━━━",
        f"💠 Credits hiện có: <b>{credits}</b>",
        "",
        "Credits dùng để check file/cookie hàng loạt và gọi API reseller.",
        "Người có gói VIP đang hoạt động được check bulk miễn phí.",
        promo_line,
        "<b>Chọn gói:</b>",
    ]
    kb = []
    for i, p in enumerate(packs):
        c_num = int(p.get("credits", 0))
        price = int(p.get("price", 0))
        kb.append([InlineKeyboardButton(
            text=f"⚡ {p.get('label', f'{c_num} lượt')} — {vnd(price)}",
            callback_data=f"buycredit:{i}",
        )])
    kb.append([InlineKeyboardButton(text="🎟️ Nhập mã giảm giá", callback_data="promo_input")])
    markup = InlineKeyboardMarkup(inline_keyboard=kb)
    txt = "\n".join(lines)
    if edit:
        try:
            await target_msg.edit_text(txt, parse_mode="HTML", reply_markup=markup)
            return
        except Exception:
            pass
    await target_msg.answer(txt, parse_mode="HTML", reply_markup=markup)


@router.message(Command("muacredit"))
async def on_muacredit(msg: Message):
    """Mua gói credits (lượt check) bằng số dư."""
    await _send_credit_packs(msg, msg.from_user.id)


@router.callback_query(F.data == "promo_input")
async def on_promo_input(cb: CallbackQuery):
    await cb.message.answer(
        "🎟️ <b>NHẬP MÃ GIẢM GIÁ</b>\n\nGõ lệnh:\n<code>/promo MÃ_CODE</code>\n\nVD: <code>/promo SALE20</code>",
        parse_mode="HTML",
    )
    await cb.answer()


@router.callback_query(F.data == "buymore")
async def on_buymore(cb: CallbackQuery):
    """Nút mua nhanh khi sắp hết credits."""
    await _send_credit_packs(cb.message, cb.from_user.id, edit=True)
    await cb.answer()


async def _maybe_low_credit_warn(bot, tg_id: int, chat_id: int):
    """Sau khi trừ credits bulk: nếu còn dưới ngưỡng thì nhắc nạp + nút mua nhanh."""
    try:
        threshold = int(db.get_setting("low_credit_warn", "50") or 50)
    except Exception:
        threshold = 50
    if threshold <= 0:
        return
    try:
        left = db.get_credits(tg_id)
    except Exception:
        return
    if left >= threshold:
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⚡ Mua thêm credits", callback_data="buymore")]
    ])
    try:
        await bot.send_message(
            chat_id,
            f"⚠️ <b>Credits sắp hết!</b>\n\n"
            f"💠 Còn lại: <b>{left}</b> credits\n"
            f"<i>Nạp thêm để không bị gián đoạn nhé 👇</i>",
            parse_mode="HTML", reply_markup=kb,
        )
    except Exception:
        pass


@router.message(Command("promo"))
async def on_promo(msg: Message):
    """User áp mã giảm giá cho lần mua gói credit tiếp theo."""
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip():
        await msg.answer("🎟️ Cú pháp: <code>/promo MÃ_CODE</code>\nVD: <code>/promo SALE20</code>", parse_mode="HTML")
        return
    code = parts[1].strip().upper()
    ok, why, row = db.promo_valid(code)
    if not ok:
        await msg.answer(f"❌ {why}")
        return
    db.set_user_promo(msg.from_user.id, code)
    exp_txt = ""
    if row["expires_at"]:
        exp_txt = f"\n⏳ Hết hạn: {time.strftime('%d/%m/%Y %H:%M', time.localtime(row['expires_at']))}"
    left_txt = ""
    if row["max_uses"]:
        left_txt = f"\n🎫 Còn lại: <b>{int(row['max_uses']) - int(row['used_count'])}</b> lượt"
    await msg.answer(
        f"✅ <b>Áp mã thành công!</b>\n\n"
        f"🎟️ Mã: <b>{code}</b>\n"
        f"💸 Giảm: <b>{int(row['pct'])}%</b> cho lần mua gói credit tiếp theo{exp_txt}{left_txt}\n\n"
        f"<i>Mở /muacredit để mua ngay.</i>",
        parse_mode="HTML",
    )


@router.message(Command("setwebhook"))
async def on_setwebhook(msg: Message):
    """Admin gán webhook cho reseller key. Cú pháp: /setwebhook <tên_key|id> <url> (url trống = xóa)"""
    if not _is_admin(msg.from_user.id):
        await msg.answer("❌ Bạn không có quyền!")
        return
    parts = (msg.text or "").split(maxsplit=2)
    if len(parts) < 2:
        await msg.answer(
            "🔔 Cú pháp: <code>/setwebhook &lt;tên_key|id&gt; &lt;url&gt;</code>\n"
            "VD: <code>/setwebhook khachA https://site.com/hook</code>\n"
            "Xóa: <code>/setwebhook khachA off</code>",
            parse_mode="HTML",
        )
        return
    key = db.get_reseller_by_name_or_id(parts[1])
    if not key:
        await msg.answer(f"❌ Không tìm thấy key <b>{parts[1]}</b>.", parse_mode="HTML")
        return
    url = "" if len(parts) < 3 or parts[2].strip().lower() in ("off", "xoa", "xóa", "-") else parts[2].strip()
    if url and not (url.startswith("http://") or url.startswith("https://")):
        await msg.answer("❌ URL phải bắt đầu bằng http:// hoặc https://")
        return
    db.set_reseller_webhook(key["id"], url)
    if url:
        await msg.answer(
            f"✅ Đã gán webhook cho key <b>{key['name']}</b>:\n<code>{url}</code>\n\n"
            f"<i>Mỗi lượt /api/v1/fb/check và /api/v1/fb/getuid sẽ POST kết quả về URL này.</i>",
            parse_mode="HTML",
        )
    else:
        await msg.answer(f"✅ Đã xóa webhook của key <b>{key['name']}</b>.", parse_mode="HTML")


@router.message(Command("taopromo"))
async def on_taopromo(msg: Message):
    """Admin tạo mã giảm giá flash sale. Cú pháp: /taopromo <CODE> <phần_trăm> [số_lượt] [số_giờ]"""
    if not _is_admin(msg.from_user.id):
        await msg.answer("❌ Bạn không có quyền!")
        return
    parts = (msg.text or "").split()
    if len(parts) < 3:
        await msg.answer(
            "🎟️ Cú pháp: <code>/taopromo &lt;CODE&gt; &lt;phần_trăm&gt; [số_lượt_dùng] [số_giờ_hiệu_lực]</code>\n"
            "VD: <code>/taopromo SALE20 20 100 24</code> — giảm 20%, tối đa 100 lượt, hiệu lực 24h\n"
            "VD: <code>/taopromo VIP50 50</code> — giảm 50%, không giới hạn",
            parse_mode="HTML",
        )
        return
    try:
        pct = int(parts[2])
    except ValueError:
        await msg.answer("❌ Phần trăm phải là số (1-90).")
        return
    max_uses = int(parts[3]) if len(parts) > 3 and parts[3].isdigit() else 0
    hours = int(parts[4]) if len(parts) > 4 and parts[4].isdigit() else 0
    ok, txt = db.create_promo(parts[1], pct, max_uses, hours)
    extra = ""
    if ok:
        if max_uses:
            extra += f"\n🎫 Giới hạn: {max_uses} lượt"
        if hours:
            extra += f"\n⏳ Hiệu lực: {hours} giờ"
        extra += f"\n\n📢 Gõ <code>/flashsale {parts[1].strip().upper()}</code> để thông báo cho toàn bộ user."
    await msg.answer(("✅ " if ok else "❌ ") + txt + extra, parse_mode="HTML")


@router.message(Command("xoapromo"))
async def on_xoapromo(msg: Message):
    if not _is_admin(msg.from_user.id):
        await msg.answer("❌ Bạn không có quyền!")
        return
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2:
        await msg.answer("Cú pháp: <code>/xoapromo MÃ_CODE</code>", parse_mode="HTML")
        return
    if db.delete_promo(parts[1]):
        await msg.answer(f"✅ Đã xóa mã <b>{parts[1].strip().upper()}</b>.", parse_mode="HTML")
    else:
        await msg.answer("❌ Mã không tồn tại.")


@router.message(Command("dspromo"))
async def on_dspromo(msg: Message):
    if not _is_admin(msg.from_user.id):
        await msg.answer("❌ Bạn không có quyền!")
        return
    rows = db.list_promos()
    if not rows:
        await msg.answer("Chưa có mã giảm giá nào. Tạo bằng /taopromo")
        return
    lines = ["🎟️ <b>DANH SÁCH MÃ GIẢM GIÁ</b>", "━━━━━━━━━━━━━━━"]
    for r in rows:
        d = dict(r)
        status = "✅" if db.promo_valid(d["code"])[0] else "⛔"
        lim = f"{int(d['used_count'])}/{int(d['max_uses'])}" if d["max_uses"] else f"{int(d['used_count'])}/∞"
        lines.append(f"{status} <code>{d['code']}</code> — giảm {int(d['pct'])}% — đã dùng {lim}")
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("flashsale"))
async def on_flashsale(msg: Message):
    """Admin broadcast thông báo flash sale cho toàn bộ user."""
    if not _is_admin(msg.from_user.id):
        await msg.answer("❌ Bạn không có quyền!")
        return
    parts = (msg.text or "").split(maxsplit=1)
    if len(parts) < 2:
        await msg.answer("Cú pháp: <code>/flashsale MÃ_CODE</code>", parse_mode="HTML")
        return
    ok, why, row = db.promo_valid(parts[1])
    if not ok:
        await msg.answer(f"❌ {why}")
        return
    d = dict(row)
    exp_txt = f"\n⏰ Kết thúc: {time.strftime('%d/%m %H:%M', time.localtime(d['expires_at']))}" if d["expires_at"] else ""
    lim_txt = f"\n🎫 Chỉ {int(d['max_uses']) - int(d['used_count'])} suất" if d["max_uses"] else ""
    text = (
        f"🔥 <b>FLASH SALE — GIẢM {int(d['pct'])}% GÓI CREDITS!</b> 🔥\n"
        f"━━━━━━━━━━━━━━━━━━━━\n\n"
        f"🎟️ Mã: <code>{d['code']}</code>\n"
        f"💸 Giảm <b>{int(d['pct'])}%</b> khi mua gói credits{lim_txt}{exp_txt}\n\n"
        f"👉 Áp mã ngay: <code>/promo {d['code']}</code>\n"
        f"⚡ Mua gói: /muacredit"
    )
    users = [r["tg_id"] for r in db.get_all_users_for_broadcast()]
    sent, failed = 0, 0
    sem = asyncio.Semaphore(20)

    async def _send(uid):
        nonlocal sent, failed
        async with sem:
            try:
                await msg.bot.send_message(uid, text, parse_mode="HTML")
                sent += 1
            except Exception:
                failed += 1

    await asyncio.gather(*[_send(u) for u in users], return_exceptions=True)
    await msg.answer(f"📢 Đã gửi flash sale <b>{d['code']}</b>: <b>{sent}</b> thành công, {failed} thất bại.", parse_mode="HTML")


@router.callback_query(F.data.startswith("buycredit:"))
async def on_buycredit(cb: CallbackQuery):
    idx = int(cb.data.split(":")[1])
    packs = _credit_packs()
    if idx < 0 or idx >= len(packs):
        await cb.answer("❌ Gói không tồn tại!", show_alert=True)
        return
    p = packs[idx]
    c_num, price = int(p.get("credits", 0)), int(p.get("price", 0))
    # Áp mã giảm giá flash sale nếu user đã nhập /promo
    final_price, used_code = db.apply_user_promo(cb.from_user.id, price)
    promo_txt = f"\n🎟️ Đã áp mã <b>{used_code}</b>!" if used_code else ""
    user = db.get_user(cb.from_user.id)
    balance = (user["balance"] or 0) if user else 0
    if balance < final_price:
        await cb.answer(
            f"❌ Số dư không đủ! Cần {vnd(final_price)}, bạn có {vnd(balance)}. Nạp thêm bằng /bank",
            show_alert=True,
        )
        return
    if not db.adjust_balance(cb.from_user.id, -final_price, f"Mua gói {c_num} credits" + (f" (promo {used_code})" if used_code else "")):
        await cb.answer("❌ Số dư không đủ.", show_alert=True)
        return
    new_credits = db.add_credits(cb.from_user.id, c_num, f"Mua gói {c_num} credits ({vnd(final_price)})")
    db.add_log("credit", f"Mua {c_num} credits (-{vnd(final_price)})", cb.from_user.id)
    await cb.answer("✅ Mua thành công!", show_alert=True)
    await cb.message.answer(
        f"✅ <b>MUA CREDITS THÀNH CÔNG!</b>{promo_txt}\n\n"
        f"⚡ Nhận: <b>{c_num}</b> credits\n"
        f"💰 Đã trừ: <b>{vnd(final_price)}</b>\n"
        f"💠 Tổng credits: <b>{new_credits}</b>",
        parse_mode="HTML",
    )


@router.message(Command("taokey"))
async def on_taokey(msg: Message):
    """Admin: tạo API key cho reseller. Cú pháp: /taokey <tên shop> [số credits]"""
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/taokey &lt;tên shop&gt; [số credits]</code>\nVí dụ: <code>/taokey shopA 1000</code>")
        return
    name = parts[1]
    credits = 0
    if len(parts) > 2:
        try:
            credits = max(0, int(parts[2]))
        except ValueError:
            await msg.answer("❌ Số credits phải là số nguyên!")
            return
    k = db.create_reseller_key(name, credits)
    await msg.answer(
        "🔑 <b>TẠO API KEY THÀNH CÔNG</b>\n\n"
        f"🏪 Shop: <b>{html.escape(name)}</b>\n"
        f"🆔 Key ID: <code>{k['id']}</code>\n"
        f"⚡ Credits: <b>{k['credits']}</b>\n\n"
        f"🔐 API Key:\n<code>{k['api_key']}</code>\n\n"
        "⚠️ <i>Gửi key này cho reseller qua kênh riêng. Key chỉ hiện 1 lần!</i>\n\n"
        "📖 Tài liệu API:\n"
        "• <code>POST /api/v1/fb/check</code> — header <code>X-API-Key</code>, body <code>{\"uid\": \"...\"}</code>\n"
        "• <code>POST /api/v1/fb/getuid</code> — body <code>{\"link\": \"...\"}</code>\n"
        "• <code>GET /api/v1/balance</code> — xem credits\n"
        "• <code>GET /api/v1/usage</code> — lịch sử dùng",
        parse_mode="HTML",
    )


@router.message(Command("napkey"))
async def on_napkey(msg: Message):
    """Admin: nạp credits cho reseller key. Cú pháp: /napkey <key_id> <số credits>"""
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 3:
        await msg.answer("⚠️ Cú pháp: <code>/napkey &lt;key_id&gt; &lt;số credits&gt;</code>")
        return
    try:
        key_id, n = int(parts[1]), int(parts[2])
    except ValueError:
        await msg.answer("❌ key_id và số credits phải là số!")
        return
    db.add_reseller_credits(key_id, n)
    await msg.answer(f"✅ Đã nạp <b>{n}</b> credits cho key <code>{key_id}</code>.", parse_mode="HTML")


@router.message(Command("khoakey"))
async def on_khoakey(msg: Message):
    """Admin: khóa/mở API key reseller. Cú pháp: /khoakey <key_id> [on|off]"""
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/khoakey &lt;key_id&gt; [on|off]</code>")
        return
    try:
        key_id = int(parts[1])
    except ValueError:
        await msg.answer("❌ key_id phải là số!")
        return
    active = True
    if len(parts) > 2:
        active = parts[2].lower() not in ("off", "0", "khoa", "khóa", "lock")
    db.set_reseller_active(key_id, active)
    await msg.answer(f"{'🟢 Đã mở' if active else '🔴 Đã khóa'} key <code>{key_id}</code>.", parse_mode="HTML")


@router.callback_query(F.data == "ck_export_xlsx")
async def on_ck_export_xlsx(cb: CallbackQuery):
    cache = check_cache_get("cookie", cb.message.chat.id, cb.from_user.id)
    if not cache:
        await cb.answer("❌ Đã hết phiên lưu trữ!", show_alert=True)
        return
    try:
        sheets = {}
        def _rows(items):
            rows = []
            for i, line in enumerate(items, 1):
                parts = line.split("|")
                uid = next((p.strip() for p in parts if p.strip().isdigit() and len(p.strip()) >= 6), "")
                rows.append([i, uid or line[:40], line])
            return rows
        if cache.get("live"):
            sheets["LIVE"] = (["STT", "UID", "Dòng gốc"], _rows(cache["live"]))
        if cache.get("checkpoint_282"):
            sheets["CP 282"] = (["STT", "UID", "Dòng gốc"], _rows(cache["checkpoint_282"]))
        if cache.get("checkpoint_956"):
            sheets["CP 956"] = (["STT", "UID", "Dòng gốc"], _rows(cache["checkpoint_956"]))
        if cache.get("checkpoint"):
            sheets["CP khac"] = (["STT", "UID", "Dòng gốc"], _rows(cache["checkpoint"]))
        if cache.get("die"):
            sheets["DIE"] = (["STT", "UID", "Dòng gốc"], _rows(cache["die"]))
        if not sheets:
            await cb.answer("❌ Không có dữ liệu!", show_alert=True)
            return
        data = _build_excel_bytes(sheets)
    except Exception as e:
        await cb.answer(f"❌ Lỗi tạo Excel: {e}", show_alert=True)
        return
    from aiogram.types import BufferedInputFile
    file = BufferedInputFile(data, filename="ket_qua_check_cookie.xlsx")
    await cb.answer()
    await cb.message.answer_document(document=file, caption="📊 File Excel kết quả check cookie (chia sheet theo loại).")


@router.message(Command("app"))
async def on_app(msg: Message):
    """Mở Mini App check UID (cần cấu hình miniapp_url là HTTPS public)."""
    from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, WebAppInfo
    url = (db.get_setting("miniapp_url", "") or "").strip()
    if not url:
        await msg.answer(
            "📱 <b>MINI APP</b>\n\n"
            "Mini App chưa được cấu hình public URL.\n"
            "<i>Admin: set setting <code>miniapp_url</code> thành địa chỉ HTTPS public của backend + <code>/miniapp</code> (VD: https://bot.ban.com/miniapp) rồi thử lại.</i>",
            parse_mode="HTML",
        )
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📱 Mở Mini App Check UID", web_app=WebAppInfo(url=url))]
    ])
    await msg.answer(
        "📱 <b>MINI APP — CHECK UID NHANH</b>\n\n"
        "Bấm nút bên dưới để mở giao diện check UID, xem số dư và lịch sử ngay trong Telegram 👇",
        parse_mode="HTML", reply_markup=kb,
    )


@router.message(Command("lichsu"))
async def on_lichsu(msg: Message):
    """Xem lại lịch sử check 7 ngày gần nhất (không tốn credit khi xem lại)."""
    rows = db.get_check_history(msg.from_user.id, days=7)
    if not rows:
        await msg.answer(
            "🕘 <b>LỊCH SỬ CHECK (7 ngày)</b>\n\n"
            "Bạn chưa có lượt check nào trong 7 ngày qua.",
            parse_mode="HTML",
        )
        return
    icon = {"live": "🟢", "die": "🔴", "checkpoint_282": "🟡", "checkpoint_956": "🟠",
            "checkpoint": "⚪", "disabled": "⛔", "dead": "🔴"}
    lines = ["🕘 <b>LỊCH SỬ CHECK (7 ngày)</b>", "━━━━━━━━━━━━━━━", ""]
    for r in rows[:20]:
        d = dict(r)
        st = (d.get("result") or "?").lower()
        ic = icon.get(st, "⚪")
        t = time.strftime("%d/%m %H:%M", time.localtime(d.get("checked_at") or 0))
        target = (d.get("target") or "")[:40]
        lines.append(f"{ic} <code>{target}</code> — <b>{st.upper()}</b> <i>({t})</i>")
    if len(rows) > 20:
        lines.append(f"\n<i>...và {len(rows) - 20} lượt khác</i>")
    lines.append("\n💡 <i>Muốn check lại UID nào, cứ gửi UID/link như bình thường.</i>")
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("accuracy"))
async def on_accuracy(msg: Message):
    """Thống kê độ chính xác / nhất quán của kết quả check (90 ngày)."""
    st = db.get_accuracy_stats(msg.from_user.id)
    if not st["total"]:
        await msg.answer(
            "📊 <b>THỐNG KÊ ĐỘ CHÍNH XÁC</b>\n\n"
            "Bạn chưa có lượt check nào được ghi nhận.\n"
            "Hãy check vài UID bằng /fb hoặc /checkfile rồi quay lại nhé!",
            parse_mode="HTML",
        )
        return
    total = st["total"]
    live_pct = st["live"] * 100 / total
    die_pct = st["die"] * 100 / total
    if st["rechecked"]:
        cons_pct = st["consistent"] * 100 / st["rechecked"]
        cons_line = (
            f"🔁 UID check lại ≥2 lần: <b>{st['rechecked']}</b>\n"
            f"🎯 Nhất quán kết quả: <b>{st['consistent']}/{st['rechecked']} ({cons_pct:.1f}%)</b>\n"
        )
    else:
        cons_line = "🔁 Chưa có UID nào được check lại để đánh giá độ nhất quán.\n"
    await msg.answer(
        "📊 <b>THỐNG KÊ ĐỘ CHÍNH XÁC (90 ngày)</b>\n"
        "━━━━━━━━━━━━━━━━━━━━\n"
        f"🔢 Tổng lượt check: <b>{total}</b>\n"
        f"🆔 Số UID riêng biệt: <b>{st['unique']}</b>\n"
        f"🟢 LIVE: <b>{st['live']}</b> ({live_pct:.1f}%)\n"
        f"🔴 DIE: <b>{st['die']}</b> ({die_pct:.1f}%)\n"
        f"⚠️ Lỗi: <b>{st['error']}</b>\n\n"
        f"{cons_line}\n"
        "<i>Độ nhất quán = tỉ lệ UID cho cùng kết quả khi check lại nhiều lần.</i>",
        parse_mode="HTML",
    )


# ─── 9. QUẢN LÝ COOKIE POOL (admin) ──────────────────────────────────────────

class CookieAddState(StatesGroup):
    waiting_for_cookie = State()


def _mask_cookie(ck: str) -> str:
    ck = ck or ""
    return (ck[:10] + "..." + ck[-6:]) if len(ck) > 20 else "***"


@router.message(Command("cookieadd"))
async def on_cookieadd(msg: Message, state: FSMContext):
    """Admin: thêm cookie vào pool xoay vòng.
    Chấp nhận cả 2 cách: /cookieadd <cookie> trong 1 tin nhắn,
    hoặc gõ /cookieadd rồi gửi cookie ở tin nhắn tiếp theo."""
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    inline_ck = (msg.text or "").partition(" ")[2].strip()
    if inline_ck:
        await _save_pool_cookie(msg, state, inline_ck)
        return
    await state.set_state(CookieAddState.waiting_for_cookie)
    await msg.answer(
        "🍪 <b>THÊM COOKIE VÀO POOL</b>\n\n"
        "Hãy gửi <b>chuỗi cookie Facebook</b> (1 tin nhắn).\n"
        "Sau khi lưu, tin nhắn chứa cookie sẽ được <b>xóa ngay</b> để bảo mật.\n\n"
        "Gõ /cancel để hủy.",
        parse_mode="HTML",
    )


async def _save_pool_cookie(msg: Message, state: FSMContext, cookie: str):
    """Lưu 1 chuỗi cookie vào pool (dùng chung cho cả 2 cách nhập)."""
    if "c_user" not in cookie and "xs" not in cookie:
        await msg.answer("❌ Chuỗi này không giống cookie Facebook (thiếu c_user/xs).")
        return
    from .fb import get_fb_cookie_pool, set_fb_cookie_pool
    pool = get_fb_cookie_pool()
    if cookie in pool:
        await msg.answer("⚠️ Cookie này đã có trong pool rồi.")
    else:
        pool.append(cookie)
        set_fb_cookie_pool(pool)
        await msg.answer(f"✅ Đã thêm cookie vào pool. Pool hiện có <b>{len(pool)}</b> cookie.", parse_mode="HTML")
    try:
        await msg.delete()
    except Exception:
        pass
    await state.clear()


@router.message(CookieAddState.waiting_for_cookie)
async def on_cookieadd_received(msg: Message, state: FSMContext):
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        await state.clear()
        return
    if (msg.text or "").strip() == "/cancel":
        await state.clear()
        await msg.answer("Đã hủy.")
        return
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("\U0001f6ab \u0110\u00e3 h\u1ee7y thao t\u00e1c \u0111ang nh\u1eadp. B\u1ea1n g\u00f5 l\u1ea1i l\u1ec7nh v\u1eeba r\u1ed3i nh\u00e9.")
        return
    cookie = (msg.text or "").strip()
    # Nếu user gõ lại lệnh /cookieadd kèm cookie trong lúc đang chờ -> bỏ prefix lệnh
    if cookie.startswith("/cookieadd"):
        cookie = cookie.partition(" ")[2].strip()
    await _save_pool_cookie(msg, state, cookie)


@router.message(Command("cookielist"))
async def on_cookielist(msg: Message):
    """Admin: xem danh sách cookie trong pool (đã che)."""
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    from .fb import get_fb_cookie_pool
    pool = get_fb_cookie_pool()
    if not pool:
        await msg.answer("🍪 Pool đang trống. Thêm bằng /cookieadd")
        return
    lines = ["🍪 <b>COOKIE POOL</b>", "━━━━━━━━━━━━", ""]
    for i, ck in enumerate(pool, 1):
        lines.append(f"{i}. <code>{html.escape(_mask_cookie(ck))}</code> ({len(ck)} ký tự)")
    lines += ["", f"Tổng: <b>{len(pool)}</b> cookie", "", "Xóa: /cookiedel &lt;số thứ tự&gt;"]
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("cookiedel"))
async def on_cookiedel(msg: Message):
    """Admin: xóa cookie khỏi pool. Cú pháp: /cookiedel <số thứ tự>"""
    from .admin_bot import is_admin
    if not is_admin(msg.chat.id, msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/cookiedel &lt;số thứ tự&gt;</code> (xem số bằng /cookielist)")
        return
    try:
        idx = int(parts[1]) - 1
    except ValueError:
        await msg.answer("❌ Số thứ tự phải là số!")
        return
    from .fb import get_fb_cookie_pool, set_fb_cookie_pool
    pool = get_fb_cookie_pool()
    if idx < 0 or idx >= len(pool):
        await msg.answer("❌ Số thứ tự không đúng!")
        return
    pool.pop(idx)
    set_fb_cookie_pool(pool)
    await msg.answer(f"✅ Đã xóa. Pool còn <b>{len(pool)}</b> cookie.", parse_mode="HTML")


# ─── 10. QUẢN LÝ THEO DÕI FB (alert mode) ───────────────────────────────────

@router.message(Command("mywatches"))
async def on_mywatches(msg: Message):
    """Liệt kê các UID đang theo dõi + chế độ báo."""
    rows = db.get_user_watches(msg.from_user.id)
    if not rows:
        await msg.answer("👁 Bạn chưa theo dõi UID nào. Check 1 UID bằng /fb rồi bấm nút 👁 Theo dõi.")
        return
    lines = ["👁 <b>DANH SÁCH ĐANG THEO DÕI</b>", "━━━━━━━━━━━━", ""]
    for r in rows:
        d = dict(r)
        mode = d.get("alert_mode") or "all"
        mode_txt = "🔔 mọi thay đổi" if mode == "all" else "🔕 chỉ khi DIE"
        st = (d.get("last_status") or "?").upper()
        lines.append(f"🆔 <code>{d['uid']}</code> — <b>{st}</b> ({mode_txt})")
    lines += ["", "Đổi chế độ báo: <code>/trackmode &lt;uid&gt; &lt;all|die&gt;</code>"]
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("trackmode"))
async def on_trackmode(msg: Message):
    """Đặt chế độ báo cho 1 watch. Cú pháp: /trackmode <uid> <all|die>"""
    parts = (msg.text or "").split()
    if len(parts) < 3:
        await msg.answer(
            "⚠️ Cú pháp: <code>/trackmode &lt;uid&gt; &lt;all|die&gt;</code>\n"
            "• <b>all</b> — báo mọi lần đổi trạng thái\n"
            "• <b>die</b> — chỉ báo khi acc chuyển sang DIE",
            parse_mode="HTML",
        )
        return
    uid, mode = parts[1], parts[2].lower()
    if mode not in ("all", "die"):
        await msg.answer("❌ Chế độ chỉ có <b>all</b> hoặc <b>die</b>!")
        return
    rows = db.get_user_watches(msg.from_user.id)
    target = next((dict(r) for r in rows if str(r["uid"]) == uid), None)
    if not target:
        await msg.answer("❌ Bạn không theo dõi UID này. Xem danh sách: /mywatches")
        return
    db.set_watch_alert_mode(target["id"], msg.from_user.id, "die_only" if mode == "die" else "all")
    await msg.answer(
        f"✅ UID <code>{html.escape(uid)}</code>: chế độ báo = <b>{'chỉ khi DIE' if mode == 'die' else 'mọi thay đổi'}</b>.",
        parse_mode="HTML",
    )


# ============================ TRUNG TÂM THEO DÕI (MENU GỘP) ============================
def _trackmenu_main_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📘 Facebook UID", callback_data="trackmenu:fb"),
         InlineKeyboardButton(text="🎵 TikTok", callback_data="trackmenu:tiktok")],
        [InlineKeyboardButton(text="📸 Instagram", callback_data="trackmenu:ig"),
         InlineKeyboardButton(text="💬 Zalo", callback_data="trackmenu:zalo")],
        [InlineKeyboardButton(text="📁 Danh sách UID", callback_data="trackmenu:lists"),
         InlineKeyboardButton(text="🔔 Cảnh báo", callback_data="trackmenu:alerts")],
    ])


def _trackmenu_sub_kb(prefix: str):
    """Submenu chung: Thêm mới + Danh sách + Quay lại."""
    rows = []
    if prefix == "fb":
        rows = [
            [InlineKeyboardButton(text="➕ Thêm UID FB", callback_data="trackmenu:fb_add")],
            [InlineKeyboardButton(text="📋 Đang theo dõi", callback_data="trackmenu:fb_list")],
        ]
    elif prefix == "tiktok":
        rows = [
            [InlineKeyboardButton(text="➕ Theo dõi tài khoản", callback_data="trackmenu:tiktok_add"),
             InlineKeyboardButton(text="🎬 Theo dõi video", callback_data="trackmenu:tiktok_video_add")],
            [InlineKeyboardButton(text="📋 DS tài khoản", callback_data="trackmenu:tiktok_list"),
             InlineKeyboardButton(text="📋 DS video", callback_data="trackmenu:tiktok_vlist")],
        ]
    elif prefix == "ig":
        rows = [
            [InlineKeyboardButton(text="➕ Theo dõi tài khoản", callback_data="trackmenu:ig_add"),
             InlineKeyboardButton(text="🎞️ Theo dõi bài viết", callback_data="trackmenu:ig_post_add")],
            [InlineKeyboardButton(text="📋 DS tài khoản", callback_data="trackmenu:ig_list"),
             InlineKeyboardButton(text="📋 DS bài viết", callback_data="trackmenu:ig_plist")],
        ]
    elif prefix == "zalo":
        rows = [
            [InlineKeyboardButton(text="➕ Theo dõi SĐT", callback_data="trackmenu:zalo_add")],
            [InlineKeyboardButton(text="📋 DS Zalo", callback_data="trackmenu:zalo_list")],
        ]
    elif prefix == "lists":
        rows = [
            [InlineKeyboardButton(text="➕ Tạo danh sách", callback_data="trackmenu:lists_new")],
            [InlineKeyboardButton(text="📋 Xem danh sách", callback_data="trackmenu:lists_view")],
            [InlineKeyboardButton(text="➕ Thêm UID vào DS", callback_data="trackmenu:lists_add")],
        ]
    elif prefix == "alerts":
        rows = [
            [InlineKeyboardButton(text="➕ Thêm cảnh báo", callback_data="trackmenu:alert_add")],
            [InlineKeyboardButton(text="📋 DS cảnh báo", callback_data="trackmenu:alert_list")],
        ]
    rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:main")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _trackmenu_text_main() -> str:
    return (
        "👁️ <b>TRUNG TÂM THEO DÕI</b>\n"
        "━━━━━━━━━━━━\n\n"
        "Chọn loại muốn quản lý — khỏi cần nhớ lệnh:\n\n"
        "📘 <b>Facebook UID</b> — theo dõi LIVE/DIE\n"
        "🎵 <b>TikTok</b> — tài khoản & video\n"
        "📸 <b>Instagram</b> — tài khoản & bài viết\n"
        "💬 <b>Zalo</b> — theo dõi SĐT\n"
        "📁 <b>Danh sách UID</b> — gom UID để quét hàng loạt\n"
        "🔔 <b>Cảnh báo</b> — báo biến động tự động"
    )


# ─── TIỆN ÍCH /tienich — menu nút cho các lệnh còn gõ tay ────────────────────

def _tienich_main_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🎁 Nhập giftcode", callback_data="tienich:code"),
         InlineKeyboardButton(text="🎟️ Áp mã giảm giá", callback_data="tienich:promo")],
        [InlineKeyboardButton(text="🎂 Ngày sinh", callback_data="tienich:birthday"),
         InlineKeyboardButton(text="✏️ Đổi mã giới thiệu", callback_data="tienich:refcode")],
        [InlineKeyboardButton(text="🔗 Lấy UID từ link FB", callback_data="tienich:getuid"),
         InlineKeyboardButton(text="💸 Rút hoa hồng", callback_data="tienich:withdraw")],
        [InlineKeyboardButton(text="↔️ Chuyển tiền", callback_data="tienich:transfer"),
         InlineKeyboardButton(text="💱 Đổi hoa hồng → số dư", callback_data="tienich:doitien")],
        [InlineKeyboardButton(text="📜 Lịch sử check", callback_data="tienich:history")],
        [InlineKeyboardButton(text="💰 Đặt cọc giữ hàng", callback_data="tienich:coc"),
         InlineKeyboardButton(text="❌ Hủy đặt cọc", callback_data="tienich:huycoc")],
    ])


def _tienich_text_main() -> str:
    return (
        "🧰 <b>TIỆN ÍCH</b>\n"
        "━━━━━━━━━━━━\n\n"
        "Các thao tác hay dùng — bấm nút, khỏi gõ lệnh tay:\n\n"
        "🎁 <b>Giftcode / Mã giảm giá</b> — nhập mã nhận thưởng\n"
        "🎂 <b>Ngày sinh</b> — nhận quà sinh nhật hằng năm\n"
        "🔗 <b>Lấy UID</b> — từ link Facebook\n"
        "💸 <b>Rút / Chuyển tiền</b> — hoa hồng & số dư\n"
        "📜 <b>Lịch sử</b> — xem nhanh có nút lọc\n"
        "💰 <b>Đặt cọc</b> — giữ hàng hot, hủy cọc"
    )


def _tienich_back_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="◀️ Quay lại Tiện ích", callback_data="tienich:main")]
    ])


@router.message(Command("tienich"))
async def on_tienich(msg: Message, state: FSMContext):
    await state.clear()
    await msg.answer(_tienich_text_main(), parse_mode="HTML", reply_markup=_tienich_main_kb())


@router.callback_query(F.data.startswith("tienich:"))
async def on_tienich_cb(cb: CallbackQuery, state: FSMContext):
    data = cb.data or ""
    action = data.split(":", 1)[1] if ":" in data else "main"
    tg_id = cb.from_user.id

    try:
        if action == "main":
            await state.clear()
            await cb.message.edit_text(_tienich_text_main(), parse_mode="HTML", reply_markup=_tienich_main_kb())
            await cb.answer()
            return

        # ── Giftcode ──
        if action == "code":
            await state.set_state(TienIchState.waiting_for_giftcode)
            await cb.message.edit_text(
                "🎁 <b>NHẬP GIFTCODE</b>\n\nGửi mã giftcode của bạn ngay tin nhắn tiếp theo.\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_tienich_back_kb())
            await cb.answer()
            return

        # ── Promo ──
        if action == "promo":
            await state.set_state(TienIchState.waiting_for_promo)
            await cb.message.edit_text(
                "🎟️ <b>ÁP MÃ GIẢM GIÁ</b>\n\nGửi mã giảm giá (VD: SALE20) ngay tin nhắn tiếp theo.\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_tienich_back_kb())
            await cb.answer()
            return

        # ── Birthday ──
        if action == "birthday":
            await state.set_state(TienIchState.waiting_for_birthday)
            await cb.message.edit_text(
                "🎂 <b>NGÀY SINH NHẬN QUÀ</b>\n\nGửi ngày sinh theo định dạng <b>ngày/tháng/năm</b>.\nVD: <code>25/12/2000</code>\n\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_tienich_back_kb())
            await cb.answer()
            return

        # ── Refcode ──
        if action == "refcode":
            await state.set_state(TienIchState.waiting_for_refcode)
            await cb.message.edit_text(
                "✏️ <b>ĐỔI MÃ GIỚI THIỆU</b>\n\nGửi mã mới (chỉ chữ và số, không dấu cách).\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_tienich_back_kb())
            await cb.answer()
            return

        # ── GetUID ──
        if action == "getuid":
            await state.set_state(TienIchState.waiting_for_getuid)
            await cb.message.edit_text(
                "🔗 <b>LẤY UID TỪ LINK FB</b>\n\nGửi link Facebook ngay tin nhắn tiếp theo.\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_tienich_back_kb())
            await cb.answer()
            return

        # ── Withdraw (ruttien) bước 1: số tiền ──
        if action == "withdraw":
            await state.set_state(TienIchState.waiting_withdraw_amount)
            await cb.message.edit_text(
                "💸 <b>RÚT HOA HỒNG</b> (bước 1/3)\n\nGửi <b>số tiền</b> muốn rút (tối thiểu 50.000đ).\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_tienich_back_kb())
            await cb.answer()
            return

        # ── Transfer (chuyentien) bước 1: user_id ──
        if action == "transfer":
            await state.set_state(TienIchState.waiting_transfer_uid)
            await cb.message.edit_text(
                "↔️ <b>CHUYỂN TIỀN</b> (bước 1/2)\n\nGửi <b>User ID</b> người nhận.\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_tienich_back_kb())
            await cb.answer()
            return

        # ── Doitien: không cần nhập, chạy luôn ──
        if action == "doitien":
            await cb.answer()
            # tái sử dụng logic on_doitien bằng cách giả lập msg
            await _do_doitien(cb.message, tg_id)
            return

        # ── History: nút lọc ──
        if action == "history":
            kb = InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="📜 Tất cả", callback_data="tienich:hist_all"),
                 InlineKeyboardButton(text="📘 FB", callback_data="tienich:hist_fb")],
                [InlineKeyboardButton(text="🎵 TikTok", callback_data="tienich:hist_tiktok"),
                 InlineKeyboardButton(text="📸 IG", callback_data="tienich:hist_ig")],
                [InlineKeyboardButton(text="💬 Zalo", callback_data="tienich:hist_zalo")],
                [InlineKeyboardButton(text="◀️ Quay lại Tiện ích", callback_data="tienich:main")],
            ])
            await cb.message.edit_text("📜 <b>LỊCH SỬ CHECK</b>\n\nChọn nền tảng muốn xem:", parse_mode="HTML", reply_markup=kb)
            await cb.answer()
            return

        if action.startswith("hist_"):
            kind = action[5:]
            kind = None if kind == "all" else kind
            await cb.answer()
            await _show_history(cb.message, tg_id, kind)
            return

        # ── Coc: chọn loại acc ──
        if action == "coc":
            cats = [dict(x) for x in db.acc_category_list() if x.get("active", 1)]
            if not cats:
                await cb.message.edit_text("⛔ Hiện chưa có loại acc nào để đặt cọc.", reply_markup=_tienich_back_kb())
                await cb.answer()
                return
            rows = []
            for c in cats[:10]:
                rows.append([InlineKeyboardButton(
                    text=f"💰 {c['name']} (#{c['id']})",
                    callback_data=f"tienich:coc_go_{c['id']}")])
            rows.append([InlineKeyboardButton(text="◀️ Quay lại Tiện ích", callback_data="tienich:main")])
            await cb.message.edit_text("💰 <b>ĐẶT CỌC GIỮ HÀNG</b>\n\nChọn loại acc muốn đặt cọc:", parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))
            await cb.answer()
            return

        if action.startswith("coc_go_"):
            cat_id = int(action[7:])
            await cb.answer()
            await _do_deposit(tg_id, cat_id, cb.message, cb.message.bot)
            return

        # ── Huycoc: chọn khoản cọc ──
        if action == "huycoc":
            deps = [dict(d) for d in db.acc_deposit_list(tg_id) if d.get("status") == "WAITING"]
            if not deps:
                await cb.message.edit_text("💰 Bạn chưa có khoản đặt cọc chờ hàng nào.", reply_markup=_tienich_back_kb())
                await cb.answer()
                return
            rows = []
            for d in deps[:10]:
                rows.append([InlineKeyboardButton(
                    text=f"❌ Hủy #{d['id']} — {d.get('cat_name') or ''} ({vnd(d['amount'])})",
                    callback_data=f"tienich:huycoc_go_{d['id']}")])
            rows.append([InlineKeyboardButton(text="◀️ Quay lại Tiện ích", callback_data="tienich:main")])
            await cb.message.edit_text("❌ <b>HỦY ĐẶT CỌC</b>\n\nChọn khoản cọc muốn hủy (hoàn tiền vào ví shop):", parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))
            await cb.answer()
            return

        if action.startswith("huycoc_go_"):
            dep_id = int(action[11:])
            await cb.answer()
            amt = db.acc_deposit_cancel(dep_id, tg_id)
            if amt is None:
                await cb.message.answer("❌ Không tìm thấy khoản cọc chờ hàng này của bạn.")
                return
            db.add_shop_balance_only(tg_id, amt, f"huy_coc:{dep_id}")
            await cb.message.answer(f"✅ Đã hủy cọc <b>#{dep_id}</b>, hoàn <b>{vnd(amt)}</b> vào ví shop.", parse_mode="HTML")
            return

        await cb.answer("⏳ Đang phát triển.")
    except Exception as e:
        log.exception("tienich cb %s: %s", action, e)
        try:
            await cb.answer("❌ Có lỗi xảy ra.", show_alert=True)
        except Exception:
            pass


def _vip_limit_for(user) -> int:
    vip_level = 0
    try:
        vip_level = dict(user).get("vip_level", 0) if user else 0
    except Exception:
        pass
    try:
        return int(db.get_setting(f"vip{vip_level}_limit", [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]))
    except Exception:
        return [5, 50, 200, 1000][vip_level if vip_level <= 3 else 3]


@router.message(Command("theodoi"))
async def on_theodoi(msg: Message, state: FSMContext):
    await state.clear()
    await msg.answer(_trackmenu_text_main(), parse_mode="HTML", reply_markup=_trackmenu_main_kb())


@router.callback_query(F.data.startswith("trackmenu:"))
async def on_trackmenu_cb(cb: CallbackQuery, state: FSMContext):
    data = cb.data or ""
    parts = data.split(":")
    action = parts[1] if len(parts) > 1 else "main"
    sub = parts[2] if len(parts) > 2 else ""
    tg_id = cb.from_user.id

    try:
        # ── Main menu ──
        if action == "main":
            await state.clear()
            await cb.message.edit_text(_trackmenu_text_main(), parse_mode="HTML", reply_markup=_trackmenu_main_kb())
            await cb.answer()
            return

        # ── Submenus ──
        if action in ("fb", "tiktok", "ig", "zalo", "lists", "alerts") and not sub:
            titles = {
                "fb": "📘 <b>FACEBOOK UID</b>\nTheo dõi biến động LIVE/DIE của UID FB.\n\nChọn thao tác:",
                "tiktok": "🎵 <b>TIKTOK</b>\nTheo dõi tài khoản & video TikTok.\n\nChọn thao tác:",
                "ig": "📸 <b>INSTAGRAM</b>\nTheo dõi tài khoản & bài viết IG.\n\nChọn thao tác:",
                "zalo": "💬 <b>ZALO</b>\nTheo dõi biến động SĐT Zalo.\n\nChọn thao tác:",
                "lists": "📁 <b>DANH SÁCH UID</b>\nGom UID thành danh sách để quét hàng loạt.\n\nChọn thao tác:",
                "alerts": "🔔 <b>CẢNH BÁO</b>\nBáo tự động khi có biến động.\n\nChọn thao tác:",
            }
            await cb.message.edit_text(titles.get(action, ""), parse_mode="HTML",
                                       reply_markup=_trackmenu_sub_kb(action))
            await cb.answer()
            return

        # ── Facebook: prompt thêm ──
        if data == "trackmenu:fb_add":
            await state.set_state(TrackMenuState.waiting_fb_uid)
            await cb.message.edit_text(
                "📘 <b>THÊM THEO DÕI UID FB</b>\n\nGửi UID Facebook cần theo dõi.\nVD: <code>1000123456789</code>\n\nGõ /huy để hủy.",
                parse_mode="HTML", reply_markup=_trackmenu_back_kb("trackmenu:fb"))
            await cb.answer()
            return

        # ── Facebook: danh sách ──
        if data == "trackmenu:fb_list":
            rows = db.get_user_watches(tg_id)
            if not rows:
                await cb.message.edit_text(
                    "📭 Bạn chưa theo dõi UID nào.\n\nBấm ➕ Thêm UID FB để bắt đầu.",
                    reply_markup=_trackmenu_sub_kb("fb"))
                await cb.answer()
                return
            lines = ["📘 <b>UID ĐANG THEO DÕI</b>", "━━━━━━━━━━━━", ""]
            kb_rows = []
            for r in rows[:10]:
                d = dict(r)
                uid = str(d.get("uid") or "")
                st = (d.get("last_status") or "?").upper()
                mode = d.get("alert_mode") or "all"
                mode_icon = "🔔" if mode == "all" else "🔕"
                lines.append(f"{mode_icon} <code>{html.escape(uid)}</code> — <b>{st}</b>")
                kb_rows.append([
                    InlineKeyboardButton(text=f"{mode_icon} {uid[:12]}", callback_data=f"trackmenu:fb_toggle:{d.get('id')}"),
                    InlineKeyboardButton(text="❌", callback_data=f"trackmenu:fb_del:{uid}"),
                ])
            if len(rows) > 10:
                lines.append(f"\n<i>…và {len(rows) - 10} UID nữa — gõ /mywatches để xem hết</i>")
            lines += ["", "🔔 = báo mọi thay đổi | 🔕 = chỉ khi DIE", "Bấm 🔔/🔕 để đổi chế độ, ❌ để hủy."]
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:fb")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))
            await cb.answer()
            return

        # ── Facebook: toggle chế độ ──
        if action == "fb_toggle" and sub:
            try:
                wid = int(sub)
            except ValueError:
                await cb.answer("❌")
                return
            rows = db.get_user_watches(tg_id)
            target = next((dict(r) for r in rows if str(r["id"]) == str(wid)), None)
            if not target:
                await cb.answer("❌ Không tìm thấy.", show_alert=True)
                return
            cur = target.get("alert_mode") or "all"
            new_mode = "die_only" if cur == "all" else "all"
            db.set_watch_alert_mode(wid, tg_id, new_mode)
            await cb.answer("✅ Đã đổi thành " + ("chỉ khi DIE 🔕" if new_mode == "die_only" else "mọi thay đổi 🔔"))
            # refresh list
            cb2 = cb.model_copy(update={"data": "trackmenu:fb_list"})
            await on_trackmenu_cb(cb2, state)
            return

        # ── Facebook: xóa ──
        if action == "fb_del" and sub:
            n = db.remove_watch(tg_id, sub)
            await cb.answer("✅ Đã hủy theo dõi." if n else "❌ Không tìm thấy.")
            cb2 = cb.model_copy(update={"data": "trackmenu:fb_list"})
            await on_trackmenu_cb(cb2, state)
            return

        # ── TikTok / IG / Zalo: prompt thêm ──
        prompts = {
            "trackmenu:tiktok_add": (TrackMenuState.waiting_tiktok_username,
                "🎵 <b>THEO DÕI TIKTOK</b>\n\nGửi username TikTok (không cần @).\nVD: <code>cristiano</code>\n\nGõ /huy để hủy."),
            "trackmenu:tiktok_video_add": (TrackMenuState.waiting_tiktok_video,
                "🎬 <b>THEO DÕI VIDEO TIKTOK</b>\n\nGửi link video TikTok.\nVD: <code>https://tiktok.com/@user/video/123</code>\n\nGõ /huy để hủy."),
            "trackmenu:ig_add": (TrackMenuState.waiting_ig_username,
                "📸 <b>THEO DÕI INSTAGRAM</b>\n\nGửi username IG (không cần @).\nVD: <code>cristiano</code>\n\nGõ /huy để hủy."),
            "trackmenu:ig_post_add": (TrackMenuState.waiting_ig_post,
                "🎞️ <b>THEO DÕI BÀI VIẾT IG</b>\n\nGửi link bài viết Instagram.\nVD: <code>https://www.instagram.com/p/C123456/</code>\n\nGõ /huy để hủy."),
            "trackmenu:zalo_add": (TrackMenuState.waiting_zalo_phone,
                "💬 <b>THEO DÕI ZALO</b>\n\nGửi số điện thoại cần theo dõi.\nVD: <code>0901234567</code>\n\nGõ /huy để hủy."),
            "trackmenu:lists_new": (TrackMenuState.waiting_list_name,
                "📁 <b>TẠO DANH SÁCH MỚI</b>\n\nGửi tên danh sách.\nVD: <code>khach-vip</code>\n\nGõ /huy để hủy."),
            "trackmenu:lists_add": (TrackMenuState.waiting_list_add,
                "📁 <b>THÊM UID VÀO DANH SÁCH</b>\n\nGửi theo dạng:\n<code>tên danh sách | uid</code>\nVD: <code>khach-vip | 1000123456789</code>\n\nGõ /huy để hủy."),
            "trackmenu:alert_add": (TrackMenuState.waiting_alert_add,
                "🔔 <b>THÊM CẢNH BÁO</b>\n\nGửi theo dạng:\n<code>platform target</code>\nVD: <code>fb_watch 1000123456789</code>\n\nGõ /huy để hủy."),
        }
        if data in prompts:
            st, txt = prompts[data]
            back_map = {
                "trackmenu:tiktok_add": "trackmenu:tiktok", "trackmenu:tiktok_video_add": "trackmenu:tiktok",
                "trackmenu:ig_add": "trackmenu:ig", "trackmenu:ig_post_add": "trackmenu:ig",
                "trackmenu:zalo_add": "trackmenu:zalo", "trackmenu:lists_new": "trackmenu:lists",
                "trackmenu:lists_add": "trackmenu:lists", "trackmenu:alert_add": "trackmenu:alerts",
            }
            await state.set_state(st)
            await cb.message.edit_text(txt, parse_mode="HTML",
                                       reply_markup=_trackmenu_back_kb(back_map.get(data, "trackmenu:main")))
            await cb.answer()
            return

        # ── TikTok: DS tài khoản ──
        if data == "trackmenu:tiktok_list":
            tracks = db.user_tracks(tg_id)
            if not tracks:
                await cb.message.edit_text("📭 Chưa theo dõi tài khoản TikTok nào.",
                                           reply_markup=_trackmenu_sub_kb("tiktok"))
                await cb.answer()
                return
            lines = ["🎵 <b>TIKTOK ĐANG THEO DÕI</b>", ""]
            kb_rows = []
            for t in tracks[:10]:
                t = dict(t)
                un = t.get("tiktok_username") or ""
                fl = t.get("last_followers") or 0
                lines.append(f"• <b>@{html.escape(un)}</b> — 👥 {fmt_num(fl)}")
                kb_rows.append([InlineKeyboardButton(text=f"❌ @{un[:18]}", callback_data=f"trackmenu:tiktok_del:{un}")])
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:tiktok")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))
            await cb.answer()
            return

        # ── TikTok: DS video ──
        if data == "trackmenu:tiktok_vlist":
            vtracks = db.user_video_tracks(tg_id)
            if not vtracks:
                await cb.message.edit_text("📭 Chưa theo dõi video TikTok nào.",
                                           reply_markup=_trackmenu_sub_kb("tiktok"))
                await cb.answer()
                return
            lines = ["🎬 <b>VIDEO TIKTOK ĐANG THEO DÕI</b>", ""]
            kb_rows = []
            for v in vtracks[:10]:
                v = dict(v)
                vid = str(v.get("video_id") or "")
                un = v.get("tiktok_username") or ""
                lines.append(f"• <a href=\"{html.escape(v.get('video_url') or '')}\">@{html.escape(un)}</a>")
                kb_rows.append([InlineKeyboardButton(text=f"❌ Video {vid[:14]}", callback_data=f"trackmenu:tiktokv_del:{vid}")])
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:tiktok")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows),
                                       disable_web_page_preview=True)
            await cb.answer()
            return

        # ── TikTok: xóa ──
        if action in ("tiktok_del", "tiktokv_del") and sub:
            if action == "tiktok_del":
                ok = db.remove_track(tg_id, sub)
            else:
                ok = db.remove_video_track(tg_id, sub)
            await cb.answer("✅ Đã hủy." if ok else "❌ Không tìm thấy.")
            cb2 = cb.model_copy(update={"data": "trackmenu:tiktok_list" if action == "tiktok_del" else "trackmenu:tiktok_vlist"})
            await on_trackmenu_cb(cb2, state)
            return

        # ── IG: DS tài khoản ──
        if data == "trackmenu:ig_list":
            tracks = db.user_ig_tracks(tg_id)
            if not tracks:
                await cb.message.edit_text("📭 Chưa theo dõi tài khoản IG nào.",
                                           reply_markup=_trackmenu_sub_kb("ig"))
                await cb.answer()
                return
            lines = ["📸 <b>IG ĐANG THEO DÕI</b>", ""]
            kb_rows = []
            for t in tracks[:10]:
                t = dict(t)
                un = t.get("ig_username") or ""
                fl = t.get("last_followers") or 0
                lines.append(f"• <b>@{html.escape(un)}</b> — 👥 {fmt_num(fl)}")
                kb_rows.append([InlineKeyboardButton(text=f"❌ @{un[:18]}", callback_data=f"trackmenu:ig_del:{un}")])
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:ig")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))
            await cb.answer()
            return

        # ── IG: DS bài viết ──
        if data == "trackmenu:ig_plist":
            vtracks = db.user_ig_video_tracks(tg_id)
            if not vtracks:
                await cb.message.edit_text("📭 Chưa theo dõi bài viết IG nào.",
                                           reply_markup=_trackmenu_sub_kb("ig"))
                await cb.answer()
                return
            lines = ["🎞️ <b>BÀI VIẾT IG ĐANG THEO DÕI</b>", ""]
            kb_rows = []
            for v in vtracks[:10]:
                v = dict(v)
                pid = str(v.get("post_id") or "")
                un = v.get("ig_username") or ""
                lines.append(f"• <a href=\"{html.escape(v.get('post_url') or '')}\">@{html.escape(un)}</a>")
                kb_rows.append([InlineKeyboardButton(text=f"❌ Post {pid[:14]}", callback_data=f"trackmenu:igp_del:{pid}")])
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:ig")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows),
                                       disable_web_page_preview=True)
            await cb.answer()
            return

        # ── IG: xóa ──
        if action in ("ig_del", "igp_del") and sub:
            if action == "ig_del":
                ok = db.remove_ig_track(tg_id, sub)
            else:
                ok = db.remove_ig_video_track(tg_id, sub)
            await cb.answer("✅ Đã hủy." if ok else "❌ Không tìm thấy.")
            cb2 = cb.model_copy(update={"data": "trackmenu:ig_list" if action == "ig_del" else "trackmenu:ig_plist"})
            await on_trackmenu_cb(cb2, state)
            return

        # ── Zalo: DS ──
        if data == "trackmenu:zalo_list":
            tracks = db.user_zalo_tracks(tg_id)
            if not tracks:
                await cb.message.edit_text("📭 Chưa theo dõi SĐT Zalo nào.",
                                           reply_markup=_trackmenu_sub_kb("zalo"))
                await cb.answer()
                return
            lines = ["💬 <b>SĐT ZALO ĐANG THEO DÕI</b>", ""]
            kb_rows = []
            for t in tracks[:10]:
                t = dict(t)
                phone = str(t.get("phone") or "")
                nm = t.get("name") or ""
                st = t.get("status") or ""
                lines.append(f"• <code>{html.escape(phone)}</code> — {html.escape(nm)} ({st})")
                kb_rows.append([InlineKeyboardButton(text=f"❌ {phone}", callback_data=f"trackmenu:zalo_del:{phone}")])
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:zalo")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))
            await cb.answer()
            return

        if action == "zalo_del" and sub:
            ok = db.remove_zalo_track(tg_id, sub)
            await cb.answer("✅ Đã hủy." if ok else "❌ Không tìm thấy.")
            cb2 = cb.model_copy(update={"data": "trackmenu:zalo_list"})
            await on_trackmenu_cb(cb2, state)
            return

        # ── Lists: xem ──
        if data == "trackmenu:lists_view":
            lists = db.get_user_lists(tg_id)
            if not lists:
                await cb.message.edit_text("📭 Chưa có danh sách nào.\nBấm ➕ Tạo danh sách để bắt đầu.",
                                           reply_markup=_trackmenu_sub_kb("lists"))
                await cb.answer()
                return
            lines = ["📁 <b>DANH SÁCH CỦA BẠN</b>", ""]
            kb_rows = []
            for l in lists[:10]:
                l = dict(l)
                name = l.get("name") or ""
                cnt = l.get("item_count", 0)
                lines.append(f"• <b>{html.escape(name)}</b> — {cnt} UID")
                kb_rows.append([
                    InlineKeyboardButton(text=f"🔍 {name[:14]}", callback_data=f"trackmenu:lists_scan:{name}"),
                    InlineKeyboardButton(text="🗑", callback_data=f"trackmenu:lists_del:{name}"),
                ])
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:lists")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))
            await cb.answer()
            return

        # ── Lists: quét nhanh ──
        if action == "lists_scan" and sub:
            items = db.get_list_items(tg_id, sub)
            uids = [str(dict(i).get("value", "")) for i in items if dict(i).get("value")]
            if not uids:
                await cb.answer("Danh sách trống.", show_alert=True)
                return
            await cb.answer(f"⏳ Đang quét {len(uids)} UID...")
            # gọi logic scan dùng chung (không dùng FakeMsg nữa)
            await _run_scanlist(cb.message.chat.id, cb.from_user.id, sub,
                                answer=cb.message.answer, bot=cb.bot)
            return

        # ── Lists: xóa ──
        if action == "lists_del" and sub:
            ok = db.delete_user_list(tg_id, sub)
            await cb.answer("🗑 Đã xóa." if ok else "❌ Không tìm thấy.")
            cb2 = cb.model_copy(update={"data": "trackmenu:lists_view"})
            await on_trackmenu_cb(cb2, state)
            return

        # ── Alerts: DS ──
        if data == "trackmenu:alert_list":
            rules = db.get_alert_rules(tg_id=str(tg_id))
            if not rules:
                await cb.message.edit_text("📭 Chưa có cảnh báo nào.",
                                           reply_markup=_trackmenu_sub_kb("alerts"))
                await cb.answer()
                return
            lines = ["🔔 <b>DANH SÁCH CẢNH BÁO</b>", ""]
            kb_rows = []
            for r in rules[:10]:
                r = dict(r)
                lines.append(f"• ID {r['id']}: [{html.escape(str(r['platform']))}] {html.escape(str(r['target']))}")
                kb_rows.append([InlineKeyboardButton(text=f"❌ Xóa #{r['id']}", callback_data=f"trackmenu:alert_del:{r['id']}")])
            kb_rows.append([InlineKeyboardButton(text="◀️ Quay lại", callback_data="trackmenu:alerts")])
            await cb.message.edit_text("\n".join(lines), parse_mode="HTML",
                                       reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))
            await cb.answer()
            return

        if action == "alert_del" and sub:
            try:
                db.delete_alert_rule(int(sub))
                await cb.answer("✅ Đã xóa.")
            except Exception:
                await cb.answer("❌ Lỗi.")
            cb2 = cb.model_copy(update={"data": "trackmenu:alert_list"})
            await on_trackmenu_cb(cb2, state)
            return

        await cb.answer("❌")
    except Exception as e:
        log.warning("trackmenu cb lỗi: %s", e)
        try:
            await cb.answer("❌ Có lỗi xảy ra.")
        except Exception:
            pass


def _trackmenu_back_kb(back_to: str):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="◀️ Quay lại", callback_data=back_to)]
    ])


def _is_cancel(text: str) -> bool:
    t = (text or "").strip().lower()
    return t in ("/huy", "/cancel", "hủy", "huỷ")




# ============================ SHOP ACC FB ============================
class AccShopState(StatesGroup):
    waiting_for_stock_file = State()
    waiting_for_cover = State()
    waiting_for_review_comment = State()
    waiting_for_custom_qty = State()


def _purchase_alert_text(buyer, orders: list) -> str:
    """Tin báo admin: thông tin khách + chi tiết acc đã mua."""
    e = html.escape
    uname = (getattr(buyer, "username", "") or "").strip()
    fname = (getattr(buyer, "full_name", "") or "").strip()
    lines = ["🛒 <b>ĐƠN MUA ACC MỚI</b>", "━━━━━━━━━━━━",
             f"👤 Khách: <b>{e(fname)}</b>"
             + (f" @{e(uname)}" if uname else "")
             + f" — <code>{buyer.id}</code>"]
    for o in orders:
        lines += ["",
                  f"🆔 Mã đơn: <code>{o['id']}</code>",
                  f"📦 {e(o.get('cat_name') or '')} | 💰 {vnd(o.get('price') or 0)}",
                  f"👤 UID: <code>{e(o.get('uid') or '')}</code>",
                  f"🔑 MK: <code>{e(o.get('password') or '')}</code>"]
        if o.get("created_date"):
            lines.append(f"📅 Ngày tạo: <code>{e(o['created_date'])}</code>")
        if o.get("backup_mail"):
            lines.append(f"📧 Mail thay: <code>{e(o['backup_mail'])}</code>")
        if o.get("totp"):
            lines.append(f"🔐 2FA: <code>{e(o['totp'])}</code>")
        if o.get("note"):
            lines.append(f"📝 Ghi chú: {e(o['note'])}")
        if o.get("cookie"):
            lines.append(f"🍪 Cookie: <code>{e(o['cookie'])}</code>")
        if o.get("token"):
            lines.append(f"🎫 Token: <code>{e(o['token'])}</code>")
    txt = "\n".join(lines)
    return txt[:3900]


def _buyer_shim(tg_id: int):
    """Object buyer toi thieu (id/username/full_name) khi khong co message User."""
    username, full_name = "", ""
    try:
        u = db.get_user(tg_id)
        if u:
            username = u["username"] or ""
            try:
                full_name = u["full_name"] or ""
            except Exception:
                full_name = ""
    except Exception:
        pass
    return type("BuyerShim", (), {"id": tg_id, "username": username,
                                  "full_name": full_name})()


async def _notify_purchase_admin(bot, buyer, orders: list):
    """Báo admin khi khách mua acc: qua bot thông báo riêng nếu có,
    fallback về bot chính. Admin phụ có quyền Đơn hàng cũng nhận tin."""
    if not orders:
        return
    try:
        text = _purchase_alert_text(buyer, orders)
        sent = await _notify_bot.manager.send_to_privileged(text)
        if not sent:
            await _notify_shop_admin(bot, text)
        try:
            priv = set(_notify_bot.privileged_ids())
        except Exception:
            priv = set()
        for tid in _perms.notify_extra_ids("orders"):
            if tid in priv:
                continue
            try:
                await bot.send_message(tid, text, parse_mode="HTML")
            except Exception:
                pass
    except Exception as e:
        log.warning("Báo admin đơn mua acc lỗi: %s", e)


_STOCK_DIE_STATUSES = {"dead", "disabled", "checkpoint", "checkpoint_282", "checkpoint_956"}


async def _check_stock_live(stock_rows):
    """Check song song 1 loạt acc trong kho.
    Trả (live_ids, die_ids, unknown_ids) theo stock id.
    - live: status == "live"
    - die: dead/disabled/checkpoint... → cách ly
    - unknown: cookie_invalid/error/... → lỗi hạ tầng check, không kết luận được"""
    sem = asyncio.Semaphore(10)
    live_ids, die_ids, unknown_ids = [], [], []

    async def _one(r):
        async with sem:
            try:
                res = await check_uid(str(r["uid"]))
                st = str(res.get("status") or "").lower()
            except Exception:
                st = "error"
            return r["id"], st

    for sid, st in await asyncio.gather(*[_one(r) for r in stock_rows]):
        if st == "live":
            live_ids.append(sid)
        elif st in _STOCK_DIE_STATUSES:
            die_ids.append(sid)
        else:
            unknown_ids.append(sid)
    return live_ids, die_ids, unknown_ids


async def _notify_die_quarantine(bot, die_rows):
    """Báo admin các acc vừa bị cách ly vì check DIE trước khi giao."""
    if not die_rows:
        return
    try:
        by_cat = {}
        for r in die_rows:
            by_cat.setdefault(r.get("cat_id"), []).append(str(r.get("uid") or ""))
        lines = ["🗑 <b>CÁCH LY ACC DIE</b>",
                 "Shop vừa check trước khi giao — đã loại khỏi kho bán:"]
        for cid, uids in by_cat.items():
            try:
                c = db.acc_category_get(cid)
                name = c["name"] if c else f"#{cid}"
            except Exception:
                name = f"#{cid}"
            show = ", ".join(f"<code>{html.escape(u)}</code>" for u in uids[:10])
            more = f" <i>(+{len(uids) - 10})</i>" if len(uids) > 10 else ""
            lines.append(f"📦 <b>{html.escape(name)}</b>: {show}{more}")
        lines.append("\n<i>Dọn hẳn: /xoadie [id_loại]</i>")
        text = "\n".join(lines)
        sent = await _notify_bot.manager.send_to_privileged(text)
        if not sent:
            await _notify_shop_admin(bot, text)
    except Exception as e:
        log.warning("Báo admin cách ly acc die lỗi: %s", e)


async def _sell_live_stock(bot, tg_id: int, price_total: int, qty: int, pick_fn,
                          check_live: bool = True):
    """Bán qty acc nhưng CHỈ giao acc đã check LIVE (nếu check_live=True).
    pick_fn(exclude_ids) -> list[{"id","uid","cat_id"}] (ứng viên AVAILABLE).
    - check_live=True: Acc check DIE → cách ly khỏi kho + báo admin.
    - check_live=False (gian hàng không phải FB): giao thẳng, không check.
    - Check lỗi hạ tầng (không kết luận được) → không giao, không cách ly.
    Trả (orders, fail): orders = list[dict order] hoặc None;
    fail ∈ {"short" (thiếu hàng live), "infra" (lỗi check), "race" (bị mua mất)}."""
    need = max(1, int(qty))
    seen, live_rows = set(), []
    infra = False
    for _ in range(4):
        cands = [r for r in pick_fn(seen) if r["id"] not in seen]
        if not cands:
            break
        seen.update(r["id"] for r in cands)
        if check_live:
            l_ids, d_ids, u_ids = await _check_stock_live(cands)
            if d_ids:
                db.acc_stock_quarantine(d_ids)
                dset = set(d_ids)
                await _notify_die_quarantine(bot, [r for r in cands if r["id"] in dset])
            if u_ids:
                infra = True
            idmap = {r["id"]: r for r in cands}
            for sid in l_ids:
                if len(live_rows) < need:
                    live_rows.append(idmap[sid])
        else:
            for r in cands:
                if len(live_rows) < need:
                    live_rows.append(r)
        if len(live_rows) >= need:
            break
    if len(live_rows) < need:
        return None, ("infra" if infra else "short")
    # Bán đúng các acc đã check live (nguyên tử theo từng loại)
    orders = []
    groups = {}
    for r in live_rows[:need]:
        groups.setdefault(r["cat_id"], []).append(r["id"])
    per = price_total // need if need else price_total
    left = price_total
    items = list(groups.items())
    for gi, (cid, sids) in enumerate(items):
        share = left if gi == len(items) - 1 else per * len(sids)
        left -= share
        sold = db.acc_sell_stock_ids(cid, tg_id, share, sids)
        if not sold:
            return None, "race"
        orders += [dict(db.acc_get_order(oid)) for oid, _ in sold]
    return orders, None


async def _notify_shop_admin(bot, text: str):
    """Báo cho admin qua bot chính."""
    try:
        aid = db.get_setting("admin_tg_id", "")
        if aid:
            await bot.send_message(int(aid), text, parse_mode="HTML")
    except Exception:
        pass


async def _notify_admin_photo(bot, file_id: str, caption: str):
    """Gui anh bang chung cho admin qua bot chinh."""
    targets = []
    try:
        aid = db.get_setting("admin_tg_id", "")
        if aid:
            targets.append(int(aid))
    except Exception:
        pass
    try:
        gid = db.get_setting("admin_tg_group_id", "")
        if gid:
            targets.append(int(gid))
    except Exception:
        pass
    for t in targets:
        try:
            await bot.send_photo(t, file_id, caption=caption, parse_mode="HTML")
        except Exception:
            pass


async def _notify_admin_smart(bot, text: str, perm: str = None):
    """Báo admin: ưu tiên bot thông báo riêng, fallback về bot chính.
    perm: nếu có, gửi thêm cho các admin phụ đang giữ quyền đó
    (tránh trùng người đã nhận qua kênh privileged)."""
    try:
        sent = await _notify_bot.manager.send_to_privileged(text)
        if sent:
            pass
        else:
            await _notify_shop_admin(bot, text)
    except Exception:
        try:
            await _notify_shop_admin(bot, text)
        except Exception:
            pass
    if perm:
        try:
            priv = set(_notify_bot.privileged_ids())
        except Exception:
            priv = set()
        for tid in _perms.notify_extra_ids(perm):
            if tid in priv:
                continue
            try:
                await bot.send_message(tid, text, parse_mode="HTML")
            except Exception:
                pass


def _totp_now(secret: str, digits: int = 6, period: int = 30) -> str:
    """Sinh mã TOTP 6 số hiện tại từ 2FA secret (RFC 6238, chu kỳ 30s)."""
    import base64
    import hashlib
    import hmac
    import struct
    import time
    s = re.sub(r"[^A-Za-z2-7]", "", (secret or "").upper())
    if not s:
        return ""
    try:
        key = base64.b32decode(s + "=" * (-len(s) % 8))
    except Exception:
        return ""
    counter = int(time.time() // period)
    mac = hmac.new(key, struct.pack(">Q", counter), hashlib.sha1).digest()
    off = mac[-1] & 0x0F
    code = struct.unpack(">I", mac[off:off + 4])[0] & 0x7FFFFFFF
    return str(code % (10 ** digits)).zfill(digits)


DEFAULT_MAIL_APP_LINK = "https://www.swisstransfer.com/d/3ec3521b-a0b7-4dc9-bd78-53298af278ac"


def _mail_app_link() -> str:
    """Link tải app mail ảo (mail của acc là mail ảo). Admin đổi bằng /setmailapp."""
    return (db.get_setting("mail_app_link", DEFAULT_MAIL_APP_LINK) or "").strip()


def _mail_app_line() -> str:
    link = _mail_app_link()
    if not link:
        return ""
    return (
        "📧 <b>Mail của acc là mail ảo</b> — tải app xem mail tại đây: "
        f"<a href=\"{html.escape(link)}\">Tải app mail</a>"
    )


def _pickup_suffix() -> str:
    """Dòng chọn cách nhận acc + link app mail (nếu có)."""
    line = _mail_app_line()
    return (line + "\n" if line else "") + "👇 <b>Chọn cách nhận acc:</b>"


def _acc_delivery_caption(o) -> str:
    e = html.escape
    parts = [
        f"🎉 <b>MUA THÀNH CÔNG!</b>",
        f"{_cat_icon(o['cat_name'])} <b>{e(o['cat_name'])}</b>",
        f"🧾 Đơn hàng: <b>#{o['id']}</b>  •  💰 <b>{vnd(o['price'])}</b>",
        "━━━━━━━━━━━━━━",
        f"👤 UID: <code>{e(o['uid'] or '')}</code>",
        f"🔑 Mật khẩu: <code>{e(o['password'] or '')}</code>",
    ]
    if o["created_date"]:
        parts.append(f"📅 Ngày tạo: <code>{e(o['created_date'])}</code>")
    if o["backup_mail"]:
        parts.append(f"📧 Mail thay: <code>{e(o['backup_mail'])}</code>")
    if o["totp"]:
        _code = _totp_now(o["totp"])
        _code_txt = f" — mã hiện tại: <code>{_code}</code>" if _code else ""
        parts.append(f"🔐 2FA: <code>{e(o['totp'])}</code>{_code_txt}")
        parts.append("   <i>Mã đổi 30s/lần — lấy mã mới trong /damua.</i>")
    if o["note"]:
        parts.append(f"📝 Ghi chú: {e(o['note'])}")
    if o["cookie"]:
        parts.append(f"🍪 Cookie: <code>{e(o['cookie'])}</code>")
    if o["token"]:
        parts.append(f"🎫 Token: <code>{e(o['token'])}</code>")
    parts += [
        "",
        "⚠️ <b>Đổi mật khẩu ngay</b> sau khi đăng nhập.",
        f"🛡 Bảo hành <b>{_fmt_warranty(o['warranty_hours'])}</b> — acc die thì vào /damua bấm \"Bảo hành\".",
        "📌 <b>Chỉ bảo hành trường hợp đăng nhập báo sai mật khẩu.</b>",
    ]
    mail_line = _mail_app_line()
    if mail_line:
        parts += ["", mail_line]
    return "\n".join(parts)


def _cat_icon(name: str) -> str:
    """Emoji theo tên loại acc cho giao diện shop đẹp hơn."""
    n = (name or "").lower()
    if "bm" in n: return "🏢"
    if "page" in n: return "📄"
    if "clone" in n: return "👥"
    if "via" in n: return "🌐"
    if "xmdt" in n or "xác minh" in n: return "✅"
    if "mail" in n: return "📧"
    return "📦"


@router.message(Command("shop"))
async def on_shop(msg: Message):
    stalls = [s for s in db.acc_stall_list() if int(s["n"] or 0) > 0]
    if len(stalls) > 1:
        await _shop_stall_picker(msg)
        return
    await _shop_list(msg, None)


@router.callback_query(F.data.startswith("shopstall:"))
async def on_shop_stall(cb: CallbackQuery):
    arg = (cb.data or "").split(":", 1)[1]
    await cb.answer()
    if arg == "__all__":
        try:
            await cb.message.delete()
        except Exception:
            pass
        await _shop_stall_picker(cb.message)
        return
    await _shop_list(cb.message, arg, edit=True, user_id=cb.from_user.id)


async def _shop_stall_picker(msg: Message):
    """Màn hình chọn gian hàng (chỉ hiện khi có >1 gian hàng)."""
    stalls = [s for s in db.acc_stall_list() if int(s["n"] or 0) > 0]
    lines = ["🏪 <b>CHỌN GIAN HÀNG</b>", "━━━━━━━━━━━━━━", ""]
    kb_rows = []
    for s in stalls:
        st = s["stall"]
        total = db.acc_stock_count_stall(st)
        lines.append(f"🏪 <b>{html.escape(st)}</b> — 📦 {total} còn bán • {s['n']} loại hàng")
        kb_rows.append([InlineKeyboardButton(
            text=f"🏪 {st} ({total})", callback_data=f"shopstall:{st}")])
    kb_rows.append([InlineKeyboardButton(text="🛒 Xem giỏ hàng", callback_data="cartview")])
    await msg.answer("\n".join(lines), parse_mode="HTML",
                     reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))


async def _shop_list(msg: Message, stall=None, edit=False, user_id=None):
    cats = db.acc_category_list()
    if stall:
        cats = [c for c in cats
                if ((c["stall"] if "stall" in c.keys() else "") or "Acc Facebook") == stall]
    if not cats:
        txt = "🛒 Gian hàng này hiện chưa có mặt hàng nào. Quay lại sau nhé!"
        if edit:
            await msg.edit_text(txt)
        else:
            await msg.answer(txt)
        return
    _is_fb = (stall or "Acc Facebook") == "Acc Facebook"
    kb_rows = []
    now = int(time.time())
    _uid = user_id or (msg.from_user.id if getattr(msg, "from_user", None) else 0)
    _u = db.get_user(_uid) if _uid else None
    _shop_bal = int(_u["shop_balance"] or 0) if _u and "shop_balance" in _u.keys() else 0
    lines = [
        f"🏪 <b>{html.escape(stall)}</b>" if stall else "🛒 <b>SHOP TÀI KHOẢN</b>",
        "━━━━━━━━━━━━━━",
        f"👛 Ví shop: <b>{vnd(_shop_bal)}</b>",
        "",
    ]
    # 4.10 Giờ vàng giảm giá
    hh_on, hh_pct, hh_cat = db.happy_hour_active()
    if hh_on:
        scope = "toàn shop" if hh_cat == 0 else f"loại #{hh_cat}"
        lines += [f"⚡ <b>GIỜ VÀNG</b> — giảm <b>{hh_pct}%</b> {scope}!",
              ""]
    # gom số liệu từng loại
    infos = []
    for c in cats:
        c = dict(c)
        up = db.shop_unit_price(c)
        n = up["stock"]
        sold = db.acc_sold_count(c["id"])
        avg, rcnt = db.acc_review_avg(c["id"])
        sample = db.get_conn().execute(
            "SELECT uid FROM acc_stock WHERE cat_id=? AND status='AVAILABLE' ORDER BY id LIMIT 1",
            (c["id"],)).fetchone()
        sample_uid = (sample["uid"] or "") if sample else ""
        infos.append((c, up, n, sold, avg, rcnt, sample_uid))
    # 💡 GỢI Ý CHO BẠN: bán chạy / rẻ nhất / đánh giá cao
    in_stock = [i for i in infos if i[2] > 0]
    sugg = []
    if in_stock:
        best = max(in_stock, key=lambda i: i[3])
        cheap = min(in_stock, key=lambda i: i[1]["price"])
        rated = [i for i in in_stock if i[5] >= 2]
        top = max(rated, key=lambda i: i[4]) if rated else None
        seen = set()
        for label, item in (("🔥 Bán chạy", best), ("💰 Rẻ nhất", cheap), ("⭐ Đánh giá cao", top)):
            if item and item[0]["id"] not in seen:
                seen.add(item[0]["id"])
                sugg.append(InlineKeyboardButton(
                    text=f"{label}: {item[0]['name']}",
                    callback_data=f"accbuy:{item[0]['id']}"))
    if sugg and any(i[3] > 0 for i in in_stock):
        lines.append("💡 <b>GỢI Ý CHO BẠN</b>")
        lines.append("")
        kb_rows.append(sugg[:2])
        if len(sugg) > 2:
            kb_rows.append(sugg[2:])
    for c, up, n, sold, avg, rcnt, sample_uid in infos:
        icon = _cat_icon(c["name"])
        price_txt = vnd(up["price"])
        badges = []
        if n == 0:
            badges.append("⛔ HẾT HÀNG")
        else:
            if sold >= 10:
                badges.append("🔥 BÁN CHẠY")
            try:
                is_new = now - int(c.get("created_at") or 0) < 7 * 86400
            except Exception:
                is_new = False
            if is_new:
                badges.append("✨ MỚI")
            if n <= 5:
                badges.append("⚠️ SẮP HẾT")
        if up["happy"]:
            badges.append(f"⚡ giờ vàng −{up['happy_pct']}%")
        badge_txt = (" " + " ".join(badges)) if badges else ""
        stock_txt = f"📦 Còn <b>{n}</b>" if n else "📦 <b>Hết hàng</b>"
        sold_txt = f"  •  🔁 Đã bán <b>{sold}</b>" if sold else ""
        lines.append(f"{icon} <b>{html.escape(c['name'])}</b>{badge_txt}")
        lines.append(f"┌ 💰 <b>{price_txt}</b>/acc")
        lines.append(f"└ {stock_txt}{sold_txt}")
        if rcnt:
            stars = "⭐" * int(round(avg)) if avg else "⭐"
            lines.append(f"    {stars} <b>{avg}/5</b> <i>({rcnt} đánh giá)</i>")
        if c["description"]:
            lines.append(f"    <i>{html.escape(c['description'])}</i>")
        if sample_uid and _is_fb:
            lines.append(f'    🔗 Check thử: <a href="https://facebook.com/{html.escape(sample_uid)}">facebook.com/{html.escape(sample_uid)}</a>')
        lines.append("")
        kb_rows.append([InlineKeyboardButton(
            text=f"{icon} {c['name']} • {price_txt}",
            callback_data=f"accbuy:{c['id']}")])
    # 4.8 Hộp mù acc
    try:
        m_price = int(db.get_setting("mystery_price", "0") or 0)
    except Exception:
        m_price = 0
    if m_price > 0 and _is_fb:
        m_cats = [dict(x) for x in db.acc_category_list()
                  if int(x["mystery_eligible"] or 0) and db.acc_stock_count(x["id"]) > 0]
        if m_cats:
            lines += ["", f"🎁 <b>HỘP MÙ ACC — {vnd(m_price)}</b>: mở ngẫu nhiên 1 acc "
                         f"từ {len(m_cats)} loại, rẻ hơn giá lẻ!"]
            kb_rows.append([InlineKeyboardButton(
                text=f"🎁 Mua hộp mù — {vnd(m_price)}",
                callback_data="accmystery")])
    lines += ["👉 <i>Chạm vào từng loại để xem chi tiết & mua.</i>",
              "",
              "🎡 <i>/quay</i> vòng quay may mắn  •  🏅 <i>/hang</i> hạng thành viên  •  🎁 <i>/doiqua</i> đổi điểm",
              "🛒 <i>/giohang</i> giỏ hàng của bạn"]
    kb_rows.append([InlineKeyboardButton(text="🛒 Xem giỏ hàng", callback_data="cartview"),
                    InlineKeyboardButton(text="💰 Nạp ví shop", callback_data="wal_nap:shop")])
    if stall:
        kb_rows.append([InlineKeyboardButton(text="◀️ Gian hàng khác",
                                             callback_data="shopstall:__all")])
    _kb = InlineKeyboardMarkup(inline_keyboard=kb_rows)
    _text = "\n".join(lines)
    if edit:
        await msg.edit_text(_text, parse_mode="HTML", reply_markup=_kb,
                            disable_web_page_preview=True)
    else:
        await msg.answer(_text, parse_mode="HTML", reply_markup=_kb,
                         disable_web_page_preview=True)


@router.message(Command("giasi"))
async def on_giasi(msg: Message):
    """Bảng giá sỉ: giá mỗi acc khi mua 1 / 5 / 10 / 20."""
    cats = db.acc_category_list()
    if not cats:
        await msg.answer("💼 Shop hiện chưa có mặt hàng nào.")
        return
    p5, p10, p20 = _shop_bulk_pcts()
    lines = ["💼 <b>BẢNG GIÁ SỈ</b> <i>(giá mỗi acc)</i>", "━━━━━━━━━━━━", ""]
    for c in cats:
        c = dict(c)
        up = db.shop_unit_price(c)
        base = up["price"]
        n = up["stock"]
        if n <= 0:
            continue
        u5 = base * (100 - p5) // 100
        u10 = base * (100 - p10) // 100
        u20 = base * (100 - p20) // 100
        lines.append(f"{_cat_icon(c['name'])} <b>{html.escape(c['name'])}</b> — còn {n}")
        lines.append(f"   Lẻ: <b>{vnd(base)}</b> | 5: <b>{vnd(u5)}</b> | "
                     f"10: <b>{vnd(u10)}</b> | 20: <b>{vnd(u20)}</b>")
        lines.append("")
    lines.append("<i>Mua càng nhiều càng rẻ. Vào /shop chọn loại rồi bấm nút số lượng.</i>")
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.callback_query(F.data.startswith("accbuy:"))
async def on_acc_buy(cb: CallbackQuery):
    await cb.answer()
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    c = db.acc_category_get(cat_id)
    if not c or not c["active"]:
        await cb.message.answer("❌ Loại acc này không còn bán.")
        return
    c = dict(c)
    n = db.acc_stock_count(cat_id)
    p5, p10, p20 = _shop_bulk_pcts()
    up = db.shop_unit_price(c)
    unit = up["price"]
    icon = _cat_icon(c["name"])
    sold = db.acc_sold_count(cat_id)
    sample = db.get_conn().execute(
        "SELECT uid, created_date FROM acc_stock WHERE cat_id=? AND status='AVAILABLE' ORDER BY id LIMIT 1",
        (cat_id,)).fetchone()
    mail_line = _mail_app_line()
    txt = (
        f"{icon} <b>{html.escape(c['name'])}</b>\n"
        f"━━━━━━━━━━━━━━\n"
        f"💰 <b>{vnd(unit)}</b>/acc"
        + (f"  <i>(gốc {vnd(up['base'])})</i>" if up["happy"] else "") + "\n"
        f"📦 Tồn kho: <b>{n}</b> acc"
        + (f"  •  🔁 Đã bán <b>{sold}</b>" if sold else "") + "\n"
        f"━━━━━━━━━━━━━━\n"
        f"🛡 <b>Bảo hành:</b> 1 đổi 1 khi acc die\n"
        + (mail_line + "\n" if mail_line else "")
    )
    if sample and sample["uid"]:
        uid_e = html.escape(sample["uid"])
        txt += f'🔍 <i>Acc mẫu: <a href="https://facebook.com/{uid_e}">facebook.com/{uid_e}</a>'
        if sample["created_date"]:
            txt += f" (tạo {html.escape(sample['created_date'])})"
        txt += "</i>\n"

    if up["happy"]:
        txt += f"⚡ <b>GIỜ VÀNG −{up['happy_pct']}%</b> <i>(giá gốc {vnd(up['base'])})</i>\n"
    avg, rcnt = db.acc_review_avg(cat_id)
    if rcnt:
        stars = "⭐" * int(round(avg)) if avg else "⭐"
        txt += f"{stars} <b>{avg}/5</b> <i>({rcnt} đánh giá)</i>\n"
        for rv in db.acc_review_list(cat_id, 2):
            q = html.escape((rv["comment"] or "")[:120])
            txt += f'   <i>"{q}"</i> {"⭐" * int(rv["stars"] or 5)}\n'
    if c["description"]:
        txt += f"\n<i>{html.escape(c['description'])}</i>\n"
    if int(c.get("credit_bonus") or 0) > 0:
        txt += f"\n🎁 Tặng <b>{c['credit_bonus']} credits</b> cho mỗi acc mua.\n"
    if p5 > 0 or p10 > 0 or p20 > 0:
        txt += f"\n🏷 <b>Mua nhiều, giảm nhiều:</b>\n"
        if p5 > 0:
            txt += f"   5 acc → <b>−{p5}%</b>\n"
        if p10 > 0:
            txt += f"   10 acc → <b>−{p10}%</b>\n"
        if p20 > 0:
            txt += f"   20 acc → <b>−{p20}%</b> <i>(giá sỉ)</i>\n"
    tier = db.member_tier_info(cb.from_user.id)
    if tier["pct"] > 0:
        txt += f"\n{tier['tier']} — bạn được <b>giảm thêm {tier['pct']}%</b> mọi đơn.\n"
    txt += "\n<i>⚡ Giao acc tự động sau thanh toán • Đổi mật khẩu ngay khi đăng nhập.</i>\n"
    kb_rows = []
    if n:
        kb_rows.append([InlineKeyboardButton(
            text=f"✅ Mua 1 — {vnd(unit)}",
            callback_data=f"accconfirm:{cat_id}:1")])
        if n >= 5 and p5 > 0:
            kb_rows.append([InlineKeyboardButton(
                text=f"🎁 Mua 5 — giảm {p5}%",
                callback_data=f"accconfirm:{cat_id}:5")])
        if n >= 10 and p10 > 0:
            kb_rows.append([InlineKeyboardButton(
                text=f"🎁 Mua 10 — giảm {p10}%",
                callback_data=f"accconfirm:{cat_id}:10")])
        if n >= 20 and p20 > 0:
            kb_rows.append([InlineKeyboardButton(
                text=f"💼 Mua 20 — giá sỉ giảm {p20}%",
                callback_data=f"accconfirm:{cat_id}:20")])
        kb_rows.append([InlineKeyboardButton(
            text="✏️ Nhập số lượng khác",
            callback_data=f"accqty:{cat_id}")])
        kb_rows.append([InlineKeyboardButton(
            text="🛒 Thêm vào giỏ",
            callback_data=f"accaddcart:{cat_id}")])
        kb_rows.append([InlineKeyboardButton(
            text="🔍 Soi acc mẫu (miễn phí)",
            callback_data=f"accpreview:{cat_id}")])
    else:
        # 4.9 Đặt cọc giữ hàng hot khi hết hàng
        try:
            dep_pct = int(db.get_setting("deposit_pct", "30") or 30)
        except Exception:
            dep_pct = 30
        dep_amt = int(c["price"]) * dep_pct // 100
        kb_rows.append([InlineKeyboardButton(
            text=f"💰 Đặt cọc {dep_pct}% ({vnd(dep_amt)}) giữ hàng",
            callback_data=f"acccoc:{cat_id}")])
        sub = db.acc_restock_sub(cb.from_user.id, cat_id)
        if sub:
            mode = f"tự mua {sub['qty']} acc" if sub["auto_buy"] else "chỉ báo hàng"
            kb_rows.append([InlineKeyboardButton(
                text=f"🔕 Hủy đăng ký ({mode})",
                callback_data=f"accunnotify:{cat_id}")])
        else:
            kb_rows.append([InlineKeyboardButton(
                text="🔔 Báo tôi khi có hàng",
                callback_data=f"accnotify:{cat_id}")])
    kb_rows.append([InlineKeyboardButton(text="🔙 Quay lại shop", callback_data="accshop_back")])
    kb = InlineKeyboardMarkup(inline_keyboard=kb_rows)
    cover = (c.get("cover_photo") or "").strip()
    if cover:
        try:
            cap = txt if len(txt) <= 1000 else txt[:990] + "…"
            await cb.message.answer_photo(cover, caption=cap, parse_mode="HTML",
                                          reply_markup=kb)
        except Exception:
            await cb.message.answer(txt, parse_mode="HTML",
                                    disable_web_page_preview=True, reply_markup=kb)
    else:
        await cb.message.answer(txt, parse_mode="HTML",
                                disable_web_page_preview=True, reply_markup=kb)


@router.callback_query(F.data.startswith("accnotify:"))
async def on_acc_notify(cb: CallbackQuery):
    """Chọn kiểu báo hàng: chỉ báo, hoặc tự động trừ ví shop mua ngay khi hàng về."""
    await cb.answer()
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    c = db.acc_category_get(cat_id)
    if not c:
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔔 Chỉ báo cho tôi khi có hàng",
                              callback_data=f"accnotify_do:{cat_id}:0:1")],
        [InlineKeyboardButton(text="⚡ Tự mua 1 acc khi có hàng về",
                              callback_data=f"accnotify_do:{cat_id}:1:1")],
        [InlineKeyboardButton(text="⚡ Tự mua 5 acc khi có hàng về",
                              callback_data=f"accnotify_do:{cat_id}:1:5")],
        [InlineKeyboardButton(text="⚡ Tự mua 10 acc khi có hàng về",
                              callback_data=f"accnotify_do:{cat_id}:1:10")],
    ])
    await cb.message.answer(
        f"📦 <b>{html.escape(c['name'])}</b> đang hết hàng.\n\n"
        "Khi hàng về, bạn muốn:\n"
        "• <b>Chỉ báo:</b> bot nhắn riêng cho bạn.\n"
        "• <b>Tự mua:</b> bot tự trừ tiền <b>ví shop</b> và giao acc ngay "
        "(acc đã check LIVE). Nếu ví không đủ tiền lúc đó, bot chỉ báo hàng về thôi.",
        parse_mode="HTML", reply_markup=kb)


@router.callback_query(F.data.startswith("accnotify_do:"))
async def on_acc_notify_do(cb: CallbackQuery):
    await cb.answer()
    try:
        _, cat_id, auto, qty = cb.data.split(":")
        cat_id, auto, qty = int(cat_id), int(auto), int(qty)
    except Exception:
        return
    c = db.acc_category_get(cat_id)
    if not c:
        return
    is_new = db.acc_sub_restock(cb.from_user.id, cat_id, qty=qty, auto_buy=auto)
    if auto:
        await cb.message.answer(
            f"⚡ Đã đặt trước <b>{qty} acc</b> <b>{html.escape(c['name'])}</b>!\n\n"
            f"Khi hàng về, bot sẽ tự trừ <b>{vnd(int(c['price']) * qty)}</b> (giá gốc, "
            "chưa tính giảm mua nhiều) từ ví shop và giao acc ngay cho bạn.\n"
            "Nhớ nạp đủ tiền vào ví shop bằng /napshop nhé!",
            parse_mode="HTML")
    else:
        await cb.message.answer(
            f"🔔 Đã đăng ký! Khi <b>{html.escape(c['name'])}</b> có hàng về, "
            f"bot sẽ nhắn riêng cho bạn ngay.",
            parse_mode="HTML")


@router.callback_query(F.data.startswith("accunnotify:"))
async def on_acc_unnotify(cb: CallbackQuery):
    await cb.answer()
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    db.acc_unsub_restock(cb.from_user.id, cat_id)
    await cb.message.answer("🔕 Đã hủy đăng ký báo hàng.")


@router.callback_query(F.data.startswith("accpreview:"))
async def on_acc_preview(cb: CallbackQuery):
    """Soi 5 acc AVAILABLE ngẫu nhiên (chỉ hiện link để khách check, không lộ pass/cookie/token)."""
    await cb.answer()
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    c = db.acc_category_get(cat_id)
    if not c:
        return
    rows = db.get_conn().execute(
        "SELECT uid, created_date FROM acc_stock "
        "WHERE cat_id=? AND status='AVAILABLE' ORDER BY RANDOM() LIMIT 5",
        (cat_id,)).fetchall()
    rows = [r for r in rows if (r["uid"] or "").strip()]
    if not rows:
        await cb.message.answer("⛔ Loại này vừa hết hàng.")
        return
    e = html.escape
    parts = [f"🔍 <b>SOI ACC MẪU</b> — {e(c['name'])}",
             f"Lấy ngẫu nhiên <b>{len(rows)}</b> acc trong kho để bạn check chất lượng:", ""]
    for i, r in enumerate(rows, 1):
        uid = (r["uid"] or "").strip()
        line = f'{i}. <a href="https://facebook.com/{e(uid)}">facebook.com/{e(uid)}</a>'
        if r["created_date"]:
            line += f" (tạo {e(r['created_date'])})"
        parts.append(line)
    parts += ["", "🔒 Mật khẩu, 2FA, cookie, token chỉ hiện sau khi bạn mua.",
              "🛡 Các acc trong kho đã qua quét kiểm tra chất lượng."]
    await cb.message.answer("\n".join(parts), parse_mode="HTML", disable_web_page_preview=True)


def _shop_bulk_pcts() -> tuple[int, int, int]:
    """% giảm giá khi mua 5 / 10 / 20 acc (giá sỉ). Đổi qua settings shop_bulk_5_pct, shop_bulk_10_pct, shop_bulk_20_pct."""
    def _g(key, default):
        try:
            return max(0, int(db.get_setting(key, str(default)) or default))
        except Exception:
            return default
    return _g("shop_bulk_5_pct", 5), _g("shop_bulk_10_pct", 10), _g("shop_bulk_20_pct", 15)


@router.callback_query(F.data == "accshop_back")
async def on_acc_shop_back(cb: CallbackQuery):
    await cb.answer()
    await on_shop(cb.message)


@router.callback_query(F.data.startswith("accqty:"))
async def on_acc_qty_prompt(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    c = db.acc_category_get(cat_id)
    if not c or not c["active"]:
        await cb.message.answer("❌ Loại acc này không còn bán.")
        return
    n = db.acc_stock_count(cat_id)
    if n <= 0:
        await cb.message.answer("⛔ Loại này vừa hết hàng.")
        return
    await state.set_state(AccShopState.waiting_for_custom_qty)
    await state.update_data(acc_qty_cat_id=cat_id)
    await cb.message.answer(
        f"✏️ Nhập <b>số lượng</b> muốn mua (1–{n} acc):\n"
        f"<i>Gõ /huy để hủy.</i>",
        parse_mode="HTML")


@router.message(AccShopState.waiting_for_custom_qty)
async def on_acc_qty_input(msg: Message, state: FSMContext):
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("🚫 Đã hủy thao tác đang nhập. Bạn gõ lại lệnh vừa rồi nhé.")
        return
    if _t.lower() in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("Đã hủy nhập số lượng.")
        return
    data = await state.get_data()
    cat_id = data.get("acc_qty_cat_id")
    if not cat_id:
        await state.clear()
        return
    c = db.acc_category_get(cat_id)
    if not c or not c["active"]:
        await state.clear()
        await msg.answer("❌ Loại acc này không còn bán.")
        return
    try:
        qty = int(_t.replace(".", "").replace(",", "").strip())
    except Exception:
        await msg.answer("❌ Số lượng không hợp lệ. Nhập số nguyên, ví dụ: 3")
        return
    n = db.acc_stock_count(cat_id)
    if qty < 1 or qty > n:
        await msg.answer(f"❌ Số lượng phải từ 1 đến {n}. Nhập lại nhé (hoặc /huy để hủy):")
        return
    await state.clear()
    c = dict(c)
    up = db.shop_unit_price(c)
    price = up["price"]
    p5, p10, p20 = _shop_bulk_pcts()
    bulk_pct = p20 if qty >= 20 else (p10 if qty >= 10 else (p5 if qty >= 5 else 0))
    tier = db.member_tier_info(msg.from_user.id)
    tier_pct = int(tier["pct"])
    total = price * qty
    final = total * (100 - bulk_pct) // 100
    final = final * (100 - tier_pct) // 100
    disc = []
    if bulk_pct:
        disc.append(f"giảm {bulk_pct}% mua nhiều")
    if tier_pct:
        disc.append(f"giảm {tier_pct}% hạng {tier['tier']}")
    disc_txt = (" (" + ", ".join(disc) + ")") if disc else ""
    txt = (
        f"🛒 <b>Xác nhận mua</b>\n"
        f"📦 {html.escape(c['name'])} × <b>{qty}</b> acc\n"
        f"💰 Tổng: <b>{vnd(final)}</b>{disc_txt}\n"
        f"👉 Bấm xác nhận để mua ngay:")
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(
            text=f"✅ Xác nhận mua {qty} — {vnd(final)}",
            callback_data=f"accconfirm:{cat_id}:{qty}")],
        [InlineKeyboardButton(text="🔙 Chọn lại", callback_data=f"accbuy:{cat_id}")],
    ])
    await msg.answer(txt, parse_mode="HTML", reply_markup=kb)


@router.callback_query(F.data.startswith("accconfirm:"))
async def on_acc_confirm(cb: CallbackQuery):
    await cb.answer()
    try:
        parts = cb.data.split(":")
        cat_id = int(parts[1])
        qty = int(parts[2]) if len(parts) > 2 else 1
        want_upsell = len(parts) > 3 and parts[3] == "upsell"
    except Exception:
        return
    if qty < 1:
        qty = 1
    if qty > 1000:
        qty = 1000
    tg_id = cb.from_user.id
    c = db.acc_category_get(cat_id)
    if not c or not c["active"]:
        await cb.message.answer("❌ Loại acc này không còn bán.")
        return
    c = dict(c)
    up = db.shop_unit_price(c)
    price = up["price"]
    # 4.6 Upsell: mua thêm trong 10 phút sau đơn trước được giảm thêm
    upsell_pct = 0
    try:
        wmin = int(db.get_setting("upsell_window_min", "10") or 10)
        upsell_pct_cfg = int(db.get_setting("upsell_pct", "5") or 5)
    except Exception:
        wmin, upsell_pct_cfg = 10, 5
    if want_upsell and upsell_pct_cfg > 0:
        last_at = db.acc_last_order_at(tg_id)
        if last_at and now() - last_at <= wmin * 60:
            upsell_pct = upsell_pct_cfg
        else:
            want_upsell = False
    p5, p10, p20 = _shop_bulk_pcts()
    bulk_pct = p20 if qty >= 20 else (p10 if qty >= 10 else (p5 if qty >= 5 else 0))
    tier = db.member_tier_info(tg_id)
    tier_pct = int(tier["pct"])
    total = price * qty
    final = total * (100 - bulk_pct) // 100
    final = final * (100 - tier_pct) // 100
    final = final * (100 - upsell_pct) // 100
    if db.acc_stock_count(cat_id) < qty:
        await cb.message.answer(
            f"⛔ Kho chỉ còn <b>{db.acc_stock_count(cat_id)}</b> acc, "
            f"không đủ {qty} acc. Bạn chọn số lượng ít hơn nhé.",
            parse_mode="HTML")
        return
    u = db.get_user(tg_id)
    balance = int(u["shop_balance"] or 0) if u else 0
    if balance < final:
        disc_txt = []
        if bulk_pct:
            disc_txt.append(f"giảm {bulk_pct}% mua nhiều")
        if tier_pct:
            disc_txt.append(f"giảm {tier_pct}% hạng {tier['tier']}")
        if upsell_pct:
            disc_txt.append(f"giảm {upsell_pct}% mua thêm trong {wmin} phút")
        await cb.message.answer(
            f"😢 <b>VÍ SHOP KHÔNG ĐỦ</b>\n"
            f"━━━━━━━━━━━━━━\n"
            f"Mua {qty} acc: <b>{vnd(final)}</b>"
            + (f"\n<i>🎉 {', '.join(disc_txt)}</i>" if disc_txt else "") + "\n"
            f"👛 Ví shop của bạn: <b>{vnd(balance)}</b>\n"
            f"💸 Còn thiếu: <b>{vnd(final - balance)}</b>\n"
            f"━━━━━━━━━━━━━━\n"
            f"Nạp thêm bằng /napshop (tự động, quét QR) rồi mua lại nhé.",
            parse_mode="HTML",
        )
        return
    # Trừ tiền ví shop trước, giao acc sau (nguyên tử ở acc_sell_many)
    if not db.adjust_shop_balance(tg_id, -final, f"mua_acc:{cat_id}x{qty}"):
        await cb.message.answer(
            "❌ <b>Ví shop không đủ!</b>\nNạp thêm bằng /napshop rồi mua lại nhé."
            + _SHOP_WALLET_HINT,
            parse_mode="HTML")
        return
    # Check LIVE trước khi giao (chỉ sạp Acc Facebook): chỉ bán acc đã check live
    _need_check = _cat_live_check(cat_id)
    if _need_check:
        await cb.message.answer("🔍 <b>Đang kiểm tra chất lượng acc...</b>", parse_mode="HTML")
    sold_orders, _sell_fail = await _sell_live_stock(
        cb.bot, tg_id, final, qty,
        lambda seen: db.acc_stock_pick_candidates(cat_id, qty + 4, seen),
        check_live=_need_check)
    if not sold_orders:
        db.add_shop_balance_only(tg_id, final, "hoan_tien_khong_du_hang_live")
        await cb.message.answer(
            "😔 <b>Rất tiếc, hiện không đủ hàng</b> để giao cho bạn.\n"
            f"Tiền <b>{vnd(final)}</b> đã được hoàn vào ví shop.\n"
            "Bạn quay lại sau hoặc chọn loại acc khác nhé!",
            parse_mode="HTML")
        return
    # Giao từng acc — khách chọn cách nhận: hiện thông tin hoặc tải file
    delivered = sold_orders
    for order in delivered:
        order_id = order["id"]
        await cb.message.answer(
            f"🎉 <b>MUA THÀNH CÔNG!</b>\n"
            f"━━━━━━━━━━━━━━\n"
            f"{icon} <b>{html.escape(order['cat_name'])}</b>\n"
            f"🧾 Đơn hàng: <b>#{order_id}</b>\n"
            f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
            f"💰 Đã thanh toán: <b>{vnd(order['price'])}</b>\n"
            f"{_live_line(cat_id)}"
            f"━━━━━━━━━━━━━━\n\n"
            f"{_pickup_suffix()}",
            parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
    # Báo admin: thông tin khách + acc đã mua
    await _notify_purchase_admin(cb.bot, cb.from_user, delivered)
    # Quà tặng kèm: credits + vé quay
    extras = []
    bonus_per = int(c.get("credit_bonus") or 0)
    if bonus_per > 0:
        new_credits = db.add_credits(tg_id, bonus_per * qty, f"combo_mua_acc:{cat_id}")
        extras.append(f"🎁 Tặng <b>{bonus_per * qty} credits</b> (số dư credits: {new_credits})")
    tickets = db.spin_add_tickets(tg_id, qty)
    extras.append(f"🎡 Nhận <b>{qty} vé quay</b> may mắn — gõ /quay để thử vận may (đang có {tickets} vé)")
    # Tích điểm loyalty: 1 điểm / loyalty_per_vnd
    try:
        per = int(db.get_setting("loyalty_per_vnd", "100000") or 100000)
    except Exception:
        per = 100000
    pts = final // per if per > 0 else 0
    if pts > 0:
        new_pts = db.loyalty_add(tg_id, pts, f"mua_acc:{cat_id}")
        extras.append(f"⭐ Tích <b>{pts} điểm</b> loyalty (tổng: <b>{new_pts}</b> điểm) — "
                      f"đủ điểm đổi acc miễn phí bằng /doiqua")
    if extras:
        rkey = f"o{min(o['id'] for o in delivered)}" if delivered else None
        await cb.message.answer("\n".join(extras), parse_mode="HTML",
                                reply_markup=_spin_kb(rkey, tg_id, final))
    # 4.6 Upsell: gợi ý mua thêm trong 10 phút được giảm thêm
    try:
        if upsell_pct_cfg > 0 and not want_upsell:
            kb = InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(
                    text=f"⚡ Mua thêm 1 acc giảm {upsell_pct_cfg}% (trong {wmin} phút)",
                    callback_data=f"accconfirm:{cat_id}:1:upsell")],
            ])
            await cb.message.answer(
                f"⚡ <b>ƯU ĐÃI NÓNG:</b> bạn vừa mua <b>{html.escape(c['name'])}</b> — "
                f"mua thêm acc cùng loại trong <b>{wmin} phút</b> được "
                f"<b>giảm thêm {upsell_pct_cfg}%</b>!",
                parse_mode="HTML", reply_markup=kb)
    except Exception:
        pass
    # 5.12 Hướng dẫn sau mua tự động
    try:
        guide = db.get_setting("postbuy_guide", "") or (
            "📋 <b>HƯỚNG DẪN SAU KHI MUA ACC</b>\n"
            "1️⃣ Đổi mật khẩu ngay sau khi đăng nhập.\n"
            "2️⃣ Bật xác thực 2 lớp (2FA) cho acc.\n"
            "3️⃣ 3 ngày đầu: đừng đổi tên/avatar vội, lướt newsfeed nhẹ nhàng.\n"
            "4️⃣ Không đăng nhập nhiều acc cùng 1 IP/proxy lạ.\n"
            "💡 Làm đúng các bước trên giúp acc sống lâu, ít bị checkpoint!")
        await cb.message.answer(guide, parse_mode="HTML")
    except Exception:
        pass
    # Hoa hồng cho người giới thiệu (F1)
    try:
        f1_id, comm = db.ref_shop_commission(tg_id, final)
        if f1_id and comm:
            try:
                await cb.bot.send_message(
                    f1_id,
                    f"🎁 <b>Hoa hồng shop acc!</b>\n"
                    f"Người bạn giới thiệu vừa mua {qty} acc "
                    f"<b>{html.escape(c['name'])}</b> — bạn nhận <b>{vnd(comm)}</b>.",
                    parse_mode="HTML")
            except Exception:
                pass
    except Exception:
        pass
    # Cảnh báo hết hàng cho admin
    try:
        warn_at = int(db.get_setting("acc_low_stock_warn", "20") or 20)
    except Exception:
        warn_at = 20
    left = db.acc_stock_count(cat_id)
    if left <= warn_at:
        await _notify_admin_smart(
            cb.bot,
            f"⚠️ <b>Sắp hết hàng:</b> {html.escape(c['name'])} chỉ còn <b>{left}</b> acc. "
            f"Nhập thêm bằng /themacc {cat_id}",
            perm="kho",
        )


# ============================ GIỎ HÀNG SHOP ACC ============================
def _line_price(cat: dict, qty: int, tg_id: int) -> dict:
    """Giá 1 dòng giỏ/mua: giữ nguyên công thức mua lẻ
    (giảm bulk theo số lượng của dòng + giảm hạng TV)."""
    up = db.shop_unit_price(cat)
    price = up["price"]
    p5, p10, p20 = _shop_bulk_pcts()
    bulk_pct = p20 if qty >= 20 else (p10 if qty >= 10 else (p5 if qty >= 5 else 0))
    tier = db.member_tier_info(tg_id)
    tier_pct = int(tier["pct"])
    total = price * qty
    final = total * (100 - bulk_pct) // 100
    final = final * (100 - tier_pct) // 100
    return {"unit": price, "bulk_pct": bulk_pct, "tier_pct": tier_pct,
            "tier_name": tier.get("tier", ""), "total": total, "final": final}


def _cart_render(tg_id: int):
    """Trả (text, keyboard) cho giỏ hàng, hoặc None nếu giỏ trống."""
    items = [it for it in db.cart_list(tg_id) if it["active"]]
    if not items:
        return None
    lines = ["🛒 <b>GIỎ HÀNG CỦA BẠN</b>", "━━━━━━━━━━━━━━"]
    kb_rows = []
    grand = 0
    saved = 0
    for i, it in enumerate(items, 1):
        cat = db.acc_category_get(it["cat_id"])
        if not cat:
            continue
        lp = _line_price(dict(cat), it["qty"], tg_id)
        grand += lp["final"]
        disc = []
        if lp["bulk_pct"]:
            disc.append(f"−{lp['bulk_pct']}% mua nhiều")
        if lp["tier_pct"]:
            disc.append(f"−{lp['tier_pct']}% {lp['tier_name']}")
        # tiền tiết kiệm so với giá gốc
        try:
            base_total = int(lp["unit"]) * int(it["qty"])
            saved += max(0, base_total - int(lp["final"]))
        except Exception:
            pass
        dtxt = f"\n   <i>🎉 {', '.join(disc)}</i>" if disc else ""
        lines.append(
            f"<b>{i}. {_cat_icon(it['name'])} {html.escape(it['name'])}</b>\n"
            f"   {it['qty']} × {vnd(lp['unit'])} = <b>{vnd(lp['final'])}</b>{dtxt}")
        kb_rows.append([
            InlineKeyboardButton(text="➖", callback_data=f"cartdec:{it['cat_id']}"),
            InlineKeyboardButton(text=f"×{it['qty']}", callback_data="cartnoop"),
            InlineKeyboardButton(text="➕", callback_data=f"cartinc:{it['cat_id']}"),
            InlineKeyboardButton(text="❌", callback_data=f"cartdel:{it['cat_id']}"),
        ])
    u = db.get_user(tg_id)
    balance = int(u["shop_balance"] or 0) if u else 0
    lines += ["━━━━━━━━━━━━━━",
              f"🧾 <b>Tổng cộng: {vnd(grand)}</b>"]
    if saved > 0:
        lines.append(f"🎉 <i>Bạn tiết kiệm được {vnd(saved)}</i>")
    lines.append(f"👛 Ví shop: <b>{vnd(balance)}</b>"
                 + ("" if balance >= grand else f"  <i>(thiếu {vnd(grand - balance)})</i>"))
    kb_rows.append([InlineKeyboardButton(
        text=f"💳 Thanh toán — {vnd(grand)}", callback_data="cartcheckout")])
    kb_rows.append([
        InlineKeyboardButton(text="🗑 Xóa giỏ", callback_data="cartclear"),
        InlineKeyboardButton(text="🛍 Tiếp tục mua", callback_data="accshop_back"),
    ])
    return "\n".join(lines), InlineKeyboardMarkup(inline_keyboard=kb_rows)


async def _cart_refresh(cb: CallbackQuery):
    r = _cart_render(cb.from_user.id)
    if not r:
        try:
            await cb.message.edit_text(
                "🛒 Giỏ hàng đang trống.\nVào /shop chọn acc rồi bấm \"🛒 Thêm vào giỏ\" nhé!")
        except Exception:
            pass
        return
    txt, kb = r
    try:
        await cb.message.edit_text(txt, parse_mode="HTML", reply_markup=kb)
    except Exception:
        pass


def _cart_validated_lines(tg_id: int):
    """Validate giỏ: bỏ loại ngừng bán/hết hàng, cắt số lượng theo tồn kho.
    Trả (lines[(cat, qty)], notes[str])."""
    lines, notes = [], []
    for it in db.cart_list(tg_id):
        c = db.acc_category_get(it["cat_id"])
        if not c or not c["active"]:
            db.cart_remove(tg_id, it["cat_id"])
            notes.append(f"❌ {it['name']}: ngừng bán → đã xóa khỏi giỏ")
            continue
        stock = db.acc_stock_count(it["cat_id"])
        if stock <= 0:
            db.cart_remove(tg_id, it["cat_id"])
            notes.append(f"⛔ {it['name']}: vừa hết hàng → đã xóa khỏi giỏ")
            continue
        qty = it["qty"]
        if qty > stock:
            db.cart_set_qty(tg_id, it["cat_id"], stock)
            notes.append(f"⚠️ {it['name']}: kho chỉ còn {stock} → đã giảm số lượng")
            qty = stock
        lines.append((dict(c), qty))
    return lines, notes


@router.callback_query(F.data.startswith("accaddcart:"))
async def on_acc_add_cart(cb: CallbackQuery):
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        await cb.answer()
        return
    tg_id = cb.from_user.id
    c = db.acc_category_get(cat_id)
    if not c or not c["active"]:
        await cb.answer("❌ Loại acc này không còn bán.", show_alert=True)
        return
    stock = db.acc_stock_count(cat_id)
    if stock <= 0:
        await cb.answer("⛔ Loại này đang hết hàng.", show_alert=True)
        return
    q = db.cart_add(tg_id, cat_id, 1)
    if q > stock:
        db.cart_set_qty(tg_id, cat_id, stock)
        q = stock
    n = db.cart_count(tg_id)
    await cb.answer(f"🛒 Đã thêm vào giỏ ({n} acc). Gõ /giohang để xem và thanh toán.")


@router.message(Command("giohang"))
async def on_giohang(msg: Message):
    r = _cart_render(msg.from_user.id)
    if not r:
        await msg.answer("🛒 Giỏ hàng đang trống.\nVào /shop chọn acc rồi bấm \"🛒 Thêm vào giỏ\" nhé!")
        return
    txt, kb = r
    await msg.answer(txt, parse_mode="HTML", reply_markup=kb)


@router.callback_query(F.data == "cartnoop")
async def on_cart_noop(cb: CallbackQuery):
    await cb.answer()


@router.callback_query(F.data.startswith("cartinc:"))
async def on_cart_inc(cb: CallbackQuery):
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        await cb.answer()
        return
    tg_id = cb.from_user.id
    cur = next((it["qty"] for it in db.cart_list(tg_id) if it["cat_id"] == cat_id), 0)
    if not cur:
        await cb.answer()
        return
    stock = db.acc_stock_count(cat_id)
    if cur + 1 > stock:
        await cb.answer(f"⛔ Kho chỉ còn {stock} acc loại này.", show_alert=True)
        return
    await cb.answer()
    db.cart_set_qty(tg_id, cat_id, cur + 1)
    await _cart_refresh(cb)


@router.callback_query(F.data.startswith("cartdec:"))
async def on_cart_dec(cb: CallbackQuery):
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        await cb.answer()
        return
    tg_id = cb.from_user.id
    cur = next((it["qty"] for it in db.cart_list(tg_id) if it["cat_id"] == cat_id), 0)
    await cb.answer()
    db.cart_set_qty(tg_id, cat_id, cur - 1)
    await _cart_refresh(cb)


@router.callback_query(F.data.startswith("cartdel:"))
async def on_cart_del(cb: CallbackQuery):
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        await cb.answer()
        return
    await cb.answer("Đã xóa khỏi giỏ.")
    db.cart_remove(cb.from_user.id, cat_id)
    await _cart_refresh(cb)


@router.callback_query(F.data == "cartclear")
async def on_cart_clear(cb: CallbackQuery):
    db.cart_clear(cb.from_user.id)
    await cb.answer("🗑 Đã xóa toàn bộ giỏ hàng.")
    await _cart_refresh(cb)


@router.callback_query(F.data == "cartview")
async def on_cart_view(cb: CallbackQuery):
    await cb.answer()
    r = _cart_render(cb.from_user.id)
    if not r:
        await cb.message.answer("🛒 Giỏ hàng đang trống.")
        return
    txt, kb = r
    await cb.message.answer(txt, parse_mode="HTML", reply_markup=kb)


@router.callback_query(F.data == "cartcheckout")
async def on_cart_checkout(cb: CallbackQuery):
    await cb.answer()
    tg_id = cb.from_user.id
    lines, notes = _cart_validated_lines(tg_id)
    if not lines:
        txt = "🛒 Giỏ hàng không còn món nào mua được."
        if notes:
            txt += "\n" + "\n".join(notes)
        await cb.message.answer(txt)
        return
    priced = []
    grand = 0
    for c, qty in lines:
        lp = _line_price(c, qty, tg_id)
        priced.append((c, qty, lp))
        grand += lp["final"]
    u = db.get_user(tg_id)
    balance = int(u["shop_balance"] or 0) if u else 0
    txt = ["🧾 <b>XÁC NHẬN THANH TOÁN</b>", "━━━━━━━━━━━━━━"]
    for c, qty, lp in priced:
        txt.append(f"{_cat_icon(c['name'])} <b>{html.escape(c['name'])}</b>\n"
                   f"   {qty} × {vnd(lp['unit'])} = <b>{vnd(lp['final'])}</b>")
    txt += ["━━━━━━━━━━━━━━",
            f"🧾 <b>Tổng cộng: {vnd(grand)}</b>",
            f"👛 Ví shop: <b>{vnd(balance)}</b>"]
    if notes:
        txt.append("")
        txt += [f"<i>{html.escape(n)}</i>" for n in notes]
    if balance < grand:
        txt.append(f"\n❌ <b>Ví shop không đủ!</b> Còn thiếu <b>{vnd(grand - balance)}</b>.\n"
                   "Nạp thêm bằng /napshop rồi quay lại nhé." + _SHOP_WALLET_HINT)
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔙 Sửa giỏ hàng", callback_data="cartview")]])
    else:
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=f"✅ Xác nhận thanh toán {vnd(grand)}",
                                  callback_data="cartconfirm")],
            [InlineKeyboardButton(text="🔙 Sửa giỏ hàng", callback_data="cartview")]])
    await cb.message.answer("\n".join(txt), parse_mode="HTML", reply_markup=kb)


@router.callback_query(F.data == "cartconfirm")
async def on_cart_confirm(cb: CallbackQuery):
    await cb.answer()
    tg_id = cb.from_user.id
    lines, _notes = _cart_validated_lines(tg_id)
    if not lines:
        await cb.message.answer("🛒 Giỏ hàng không còn món nào mua được.")
        return
    priced = []
    grand = 0
    for c, qty in lines:
        stock = db.acc_stock_count(c["id"])
        if stock < qty:
            await cb.message.answer(
                f"⛔ <b>{html.escape(c['name'])}</b> chỉ còn {stock} acc, "
                f"không đủ {qty}.\nBạn sửa lại giỏ hàng rồi thanh toán lại nhé.",
                parse_mode="HTML")
            return
        lp = _line_price(c, qty, tg_id)
        priced.append((c, qty, lp))
        grand += lp["final"]
    u = db.get_user(tg_id)
    balance = int(u["shop_balance"] or 0) if u else 0
    if balance < grand:
        await cb.message.answer(
            f"❌ Ví shop không đủ ({vnd(balance)} < {vnd(grand)}). "
            "Nạp thêm bằng /napshop nhé." + _SHOP_WALLET_HINT, parse_mode="HTML")
        return
    if not db.adjust_shop_balance(tg_id, -grand, f"mua_giohang:{len(priced)}mon"):
        await cb.message.answer("❌ Ví shop không đủ. Nạp thêm bằng /napshop nhé."
                                + _SHOP_WALLET_HINT,
                                parse_mode="HTML")
        return
    await cb.message.answer("🔍 <b>Đang kiểm tra chất lượng acc...</b>", parse_mode="HTML")
    delivered, paid_total = [], 0
    failed = []
    for c, qty, lp in priced:
        cat_id = c["id"]
        sold, _fail = await _sell_live_stock(
            cb.bot, tg_id, lp["final"], qty,
            lambda seen, _cid=cat_id, _q=qty: db.acc_stock_pick_candidates(_cid, _q + 4, seen),
            check_live=_cat_live_check(cat_id))
        if sold:
            delivered += sold
            paid_total += lp["final"]
        else:
            failed.append((c, qty))
            db.add_shop_balance_only(tg_id, lp["final"], "hoan_tien_giohang_thieu_live")
    db.cart_clear(tg_id)
    if not delivered:
        await cb.message.answer(
            "😔 <b>Rất tiếc, hiện không đủ hàng</b> để giao.\n"
            f"Tiền <b>{vnd(grand)}</b> đã được hoàn vào ví shop.",
            parse_mode="HTML")
        return
    for order in delivered:
        order_id = order["id"]
        await cb.message.answer(
            f"🎉 <b>MUA THÀNH CÔNG!</b>\n"
            f"━━━━━━━━━━━━━━\n"
            f"{_cat_icon(order['cat_name'])} <b>{html.escape(order['cat_name'])}</b>\n"
            f"🧾 Đơn hàng: <b>#{order_id}</b>\n"
            f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
            f"💰 Đã thanh toán: <b>{vnd(order['price'])}</b>\n"
            f"{_live_line(cat_id)}"
            f"━━━━━━━━━━━━━━\n\n"
            f"{_pickup_suffix()}",
            parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
    by_cat = {}
    for o in delivered:
        by_cat.setdefault(o["cat_id"], []).append(o)
    summary = [f"🧾 <b>Xong giỏ hàng:</b> giao <b>{len(delivered)}</b> acc, "
               f"thanh toán <b>{vnd(paid_total)}</b>."]
    if failed:
        fnames = ", ".join(f"{html.escape(c['name'])}×{q}" for c, q in failed)
        summary.append(f"⚠️ {fnames}: không đủ acc LIVE → đã hoàn tiền dòng này.")
    extras = []
    for cid, olist in by_cat.items():
        crow = db.acc_category_get(cid)
        c = dict(crow) if crow else None
        bonus_per = int((c.get("credit_bonus") or 0)) if c else 0
        if bonus_per > 0:
            db.add_credits(tg_id, bonus_per * len(olist), f"combo_mua_giohang:{cid}")
            extras.append(f"🎁 Tặng <b>{bonus_per * len(olist)} credits</b> "
                          f"({html.escape(c['name']) if c else ''})")
    tickets = db.spin_add_tickets(tg_id, len(delivered))
    extras.append(f"🎡 Nhận <b>{len(delivered)} vé quay</b> may mắn — gõ /quay "
                  f"(đang có {tickets} vé)")
    try:
        per = int(db.get_setting("loyalty_per_vnd", "100000") or 100000)
    except Exception:
        per = 100000
    pts = paid_total // per if per > 0 else 0
    if pts > 0:
        new_pts = db.loyalty_add(tg_id, pts, "mua_giohang")
        extras.append(f"⭐ Tích <b>{pts} điểm</b> loyalty (tổng: <b>{new_pts}</b> điểm) — "
                      "đủ điểm đổi acc miễn phí bằng /doiqua")
    summary.append("")
    summary += extras
    rkey = f"c{min(o['id'] for o in delivered)}" if delivered else None
    await cb.message.answer("\n".join(summary), parse_mode="HTML",
                            reply_markup=_spin_kb(rkey, tg_id, paid_total))
    try:
        guide = db.get_setting("postbuy_guide", "") or (
            "📋 <b>HƯỚNG DẪN SAU KHI MUA ACC</b>\n"
            "1️⃣ Đổi mật khẩu ngay sau khi đăng nhập.\n"
            "2️⃣ Bật xác thực 2 lớp (2FA) cho acc.\n"
            "3️⃣ 3 ngày đầu: đừng đổi tên/avatar vội, lướt newsfeed nhẹ nhàng.\n"
            "4️⃣ Không đăng nhập nhiều acc cùng 1 IP/proxy lạ.\n"
            "💡 Làm đúng các bước trên giúp acc sống lâu, ít bị checkpoint!")
        await cb.message.answer(guide, parse_mode="HTML")
    except Exception:
        pass
    try:
        f1_id, comm = db.ref_shop_commission(tg_id, paid_total)
        if f1_id and comm:
            try:
                await cb.bot.send_message(
                    f1_id,
                    "🎁 <b>Hoa hồng shop acc!</b>\n"
                    f"Người bạn giới thiệu vừa mua giỏ hàng {len(delivered)} acc — "
                    f"bạn nhận <b>{vnd(comm)}</b>.",
                    parse_mode="HTML")
            except Exception:
                pass
    except Exception:
        pass
    await _notify_purchase_admin(cb.bot, cb.from_user, delivered)
    try:
        warn_at = int(db.get_setting("acc_low_stock_warn", "20") or 20)
    except Exception:
        warn_at = 20
    for cid in by_cat:
        left = db.acc_stock_count(cid)
        if left <= warn_at:
            c = db.acc_category_get(cid)
            if c:
                await _notify_admin_smart(
                    cb.bot,
                    f"⚠️ <b>Sắp hết hàng:</b> {html.escape(c['name'])} chỉ còn "
                    f"<b>{left}</b> acc. Nhập thêm bằng /themacc {cid}",
                    perm="kho")




@router.message(Command("quay"))
async def on_quay(msg: Message):
    """Vòng quay may mắn — mỗi acc mua = 1 vé."""
    tg_id = msg.from_user.id
    tickets = db.spin_get_tickets(tg_id)
    if tickets <= 0:
        await msg.answer(
            "🎡 <b>VÒNG QUAY MAY MẮN</b>\n\n"
            "Bạn chưa có vé quay nào.\n"
            "👉 Mua acc ở /shop — mỗi acc = 1 vé quay, trúng credits và tiền mặt vào số dư!",
            parse_mode="HTML")
        return
    await _do_spin(msg, tg_id, tickets)


@router.callback_query(F.data == "spin_now")
async def on_spin_now(cb: CallbackQuery):
    await cb.answer()
    tickets = db.spin_get_tickets(cb.from_user.id)
    if tickets <= 0:
        await cb.message.answer("😅 Bạn đã hết vé quay. Mua acc ở /shop để nhận thêm vé nhé!")
        return
    await _do_spin(cb.message, cb.from_user.id, tickets)


# ---- Điểm loyalty ngẫu nhiên sau khi mua acc ----
def _loyalty_random_cfg():
    """(min, max, cap_user, min_order, daily_max) cho điểm ngẫu nhiên sau mua.
    None = đang tắt. min_order: đơn phải >= X tiền mới được random (0 = không
    giới hạn). daily_max: mỗi user tối đa N lượt random/ngày (0 = không giới hạn)."""
    try:
        lo = int(db.get_setting("loyalty_random_min", "1") or 1)
        hi = int(db.get_setting("loyalty_random_max", "5") or 5)
        cap = int(db.get_setting("loyalty_random_cap", "0") or 0)
        min_order = int(db.get_setting("loyalty_random_min_order", "0") or 0)
        daily_max = int(db.get_setting("loyalty_random_daily_max", "0") or 0)
    except Exception:
        return None
    if cap <= 0 or lo <= 0 or hi < lo:
        return None
    return (lo, hi, cap, max(0, min_order), max(0, daily_max))


def _loyalty_random_received(tg_id: int) -> int:
    """Tổng điểm ngẫu nhiên user đã nhận (để so với trần cap_user)."""
    r = db.get_conn().execute(
        "SELECT COALESCE(SUM(delta),0) AS s FROM loyalty_history "
        "WHERE tg_id=? AND reason LIKE 'randpts:%'", (tg_id,)).fetchone()
    try:
        return int(r["s"] or 0)
    except Exception:
        return 0


def _loyalty_random_today(tg_id: int) -> int:
    """Số lượt random user đã bấm trong hôm nay (chống cày điểm)."""
    try:
        import datetime as _dt
        start = int(_dt.datetime.now().replace(
            hour=0, minute=0, second=0, microsecond=0).timestamp())
        r = db.get_conn().execute(
            "SELECT COUNT(*) AS c FROM loyalty_history "
            "WHERE tg_id=? AND reason LIKE 'randpts:%' AND created_at>=?",
            (tg_id, start)).fetchone()
        return int(r["c"] or 0)
    except Exception:
        return 0


def _order_total_for_randkey(key: str) -> int:
    """Tổng tiền của đơn từ rand key (o<order_id> / c<order_id>). 0 nếu không rõ."""
    try:
        oid = int((key or "")[1:])
        r = db.get_conn().execute(
            "SELECT price FROM acc_orders WHERE id=?", (oid,)).fetchone()
        return int(r["price"] or 0) if r else 0
    except Exception:
        return 0


def _spin_kb(order_key=None, tg_id: int = 0, order_total: int = 0):
    """Bàn phím tin quà tặng sau mua: nút quay + nút điểm ngẫu nhiên (nếu bật
    và đơn/user đủ điều kiện chống cày điểm). Nhãn nút hiện tiến độ (x/y)."""
    rows = [[InlineKeyboardButton(text="🎡 Quay ngay", callback_data="spin_now")]]
    if order_key:
        cfg = _loyalty_random_cfg()
        if cfg:
            lo, hi, cap, min_order, daily_max = cfg
            ok = True
            if min_order > 0 and order_total < min_order:
                ok = False
            if ok and daily_max > 0 and tg_id:
                ok = _loyalty_random_today(tg_id) < daily_max
            if ok:
                prog = ""
                if tg_id and cap > 0:
                    prog = f" ({_loyalty_random_received(tg_id)}/{cap})"
                rows.append([InlineKeyboardButton(
                    text=f"🎲 Nhận điểm ngẫu nhiên{prog}",
                    callback_data=f"randpts:{order_key}")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


@router.callback_query(F.data.startswith("randpts:"))
async def on_randpts(cb: CallbackQuery):
    tg_id = cb.from_user.id
    key = (cb.data or "").split(":", 1)[1].strip()[:40]
    cfg = _loyalty_random_cfg()
    if not cfg or not key:
        await cb.answer("Tính năng đang tắt.", show_alert=True)
        return
    lo, hi, cap, min_order, daily_max = cfg
    # Chống cày điểm: kiểm tra lại điều kiện lúc bấm (đơn tối thiểu + lượt/ngày)
    if min_order > 0 and _order_total_for_randkey(key) < min_order:
        await cb.answer(
            f"Đơn này chưa đủ {vnd(min_order)} để nhận điểm ngẫu nhiên.",
            show_alert=True)
        return
    if daily_max > 0 and _loyalty_random_today(tg_id) >= daily_max:
        await cb.answer(
            f"Hôm nay bạn đã nhận đủ {daily_max} lượt điểm ngẫu nhiên rồi!",
            show_alert=True)
        return
    import random as _random
    with db._lock:
        c = db.get_conn()
        dup = c.execute("SELECT 1 FROM loyalty_history WHERE tg_id=? AND reason=?",
                        (tg_id, f"randpts:{key}")).fetchone()
        r = c.execute("SELECT COALESCE(SUM(delta),0) AS s FROM loyalty_history "
                      "WHERE tg_id=? AND reason LIKE 'randpts:%'", (tg_id,)).fetchone()
        try:
            received = int(r["s"] or 0)
        except Exception:
            received = 0
        if dup:
            outcome, pts, total = "dup", 0, 0
        elif received >= cap:
            outcome, pts, total = "capped", 0, 0
        else:
            pts = min(_random.randint(lo, hi), cap - received)
            total = db.loyalty_add(tg_id, pts, f"randpts:{key}")
            outcome = "ok"
    if outcome == "dup":
        await cb.answer("Đơn này bạn đã nhận điểm rồi!", show_alert=True)
        return
    if outcome == "capped":
        await cb.answer(f"Bạn đã nhận đủ tối đa {cap} điểm ngẫu nhiên!",
                        show_alert=True)
        return
    try:
        await cb.message.edit_reply_markup(reply_markup=_spin_kb())
    except Exception:
        pass
    new_recv = received + pts
    await cb.answer(f"🎲 Bạn nhận được {pts} điểm ngẫu nhiên! ({new_recv}/{cap})",
                    show_alert=True)
    try:
        await cb.message.answer(
            f"🎲 <b>ĐIỂM NGẪU NHIÊN</b>\n"
            f"Chúc mừng! Bạn nhận được <b>{pts}</b> điểm loyalty "
            f"(đã nhận <b>{new_recv}/{cap}</b> điểm ngẫu nhiên, "
            f"tổng: <b>{total}</b> điểm) — đủ điểm đổi acc miễn phí bằng /doiqua.",
            parse_mode="HTML")
    except Exception:
        pass


async def _do_spin(msg, tg_id: int, tickets: int):
    if not db.spin_consume_ticket(tg_id):
        await msg.answer("😅 Vé vừa được dùng hết. Mua acc ở /shop để nhận thêm vé nhé!")
        return
    prize = db.spin_roll()
    db.spin_award(tg_id, prize)
    left = db.spin_get_tickets(tg_id)
    label = prize.get("label", "")
    kind = prize.get("kind", "none")
    value = int(prize.get("value", 0) or 0)
    if kind == "none" or value <= 0:
        txt = f"🎡 <b>VÒNG QUAY MAY MẮN</b>\n\n{label}\n\nCòn <b>{left}</b> vé — chúc bạn may mắn lần sau!"
    else:
        txt = (f"🎡 <b>VÒNG QUAY MAY MẮN</b>\n\n"
               f"🎉 <b>TRÚNG: {label}!</b>\n"
               f"Quà đã cộng vào tài khoản của bạn.\n\n"
               f"Còn <b>{left}</b> vé.")
    kb = None
    if left > 0:
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🎡 Quay tiếp", callback_data="spin_now")],
        ])
    await msg.answer(txt, parse_mode="HTML", reply_markup=kb)


@router.message(Command("hang"))
async def on_hang(msg: Message):
    """Xem hạng thành viên shop acc."""
    t = db.member_tier_info(msg.from_user.id)
    lines = ["🏅 <b>HẠNG THÀNH VIÊN SHOP ACC</b>", "━━━━━━━━━━━━",
             f"Hạng hiện tại: <b>{t['tier']}</b>",
             f"Tổng đã mua: <b>{vnd(t['spent'])}</b>"]
    if t["pct"]:
        lines.append(f"Ưu đãi: <b>giảm {t['pct']}%</b> mọi đơn mua acc")
    lines += ["",
              f"🥈 Bạc: mua từ {vnd(t['silver_min'])} — giảm {t['silver_pct']}%",
              f"🥇 Vàng: mua từ {vnd(t['gold_min'])} — giảm {t['gold_pct']}%"]
    if t["next_tier"]:
        need = t["next_min"] - t["spent"]
        lines.append(f"\n👉 Mua thêm <b>{vnd(need)}</b> để lên {t['next_tier']}!")
    else:
        lines.append("\n🌟 Bạn đang ở hạng cao nhất!")
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("doiqua"))
async def on_doiqua(msg: Message):
    """Đổi điểm loyalty lấy acc miễn phí."""
    tg_id = msg.from_user.id
    try:
        need = int(db.get_setting("loyalty_redeem_points", "10") or 10)
    except Exception:
        need = 10
    pts = db.loyalty_get(tg_id)
    cat_id = db.loyalty_redeem_cat()
    c = db.acc_category_get(cat_id) if cat_id else None
    if not c or not c["active"]:
        await msg.answer(
            "🎁 <b>ĐỔI QUÀ LOYALTY</b>\n\n"
            "Shop chưa mở quà đổi điểm. Quay lại sau nhé!",
            parse_mode="HTML")
        return
    c = dict(c)
    if pts < need:
        await msg.answer(
            f"🎁 <b>ĐỔI QUÀ LOYALTY</b>\n\n"
            f"Điểm của bạn: <b>{pts}</b> / cần <b>{need}</b> điểm.\n"
            f"Quà: 1 acc <b>{html.escape(c['name'])}</b> miễn phí.\n\n"
            f"👉 Mua acc ở /shop để tích thêm điểm (1 điểm / 100k).",
            parse_mode="HTML")
        return
    if db.acc_stock_count(cat_id) <= 0:
        await msg.answer("😅 Quà đổi điểm hiện hết hàng, bạn quay lại sau nhé!")
        return
    if not db.loyalty_consume(tg_id, need):
        await msg.answer("😅 Điểm của bạn vừa thay đổi, thử lại nhé!")
        return
    sold_orders, _sell_fail = await _sell_live_stock(
        msg.bot, tg_id, 0, 1,
        lambda seen: db.acc_stock_pick_candidates(cat_id, 5, seen),
        check_live=_cat_live_check(cat_id))
    if not sold_orders:
        db.loyalty_add(tg_id, need, "hoan_diem_khong_du_hang_live")
        await msg.answer(
            "😔 Quà đổi điểm hiện hết acc <b>LIVE</b> (shop vừa kiểm tra lại), "
            "điểm đã được hoàn lại. Bạn quay lại sau nhé!",
            parse_mode="HTML")
        return
    order = sold_orders[0]
    order_id = order["id"]
    tickets = db.spin_add_tickets(tg_id, 1)
    await msg.answer(
        f"🎉 <b>ĐỔI QUÀ THÀNH CÔNG!</b> (−{need} điểm)\n"
        f"🎡 +1 vé quay may mắn (đang có {tickets} vé — gõ /quay)\n"
        f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
        f"{_live_line(cat_id)}\n"
        f"{_pickup_suffix()}",
        parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
    await _notify_purchase_admin(msg.bot, msg.from_user, [order])


@router.message(Command("damua"))
async def on_damua(msg: Message):
    orders = db.acc_user_orders(msg.from_user.id, 10)
    if not orders:
        await msg.answer("🧾 Bạn chưa mua acc nào. Xem shop: /shop")
        return
    t = db.member_tier_info(msg.from_user.id)
    pts = db.loyalty_get(msg.from_user.id)
    lines = ["🧾 <b>ACC ĐÃ MUA</b>", "━━━━━━━━━━━━",
             f"🏅 Hạng: <b>{t['tier']}</b> — ⭐ Điểm loyalty: <b>{pts}</b>",
             f"💰 Tổng đã mua: <b>{vnd(t['spent'])}</b>", ""]
    kb_rows = []
    oids = [dict(o)["id"] for o in orders]
    totp_orders = set()
    if oids:
        try:
            for r in db.get_conn().execute(
                    f"SELECT o.id FROM acc_orders o WHERE o.id IN ({','.join('?' * len(oids))}) "
                    "AND o.totp IS NOT NULL AND o.totp != ''",
                    tuple(oids)).fetchall():
                totp_orders.add(int(r["id"]))
        except Exception:
            pass
    for o in orders:
        o = dict(o)
        lines.append(
            f"#{o['id']} — <b>{html.escape(o['cat_name'])}</b> — {vnd(o['price'])} "
            f"({vn_time_str(ts=o['created_at'])})"
        )
        kb_rows.append([InlineKeyboardButton(
            text=f"🛡 Bảo hành đơn #{o['id']}",
            callback_data=f"accwarranty:{o['id']}")])
        if o["id"] in totp_orders:
            kb_rows.append([InlineKeyboardButton(
                text=f"🔑 Mã 2FA đơn #{o['id']}",
                callback_data=f"acc2fa:{o['id']}")])
        # 4.11 Đánh giá có thưởng (mỗi đơn 1 lần)
        if not db.acc_order_reviewed(o["id"]):
            kb_rows.append([InlineKeyboardButton(
                text=f"⭐ Đánh giá đơn #{o['id']} (+credits)",
                callback_data=f"accreview:{o['id']}")])
    lines += ["", "Acc die trong thời gian bảo hành thì bấm nút bên dưới."]
    await msg.answer("\n".join(lines), parse_mode="HTML",
                     reply_markup=InlineKeyboardMarkup(inline_keyboard=kb_rows))


@router.callback_query(F.data.startswith("acc2fa:"))
async def on_acc_2fa(cb: CallbackQuery):
    """Khách lấy mã 2FA hiện tại của đơn mình đã mua."""
    try:
        order_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    order = db.acc_get_order(order_id)
    if not order or int(order["tg_id"]) != cb.from_user.id:
        await cb.answer("❌ Không tìm thấy đơn hàng.", show_alert=True)
        return
    order = dict(order)
    code = _totp_now(order.get("totp") or "")
    if not code:
        await cb.answer("❌ Đơn này không có 2FA.", show_alert=True)
        return
    import time as _t
    remain = 30 - int(_t.time() % 30)
    await cb.answer()
    await cb.message.answer(
        f"🔐 <b>MÃ 2FA — ĐƠN #{order_id}</b>\n"
        f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
        f"🔢 Mã hiện tại: <code>{code}</code> <i>(hết hạn sau {remain}s)</i>\n"
        f"⚠️ Mã đổi 30s/lần — cần mã mới thì bấm lại nút này.",
        parse_mode="HTML",
    )


def _acc_delivery_kb(order_id: int) -> InlineKeyboardMarkup:
    """Bàn phím chọn cách nhận acc: hiện thông tin để sao chép / tải file."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📋 Hiện thông tin (sao chép)",
                              callback_data=f"accshow:{order_id}")],
        [InlineKeyboardButton(text="📄 Tải file .txt",
                              callback_data=f"accfile:{order_id}:txt"),
         InlineKeyboardButton(text="📊 Tải file .xlsx",
                              callback_data=f"accfile:{order_id}:xlsx")],
    ])


@router.callback_query(F.data.startswith("accshow:"))
async def on_acc_show(cb: CallbackQuery):
    """Hiện toàn bộ thông tin acc để sao chép."""
    await cb.answer()
    try:
        order_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    order = db.acc_get_order(order_id)
    if not order or int(order["tg_id"]) != cb.from_user.id:
        await cb.answer("❌ Không tìm thấy đơn hàng.", show_alert=True)
        return
    await cb.message.answer(_acc_delivery_caption(dict(order)), parse_mode="HTML")


@router.callback_query(F.data.startswith("accfile:"))
async def on_acc_file(cb: CallbackQuery):
    """Xuất acc đã mua ra file .txt / .xlsx."""
    await cb.answer()
    try:
        _, oid, fmt = cb.data.split(":")
        order_id = int(oid)
    except Exception:
        return
    order = db.acc_get_order(order_id)
    if not order or int(order["tg_id"]) != cb.from_user.id:
        await cb.answer("❌ Không tìm thấy đơn hàng.", show_alert=True)
        return
    order = dict(order)
    uid = (order.get("uid") or "").strip()
    fields = [order.get("uid") or "", order.get("password") or "",
              order.get("created_date") or "", order.get("backup_mail") or "",
              order.get("note") or "", order.get("totp") or "",
              order.get("cookie") or "", order.get("token") or ""]
    if fmt == "xlsx":
        import openpyxl
        import io as _io
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "acc"
        ws.append(["UID", "MK", "Ngày tạo", "Mail thay", "Ghi chú",
                   "2FA", "Cookie", "Token"])
        ws.append(fields)
        buf = _io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        fname = f"acc_don_{order_id}.xlsx"
    else:
        data = "|".join(fields).encode("utf-8")
        fname = f"acc_don_{order_id}.txt"
    await cb.message.answer_document(
        BufferedInputFile(data, filename=fname),
        caption=f"📥 File acc đơn <b>#{order_id}</b> — "
                f"UID <code>{html.escape(uid)}</code>",
        parse_mode="HTML")


@router.callback_query(F.data.startswith("accwarranty:"))
async def on_acc_warranty(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    try:
        order_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    order = db.acc_get_order(order_id)
    if not order or int(order["tg_id"]) != cb.from_user.id:
        await cb.message.answer("❌ Không tìm thấy đơn hàng.")
        return
    order = dict(order)
    wh = int(order["warranty_hours"] or 0)  # phút
    age_m = (now() - int(order["created_at"])) / 60
    if age_m > wh:
        await cb.message.answer(
            f"❌ Đơn <b>#{order_id}</b> đã hết thời gian bảo hành ({_fmt_warranty(wh)}).",
            parse_mode="HTML",
        )
        return
    if db.acc_warranty_has_pending(order_id):
        await cb.message.answer("⏳ Đơn này đang chờ admin xử lý bảo hành rồi.")
        return
    # Bot tự check UID để hỗ trợ admin (không tự đổi acc nữa)
    check_txt = "không check được"
    try:
        from . import fb as fb_mod
        r = await fb_mod.check_uid(str(order["uid"]))
        st = str(r.get("status", "")).lower()
        check_txt = "DIE" if st in ("die", "dead") else ("LIVE" if st == "live" else st.upper())
    except Exception:
        pass
    await state.set_state(WarrantyClaimState.waiting_for_evidence)
    await state.update_data(order_id=order_id, check_txt=check_txt)
    await cb.message.answer(
        "🛡 <b>BẢO HÀNH — CẦN BẰNG CHỨNG</b>\n\n"
        "Chỉ bảo hành trường hợp <b>log sai mật khẩu</b> được giao.\n"
        "📸 Vui lòng gửi <b>ảnh chụp màn hình</b> lỗi đăng nhập sai mk để admin xem xét.\n\n"
        "Gõ /huy để hủy yêu cầu.",
        parse_mode="HTML",
    )


@router.message(WarrantyClaimState.waiting_for_evidence)
async def on_warranty_evidence(msg: Message, state: FSMContext):
    if msg.text and msg.text.strip().lower() == "/huy":
        await state.clear()
        await msg.answer("Đã hủy yêu cầu bảo hành.")
        return
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("\U0001f6ab \u0110\u00e3 h\u1ee7y thao t\u00e1c \u0111ang nh\u1eadp. B\u1ea1n g\u00f5 l\u1ea1i l\u1ec7nh v\u1eeba r\u1ed3i nh\u00e9.")
        return
    if not msg.photo:
        await msg.answer(
            "📸 Vui lòng gửi <b>ảnh chụp màn hình</b> lỗi log sai mk (gõ /huy để hủy).",
            parse_mode="HTML")
        return
    data = await state.get_data()
    order_id = int(data.get("order_id") or 0)
    check_txt = data.get("check_txt") or ""
    await state.clear()
    order = db.acc_get_order(order_id)
    if not order or int(order["tg_id"]) != msg.from_user.id:
        await msg.answer("❌ Không tìm thấy đơn hàng.")
        return
    order = dict(order)
    claim_id = db.acc_warranty_claim(order_id, msg.from_user.id, 0, check_txt)
    if claim_id == -1:
        await msg.answer("⏳ Đơn này đang chờ admin xử lý bảo hành rồi.")
        return
    db.acc_warranty_set_evidence(claim_id, msg.photo[-1].file_id)
    # Chống lạm dụng bảo hành: quá số lần/tuần -> NEEDS_REVIEW
    try:
        maxw = int(db.get_setting("warranty_max_week", "3") or 3)
    except Exception:
        maxw = 3
    week_cnt = db.acc_warranty_week_count(msg.from_user.id)
    warn = ""
    if week_cnt > maxw:
        db.acc_warranty_set_status(claim_id, "NEEDS_REVIEW")
        warn = (f"\n\n⚠️ Bạn đã gửi <b>{week_cnt}</b> yêu cầu trong 7 ngày "
                f"(giới hạn {maxw}) — yêu cầu này sẽ được admin kiểm tra kỹ.")
    uname = (msg.from_user.username or msg.from_user.full_name or "").strip()
    caption = (
        f"🛡 <b>Yêu cầu bảo hành #{claim_id}</b>\n"
        f"👤 {html.escape(uname)} (<code>{msg.from_user.id}</code>)\n"
        f"🧾 Đơn #{order_id} — {html.escape(order['cat_name'])} — {vnd(order['price'])}\n"
        f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
        f"🔍 Bot tự check: <b>{html.escape(check_txt)}</b>\n"
        f"📸 Ảnh bằng chứng log sai mk đính kèm.\n\n"
        f"Xem: /bhdon | Xong: /bhdone {claim_id}"
    )
    await _notify_admin_photo(msg.bot, msg.photo[-1].file_id, caption)
    await msg.answer(
        f"✅ <b>Đã ghi nhận yêu cầu bảo hành #{claim_id}</b>\n\n"
        f"🧾 Đơn: <b>#{order_id}</b> — {html.escape(order['cat_name'])}\n"
        f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n\n"
        f"Admin sẽ xem ảnh bằng chứng và xử lý sớm nhất.{warn}",
        parse_mode="HTML",
    )

def _strip_acc(s: str) -> str:
    """Bỏ dấu tiếng Việt để so khớp đơn vị."""
    import unicodedata as _ud
    t = _ud.normalize("NFD", s or "")
    t = "".join(c for c in t if _ud.category(c) != "Mn")
    return t.replace("đ", "d").replace("Đ", "D")


def _parse_warranty(text: str):
    """Parse th\u1eddi gian b\u1ea3o h\u00e0nh: 30p | 24h | 2 ng\u00e0y | 1 tu\u1ea7n | 24 (m\u1eb7c \u0111\u1ecbnh = gi\u1edd). Tr\u1ea3 v\u1ec1 s\u1ed1 PH\u00daT, None n\u1ebfu sai."""
    t = _strip_acc((text or "").strip().lower()).replace(",", ".")
    m = re.match(r"^(\d+(?:\.\d+)?)\s*([a-z ]*)$", t)
    if not m:
        return None
    num = float(m.group(1))
    unit = m.group(2).strip()
    if unit in ("", "h", "g", "gio", "tieng"):
        mins = num * 60
    elif unit in ("p", "m", "ph", "phut", "min", "mins"):
        mins = num
    elif unit in ("d", "ngay", "day", "days"):
        mins = num * 1440
    elif unit in ("w", "tuan", "week", "weeks"):
        mins = num * 10080
    else:
        return None
    mins = int(round(mins))
    return mins if mins > 0 else None


def _fmt_warranty(mins) -> str:
    """Hi\u1ec3n th\u1ecb th\u1eddi gian b\u1ea3o h\u00e0nh (ph\u00fat) d\u1ea1ng \u0111\u1eb9p: 30 ph\u00fat | 24h | 1h30."""
    try:
        m = int(mins or 0)
    except Exception:
        m = 0
    if m <= 0:
        return "0"
    if m < 60:
        return f"{m} ph\u00fat"
    if m % 60 == 0:
        return f"{m // 60}h"
    return f"{m // 60}h{m % 60}"


@router.message(Command("themloai"))
async def on_themloai(msg: Message):
    """Thêm loại acc. Cú pháp: /themloai Tên loại | giá | bảo_hành (30p|24h|2 ngày|1 tuần) | mô tả"""
    if not _is_admin(msg.from_user.id):
        return
    raw = (msg.text or "").split(maxsplit=1)
    if len(raw) < 2 or "|" not in raw[1]:
        await msg.answer(
            "⚠️ Cú pháp: <code>/themloai Tên loại | giá | bảo_hành | mô tả</code>\n"
            "Bảo hành: <code>30p</code> (30 phút) | <code>24h</code> | <code>2 ngày</code> | <code>1 tuần</code>\n"
            "VD: <code>/themloai Via Việt | 25000 | 24h | Via VN 50-500 bạn</code>",
            parse_mode="HTML",
        )
        return
    parts = [p.strip() for p in raw[1].split("|")]
    try:
        name = parts[0]
        price = int(parts[1].replace(".", "").replace(",", "").replace(" ", ""))
        desc = parts[3] if len(parts) > 3 else ""
    except Exception:
        await msg.answer("❌ Giá phải là số.")
        return
    wh = _parse_warranty(parts[2])
    if wh is None:
        await msg.answer("❌ Bảo hành không hợp lệ. Nhập dạng: <code>30p</code> (30 phút) | <code>24h</code> | <code>2 ngày</code> | <code>1 tuần</code>",
                         parse_mode="HTML")
        return
    cid = db.acc_category_add(name, price, wh, desc)
    if cid == -1:
        await msg.answer("❌ Tên loại này đã tồn tại.")
        return
    await msg.answer(
        f"✅ Đã thêm loại acc <b>#{cid} — {html.escape(name)}</b>\n"
        f"💰 Giá: {vnd(price)} | 🛡 BH: {_fmt_warranty(wh)}\n"
        f"Nhập hàng: <code>/themacc {cid}</code>",
        parse_mode="HTML",
    )


def _cat_live_check(cat_id: int) -> bool:
    """Loại acc này có cần check LIVE trước khi giao không?
    Chỉ sạp 'Acc Facebook' (live_check=1) mới check."""
    try:
        c = db.acc_category_get(cat_id)
        if c and "live_check" in c.keys():
            return int(c["live_check"] or 0) == 1
    except Exception:
        pass
    return True


def _live_line(cat_id: int) -> str:
    """Dòng 'Đã kiểm tra LIVE' trong tin giao hàng — chỉ hiện khi có check thật."""
    return "🟢 <i>Đã kiểm tra LIVE trước khi giao</i>\n" if _cat_live_check(cat_id) else ""


@router.message(Command("themstall"))
async def on_themstall(msg: Message):
    """Thêm gian hàng mới + loại hàng đầu tiên.
    Cú pháp: /themstall Tên gian hàng | Tên loại | giá | bảo_hành | mô tả
    VD: /themstall Gmail | Gmail Edu | 15000 | 24h | Gmail edu giá rẻ"""
    if not _is_admin(msg.from_user.id):
        return
    raw = (msg.text or "").split(maxsplit=1)
    if len(raw) < 2 or "|" not in raw[1]:
        await msg.answer(
            "⚠️ Cú pháp: <code>/themstall Tên gian hàng | Tên loại | giá | bảo_hành | mô tả</code>\n"
            "VD: <code>/themstall Gmail | Gmail Edu | 15000 | 24h | Gmail edu</code>\n"
            "<i>Gian hàng mới KHÔNG check live/die (chỉ sạp Acc Facebook mới check).</i>",
            parse_mode="HTML",
        )
        return
    parts = [p.strip() for p in raw[1].split("|")]
    if len(parts) < 4:
        await msg.answer("❌ Thiếu tham số. Cần: <code>Tên gian hàng | Tên loại | giá | bảo_hành</code>",
                         parse_mode="HTML")
        return
    try:
        stall, name = parts[0], parts[1]
        price = int(parts[2].replace(".", "").replace(",", "").replace(" ", ""))
        desc = parts[4] if len(parts) > 4 else ""
    except Exception:
        await msg.answer("❌ Giá phải là số.")
        return
    if not stall or not name:
        await msg.answer("❌ Tên gian hàng và tên loại không được trống.")
        return
    wh = _parse_warranty(parts[3])
    if wh is None:
        await msg.answer("❌ Bảo hành không hợp lệ (<code>30p</code>|<code>24h</code>|<code>2 ngày</code>|<code>1 tuần</code>).",
                         parse_mode="HTML")
        return
    cid = db.acc_category_add(name, price, wh, desc, stall=stall, live_check=0)
    if cid == -1:
        await msg.answer("❌ Tên loại này đã tồn tại.")
        return
    # Tự tạo tab Sheet cho gian hàng mới (chạy nền, lỗi thì báo sau)
    sid = db.get_setting("sheet_import_id") or ""
    default_tab = db.get_setting("sheet_import_tab") or "NhapKho"
    from . import sheet_import as _si
    tab = _si.tab_for_stall(stall, default_tab)

    async def _mk_tab():
        try:
            ok = await _si.ensure_tab(sid, tab) if sid else False
            if not ok and sid:
                await msg.answer(f"⚠️ Đã tạo gian hàng nhưng <b>chưa tạo được tab Sheet</b> "
                                 f"(<code>{html.escape(tab)}</code>). Vào /shopadm → Nhập kho Sheet để thử lại.",
                                 parse_mode="HTML")
        except Exception as e:
            log.warning("themstall ensure_tab lỗi: %s", e)

    await msg.answer(
        f"🏪 Đã tạo gian hàng <b>{html.escape(stall)}</b>\n"
        f"✅ Loại hàng <b>#{cid} — {html.escape(name)}</b> | 💰 {vnd(price)} | 🛡 {_fmt_warranty(wh)}\n"
        f"📊 Tab Sheet: <code>{html.escape(tab)}</code> (đang tạo...)\n"
        f"Nhập hàng: <code>/themacc {cid}</code> hoặc /shopadm → 📊 Nhập kho Sheet",
        parse_mode="HTML",
    )
    asyncio.create_task(_mk_tab())


@router.message(Command("themacc"))
async def on_themacc(msg: Message, state: FSMContext):
    """Nhập kho: /themacc <id_loại> [ncc_id] [giá_vốn] rồi gửi file .txt"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/themacc &lt;id_loại&gt; [ncc_id] [giá_vốn]</code> (xem id: /kho)",
                         parse_mode="HTML")
        return
    try:
        cat_id = int(parts[1])
        ncc_id = int(parts[2]) if len(parts) > 2 else 0
        cost = int(parts[3].replace(".", "").replace(",", "")) if len(parts) > 3 else 0
    except Exception:
        await msg.answer("❌ ID loại/NCC/giá vốn phải là số.")
        return
    c = db.acc_category_get(cat_id)
    if not c:
        await msg.answer("❌ Không có loại acc này. Xem: /kho")
        return
    if ncc_id and not db.supplier_get(ncc_id):
        await msg.answer("❌ Không có NCC này. Xem: /ncc")
        return
    await state.update_data(acc_cat_id=cat_id, acc_ncc_id=ncc_id, acc_cost=cost)
    await state.set_state(AccShopState.waiting_for_stock_file)
    extra = ""
    if ncc_id:
        extra += f"\n🏭 NCC: <b>#{ncc_id}</b>"
    if cost:
        extra += f" — 💰 giá vốn <b>{vnd(cost)}</b>/acc"
    await msg.answer(
        f"📤 Gửi file <b>.txt</b> hoặc <b>.xlsx</b> chứa acc loại <b>{html.escape(c['name'])}</b>.{extra}\n"
        f"Mỗi dòng 1 acc — .txt dùng định dạng:\n"
        f"<code>uid|mk|ngày tạo|mail thay|ghi chú|2fa|cookie|token</code>\n"
        f"💡 Cột đầu có thể là <b>link Facebook</b> (profile/share/fb.watch) — bot tự giải ra UID.\n"
        f"📊 .xlsx: 8 cột đầu theo thứ tự trên (dòng đầu có thể là tiêu đề).\n"
        f"🔑 Cột <b>mk</b> để trống → tự điền <code>khai2006</code>.",
        parse_mode="HTML",
    )


def _smart_stock_fields(fields: list) -> dict:
    """Nhận diện trường acc từ 1 dòng .txt theo MẪU nội dung (không phụ thuộc thứ tự cột):
    link FB | email | mã 2FA (base32) | ngày tạo | cookie | token EAAG | UID số.
    Trường còn lại: trường đầu -> mật khẩu, các trường sau -> ghi chú."""
    import re as _re
    r = {"uid": "", "password": "", "created_date": "", "backup_mail": "",
         "note": "", "totp": "", "cookie": "", "token": ""}
    others = []
    for _f in fields:
        f = (_f or "").strip()
        if not f:
            continue
        fl = f.lower()
        if "facebook.com" in fl or "fb.com" in fl or "fb.watch" in fl:
            if not r["uid"]:
                r["uid"] = f
            else:
                others.append(f)
        elif _re.search(r"\S+@\S+\.\S+", f):
            if not r["backup_mail"]:
                r["backup_mail"] = f
            else:
                others.append(f)
        elif (f.startswith("EAAG") or f.startswith("EAAB")) and not r["token"]:
            r["token"] = f
        elif ("c_user" in fl or ("=" in f and len(f) > 80)) and not r["cookie"]:
            r["cookie"] = f
        elif (_re.fullmatch(r"[A-Z2-7][A-Z2-7 ]+", f.upper())
              and not r["totp"]
              and ((" " in f and len(_re.sub(r"\s+", "", f)) >= 12)
                   or len(f) >= 16)):
            r["totp"] = _re.sub(r"\s+", "", f).upper()
        elif (_re.search(r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2}", f)
              and not r["created_date"]):
            r["created_date"] = f
        elif f.isdigit() and len(f) >= 6 and not r["uid"]:
            r["uid"] = f
        else:
            others.append(f)
    if others:
        r["password"] = others[0]
        if len(others) > 1:
            r["note"] = " | ".join(others[1:])
    return r



async def _import_stock_rows(rows, cat_id, ncc_id, cost, c, msg, wait, sheet_ctx=None):
    """Pipeline nhập kho dùng chung cho /themacc (file) và /nhapkhosheet (Google Sheet).
    rows: list dict {uid,password,created_date,backup_mail,note,totp,cookie,token}.
    sheet_ctx: None hoặc {"sheet_id","tab"} — ghi trạng thái ngược vào sheet."""
    # Cột mk trống -> dùng mặc định "khai2006" (các cột khác giữ đúng vị trí)
    for r in rows:
        if not (r.get("password") or "").strip():
            r["password"] = "khai2006"
    # Tự giải link FB ở cột đầu -> UID số (nếu file ghi link thay vì UID)
    link_idx = [i for i, r in enumerate(rows)
                if re.search(r'facebook\.com|fb\.com|fb\.watch|^https?://', (r.get("uid") or ""), re.I)]
    resolved = failed = 0
    failed_lines = []
    if link_idx:
        await wait.edit_text(f"⏳ Đang nhập kho... 🔗 Giải {len(link_idx)} link FB → UID...")
        from .fb import resolve_fb_uid
        _cache = {}
        _sem = asyncio.Semaphore(4)
        async def _res(i):
            nonlocal resolved, failed
            link = (rows[i].get("uid") or "").strip()
            rows[i]["_orig_link"] = link  # giữ link gốc để vòng thử lại dùng
            if link in _cache:
                uid = _cache[link]
            else:
                async with _sem:
                    try:
                        uid, _nm, _md = await resolve_fb_uid(link)
                    except Exception:
                        uid = ""
                _cache[link] = uid
            if uid and str(uid).isdigit():
                rows[i]["uid"] = str(uid)
                resolved += 1
            else:
                rows[i]["uid"] = ""
                failed += 1
                failed_lines.append(f"Dòng {i + 1}: {link[:60]}")
        await asyncio.gather(*[_res(i) for i in link_idx])
        if failed:
            # Vòng 2: thử lại TUẦN TỰ các link lỗi, cách nhau vài giây.
            # (Vòng 1 quét song song dễ bị API giới hạn tốc độ -> báo lỗi oan.)
            try:
                await wait.edit_text(
                    f"🔁 Thử lại {failed} link chưa giải được (chậm để tránh bị giới hạn)...")
            except Exception:
                pass
            still_failed = []
            for i in link_idx:
                if rows[i].get("uid"):
                    continue
                link = (rows[i].get("_orig_link") or "").strip()
                if not link:
                    continue
                try:
                    uid, _nm, _md = await resolve_fb_uid(link)
                except Exception:
                    uid = ""
                if uid and str(uid).isdigit():
                    rows[i]["uid"] = str(uid)
                    resolved += 1
                    failed -= 1
                else:
                    still_failed.append(f"Dòng {i + 1}: {link[:60]}")
                await asyncio.sleep(2)
            failed_lines = still_failed
    # Map uid -> link gốc (hiển thị danh sách acc die sau khi quét)
    uid_to_link = {}
    for i in link_idx:
        if rows[i]["uid"] and str(rows[i]["uid"]).isdigit():
            uid_to_link[rows[i]["uid"]] = (rows[i].get("_orig_link") or "")
    # Chống trùng: bỏ dòng trùng trong file và UID đang còn trong kho
    # (chỉ tính acc AVAILABLE/DIE; acc đã bán SOLD được nhập lại bình thường)
    _seen = set()
    try:
        _existing = {r[0] for r in db.get_conn().execute(
            "SELECT uid FROM acc_stock WHERE cat_id=? AND status IN ('AVAILABLE','DIE')",
            (cat_id,)).fetchall()}
    except Exception:
        _existing = set()
    _uniq, empty, dup_in_file, dup_in_stock = [], 0, 0, 0
    all_rows = rows  # giữ để ghi trạng thái sheet cho cả dòng bị loại
    for r in rows:
        uid = (r.get("uid") or "").strip()
        if not uid:
            empty += 1
            if sheet_ctx is not None:
                r["_sheet_status"] = "🔗 LỖI LINK" if r.get("_orig_link") else "⏭ BỎ QUA"
            continue
        if uid in _seen:
            dup_in_file += 1
            if sheet_ctx is not None:
                r["_sheet_status"] = "🗑 TRÙNG"
            continue
        if uid in _existing:
            dup_in_stock += 1
            if sheet_ctx is not None:
                r["_sheet_status"] = "🗑 TRÙNG"
            continue
        _seen.add(uid)
        _uniq.append(r)
    rows = _uniq
    if sheet_ctx is not None:
        for r in _uniq:
            r["_sheet_status"] = "✅ OK"
        sheet_ctx["uid_to_row"] = {
            str(r["uid"]): r.get("_sheet_row")
            for r in _uniq if str(r.get("uid") or "").isdigit() and r.get("_sheet_row")
        }
    skipped = empty + dup_in_file + dup_in_stock
    stock_since = int(time.time())
    added, _ = db.acc_stock_add_batch(
        cat_id, rows,
        batch="Lô " + time.strftime("%d/%m %H:%M", time.localtime()),
        supplier_id=ncc_id, cost_per_acc=cost)
    total = db.acc_stock_count(cat_id)
    extra_info = ""
    if ncc_id:
        extra_info += f"\n🏭 NCC: <b>#{ncc_id}</b>"
    if cost:
        extra_info += f" — 💰 Vốn: <b>{vnd(cost)}</b>/acc"
    dup_txt = ""
    if dup_in_file or dup_in_stock:
        dup_txt = f"🗑 Trùng loại: {dup_in_file} trong file, {dup_in_stock} đã có trong kho\n"
    link_txt = ""
    if link_idx:
        link_txt = f"\n🔗 Giải link FB: <b>{resolved}</b> OK" + (f", <b>{failed}</b> lỗi" if failed else "")
    await wait.edit_text(
        f"✅ <b>Nhập kho xong:</b> {html.escape(c['name'])}\n"
        f"➕ Thêm: <b>{added}</b> acc | ⏭ Bỏ qua: <b>{skipped}</b> dòng{extra_info}\n"
        f"{dup_txt}"
        f"📊 Tồn kho hiện tại: <b>{total}</b> acc"
        f"{link_txt}"
        "\n\n🔍 Đang tự quét kiểm tra chất lượng...",
        parse_mode="HTML",
    )
    if sheet_ctx is not None:
        # Ghi đánh dấu ngược vào sheet: cột I trạng thái + ghi đè link bằng UID (cách 1)
        try:
            from . import sheet_import as _si
            _updates = []
            for r in all_rows:
                sr = r.get("_sheet_row")
                if not sr:
                    continue
                _updates.append((sr, 9, r.get("_sheet_status") or "⏭ BỎ QUA"))
                if r.get("_orig_link") and str(r.get("uid") or "").isdigit():
                    _updates.append((sr, 1, str(r["uid"])))
            await _si.write_updates(sheet_ctx["sheet_id"], sheet_ctx["tab"], _updates)
        except Exception as e:
            log.warning("sheet write-back: %s", e)
    if failed_lines:
        try:
            txt = ("⚠️ <b>Các link không giải được UID</b> (đã thử lại nhiều lần, đã bỏ qua):\n"
                   "<i>Thường do link đã chết hoặc bài viết không để công khai — kiểm tra lại các link này.</i>\n"
                   ) + "\n".join(
                html.escape(x) for x in failed_lines[:15])
            if len(failed_lines) > 15:
                txt += f"\n...và {len(failed_lines) - 15} dòng nữa"
            await msg.answer(txt, parse_mode="HTML")
        except Exception:
            pass
    # Quét chất lượng nền: acc die -> loại khỏi kho bán
    async def _scan(stock_since: int):
        from . import fb as fb_mod
        import asyncio
        ids = db.get_conn().execute(
            "SELECT id, uid FROM acc_stock WHERE cat_id=? AND status='AVAILABLE' "
            "AND added_at >= ? ORDER BY id", (cat_id, stock_since)).fetchall()
        sem = asyncio.Semaphore(10)
        dead = 0
        dead_list = []
        dead_uids = []
        async def _one(sid, uid):
            nonlocal dead
            async with sem:
                try:
                    r = await fb_mod.check_uid(str(uid))
                    if str(r.get("status", "")).lower() in ("die", "dead"):
                        db.acc_mark_status(sid, "DEAD")
                        dead += 1
                        dead_uids.append(str(uid))
                        lk = (uid_to_link.get(str(uid)) or "").strip()
                        dead_list.append(f"• {html.escape(str(uid))}" + (f" — {html.escape(lk)}" if lk else ""))
                except Exception:
                    pass
        await asyncio.gather(*[_one(r["id"], r["uid"]) for r in ids])
        left = db.acc_stock_count(cat_id)
        dead_txt = ""
        if dead_list:
            dead_txt = "\n⚠️ <b>Danh sách acc die:</b>\n" + "\n".join(dead_list[:20])
            if len(dead_list) > 20:
                dead_txt += f"\n...và {len(dead_list) - 20} acc nữa"
        if sheet_ctx is not None and dead_uids:
            # Đánh dấu DIE ngược vào sheet
            try:
                from . import sheet_import as _si
                _u = [(r, 9, "☠️ DIE") for u, r in sheet_ctx.get("uid_to_row", {}).items()
                      if u in dead_uids]
                await _si.write_updates(sheet_ctx["sheet_id"], sheet_ctx["tab"], _u)
            except Exception as e:
                log.warning("sheet die mark: %s", e)
        await msg.answer(
            f"🔍 <b>Quét xong {len(ids)} acc mới nhập:</b> "
            f"<b>{dead}</b> acc die đã loại khỏi kho.\n"
            f"📊 Tồn kho bán được: <b>{left}</b> acc.{dead_txt}",
            parse_mode="HTML",
        )
        # Báo / tự giao cho những ai đã đăng ký "có hàng nhắn tôi"
        if left > 0:
            try:
                await _fulfill_restock_subs(cat_id, msg.bot, dict(c))
            except Exception as e:
                log.warning("fulfill restock subs: %s", e)
            # 4.9 Giao hàng cho người đã đặt cọc (FIFO)
            try:
                await _fulfill_deposits(cat_id, msg.bot, dict(c))
            except Exception as e:
                log.warning("fulfill deposits: %s", e)
    asyncio.create_task(_scan(int(now()) - 600))

@router.message(AccShopState.waiting_for_stock_file, F.document)
async def on_stock_file(msg: Message, state: FSMContext):
    if not _is_admin(msg.from_user.id):
        await state.clear()
        return
    data = await state.get_data()
    await state.clear()
    cat_id = data.get("acc_cat_id")
    ncc_id = int(data.get("acc_ncc_id") or 0)
    cost = int(data.get("acc_cost") or 0)
    c = db.acc_category_get(cat_id)
    if not c:
        await msg.answer("❌ Loại acc không tồn tại.")
        return
    doc = msg.document
    file_name = (doc.file_name or "").lower()
    if not (file_name.endswith(".txt") or file_name.endswith(".xlsx")):
        await msg.answer("❌ Chỉ nhận file .txt hoặc .xlsx.")
        return
    wait = await msg.answer("⏳ Đang nhập kho...")
    try:
        file_info = await msg.bot.get_file(doc.file_id)
        raw = await msg.bot.download_file(file_info.file_path)
    except Exception as e:
        await wait.edit_text(f"❌ Không đọc được file: {e}")
        return
    rows = []
    if file_name.endswith(".xlsx"):
        # .xlsx: 8 cột đầu = uid/link|mk|ngày tạo|mail thay|ghi chú|2fa|cookie|token
        from openpyxl import load_workbook
        import io as _io
        try:
            wb = load_workbook(_io.BytesIO(raw.getvalue()), read_only=True, data_only=True)
            ws = wb.active
        except Exception as e:
            await wait.edit_text(f"❌ Không đọc được file Excel: {e}")
            return
        first_row = True
        for vals in ws.iter_rows(values_only=True):
            cells = [(str(c).strip() if c is not None else "") for c in vals[:8]]
            while len(cells) < 8:
                cells.append("")
            if not any(cells):
                continue
            c0 = cells[0].lower()
            if first_row:
                first_row = False
                # Bỏ dòng tiêu đề (UID, Link, Url, STT...)
                if not (c0.isdigit() or "http" in c0 or "facebook.com" in c0 or "fb.com" in c0):
                    continue
            rows.append({
                "uid": cells[0], "password": cells[1], "created_date": cells[2],
                "backup_mail": cells[3], "note": cells[4], "totp": cells[5],
                "cookie": cells[6], "token": cells[7],
            })
    else:
        text = raw.getvalue().decode("utf-8", errors="ignore")
        for line in text.splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            p = [x.strip() for x in line.split("|")]
            rows.append(_smart_stock_fields(p))
    if not rows:
        await wait.edit_text("❌ File không có dòng acc nào hợp lệ.")
        return
    await _import_stock_rows(rows, cat_id, ncc_id, cost, c, msg, wait)



@router.message(Command("setsheet"))
async def on_setsheet(msg: Message):
    """Cài đặt Google Sheet dùng để nhập kho: /setsheet <link|id> [tab]"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) == 1:
        cur = db.get_setting("sheet_import_id") or ""
        tab = db.get_setting("sheet_import_tab") or "NhapKho"
        if cur:
            await msg.answer(
                f"📊 <b>Sheet nhập kho hiện tại:</b>\n"
                f"🔗 <code>{html.escape(cur)}</code>\n"
                f"📑 Tab: <b>{html.escape(tab)}</b>\n\n"
                f"Đổi: <code>/setsheet &lt;link_sheet&gt; [tab]</code>",
                parse_mode="HTML")
        else:
            await msg.answer("📊 Chưa cài đặt sheet nhập kho.\n"
                             "Dùng: <code>/setsheet &lt;link_google_sheet&gt; [tab]</code>")
        return
    raw = parts[1]
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9-_]+)", raw)
    sid = m.group(1) if m else raw.strip()
    tab = parts[2].strip() if len(parts) > 2 else (db.get_setting("sheet_import_tab") or "NhapKho")
    db.set_setting("sheet_import_id", sid)
    db.set_setting("sheet_import_tab", tab)
    await msg.answer(
        f"✅ <b>Đã cài đặt sheet nhập kho</b>\n📑 Tab: <b>{html.escape(tab)}</b>\n\n"
        f"Nhập kho: <code>/nhapkhosheet &lt;id_loại&gt; [ncc_id] [giá_vốn]</code>",
        parse_mode="HTML")


@router.message(Command("nhapkhosheet"))
async def on_nhapkhosheet(msg: Message):
    """Nhập kho từ Google Sheet: quét dòng chưa đánh dấu -> giải link -> nhập kho."""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/nhapkhosheet &lt;id_loại&gt; [ncc_id] [giá_vốn]</code> (xem id: /kho)",
                         parse_mode="HTML")
        return
    try:
        cat_id = int(parts[1])
        ncc_id = int(parts[2]) if len(parts) > 2 else 0
        cost = int(parts[3].replace(".", "").replace(",", "")) if len(parts) > 3 else 0
    except Exception:
        await msg.answer("❌ ID loại/NCC/giá vốn phải là số.")
        return
    await _run_sheet_import(msg, cat_id, ncc_id, cost)


@router.message(Command("kho"))
async def on_kho(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    cats = db.acc_category_list(active_only=False, include_hidden=True)
    if not cats:
        await msg.answer("📦 Chưa có loại acc nào. Thêm: /themloai")
        return
    lines = ["📦 <b>KHO ACC</b>", "━━━━━━━━━━━━", ""]
    for c in cats:
        c = dict(c)
        n = db.acc_stock_count(c["id"])
        n_die = db.acc_stock_die_count(c["id"])
        n_sold = db.acc_stock_sold_count(c["id"])
        st = "✅ đang bán" if c["active"] else "⏸ đã ẩn"
        if int(c.get("hidden") or 0):
            st += " [ẨN KHỎI SHOP]"
        bonus = int(c.get("credit_bonus") or 0)
        bonus_txt = f" — 🎁+{bonus}cr" if bonus else ""
        lines.append(
            f"#{c['id']} <b>{html.escape(c['name'])}</b> — {vnd(c['price'])} — "
            f"BH {_fmt_warranty(c['warranty_hours'])} — kho: <b>{n}</b>"
            + (f" 🗑<b>{n_die}</b> die" if n_die else "")
            + (f" 🛒<b>{n_sold}</b> đã bán" if n_sold else "")
            + f"{bonus_txt} — {st}"
        )
    lines += ["", "Đổi giá: <code>/gia &lt;id&gt; &lt;giá&gt;</code>",
              "Đổi giờ BH: <code>/suabh &lt;id&gt; &lt;giờ&gt;</code>",
              "Tặng credits: <code>/creditbonus &lt;id&gt; &lt;số&gt;</code>",
              "Quà đổi điểm: <code>/quadoi &lt;id&gt; [điểm]</code>",
              "Điểm ngẫu nhiên sau mua: <code>/loyaltyrandom &lt;min&gt; &lt;max&gt; &lt;cap_user&gt;</code>",
              "Giờ vàng: <code>/giovang &lt;id|0&gt; [giờ_bd-giờ_kt] [%]</code>",
              "Hộp mù: <code>/hopmu &lt;id&gt;</code> — giá: <code>/hopmugia &lt;giá|0&gt;</code>",
              "NCC: <code>/themncc</code> / <code>/ncc</code>",
              "Nhập hàng: <code>/themacc &lt;id&gt; [ncc_id] [giá_vốn]</code>",
              "Sao lưu toàn kho: <code>/xuatkho [id_loại]</code>"]
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("xuatkho"))
async def on_xuatkho(msg: Message):
    """Admin: xuất toàn bộ acc trong kho ra file .xlsx (sao lưu dự phòng).
    Mỗi loại acc = 1 sheet, 8 cột đúng chuẩn nhập kho, hàng vẫn nằm trong kho."""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    cat_ids = None
    if len(parts) > 1:
        try:
            cat_ids = {int(parts[1])}
        except ValueError:
            await msg.answer("❌ ID loại phải là số (xem id: /kho).")
            return
    cats = db.acc_category_list(active_only=False, include_hidden=True)
    if cat_ids:
        cats = [c for c in cats if int(c["id"]) in cat_ids]
    if not cats:
        await msg.answer("📦 Không có loại acc nào.")
        return
    sheets = {}
    total = 0
    for c in cats:
        c = dict(c)
        rows = db.get_conn().execute(
            "SELECT uid, password, created_date, backup_mail, note, totp, cookie, token "
            "FROM acc_stock WHERE cat_id=? AND status='AVAILABLE' ORDER BY id",
            (c["id"],)).fetchall()
        if not rows:
            continue
        sheet_name = "".join(ch for ch in c["name"] if ch not in r'\/\?*[]:')[:25] or f"Loai_{c['id']}"
        sheets[f"#{c['id']} {sheet_name}"] = (
            ["UID", "MK", "Ngày tạo", "Mail thay", "Ghi chú", "2FA", "Cookie", "Token"],
            [[r["uid"], r["password"], r["created_date"], r["backup_mail"],
              r["note"], r["totp"], r["cookie"], r["token"]] for r in rows],
        )
        total += len(rows)
    if not sheets:
        await msg.answer("📦 Kho đang trống, không có gì để xuất.")
        return
    from aiogram.types import BufferedInputFile
    import datetime
    data = _build_excel_bytes(sheets)
    fname = f"kho_acc_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    await msg.answer_document(
        document=BufferedInputFile(data, filename=fname),
        caption=f"📦 <b>SAO LƯU KHO ACC</b>\n━━━━━━━━━━━━\n"
                f"Tổng: <b>{total}</b> acc / {len(sheets)} loại\n"
                f"⚠️ Hàng vẫn còn trong kho (không bị xóa).\n"
                f"Gửi lại file này vào <code>/themacc</code> để khôi phục kho nếu hệ thống có vấn đề.",
        parse_mode="HTML",
    )


@router.message(Command("gia"))
async def on_gia(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 3:
        await msg.answer("⚠️ Cú pháp: <code>/gia &lt;id_loại&gt; &lt;giá_mới&gt;</code>",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
        price = int(parts[2].replace(".", "").replace(",", ""))
    except Exception:
        await msg.answer("❌ ID và giá phải là số.")
        return
    cat = db.acc_category_get(cid)
    if not cat:
        await msg.answer("❌ Không tìm thấy loại này.")
        return
    old = int(cat["price"] or 0)
    if not _price_change_ok(old, price):
        # Giá chênh lệch bất thường -> hỏi xác nhận trước khi lưu
        kb = InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="✅ Vẫn đổi giá",
                                 callback_data=f"giacf:ok:{cid}:{price}"),
            InlineKeyboardButton(text="❌ Huỷ",
                                 callback_data=f"giacf:no:{cid}:{price}"),
        ]])
        await msg.answer(
            f"⚠️ <b>GIÁ BẤT THƯỜNG!</b>\n\n"
            f"Loại <b>#{cid}</b> {html.escape(cat['name'] or '')}\n"
            f"Giá cũ: <b>{vnd(old)}</b> → Giá mới: <b>{vnd(price)}</b>\n\n"
            f"Chênh lệch quá {_price_warn_pct()}% so với giá hiện tại. "
            f"Bạn có chắc không nhập nhầm?",
            parse_mode="HTML", reply_markup=kb)
        return
    if db.acc_category_update(cid, price=price):
        await msg.answer(f"✅ Loại <b>#{cid}</b> đổi giá thành <b>{vnd(price)}</b>.",
                         parse_mode="HTML")
    else:
        await msg.answer("❌ Không tìm thấy loại này.")


def _price_warn_pct() -> int:
    try:
        return max(1, int(db.get_setting("price_warn_pct", "50") or 50))
    except Exception:
        return 50


def _price_change_ok(old: int, new: int) -> bool:
    """True nếu mức đổi giá bình thường (không cần xác nhận)."""
    if old <= 0 or new < 0:
        return True
    pct = _price_warn_pct()
    lo = old * (100 - pct) / 100
    hi = old * (100 + pct) / 100
    return lo <= new <= hi


@router.callback_query(F.data.startswith("giacf:"))
async def on_giacf_cb(cb: CallbackQuery):
    """Xác nhận đổi giá khi giá mới chênh lệch bất thường."""
    if not (_perms.is_super(cb.from_user.id)
            or _perms.has_perm(cb.from_user.id, "price")):
        await cb.answer("🚫 Bạn không có quyền đổi giá.", show_alert=True)
        return
    parts = (cb.data or "").split(":")
    if len(parts) != 4:
        await cb.answer()
        return
    _, verdict, cid_s, price_s = parts
    try:
        cid, price = int(cid_s), int(price_s)
    except Exception:
        await cb.answer("Lỗi dữ liệu.", show_alert=True)
        return
    if verdict == "ok":
        cat = db.acc_category_get(cid)
        old = int(cat["price"] or 0) if cat else 0
        if db.acc_category_update(cid, price=price):
            db.admin_audit_add(cb.from_user.id, cb.from_user.full_name,
                               "doi_gia_xac_nhan",
                               f"loai #{cid} {vnd(old)} -> {vnd(price)}")
            await cb.message.edit_text(
                f"✅ Loại <b>#{cid}</b> đổi giá thành <b>{vnd(price)}</b> "
                f"(đã xác nhận giá bất thường).",
                parse_mode="HTML")
        else:
            await cb.message.edit_text("❌ Không tìm thấy loại này.")
    else:
        await cb.message.edit_text("❌ Đã huỷ đổi giá.")
    await cb.answer()


@router.message(Command("setmailapp"))
async def on_setmailapp(msg: Message):
    """Admin: đặt link tải app mail ảo hiện cho khách sau khi mua acc.
    /setmailapp <link> — đặt link mới | /setmailapp xoa — tắt | không tham số — xem link hiện tại."""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split(None, 1)
    if len(parts) < 2:
        cur = _mail_app_link()
        await msg.answer(
            "📧 <b>Link app mail hiện tại:</b>\n"
            f"{html.escape(cur) if cur else '<i>chưa đặt (khách sẽ không thấy dòng tải app)</i>'}\n\n"
            "Đổi link mới: <code>/setmailapp &lt;link&gt;</code>\n"
            "Tắt hẳn: <code>/setmailapp xoa</code>",
            parse_mode="HTML")
        return
    arg = parts[1].strip()
    if arg.lower() in ("xoa", "xóa", "off", "none"):
        db.set_setting("mail_app_link", "")
        await msg.answer("✅ Đã tắt link app mail — khách sẽ không thấy dòng tải app nữa.")
        return
    if not arg.lower().startswith(("http://", "https://")):
        await msg.answer("❌ Link phải bắt đầu bằng http:// hoặc https://")
        return
    db.set_setting("mail_app_link", arg)
    await msg.answer(f"✅ Đã đặt link app mail mới:\n{html.escape(arg)}", parse_mode="HTML")


@router.message(Command("anhbia"))
async def on_anhbia(msg: Message, state: FSMContext):
    """Đặt ảnh bìa cho loại acc: /anhbia <id> rồi gửi ảnh; /anhbia <id> xoa để gỡ."""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/anhbia &lt;id_loại&gt;</code> rồi gửi ảnh bìa.\n"
                         "Gỡ ảnh: <code>/anhbia &lt;id_loại&gt; xoa</code>", parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID loại phải là số.")
        return
    c = db.acc_category_get(cid)
    if not c:
        await msg.answer("❌ Không tìm thấy loại này.")
        return
    if len(parts) >= 3 and parts[2].lower() in ("xoa", "xóa", "del"):
        db.acc_category_update(cid, cover_photo="")
        await msg.answer(f"🗑 Đã gỡ ảnh bìa của <b>{html.escape(c['name'])}</b>.", parse_mode="HTML")
        return
    await state.update_data(cover_cat_id=cid)
    await state.set_state(AccShopState.waiting_for_cover)
    await msg.answer(f"📸 Gửi ảnh bìa cho <b>{html.escape(c['name'])}</b> (gửi 1 ảnh bất kỳ).",
                     parse_mode="HTML")


@router.message(AccShopState.waiting_for_cover, F.photo)
async def on_anhbia_photo(msg: Message, state: FSMContext):
    data = await state.get_data()
    cid = data.get("cover_cat_id")
    await state.clear()
    if not _is_admin(msg.from_user.id):
        return
    if not cid or not db.acc_category_get(cid):
        await msg.answer("❌ Loại acc không còn tồn tại.")
        return
    file_id = msg.photo[-1].file_id
    db.acc_category_update(cid, cover_photo=file_id)
    await msg.answer(f"✅ Đã đặt ảnh bìa cho loại <b>#{cid}</b>. Khách vào xem chi tiết sẽ thấy ảnh này.",
                     parse_mode="HTML")


@router.message(Command("suabh"))
async def on_suabh(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 3:
        await msg.answer("⚠️ Cú pháp: <code>/suabh &lt;id_loại&gt; &lt;bảo_hành: 30p|24h|2 ngày|1 tuần&gt;</code>",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    wh = _parse_warranty(parts[2])
    if wh is None:
        await msg.answer("❌ Bảo hành không hợp lệ. Nhập dạng: <code>30p</code> | <code>24h</code> | <code>2 ngày</code> | <code>1 tuần</code>",
                         parse_mode="HTML")
        return
    if db.acc_category_update(cid, warranty_hours=wh):
        await msg.answer(f"✅ Loại <b>#{cid}</b> đổi bảo hành thành <b>{_fmt_warranty(wh)}</b>.",
                         parse_mode="HTML")
    else:
        await msg.answer("❌ Không tìm thấy loại này.")


@router.message(Command("creditbonus"))
async def on_creditbonus(msg: Message):
    """Admin: tặng credits khi khách mua acc. Cú pháp: /creditbonus <id_loại> <số_credits/acc>"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 3:
        await msg.answer("⚠️ Cú pháp: <code>/creditbonus &lt;id_loại&gt; &lt;số_credits&gt;</code>\n"
                         "VD: <code>/creditbonus 3 50</code> — mua 1 acc loại #3 tặng 50 credits.",
                         parse_mode="HTML")
        return
    try:
        cid, n = int(parts[1]), int(parts[2])
    except Exception:
        await msg.answer("❌ ID và số credits phải là số.")
        return
    if n < 0:
        await msg.answer("❌ Số credits không được âm.")
        return
    if db.acc_category_update(cid, credit_bonus=n):
        await msg.answer(f"✅ Loại <b>#{cid}</b>: mua 1 acc tặng <b>{n} credits</b>.",
                         parse_mode="HTML")
    else:
        await msg.answer("❌ Không tìm thấy loại này.")


@router.message(Command("quadoi"))
async def on_quadoi(msg: Message):
    """Admin: chọn loại acc làm quà đổi điểm loyalty. /quadoi <id_loại> [số_điểm] | /quadoi 0 để tắt"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        cid = db.loyalty_redeem_cat()
        c = db.acc_category_get(cid) if cid else None
        pts = db.get_setting("loyalty_redeem_points", "10") or "10"
        await msg.answer(
            f"🎁 Quà đổi điểm hiện tại: <b>{html.escape(c['name']) if c else 'chưa cài'}</b> — <b>{html.escape(str(pts))}</b> điểm.\n"
            f"Cú pháp: <code>/quadoi &lt;id_loại&gt; [số_điểm]</code> (xem id: /kho)",
            parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    if cid and not db.acc_category_get(cid):
        await msg.answer("❌ Không tìm thấy loại này.")
        return
    db.set_setting("loyalty_redeem_cat", str(cid))
    pts_txt = ""
    if len(parts) >= 3:
        try:
            pts = int(parts[2])
            if pts > 0:
                db.set_setting("loyalty_redeem_points", str(pts))
                pts_txt = f" — <b>{pts}</b> điểm"
        except Exception:
            pass
    c = db.acc_category_get(cid) if cid else None
    await msg.answer(f"✅ Quà đổi điểm: <b>{html.escape(c['name']) if c else 'đã tắt'}</b>{pts_txt}.",
                     parse_mode="HTML")


@router.message(Command("loyaltyrandom"))
async def on_loyaltyrandom(msg: Message):
    """Admin: cài điểm loyalty ngẫu nhiên sau mua.
    /loyaltyrandom <min> <max> <cap_user> [min_order] [daily_max] (cap=0 tắt)"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 4:
        cfg = _loyalty_random_cfg()
        if not cfg:
            cur = "đang tắt"
        else:
            cur = (f"{cfg[0]}–{cfg[1]} điểm/lần, tối đa {cfg[2]} điểm/user"
                   f"{f', đơn ≥ {vnd(cfg[3])}' if cfg[3] > 0 else ''}"
                   f"{f', tối đa {cfg[4]} lượt/ngày' if cfg[4] > 0 else ''}")
        await msg.answer(
            f"🎲 Điểm ngẫu nhiên hiện tại: <b>{html.escape(cur)}</b>\n"
            f"Cú pháp: <code>/loyaltyrandom &lt;min&gt; &lt;max&gt; &lt;cap_user&gt; "
            f"[đơn_tối_thiểu] [lượt_tối_đa_ngày]</code>\n"
            f"<i>cap_user = 0 để tắt tính năng</i>",
            parse_mode="HTML")
        return
    try:
        lo, hi, cap = int(parts[1]), int(parts[2]), int(parts[3])
        min_order = int(parts[4]) if len(parts) > 4 else int(
            db.get_setting("loyalty_random_min_order", "0") or 0)
        daily_max = int(parts[5]) if len(parts) > 5 else int(
            db.get_setting("loyalty_random_daily_max", "0") or 0)
    except Exception:
        await msg.answer("❌ Các giá trị phải là số.")
        return
    if lo <= 0 or hi < lo or cap < 0 or min_order < 0 or daily_max < 0:
        await msg.answer("❌ Giá trị không hợp lệ (cần 0 < min ≤ max, cap ≥ 0).")
        return
    db.set_setting("loyalty_random_min", str(lo))
    db.set_setting("loyalty_random_max", str(hi))
    db.set_setting("loyalty_random_cap", str(cap))
    db.set_setting("loyalty_random_min_order", str(min_order))
    db.set_setting("loyalty_random_daily_max", str(daily_max))
    db.admin_audit_add(msg.from_user.id, msg.from_user.full_name, "loyaltyrandom",
                       f"min={lo} max={hi} cap={cap} min_order={min_order} daily_max={daily_max}")
    extra = ""
    if min_order > 0:
        extra += f", đơn ≥ {vnd(min_order)}"
    if daily_max > 0:
        extra += f", tối đa {daily_max} lượt/ngày"
    await msg.answer(
        f"✅ Điểm ngẫu nhiên: <b>{lo}–{hi}</b> điểm/lần, tối đa <b>{cap}</b> điểm/user"
        f"{extra}"
        + (" (đang tắt)" if cap <= 0 else "") + ".",
        parse_mode="HTML")


@router.message(Command("xoaloai"))
async def on_xoaloai(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/xoaloai &lt;id_loại&gt;</code>",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    if db.acc_category_update(cid, active=0):
        await msg.answer(f"✅ Đã ẩn loại <b>#{cid}</b> khỏi shop.", parse_mode="HTML")
    else:
        await msg.answer("❌ Không tìm thấy loại này.")


@router.message(Command("hienloai"))
async def on_hienloai(msg: Message):
    """Admin: hiện lại loại acc đã ẩn bằng /xoaloai. Cú pháp: /hienloai <id_loại>"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/hienloai &lt;id_loại&gt;</code> (xem id: /kho)",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    if db.acc_category_update(cid, active=1, hidden=0):
        c = db.acc_category_get(cid)
        nm = html.escape(c["name"]) if c else f"#{cid}"
        await msg.answer(f"✅ Đã hiện lại loại <b>{nm}</b> trên shop.\n"
                         f"Nhập hàng: <code>/themacc {cid}</code>",
                         parse_mode="HTML")
    else:
        await msg.answer("❌ Không tìm thấy loại này.")


@router.message(Command("recheck"))
async def on_recheck(msg: Message):
    """Admin: chạy re-check LIVE toàn bộ kho ngay (không đợi lịch định kỳ).
    Acc DIE → cách ly + báo cáo. Đổi chu kỳ: /setrecheck <số_ngày>."""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) > 1 and parts[1].isdigit():
        days = max(1, min(30, int(parts[1])))
        db.set_setting("stock_recheck_days", str(days))
        await msg.answer(f"✅ Đã đặt re-check kho mỗi <b>{days} ngày</b> (3h sáng).", parse_mode="HTML")
        return
    n = db.get_conn().execute(
        "SELECT COUNT(*) n FROM acc_stock WHERE status='AVAILABLE'").fetchone()["n"]
    await msg.answer(f"🔄 Bắt đầu re-check <b>{n}</b> acc trong kho... Xong mik báo kết quả.",
                     parse_mode="HTML")
    asyncio.create_task(poller._run_stock_recheck(manual=True))


@router.message(Command("xoadie"))
async def on_xoadie(msg: Message):
    """Admin: xóa hẳn các acc đã bị cách ly DIE. /xoadie [id_loại] yes"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    cid = None
    if len(parts) > 1 and parts[1].isdigit():
        cid = int(parts[1])
    confirm = (parts[-1].lower() == "yes")
    if cid is not None:
        total = db.acc_stock_die_count(cid)
        scope = f"loại #{cid}"
    else:
        total = db.get_conn().execute(
            "SELECT COUNT(*) n FROM acc_stock WHERE status='DIE'").fetchone()["n"]
        scope = "toàn bộ các loại"
    if total <= 0:
        await msg.answer("✅ Không có acc DIE nào cần dọn.")
        return
    if not confirm:
        await msg.answer(
            f"🗑 Có <b>{total}</b> acc DIE đang bị cách ly ({scope}).\n"
            f"Xóa hẳn: <code>/xoadie{f' {cid}' if cid else ''} yes</code>\n"
            "⚠️ Không khôi phục được.",
            parse_mode="HTML")
        return
    n = db.acc_stock_delete_die(cid)
    await msg.answer(f"✅ Đã xóa hẳn <b>{n}</b> acc DIE ({scope}).", parse_mode="HTML")


@router.message(Command("xoakho"))
async def on_xoakho(msg: Message):
    """Admin: XÓA TOÀN BỘ acc CHƯA BÁN trong kho của 1 loại (không khôi phục được).
    Cú pháp: /xoakho <id_loại> yes"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/xoakho &lt;id_loại&gt; yes</code>\n"
                         "⚠️ Xóa toàn bộ acc <b>chưa bán</b> của loại đó, "
                         "<b>không khôi phục được</b>.",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    n_av = db.get_conn().execute(
        "SELECT COUNT(*) n FROM acc_stock WHERE cat_id=? AND status='AVAILABLE'",
        (cid,)).fetchone()["n"]
    if not n_av:
        await msg.answer("📦 Loại này không còn acc chưa bán nào.")
        return
    if len(parts) < 3 or parts[2].lower() != "yes":
        c = db.acc_category_get(cid)
        nm = html.escape(c["name"]) if c else f"#{cid}"
        await msg.answer(
            f"⚠️ Bạn chắc chắn muốn <b>XÓA {n_av} ACC CHƯA BÁN</b> của loại "
            f"<b>{nm}</b>?\n"
            f"Nhớ <code>/xuatkho {cid}</code> sao lưu trước nếu cần.\n"
            f"Gõ lại: <code>/xoakho {cid} yes</code> để xác nhận.",
            parse_mode="HTML")
        return
    n = db.acc_stock_delete_available(cid)
    await msg.answer(f"🗑 Đã xóa <b>{n}</b> acc chưa bán của loại <b>#{cid}</b>.",
                     parse_mode="HTML")


@router.message(Command("xoahan"))
async def on_xoahan(msg: Message):
    """Admin: XÓA HẲN loại acc khỏi DB (không khôi phục được).
    Cú pháp: /xoahan <id_loại> yes"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/xoahan &lt;id_loại&gt; yes</code>\n"
                         "⚠️ Xóa loại acc, <b>không khôi phục được</b>. "
                         "Chỉ xóa được khi không còn acc <b>chưa bán</b> "
                         "(xóa kho trước: <code>/xoakho &lt;id&gt; yes</code>).",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    if len(parts) < 3 or parts[2].lower() != "yes":
        c = db.acc_category_get(cid)
        nm = html.escape(c["name"]) if c else f"#{cid}"
        await msg.answer(
            f"⚠️ Bạn chắc chắn muốn <b>XÓA HẲN</b> loại <b>{nm}</b>?\n"
            f"Gõ lại: <code>/xoahan {cid} yes</code> để xác nhận.",
            parse_mode="HTML")
        return
    ok, info = db.acc_category_delete_hard(cid)
    await msg.answer(f"{'🗑' if ok else '❌'} {info}", parse_mode="HTML")


@router.message(Command("donhang"))
async def on_donhang(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    limit = int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else 10
    rev = db.acc_revenue()
    orders = db.acc_recent_orders(limit)
    lines = ["🧾 <b>ĐƠN HÀNG SHOP ACC</b>", "━━━━━━━━━━━━",
             f"Tổng: <b>{rev['count']}</b> đơn — doanh thu <b>{vnd(rev['total'])}</b>", ""]
    for cat in rev["by_cat"]:
        lines.append(f"• {html.escape(cat['name'])}: {cat['n']} đơn — {vnd(cat['s'])}")
    lines.append("")
    for o in orders:
        o = dict(o)
        who = f"@{o['username']}" if o["username"] else f"<code>{o['tg_id']}</code>"
        lines.append(f"#{o['id']} {html.escape(o['cat_name'])} — {vnd(o['price'])} — {who} "
                     f"({vn_time_str(ts=o['created_at'])})")
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("bhdon"))
async def on_bhdon(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    claims = db.acc_warranty_pending()
    if not claims:
        await msg.answer("🛡 Không có yêu cầu bảo hành nào đang chờ.")
        return
    lines = ["🛡 <b>BẢO HÀNH ĐANG CHỜ</b>", "━━━━━━━━━━━━", ""]
    for w in claims:
        w = dict(w)
        lines.append(
            f"#{w['id']} — đơn <b>#{w['order_id']}</b> — {html.escape(w['cat_name'])}\n"
            f"👤 UID <code>{html.escape(w['uid'] or '')}</code> | khách <code>{w['tg_id']}</code>\n"
            f"🔍 Bot check: <b>{html.escape(w['check_result'] or '?')}</b> | "
            f"{vn_time_str(ts=w['created_at'])}{' | 📸 có ảnh' if w.get('evidence') else ''}\n"
            f"Xong: <code>/bhdone {w['id']}</code>"
        )
    await msg.answer("\n\n".join(lines), parse_mode="HTML")


@router.message(Command("bhdone"))
async def on_bhdone(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2 or not parts[1].isdigit():
        await msg.answer("⚠️ Cú pháp: <code>/bhdone &lt;id_khiếu_nại&gt;</code>",
                         parse_mode="HTML")
        return
    cid = int(parts[1])
    cl = db.acc_warranty_get(cid)
    if not cl:
        await msg.answer(f"❌ Không tìm thấy khiếu nại <b>#{cid}</b>.", parse_mode="HTML")
        return
    db.acc_warranty_set_status(cid, "DONE", handled_by=msg.from_user.id)
    handler_name = html.escape(msg.from_user.full_name or "")
    await msg.answer(
        f"✅ Đã đánh dấu xong khiếu nại <b>#{cid}</b>.\n"
        f"👤 Người duyệt: <b>{handler_name}</b> <code>{msg.from_user.id}</code> — "
        f"{vn_time_str()}",
        parse_mode="HTML")


@router.message(Command("accinfo"))
async def on_accinfo(msg: Message):
    """5.9 Admin: truy xuất hành trình 1 acc. Cú pháp: /accinfo <uid>"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/accinfo &lt;uid&gt;</code>",
                         parse_mode="HTML")
        return
    e = html.escape
    s = db.acc_stock_by_uid(parts[1])
    if not s:
        # acc đã bán (xóa khỏi kho) -> tra trong đơn hàng
        o = db.get_conn().execute(
            "SELECT o.*, c.name AS cat_name, u.username FROM acc_orders o "
            "LEFT JOIN acc_categories c ON c.id=o.cat_id "
            "LEFT JOIN tg_users u ON u.tg_id=o.tg_id "
            "WHERE o.uid=? ORDER BY o.id DESC LIMIT 1", (parts[1].strip(),)).fetchone()
        if not o:
            await msg.answer("❌ Không tìm thấy acc này trong kho.")
            return
        o = dict(o)
        who = f"@{o['username']}" if o.get("username") else ""
        await msg.answer(
            "🔎 <b>HÀNH TRÌNH ACC</b>\n━━━━━━━━━━━━\n"
            f"👤 UID: <code>{e(o['uid'] or '')}</code>\n"
            f"📦 Loại: <b>{e(o['cat_name'] or '')}</b>\n"
            f"📌 Trạng thái: <b>ĐÃ BÁN (xóa khỏi kho)</b>\n\n"
            f"🧾 Đơn #{o['id']} — khách <code>{o['tg_id']}</code> {e(who)} — "
            f"{vnd(o['price'])} — {vn_time_str(ts=o['created_at'])}",
            parse_mode="HTML")
        return
    lines = ["🔎 <b>HÀNH TRÌNH ACC</b>", "━━━━━━━━━━━━",
             f"👤 UID: <code>{e(s['uid'] or '')}</code>",
             f"📦 Loại: <b>{e(s['cat_name'] or '')}</b>",
             f"📥 Nhập kho: {vn_time_str(ts=s['added_at'])}"
             + (f" — {e(s['batch'])}" if s.get("batch") else ""),
             f"📌 Trạng thái: <b>{e(s['status'] or '')}</b>", ""],
    orders = db.acc_stock_orders(s["id"])
    if orders:
        lines.append("🧾 <b>Đã bán:</b>")
        for o in orders:
            who = f"@{o['username']}" if o.get("username") else ""
            lines.append(f"• Đơn #{o['id']} — khách <code>{o['tg_id']}</code> {e(who)} — "
                         f"{vnd(o['price'])} — {vn_time_str(ts=o['created_at'])}")
    else:
        lines.append("🆕 Chưa bán cho ai.")
    n_claim = db.acc_stock_claim_count(s["id"])
    lines.append(f"\n🛡 Số lần bảo hành: <b>{n_claim}</b>")
    await msg.answer("\n".join(lines), parse_mode="HTML")


# ------------------------- 5.11 FAQ TỰ ĐỘNG -------------------------
@router.message(Command("faq"))
async def on_faq(msg: Message):
    """Admin: xem danh sách câu hỏi tự động."""
    if not _is_admin(msg.from_user.id):
        return
    items = db.shop_faq_list()
    n_default = len(db._DEFAULT_FAQ)
    lines = ["❓ <b>FAQ TỰ ĐỘNG</b> — khách nhắn đúng từ khóa → bot tự trả lời",
             "━━━━━━━━━━━━", ""]
    for i, item in enumerate(items, 1):
        tag = "mặc định" if i <= n_default else f"tự thêm #{i - n_default}"
        kws = ", ".join(item.get("kw", []))
        lines.append(f"{i}. [{tag}] <code>{html.escape(kws)}</code>")
    lines += ["",
              "Thêm: <code>/themcauhoi từ khóa, từ khóa 2 | câu trả lời</code>",
              "Xóa câu tự thêm: <code>/xoacauhoi &lt;số&gt;</code>"]
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("themcauhoi"))
async def on_themcauhoi(msg: Message):
    """Admin: thêm câu hỏi tự động. /themcauhoi từ khóa, từ khóa 2 | câu trả lời"""
    if not _is_admin(msg.from_user.id):
        return
    raw = (msg.text or "").split(maxsplit=1)
    if len(raw) < 2 or "|" not in raw[1]:
        await msg.answer(
            "⚠️ Cú pháp: <code>/themcauhoi từ khóa, từ khóa 2 | câu trả lời</code>",
            parse_mode="HTML")
        return
    kw_part, ans = raw[1].split("|", 1)
    n = db.shop_faq_add(kw_part, ans)
    if n == -1:
        await msg.answer("❌ Thiếu từ khóa hoặc câu trả lời.")
        return
    await msg.answer(f"✅ Đã thêm câu hỏi tự động #{n}. Xem: /faq")


@router.message(Command("xoacauhoi"))
async def on_xoacauhoi(msg: Message):
    """Admin: xóa câu hỏi tự thêm. /xoacauhoi <số>"""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2 or not parts[1].isdigit():
        await msg.answer(
            "⚠️ Cú pháp: <code>/xoacauhoi &lt;số_câu_tự_thêm&gt;</code> (xem số ở /faq)",
            parse_mode="HTML")
        return
    if db.shop_faq_del(int(parts[1])):
        await msg.answer(f"✅ Đã xóa câu hỏi #{parts[1]}.")
    else:
        await msg.answer("❌ Không tìm thấy câu này (chỉ xóa được câu tự thêm).")


@router.message(StateFilter(None), F.text)
async def on_shop_faq_auto(msg: Message):
    """5.11 Trả lời tự động khi khách hỏi về shop.
    Chỉ chạy ở trạng thái thường (không chen vào lúc đang nhập số tiền/file...),
    và chỉ khi không có handler nào khác bắt tin nhắn."""
    text = msg.text or ""
    if text.startswith("/"):
        return
    ans = db.shop_faq_match(text)
    if ans:
        await msg.answer(ans, parse_mode="HTML")


# ================= SHOP MỞ RỘNG — GĐ4/GĐ5 đợt 4 =================
# 4.3 giá khan hiếm, 4.6 upsell, 4.8 hộp mù, 4.9 đặt cọc, 4.10 giờ vàng,
# 4.11 đánh giá, 5.2/5.15 NCC + chấm điểm, 5.14 lãi theo lô, 5.4 nhập kho NCC.

async def _fulfill_restock_subs(cat_id: int, bot, cat: dict):
    """Hàng về -> xử lý đăng ký báo hàng theo FIFO:
    - auto_buy=1: tự trừ ví shop, check LIVE rồi giao acc ngay.
    - auto_buy=0: chỉ nhắn báo hàng về.
    Sub nào ví không đủ tiền / chưa giao được thì GIỮ LẠI để thử đợt sau."""
    subs = db.acc_restock_subscribers(cat_id)
    if not subs:
        return
    up = db.shop_unit_price(cat)
    unit = up["price"]
    p5, p10, p20 = _shop_bulk_pcts()
    done_auto = 0
    for s in subs:
        tg_id = int(s["tg_id"])
        qty = max(1, int(s.get("qty") or 1))
        auto = int(s.get("auto_buy") or 0)
        if not auto:
            try:
                await bot.send_message(
                    tg_id,
                    f"🔔 <b>CÓ HÀNG RỒI!</b>\n\n"
                    f"📦 <b>{html.escape(cat['name'])}</b> vừa về "
                    f"<b>{db.acc_stock_count(cat_id)}</b> acc — "
                    f"giá {vnd(unit)}/acc.\n"
                    f"👉 Vào /shop mua ngay kẻo hết!",
                    parse_mode="HTML")
            except Exception:
                pass
            db.acc_unsub_restock(tg_id, cat_id)
            continue
        # Tự mua: giá như shop (giảm mua nhiều theo số lượng thực mua)
        stock = db.acc_stock_count(cat_id)
        buy_qty = min(qty, stock)
        if buy_qty <= 0:
            try:
                await bot.send_message(
                    tg_id,
                    f"😔 <b>Hàng bạn đặt trước:</b> <b>{html.escape(cat['name'])}</b> "
                    f"vừa về nhưng đã hết trước khi bot kịp giao. "
                    f"Bot giữ đơn đặt trước cho đợt hàng sau nhé!",
                    parse_mode="HTML")
            except Exception:
                pass
            continue  # giữ sub cho đợt sau
        bulk_pct = p20 if buy_qty >= 20 else (p10 if buy_qty >= 10 else (p5 if buy_qty >= 5 else 0))
        final = unit * buy_qty * (100 - bulk_pct) // 100
        u = db.get_user(tg_id)
        bal = int(u["shop_balance"] or 0) if u else 0
        if bal < final:
            try:
                await bot.send_message(
                    tg_id,
                    f"🔔 <b>HÀNG BẠN ĐẶT TRƯỚC ĐÃ VỀ!</b>\n\n"
                    f"📦 <b>{html.escape(cat['name'])}</b> — {buy_qty} acc = "
                    f"<b>{vnd(final)}</b> (ví shop bạn: {vnd(bal)}).\n"
                    f"👉 Nạp thêm bằng /napshop để bot tự giao acc ngay nhé! "
                    f"Bot giữ đơn đặt trước cho bạn.",
                    parse_mode="HTML")
            except Exception:
                pass
            continue  # giữ sub cho đợt sau
        if not db.adjust_shop_balance(tg_id, -final, f"dat_truoc:{cat_id}x{buy_qty}"):
            continue
        sold_orders, _sell_fail = await _sell_live_stock(
            bot, tg_id, final, buy_qty,
            lambda seen: db.acc_stock_pick_candidates(cat_id, buy_qty + 4, seen),
            check_live=_cat_live_check(cat_id))
        if not sold_orders:
            db.add_shop_balance_only(tg_id, final, "hoan_tien_dat_truoc")
            try:
                await bot.send_message(
                    tg_id,
                    "😔 <b>Hàng bạn đặt trước:</b> shop vừa kiểm tra lại, hiện chưa có "
                    "acc <b>LIVE</b> để giao. Tiền đã được hoàn vào ví shop — "
                    "có hàng live bot sẽ giao ngay cho bạn!",
                    parse_mode="HTML")
            except Exception:
                pass
            continue  # giữ sub cho đợt sau
        db.acc_unsub_restock(tg_id, cat_id)
        done_auto += 1
        for order in sold_orders:
            order_id = order["id"]
            try:
                await bot.send_message(
                    tg_id,
                    f"⚡ <b>ĐẶT TRƯỚC THÀNH CÔNG — ĐÃ GIAO ACC</b>\n\n"
                    f"📦 Loại: <b>{html.escape(cat['name'])}</b>\n"
                    f"🧾 Đơn hàng: <b>#{order_id}</b> | 💰 {vnd(order['price'])}\n"
                    f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
                    f"{_live_line(cat_id)}\n"
                    f"{_pickup_suffix()}",
                    parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
            except Exception:
                pass
        try:
            tickets = db.spin_add_tickets(tg_id, buy_qty)
            per = int(db.get_setting("loyalty_per_vnd", "100000") or 100000)
            pts = final // per if per > 0 else 0
            if pts > 0:
                db.loyalty_add(tg_id, pts, f"dat_truoc:{cat_id}")
            await bot.send_message(
                tg_id,
                f"🎁 Đặt trước {buy_qty} acc: nhận <b>{buy_qty} vé quay</b> /quay"
                + (f" + <b>{pts} điểm</b> loyalty" if pts > 0 else "")
                + f" (đang có {tickets} vé).",
                parse_mode="HTML")
        except Exception:
            pass
        try:
            await _notify_purchase_admin(bot, _buyer_shim(tg_id), sold_orders)
        except Exception:
            pass
    if done_auto:
        try:
            await _notify_admin_smart(
                bot,
                f"⚡ <b>Tự giao hàng đặt trước:</b> {html.escape(cat['name'])} — "
                f"đã giao <b>{done_auto}</b> đơn đặt trước.")
        except Exception:
            pass


async def _fulfill_deposits(cat_id: int, bot, cat: dict):
    """4.9 Hàng về -> giao acc cho người đặt cọc theo FIFO, trừ nốt tiền."""
    deps = db.acc_deposit_waiting(cat_id)
    if not deps:
        return
    up = db.shop_unit_price(cat)
    unit = up["price"]
    done = 0
    for d in deps:
        d = dict(d)
        tg_id = int(d["tg_id"])
        if db.acc_stock_count(cat_id) <= 0:
            break
        rest = unit - int(d["amount"])
        u = db.get_user(tg_id)
        bal = int(u["shop_balance"] or 0) if u else 0
        if rest > 0 and bal < rest:
            try:
                await bot.send_message(
                    tg_id,
                    f"🔔 <b>HÀNG BẠN ĐẶT CỌC ĐÃ VỀ!</b>\n\n"
                    f"📦 <b>{html.escape(cat['name'])}</b> — còn thiếu <b>{vnd(rest)}</b> "
                    f"(ví shop bạn: {vnd(bal)}).\n"
                    f"👉 Nạp thêm bằng /napshop để bot tự giao acc ngay nhé!",
                    parse_mode="HTML")
            except Exception:
                pass
            continue
        if rest > 0:
            if not db.adjust_shop_balance(tg_id, -rest, f"dat_coc_giao:{cat_id}"):
                continue
        sold_orders, _sell_fail = await _sell_live_stock(
            bot, tg_id, unit, 1,
            lambda seen: db.acc_stock_pick_candidates(cat_id, 5, seen),
            check_live=_cat_live_check(cat_id))
        if not sold_orders:
            if rest > 0:
                db.add_shop_balance_only(tg_id, rest, "hoan_tien_dat_coc")
            try:
                await bot.send_message(
                    tg_id,
                    "😔 <b>Hàng bạn đặt cọc:</b> shop vừa kiểm tra lại, hiện chưa có "
                    "acc <b>LIVE</b> để giao. Tiền cọc vẫn được giữ — có hàng live "
                    "bot sẽ giao ngay cho bạn!",
                    parse_mode="HTML")
            except Exception:
                pass
            break
        order = sold_orders[0]
        order_id = order["id"]
        db.acc_deposit_set_status(d["id"], "DONE", order_id)
        done += 1
        try:
            await bot.send_message(
                tg_id,
                f"✅ <b>ĐẶT CỌC THÀNH CÔNG — ĐÃ GIAO ACC</b>\n\n"
                f"📦 Loại: <b>{html.escape(cat['name'])}</b>\n"
                f"💰 Đã cọc: {vnd(d['amount'])} + trừ thêm {vnd(rest)}\n"
                f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
                f"{_live_line(cat_id)}\n"
                f"{_pickup_suffix()}",
                parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
        except Exception:
            pass
        await _notify_purchase_admin(bot, _buyer_shim(tg_id), [order])
    if done:
        await _notify_admin_smart(
            bot,
            f"💰 <b>Tự giao hàng đặt cọc:</b> {html.escape(cat['name'])} — "
            f"đã giao <b>{done}</b> acc cho khách đặt cọc.")


async def _do_deposit(tg_id: int, cat_id: int, msg, bot):
    c = db.acc_category_get(cat_id)
    if not c:
        await msg.answer("❌ Loại acc không tồn tại.")
        return
    c = dict(c)
    if db.acc_stock_count(cat_id) > 0:
        await msg.answer("✅ Loại này đang còn hàng — bạn mua luôn ở /shop nhé!")
        return
    try:
        dep_pct = int(db.get_setting("deposit_pct", "30") or 30)
    except Exception:
        dep_pct = 30
    amt = int(c["price"]) * dep_pct // 100
    u = db.get_user(tg_id)
    bal = int(u["shop_balance"] or 0) if u else 0
    if bal < amt:
        await msg.answer(
            f"❌ Ví shop không đủ đặt cọc.\nCọc {dep_pct}% = <b>{vnd(amt)}</b>, "
            f"ví shop bạn: <b>{vnd(bal)}</b>.\nNạp thêm bằng /napshop nhé."
            + _SHOP_WALLET_HINT,
            parse_mode="HTML")
        return
    if not db.adjust_shop_balance(tg_id, -amt, f"dat_coc:{cat_id}"):
        await msg.answer("❌ Ví shop không đủ đặt cọc." + _SHOP_WALLET_HINT, parse_mode="HTML")
        return
    dep_id = db.acc_deposit_create(tg_id, cat_id, amt)
    await msg.answer(
        f"✅ <b>ĐẶT CỌC THÀNH CÔNG #{dep_id}</b>\n\n"
        f"📦 Loại: <b>{html.escape(c['name'])}</b>\n"
        f"💰 Đã cọc: <b>{vnd(amt)}</b> ({dep_pct}% giá {vnd(c['price'])})\n\n"
        f"👉 Khi có hàng về, bot <b>tự giao acc ngay</b> và trừ nốt "
        f"<b>{vnd(int(c['price']) - amt)}</b> từ ví shop của bạn.\n"
        f"Xem/hủy cọc: /coclist",
        parse_mode="HTML")


@router.callback_query(F.data.startswith("acccoc:"))
async def on_acc_deposit_cb(cb: CallbackQuery):
    await cb.answer()
    try:
        cat_id = int(cb.data.split(":", 1)[1])
    except Exception:
        return
    await _do_deposit(cb.from_user.id, cat_id, cb.message, cb.bot)


@router.message(Command("coc"))
async def on_coc(msg: Message):
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/coc &lt;id_loại&gt;</code> (xem id: /shop)",
                         parse_mode="HTML")
        return
    try:
        cat_id = int(parts[1])
    except Exception:
        await msg.answer("❌ ID loại phải là số.")
        return
    await _do_deposit(msg.from_user.id, cat_id, msg, msg.bot)


@router.message(Command("coclist"))
async def on_coclist(msg: Message):
    deps = db.acc_deposit_list(msg.from_user.id)
    if not deps:
        await msg.answer("💰 Bạn chưa đặt cọc loại acc nào.")
        return
    lines = ["💰 <b>ĐẶT CỌC CỦA BẠN</b>", "━━━━━━━━━━━━", ""]
    st_map = {"WAITING": "⏳ chờ hàng", "DONE": "✅ đã giao", "CANCELLED": "❌ đã hủy"}
    for d in deps:
        d = dict(d)
        lines.append(
            f"#{d['id']} — <b>{html.escape(d['cat_name'] or '')}</b>\n"
            f"   Cọc: {vnd(d['amount'])} — {st_map.get(d['status'], d['status'])}")
    lines += ["", "Hủy cọc (hoàn tiền): <code>/huycoc &lt;id_cọc&gt;</code>"]
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("huycoc"))
async def on_huycoc(msg: Message):
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/huycoc &lt;id_cọc&gt;</code>",
                         parse_mode="HTML")
        return
    try:
        dep_id = int(parts[1])
    except Exception:
        await msg.answer("❌ ID cọc phải là số.")
        return
    amt = db.acc_deposit_cancel(dep_id, msg.from_user.id)
    if amt is None:
        await msg.answer("❌ Không tìm thấy khoản cọc chờ hàng này của bạn.")
        return
    db.add_shop_balance_only(msg.from_user.id, amt, f"huy_coc:{dep_id}")
    await msg.answer(f"✅ Đã hủy cọc <b>#{dep_id}</b>, hoàn <b>{vnd(amt)}</b> vào ví shop.",
                     parse_mode="HTML")


# ---- 4.8 Hộp mù acc ----
@router.callback_query(F.data == "accmystery")
async def on_mystery(cb: CallbackQuery):
    await cb.answer()
    try:
        m_price = int(db.get_setting("mystery_price", "0") or 0)
    except Exception:
        m_price = 0
    if m_price <= 0:
        await cb.message.answer("🎁 Hộp mù hiện đang tắt.")
        return
    m_cats = [dict(x) for x in db.acc_category_list()
              if int(x["mystery_eligible"] or 0) and db.acc_stock_count(x["id"]) > 0]
    if not m_cats:
        await cb.message.answer("⛔ Hộp mù hiện hết hàng, quay lại sau nhé!")
        return
    names = ", ".join(html.escape(x["name"]) for x in m_cats[:10])
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"🎁 Mở hộp — {vnd(m_price)}",
                              callback_data="accmysterybuy")],
        [InlineKeyboardButton(text="🔙 Quay lại shop", callback_data="accshop_back")],
    ])
    await cb.message.answer(
        f"🎁 <b>HỘP MÙ ACC — {vnd(m_price)}</b>\n\n"
        f"Mở ngẫu nhiên <b>1 acc</b> từ các loại:\n{names}\n\n"
        f"🍀 Hên xui trúng acc xịn giá rẻ!",
        parse_mode="HTML", reply_markup=kb)


@router.callback_query(F.data == "accmysterybuy")
async def on_mystery_buy(cb: CallbackQuery):
    await cb.answer()
    tg_id = cb.from_user.id
    try:
        m_price = int(db.get_setting("mystery_price", "0") or 0)
    except Exception:
        m_price = 0
    if m_price <= 0:
        return
    u = db.get_user(tg_id)
    bal = int(u["shop_balance"] or 0) if u else 0
    if bal < m_price:
        await cb.message.answer(
            f"❌ Ví shop không đủ! Hộp mù {vnd(m_price)}, bạn có {vnd(bal)}.\n"
            f"Nạp thêm bằng /napshop nhé." + _SHOP_WALLET_HINT,
            parse_mode="HTML")
        return
    if not db.adjust_shop_balance(tg_id, -m_price, "mua_hop_mu"):
        await cb.message.answer("❌ Ví shop không đủ." + _SHOP_WALLET_HINT, parse_mode="HTML")
        return
    await cb.message.answer("🔍 <b>Đang kiểm tra chất lượng acc...</b>", parse_mode="HTML")
    sold_orders, _sell_fail = await _sell_live_stock(
        cb.bot, tg_id, m_price, 1,
        lambda seen: db.acc_mystery_pick_candidates(5, seen))
    if not sold_orders:
        db.add_shop_balance_only(tg_id, m_price, "hoan_tien_hop_mu")
        await cb.message.answer(
            "😔 Hộp mù hiện hết acc <b>LIVE</b> (shop vừa kiểm tra lại), "
            "tiền đã được hoàn lại. Bạn quay lại sau nhé!",
            parse_mode="HTML")
        return
    order = sold_orders[0]
    order_id = order["id"]
    cat_name = order.get("cat_name") or ""
    db.spin_add_tickets(tg_id, 1)
    try:
        per = int(db.get_setting("loyalty_per_vnd", "100000") or 100000)
    except Exception:
        per = 100000
    if per > 0 and m_price // per:
        db.loyalty_add(tg_id, m_price // per, "mua_hop_mu")
    order = dict(order)
    await cb.message.answer(
        f"🎁 <b>MỞ HỘP THÀNH CÔNG!</b>\n"
        f"🍀 Bạn trúng acc loại: <b>{html.escape(cat_name)}</b>\n"
        f"👤 UID: <code>{html.escape(order['uid'] or '')}</code>\n"
        f"🟢 <i>Đã kiểm tra LIVE trước khi giao</i>\n\n"
        f"{_pickup_suffix()}",
        parse_mode="HTML", reply_markup=_acc_delivery_kb(order_id))
    await _notify_purchase_admin(cb.bot, cb.from_user, [order])


@router.message(Command("hopmu"))
async def on_hopmu(msg: Message):
    """Admin: /hopmu <id> bật/tắt loại acc tham gia hộp mù."""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/hopmu &lt;id_loại&gt;</code> (bật/tắt)",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    c = db.acc_category_get(cid)
    if not c:
        await msg.answer("❌ Không có loại này.")
        return
    new = 0 if int(c["mystery_eligible"] or 0) else 1
    db.acc_category_update(cid, mystery_eligible=new)
    await msg.answer(f"✅ Loại <b>{html.escape(c['name'])}</b> đã "
                     f"{'<b>THAM GIA</b> hộp mù' if new else 'rời khỏi hộp mù'}.",
                     parse_mode="HTML")


@router.message(Command("hopmugia"))
async def on_hopmugia(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/hopmugia &lt;giá&gt;</code> (0 = tắt hộp mù)",
                         parse_mode="HTML")
        return
    try:
        price = int(parts[1].replace(".", "").replace(",", ""))
    except Exception:
        await msg.answer("❌ Giá phải là số.")
        return
    db.set_setting("mystery_price", str(max(0, price)))
    await msg.answer(f"✅ Giá hộp mù: <b>{vnd(max(0, price))}</b>"
                     + (" (đang tắt)" if price <= 0 else ""), parse_mode="HTML")


# ---- 4.10 Giờ vàng (admin) ----
@router.message(Command("giovang"))
async def on_giovang(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        cur = db.get_setting("happy_hour_range", "20-22")
        pct = db.get_setting("happy_hour_pct", "10")
        hc = db.get_setting("happy_hour_cat", "0")
        on, _, _ = db.happy_hour_active()
        await msg.answer(
            f"⚡ <b>GIỜ VÀNG HIỆN TẠI</b>\nKhung giờ: <b>{cur}</b>\nGiảm: <b>{pct}%</b>\n"
            f"Áp dụng: {'toàn shop' if hc == '0' else f'loại #{hc}'}\n"
            f"Trạng thái: {'🟢 ĐANG DIỄN RA' if on else '⚪ ngoài khung giờ'}\n\n"
            f"Đặt: <code>/giovang &lt;id_loại|0=tất cả&gt; [giờ_bd-giờ_kt] [%giảm]</code>\n"
            f"Tắt: <code>/giovang off</code>",
            parse_mode="HTML")
        return
    if parts[1].lower() == "off":
        db.set_setting("happy_hour_pct", "0")
        await msg.answer("✅ Đã tắt giờ vàng.")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID loại phải là số (0 = tất cả).")
        return
    rng = parts[2] if len(parts) > 2 else db.get_setting("happy_hour_range", "20-22")
    pct = parts[3] if len(parts) > 3 else db.get_setting("happy_hour_pct", "10")
    db.set_setting("happy_hour_cat", str(cid))
    db.set_setting("happy_hour_range", rng)
    db.set_setting("happy_hour_pct", str(pct))
    await msg.answer(f"✅ Giờ vàng: giảm <b>{pct}%</b> khung <b>{rng}</b> cho "
                     f"{'toàn shop' if cid == 0 else f'loại #{cid}'}.", parse_mode="HTML")


# ---- 4.3 Giá khan hiếm (admin) ----
# ---- 4.11 Đánh giá có thưởng ----
@router.callback_query(F.data.startswith("accreview:"))
async def on_acc_review(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    parts = cb.data.split(":")
    try:
        order_id = int(parts[1])
    except Exception:
        return
    order = db.acc_get_order(order_id)
    if not order or int(order["tg_id"]) != cb.from_user.id:
        await cb.message.answer("❌ Không tìm thấy đơn.")
        return
    if len(parts) == 2:
        if db.acc_order_reviewed(order_id):
            await cb.message.answer("Bạn đã đánh giá đơn này rồi.")
            return
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⭐" * i, callback_data=f"accreview:{order_id}:{i}")]
            for i in range(1, 6)
        ])
        await cb.message.answer(
            f"⭐ <b>ĐÁNH GIÁ ĐƠN #{order_id}</b>\n"
            f"Chất lượng acc bạn nhận được thế nào?\n"
            f"Đánh giá xong được tặng credits ngay!",
            parse_mode="HTML", reply_markup=kb)
        return
    try:
        stars = int(parts[2])
    except Exception:
        return
    if not db.acc_review_add(order_id, cb.from_user.id, int(order["cat_id"]), stars):
        await cb.message.answer("Bạn đã đánh giá đơn này rồi.")
        return
    try:
        bonus = int(db.get_setting("review_bonus", "20") or 20)
    except Exception:
        bonus = 20
    new_cr = db.add_credits(cb.from_user.id, bonus, f"danh_gia:{order_id}")
    kb2 = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⏭ Bỏ qua", callback_data=f"accreview_skip:{order_id}")]
    ])
    await cb.message.answer(
        f"✅ Cảm ơn bạn đã đánh giá {'⭐' * stars}!\n"
        f"🎁 Tặng <b>{bonus} credits</b> (số dư credits: {new_cr}).\n\n"
        f"💬 Bạn có muốn <b>viết vài lời nhận xét</b> về acc không? "
        f"Nhận xét hay sẽ hiện trên shop cho khách khác tham khảo.",
        parse_mode="HTML", reply_markup=kb2)
    await state.update_data(review_order_id=order_id)
    await state.set_state(AccShopState.waiting_for_review_comment)


@router.message(AccShopState.waiting_for_review_comment, F.text)
async def on_review_comment(msg: Message, state: FSMContext):
    _t = (msg.text or "").strip()
    if _t.startswith("/") and _t.split()[0].lower() not in ("/huy", "/cancel"):
        await state.clear()
        await msg.answer("\U0001f6ab \u0110\u00e3 h\u1ee7y thao t\u00e1c \u0111ang nh\u1eadp. B\u1ea1n g\u00f5 l\u1ea1i l\u1ec7nh v\u1eeba r\u1ed3i nh\u00e9.")
        return
    data = await state.get_data()
    order_id = data.get("review_order_id")
    await state.clear()
    if not order_id or not db.acc_order_reviewed(order_id):
        return
    db.acc_review_set_comment(order_id, msg.text or "")
    await msg.answer("💬 Cảm ơn nhận xét của bạn! Lời nhận xét sẽ hiện trên shop cho khách khác tham khảo. 🙏")


@router.callback_query(F.data.startswith("accreview_skip:"))
async def on_acc_review_skip(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    await state.clear()
    await cb.message.answer("👌 Không sao! Cảm ơn bạn đã đánh giá. 🙏")


# ---- 5.15 Sổ NCC ----
@router.message(Command("themncc"))
async def on_themncc(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    rest = (msg.text or "").split(None, 1)
    if len(rest) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/themncc &lt;tên&gt; | &lt;liên hệ&gt;</code>",
                         parse_mode="HTML")
        return
    name, _, contact = rest[1].partition("|")
    sid = db.supplier_add(name.strip(), contact.strip())
    await msg.answer(f"✅ Đã thêm NCC <b>#{sid}</b>: {html.escape(name.strip())}",
                     parse_mode="HTML")


@router.message(Command("ncc"))
async def on_ncc(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    sups = db.supplier_list()
    if not sups:
        await msg.answer("🏭 Chưa có NCC nào. Thêm: /themncc")
        return
    lines = ["🏭 <b>SỔ NHÀ CUNG CẤP</b>", "━━━━━━━━━━━━", ""]
    for s in sups:
        s = dict(s)
        pct, alive, total = db.supplier_live_rate(s["id"])
        rate_txt = f"{pct}% sống ({alive}/{total})" if total else "chưa chấm điểm"
        stars = "⭐" * int(s["rating"] or 0) if s["rating"] else "—"
        lines.append(
            f"#{s['id']} <b>{html.escape(s['name'])}</b> {stars}\n"
            f"   📞 {html.escape(s['contact'] or '—')} | 📊 {rate_txt}")
    lines += ["", "Đánh giá tay: <code>/danhgiancc &lt;id&gt; &lt;sao 1-5&gt;</code>",
              "Chấm tỉ lệ sống: <code>/chamdiem &lt;id&gt; [số_ngày=7]</code>"]
    await msg.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("danhgiancc"))
async def on_danhgiancc(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 3:
        await msg.answer("⚠️ Cú pháp: <code>/danhgiancc &lt;id&gt; &lt;sao 1-5&gt;</code>",
                         parse_mode="HTML")
        return
    try:
        sid, stars = int(parts[1]), int(parts[2])
    except Exception:
        await msg.answer("❌ ID và sao phải là số.")
        return
    if db.supplier_rate(sid, stars):
        await msg.answer(f"✅ NCC #{sid} được đánh giá {'⭐' * max(1, min(5, stars))}.")
    else:
        await msg.answer("❌ Không tìm thấy NCC.")


@router.message(Command("chamdiem"))
async def on_chamdiem(msg: Message):
    """5.2 Chấm điểm NCC: check mẫu acc các lô nhập >= số_ngày trước."""
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/chamdiem &lt;id_ncc&gt; [số_ngày=7]</code>",
                         parse_mode="HTML")
        return
    try:
        sid = int(parts[1])
        days = int(parts[2]) if len(parts) > 2 else 7
    except Exception:
        await msg.answer("❌ ID và số ngày phải là số.")
        return
    sup = db.supplier_get(sid)
    if not sup:
        await msg.answer("❌ Không có NCC này.")
        return
    cutoff = int(time.time()) - days * 86400
    batches = db.get_conn().execute(
        "SELECT batch FROM acc_batches WHERE supplier_id=? AND created_at<=?",
        (sid, cutoff)).fetchall()
    if not batches:
        await msg.answer(f"🏭 NCC <b>{html.escape(sup['name'])}</b> chưa có lô nào nhập "
                         f"trước {days} ngày để chấm.", parse_mode="HTML")
        return
    bnames = [b["batch"] for b in batches]
    q = ("SELECT uid, batch FROM acc_stock WHERE batch IN (%s) AND uid != '' "
         "AND status='AVAILABLE' "
         "ORDER BY RANDOM() LIMIT 20" % ",".join("?" * len(bnames)))
    sample = db.get_conn().execute(q, tuple(bnames)).fetchall()
    if not sample:
        await msg.answer("❌ Không lấy được mẫu acc để chấm.")
        return
    wait = await msg.answer(f"🔍 Đang chấm <b>{len(sample)}</b> acc mẫu của NCC "
                            f"<b>{html.escape(sup['name'])}</b>...", parse_mode="HTML")
    from . import fb as fb_mod
    import asyncio
    sem = asyncio.Semaphore(5)
    alive = 0

    async def _one(uid):
        nonlocal alive
        async with sem:
            try:
                r = await fb_mod.check_uid(str(uid))
                if str(r.get("status", "")).lower() == "live":
                    alive += 1
            except Exception:
                pass

    await asyncio.gather(*[_one(s["uid"]) for s in sample])
    total = len(sample)
    pct = round(alive * 100 / total, 1)
    batch_txt = ", ".join(bnames[:5]) + ("..." if len(bnames) > 5 else "")
    db.supplier_score_add(sid, batch_txt, total, alive,
                          f"chấm sau {days} ngày, mẫu {total} acc")
    verdict = ("Hàng ngon, nên nhập tiếp 👍" if pct >= 80
               else "Hàng trung bình, cân nhắc ⚠️" if pct >= 50
               else "Hàng dở, nên đổi NCC 👎")
    await wait.edit_text(
        f"📊 <b>CHẤM ĐIỂM NCC #{sid} — {html.escape(sup['name'])}</b>\n\n"
        f"🧪 Mẫu: <b>{total}</b> acc (các lô: {html.escape(batch_txt)})\n"
        f"🟢 Sống: <b>{alive}</b> | 🔴 Die: <b>{total - alive}</b>\n"
        f"⭐ <b>Tỉ lệ sống: {pct}%</b>\n\n"
        f"<i>{verdict}</i>",
        parse_mode="HTML")


# ---- 5.14 Lãi theo lô ----
@router.message(Command("lo"))
async def on_lo(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    # Giá vốn/lãi chỉ chủ shop được xem — admin phụ không thấy
    if not _perms.is_super(msg.from_user.id):
        await msg.answer("🚫 Chỉ chủ shop mới xem được giá vốn/lãi.",
                         parse_mode="HTML")
        return
    parts = (msg.text or "").split()
    if len(parts) < 2:
        await msg.answer("⚠️ Cú pháp: <code>/lo &lt;id_loại&gt;</code>",
                         parse_mode="HTML")
        return
    try:
        cid = int(parts[1])
    except Exception:
        await msg.answer("❌ ID phải là số.")
        return
    batches = db.acc_batches_by_cat(cid)
    if not batches:
        await msg.answer("📦 Loại này chưa ghi nhận lô nào có giá vốn.\n"
                         "Nhập kho kèm giá vốn: <code>/themacc &lt;id&gt; [ncc_id] [giá_vốn]</code>",
                         parse_mode="HTML")
        return
    lines = ["💰 <b>LÃI THEO LÔ</b>", "━━━━━━━━━━━━", ""]
    t_cost = t_rev = 0
    for b in batches:
        p = db.acc_profit_by_batch(b["batch"])
        t_cost += p["cost"]
        t_rev += p["revenue"]
        emo = "🟢" if p["profit"] >= 0 else "🔴"
        lines.append(
            f"{emo} <b>{html.escape(p['batch'])}</b> ({p['count']} acc)\n"
            f"   Vốn: {vnd(p['cost'])} | Thu: {vnd(p['revenue'])} | "
            f"Lãi: <b>{vnd(p['profit'])}</b>")
    lines += ["", f"📊 <b>Tổng:</b> vốn {vnd(t_cost)} — thu {vnd(t_rev)} — "
                 f"<b>lãi {vnd(t_rev - t_cost)}</b>"]
    await msg.answer("\n".join(lines), parse_mode="HTML")


# ---- 5.4 Nhập kho tự động từ NCC ----
@router.message(Command("nccauto"))
async def on_nccauto(msg: Message):
    if not _is_admin(msg.from_user.id):
        return
    parts = (msg.text or "").split()
    if len(parts) >= 2 and parts[1].lower() == "off":
        db.set_setting("supplier_auto_url", "")
        await msg.answer("✅ Đã tắt nhập kho tự động từ NCC.")
        return
    if len(parts) < 3:
        cur = db.get_setting("supplier_auto_url", "")
        await msg.answer(
            "⚠️ Cú pháp: <code>/nccauto &lt;url_file_txt&gt; &lt;id_loại&gt; [ncc_id]</code>\n"
            "Tắt: <code>/nccauto off</code>\n\n"
            f"Cấu hình hiện tại: <code>{html.escape(cur or 'chưa có')}</code>\n"
            "Bot sẽ tự tải file mỗi sáng 6h, nhập kho + quét chất lượng.",
            parse_mode="HTML")
        return
    url = parts[1]
    if not (url.startswith("http://") or url.startswith("https://")):
        await msg.answer("❌ URL phải bắt đầu bằng http:// hoặc https://")
        return
    try:
        cid = int(parts[2])
        sid = int(parts[3]) if len(parts) > 3 else 0
    except Exception:
        await msg.answer("❌ ID loại/NCC phải là số.")
        return
    if not db.acc_category_get(cid):
        await msg.answer("❌ Không có loại acc này.")
        return
    db.set_setting("supplier_auto_url", url)
    db.set_setting("supplier_auto_cat", str(cid))
    db.set_setting("supplier_auto_supplier", str(sid))
    await msg.answer(f"✅ Đã bật nhập kho tự động mỗi sáng 6h:\n🔗 {html.escape(url)}\n"
                     f"📦 Loại #{cid}" + (f" — NCC #{sid}" if sid else ""),
                     parse_mode="HTML")


@router.callback_query(F.data.startswith("accok:"))
async def on_acc_ok(cb: CallbackQuery):
    """5.13 Khách xác nhận acc ổn sau tin hỏi thăm 24h."""
    await cb.answer("Cảm ơn bạn! Chúc bạn dùng acc vui vẻ 🍀")
